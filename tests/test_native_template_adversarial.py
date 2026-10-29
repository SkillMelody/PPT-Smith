"""Independent native-template adversarial/integration regressions.

No provider calls or fabricated visual acceptance. Fixtures and temporary files
stay inside the registered repository worktree. Real rendering is separately
exercised by the migration acceptance run, not mocked into a visual pass here.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import warnings
import zipfile

from lxml import etree
from pptx import Presentation

from native_template_fixture import make_template
from engine.native_template import runtime
from engine.native_template import review
from engine.native_template.safety import package, strict_json


ROOT = Path(__file__).resolve().parents[1]
SCRATCH = ROOT / ".pytest-tmp"


class NativeTemplateAdversarialTest(unittest.TestCase):
    def setUp(self):
        SCRATCH.mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix="native-independent-", dir=SCRATCH)
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = make_template(self.root / "source.pptx", second_theme=True)
        self.workspace = self.root / "work"
        runtime.init(self.source, self.workspace)
        self.inspection = runtime.inspect(self.workspace)
        self.first = self.inspection["sourceSlides"][0]
        self.second = self.inspection["sourceSlides"][1]

    def add_page(self, page="first", source=None, revision=0):
        return runtime.apply(self.workspace, {"expectedRevision": revision, "operations": [{
            "op": "add_page", "id": page,
            "sourceSlideId": source or self.first["sourceSlideId"],
            "topicIds": ["test-topic"], "notes": "Synthetic fixture. Source: test fixture, not business data.",
        }]})

    def state_bytes(self):
        return (self.workspace / "state.json").read_bytes()

    def test_wrong_operation_is_atomic_even_after_valid_page_add(self):
        before = self.state_bytes()
        with self.assertRaises(ValueError):
            runtime.apply(self.workspace, {"expectedRevision": 0, "operations": [
                {"op": "add_page", "id": "one", "sourceSlideId": self.first["sourceSlideId"]},
                {"op": "edit_native", "pageId": "one", "objectId": "9999", "text": "Should not persist"},
            ]})
        self.assertEqual(self.state_bytes(), before)
        self.assertEqual(runtime.status(self.workspace)["revision"], 0)

    def test_stale_revision_cannot_mutate_valid_work(self):
        self.add_page()
        before = self.state_bytes()
        with self.assertRaises(ValueError):
            runtime.apply(self.workspace, {"expectedRevision": 0, "operations": [
                {"op": "remove_page", "pageId": "first"},
            ]})
        self.assertEqual(before, self.state_bytes())

    def test_author_script_and_unrecognized_fields_are_not_executed(self):
        before = self.state_bytes()
        for operation in [
            {"op": "run_code", "script": "raise RuntimeError('must not run')"},
            {"op": "add_page", "id": "one", "sourceSlideId": self.first["sourceSlideId"],
             "author_script": str(self.root / "untrusted.py")},
        ]:
            with self.subTest(operation=operation["op"]):
                with self.assertRaises(ValueError):
                    runtime.apply(self.workspace, {"expectedRevision": 0, "operations": [operation]})
                self.assertEqual(before, self.state_bytes())

    def test_reorder_requires_every_page_exactly_once(self):
        self.add_page()
        self.add_page("second", revision=1)
        before = self.state_bytes()
        for ids in [["first"], ["first", "first"], ["first", "not-present"]]:
            with self.assertRaises(ValueError):
                runtime.apply(self.workspace, {"expectedRevision": 2, "operations": [{"op": "reorder_pages", "ids": ids}]})
            self.assertEqual(before, self.state_bytes())
        runtime.apply(self.workspace, {"expectedRevision": 2, "operations": [{"op": "reorder_pages", "ids": ["second", "first"]}]})
        self.assertEqual([p["id"] for p in runtime.status(self.workspace)["pages"]], ["second", "first"])

    def test_chart_clones_have_independent_native_workbooks_and_keep_source(self):
        original = self.source.read_bytes()
        chart = next(o for o in self.first["objects"] if o["type"] == "chart")
        operations = [
            {"op": "add_page", "id": "a", "sourceSlideId": self.first["sourceSlideId"], "notes": "Test A"},
            {"op": "add_page", "id": "b", "sourceSlideId": self.first["sourceSlideId"], "notes": "Test B"},
            {"op": "edit_native", "pageId": "a", "objectId": chart["objectId"], "chart": {
                "categories": ["Test A", "Test B", "Test C"],
                "series": [{"name": "Changed fixture values", "values": [3, 5, 9]}],
            }},
        ]
        runtime.apply(self.workspace, {"expectedRevision": 0, "operations": operations})
        exported = runtime.export(self.workspace)
        prs = Presentation(exported["candidate"])
        charts = [next(shape.chart for shape in slide.shapes if shape.has_chart) for slide in prs.slides]
        self.assertEqual(list(charts[0].series[0].values), [3, 5, 9])
        self.assertEqual(list(charts[1].series[0].values), [10, 20, 30])
        self.assertNotEqual(charts[0].part.partname, charts[1].part.partname)
        self.assertNotEqual(charts[0].part.chart_workbook.xlsx_part.partname,
                            charts[1].part.chart_workbook.xlsx_part.partname)
        self.assertEqual(self.source.read_bytes(), original)
        with zipfile.ZipFile(exported["candidate"]) as archive:
            self.assertTrue(any(name.endswith(".xlsx") for name in archive.namelist()))

    def test_unselected_source_page_is_absent_even_from_orphan_package_parts(self):
        self.add_page()
        exported = runtime.export(self.workspace)
        with zipfile.ZipFile(exported["candidate"]) as archive:
            leaked = [name for name in archive.namelist()
                      if b"UNSELECTED_SOURCE_SENTINEL" in archive.read(name)]
        self.assertEqual(leaked, [], "Unselected template text must not remain in unused slide XML")
        self.assertEqual(len(Presentation(exported["candidate"]).slides), 1)

    def test_frozen_template_change_invalidates_workspace(self):
        self.add_page()
        runtime.export(self.workspace)
        frozen = self.workspace / "template.pptx"
        frozen.write_bytes(frozen.read_bytes() + b"tamper")
        with self.assertRaises(ValueError):
            runtime.status(self.workspace)

    def test_frozen_workspace_survives_original_source_move_and_change(self):
        self.add_page()
        frozen = (self.workspace / "template.pptx").read_bytes()
        self.source.write_bytes(b"The original path is no longer the frozen input")
        runtime.status(self.workspace)
        self.source.unlink()
        exported = runtime.export(self.workspace)
        self.assertEqual(len(Presentation(exported["candidate"]).slides), 1)
        self.assertEqual((self.workspace / "template.pptx").read_bytes(), frozen)

    def test_candidate_change_is_not_accepted_as_current_export(self):
        self.add_page()
        candidate = runtime.export(self.workspace)
        p = Path(candidate["candidate"])
        p.write_bytes(p.read_bytes() + b"tamper")
        self.assertNotEqual(runtime.status(self.workspace)["status"], "candidate_unreviewed")
        with self.assertRaises((ValueError, FileNotFoundError)):
            review.verify(self.workspace)

    def test_missing_render_and_missing_external_review_do_not_pass(self):
        self.add_page()
        runtime.export(self.workspace)
        with self.assertRaises((ValueError, FileNotFoundError)):
            review.verify(self.workspace)
        fake = {"schema_version": "1.0.0", "status": "approved", "candidateHash": "sha256:fake",
                "templateHash": "sha256:fake", "renderReportHash": "sha256:fake",
                "reviewer": {"provider": "test", "model": "synthetic", "method": "not an actual review"},
                "pages": [], "deck_rhythm": "not actual visual evidence"}
        with self.assertRaises((ValueError, FileNotFoundError)):
            review.review(self.workspace, fake)

    def test_workspace_symlink_cannot_redirect_tool_writes(self):
        link = self.root / "linked"
        link.symlink_to(self.workspace, target_is_directory=True)
        with self.assertRaises(ValueError):
            runtime.status(link)
        self.assertEqual(runtime.status(self.workspace)["revision"], 0)

    def test_duplicate_json_keys_and_nonfinite_values_are_rejected(self):
        for text in [b'{"expectedRevision":0,"expectedRevision":1}', b'{"x":NaN}', b'{"x":Infinity}']:
            with self.subTest(text=text):
                with self.assertRaises(ValueError):
                    strict_json(text)

    def test_zip_traversal_duplicate_and_dtd_are_rejected(self):
        with zipfile.ZipFile(self.source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        for mode in ["traversal", "duplicate", "dtd"]:
            stream = BytesIO()
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                    for name, data in entries.items():
                        if mode == "dtd" and name == "ppt/slides/slide1.xml":
                            data = b'<!DOCTYPE p:sld [<!ENTITY e SYSTEM "file:///etc/passwd">]>' + data
                        archive.writestr(name, data)
                    if mode == "traversal": archive.writestr("../outside.xml", b"unsafe")
                    if mode == "duplicate": archive.writestr("ppt/slides/slide1.xml", b"duplicate")
            with self.subTest(mode=mode):
                with self.assertRaises(ValueError):
                    package(stream.getvalue())

    def test_renamed_active_payload_is_rejected_by_content_type_and_relation(self):
        with zipfile.ZipFile(self.source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        relation_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        document_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
        for active in ["oleObject", "control", "vbaProject"]:
            modified = dict(entries)
            rels = etree.fromstring(modified["ppt/slides/_rels/slide1.xml.rels"])
            etree.SubElement(rels, "{%s}Relationship" % relation_ns,
                             Id="rIdDisguisedActive", Type=document_ns + "/" + active,
                             Target="../media/innocent.bin")
            modified["ppt/slides/_rels/slide1.xml.rels"] = etree.tostring(rels)
            types = etree.fromstring(modified["[Content_Types].xml"])
            etree.SubElement(types, "{%s}Override" % content_ns,
                             PartName="/ppt/media/innocent.bin", ContentType="application/vnd.ms-office.oleObject")
            modified["[Content_Types].xml"] = etree.tostring(types)
            modified["ppt/media/innocent.bin"] = b"Synthetic inert payload; never executable"
            stream = BytesIO()
            with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for name, data in modified.items(): archive.writestr(name, data)
            with self.subTest(active=active):
                with self.assertRaises(ValueError):
                    package(stream.getvalue())

    def test_internal_link_cannot_smuggle_unselected_source_page(self):
        source = self.root / "internal-link.pptx"
        with zipfile.ZipFile(self.source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        package_r = "http://schemas.openxmlformats.org/package/2006/relationships"
        slide = etree.fromstring(entries["ppt/slides/slide1.xml"])
        run = slide.find(".//{%s}rPr" % a)
        etree.SubElement(run, "{%s}hlinkClick" % a,
                         attrib={"{%s}id" % r: "rIdHiddenSource"}, action="ppaction://hlinksldjump")
        entries["ppt/slides/slide1.xml"] = etree.tostring(slide)
        rels = etree.fromstring(entries["ppt/slides/_rels/slide1.xml.rels"])
        etree.SubElement(rels, "{%s}Relationship" % package_r,
                         Id="rIdHiddenSource", Type=r + "/slide", Target="slide3.xml")
        entries["ppt/slides/_rels/slide1.xml.rels"] = etree.tostring(rels)
        with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items(): archive.writestr(name, data)
        workspace = self.root / "link-work"
        # Conservatively rejecting unsupported jump links is also safe.
        try:
            state = runtime.init(source, workspace)
            runtime.apply(workspace, {"expectedRevision": 0, "operations": [{
                "op": "add_page", "id": "only-a", "sourceSlideId": state["sourceSlideIds"][0],
                "notes": "Synthetic hyperlink privacy fixture",
            }]})
            exported = runtime.export(workspace)
        except ValueError as error:
            self.assertRegex(str(error).lower(), "link|relation|slide")
            return
        with zipfile.ZipFile(exported["candidate"]) as archive:
            leaked = [name for name in archive.namelist()
                      if b"UNSELECTED_SOURCE_SENTINEL" in archive.read(name)]
        self.assertEqual(leaked, [], "Internal jump links must not export unselected source-page text")

    def test_real_presentation_order_not_filename_order_is_used(self):
        source = self.root / "reordered.pptx"
        with zipfile.ZipFile(self.source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        presentation = etree.fromstring(entries["ppt/presentation.xml"])
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        listing = presentation.find("p:sldIdLst", ns)
        old_ids = [item.get("id") for item in listing]
        listing[:] = list(reversed(list(listing)))
        entries["ppt/presentation.xml"] = etree.tostring(presentation, encoding="UTF-8", xml_declaration=True)
        with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items(): archive.writestr(name, data)
        workspace = self.root / "order-work"
        result = runtime.init(source, workspace)
        self.assertEqual(result["sourceSlideIds"], list(reversed(old_ids)))
        pages = runtime.inspect(workspace)["sourceSlides"]
        self.assertEqual(pages[0]["sourceSlideId"], old_ids[-1])
        self.assertTrue(any("UNSELECTED_SOURCE_SENTINEL" in str(o) for o in pages[0]["objects"]))


if __name__ == "__main__":
    unittest.main()
