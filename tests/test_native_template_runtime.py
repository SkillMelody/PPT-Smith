"""Local synthetic regression tests for the fixed native-template runtime."""
import io
import json
from pathlib import Path
import zipfile
import pytest
from pptx import Presentation
from native_template_fixture import make_template
from engine.native_template import runtime as rt
from engine.native_template import review as gate
from engine.native_template.safety import sha, strict_json, package, encode


@pytest.fixture
def fixture(tmp_path):
    template = make_template(tmp_path / 'template.pptx', second_theme=True)
    workspace = tmp_path / 'work'
    result = rt.init(template, workspace)
    return template, workspace, result


def apply(work, revision, *ops):
    for op in ops:
        if op.get('op') == 'add_page': op.setdefault('notes', 'Synthetic task-specific notes; not factual evidence.')
    return rt.apply(work, {'expectedRevision': revision, 'operations': list(ops)})


def test_clone_chart_table_group_and_theme(fixture):
    template, work, init = fixture
    original = template.read_bytes()
    inventory = rt.inspect(work)
    ids = init['sourceSlideIds']; assert len(ids) == 3
    chart_id = next(o['objectId'] for o in inventory['sourceSlides'][0]['objects'] if o['type'] == 'chart')
    group_text = next(o['objectId'] for o in inventory['sourceSlides'][0]['objects'] if o['parentObjectId'])
    table_id = next(o['objectId'] for o in inventory['sourceSlides'][1]['objects'] if o['type'] == 'table')
    apply(work, 0, {'op': 'add_page', 'id': 'a', 'sourceSlideId': ids[0]},
          {'op': 'add_page', 'id': 'b', 'sourceSlideId': ids[0]},
          {'op': 'add_page', 'id': 'table', 'sourceSlideId': ids[1]})
    apply(work, 1, {'op': 'edit_native', 'pageId': 'a', 'objectId': chart_id,
                   'chart': {'categories': ['New A', 'New B'], 'series': [{'name': 'New', 'values': [99, 88]}]}},
          {'op': 'edit_native', 'pageId': 'a', 'objectId': group_text, 'text': 'Nested change'},
          {'op': 'edit_native', 'pageId': 'table', 'objectId': table_id, 'rows': [['A','B','C'],['1','2','3'],['4','5','6']]},
          {'op': 'set_page', 'pageId': 'a', 'notes': 'New speaker notes'})
    report = rt.export(work); assert report['status'] == 'candidate_unreviewed'
    out = Presentation(work / 'candidate.pptx')
    charts = [next(s.chart for s in slide.shapes if s.has_chart) for slide in list(out.slides)[:2]]
    assert list(charts[0].series[0].values) == [99,88]
    assert list(charts[1].series[0].values) == [10,20,30]
    assert charts[0].part.chart_workbook.xlsx_part is not charts[1].part.chart_workbook.xlsx_part
    assert 'New speaker notes' in out.slides[0].notes_slide.notes_text_frame.text
    assert out.slides[2].slide_layout.slide_master.part is not out.slides[0].slide_layout.slide_master.part
    with zipfile.ZipFile(work / 'candidate.pptx') as z:
        xmls = b''.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
        assert b'UNSELECTED_SOURCE_SENTINEL' not in xmls
    assert template.read_bytes() == original
    assert rt.export(work)['candidateHash'] == report['candidateHash']


def test_add_native_objects_and_isolate_chart(fixture):
    _, work, init = fixture
    objects = rt.inspect(work)['sourceSlides'][0]['objects']
    chart = next(o for o in objects if o['type'] == 'chart')
    text = next(o for o in objects if o['type'] == 'text' and not o['parentObjectId'])
    apply(work, 0, {'op': 'add_page', 'id':'a', 'sourceSlideId':init['sourceSlideIds'][0]},
          {'op':'add_object','pageId':'a','objectId':'100','type':'text','styleSourceObjectId':text['objectId'], 'text':'New same style', 'options': {'x':1,'y':1,'width':3,'height':1}},
          {'op':'add_object','pageId':'a','objectId':'101','type':'chart','styleSourceObjectId':chart['objectId'],
           'chart': {'categories':['X'], 'series':[{'name':'Added','values':[42]}]}})
    rt.export(work)
    out = Presentation(work / 'candidate.pptx')
    charts = [s.chart for s in out.slides[0].shapes if s.has_chart]
    assert [list(c.series[0].values) for c in charts] == [[10,20,30],[42]]
    assert len(rt.inspect(work)['pages'][0]['objects']) == len(objects) + 2


@pytest.mark.parametrize('bad', [
    {'op':'script','script':'print(1)'},
    {'op':'add_page','id':'bad','sourceSlideId':'999'},
    {'op':'add_page','id':'bad','sourceSlideId':'256','unknown':True},
    {'op':'edit_native','pageId':'a','objectId':'99999','text':'x'},
    {'op':'edit_native','pageId':'a','objectId':'2','options': {'x':-1}},
    {'op':'add_object','pageId':'a','objectId':'100','type':'text','styleSourceObjectId':'999','text':'x'},
])
def test_transaction_rollback(fixture,bad):
    _, work, init = fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    before = (work/'state.json').read_bytes()
    with pytest.raises((ValueError, KeyError)):
        apply(work,1,{'op':'set_page','pageId':'a','notes':'should rollback'},bad)
    assert (work/'state.json').read_bytes() == before
    with pytest.raises(ValueError,match='REVISION_CONFLICT'): apply(work,0,{'op':'remove_page','pageId':'a'})


def test_reorder_delete_constraints(fixture):
    _,work,init=fixture
    apply(work,0,*[{'op':'add_page','id':p,'sourceSlideId':init['sourceSlideIds'][i]} for i,p in enumerate(['a','b'])])
    apply(work,1,{'op':'reorder_pages','ids':['b','a']})
    assert rt.export(work)['pageIds']==['b','a']
    apply(work,2,{'op':'remove_page','pageId':'a'})
    assert rt.export(work)['pageIds']==['b']
    assert rt.status(work)['lastChange']['reviewReuse'] is False


def test_safety_inputs_and_source_change(fixture,tmp_path):
    template,work,init=fixture
    with pytest.raises(ValueError): strict_json(b'{"a":NaN}')
    with pytest.raises(ValueError): strict_json(b'{"a":1,"a":2}')
    with pytest.raises(ValueError): strict_json(b'{"a":1e999}')
    alias=tmp_path/'alias'; alias.symlink_to(work,target_is_directory=True)
    with pytest.raises(ValueError,match='SYMLINK'): rt.status(alias)
    template.write_bytes(template.read_bytes()+b'x')
    assert rt.status(work)['revision'] == 0  # Original path is provenance, not the frozen source.
    frozen=work/'template.pptx';frozen.write_bytes(frozen.read_bytes()+b'x')
    with pytest.raises(ValueError,match='FROZEN_TEMPLATE_HASH_MISMATCH'): rt.status(work)


def test_unavailable_render_and_unreviewed_verify(fixture,monkeypatch):
    _,work,init=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    rt.export(work)
    with pytest.raises(ValueError): gate.verify(work)
    monkeypatch.setenv('PPTSMITH_TEST_RENDERER_UNAVAILABLE','1')
    assert gate.render(work)['status']=='unavailable'
    with pytest.raises(ValueError): gate.review(work,{})


def test_zip_unsafe_and_entities(tmp_path):
    import warnings
    def z(entries):
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as archive:
            for n,v in entries: archive.writestr(n,v)
        return out.getvalue()
    for entries in [[('../escape','x')],[('/absolute','x')],[('a.xml','<!DOCTYPE a [<!ENTITY x "y">]><a/>')],[('ppt/vbaProject.bin','x')]]:
        with pytest.raises(ValueError): package(z(entries))
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        bad=z([('a.xml','<a/>'),('a.xml','<b/>')])
    with pytest.raises(ValueError,match='DUPLICATE'): package(bad)


def _mock_render(monkeypatch):
    # Unit-only renderer double. Real LibreOffice is exercised separately, never
    # label these synthetic observations as an actual independent visual review.
    from PIL import Image
    import engine.native_template.rendering
    def render(path, output, **kwargs):
        output.mkdir(parents=True,exist_ok=True)
        pages=[]
        for i in range(1,kwargs['expected_slides']+1):
            p=output/('page%d.png'%i); Image.new('RGB',(320,180),(i*50,30,60)).save(p)
            pages.append({'slide_index':i,'image':str(p),'status':'passed'})
        return {'status':'passed','slides':pages}
    monkeypatch.setattr(engine.native_template.rendering,'run_render_report',render)


def _review_doc(work, rendered):
    return {'schema_version':'1.0.0','status':'approved',
            'candidateHash':rendered['candidateHash'],'templateHash':rendered['templateHash'],
            'renderReportHash':sha((work/'render-report.json').read_bytes()),
            'reviewer':{'provider':'test-double','model':'synthetic','method':'unit-test only'},
            'deck_rhythm':'Unit-test synthetic observation; not real external review.',
            'pages':[{'slide_index':p['slide_index'],'pageId':p['pageId'],'sourceSlideId':p['sourceSlideId'],
                      'preview_sha256':p['sha256'],'reference_sha256':p['referenceHash'],
                      **{k:{'status':'pass','observation':'Unit-test observation.'} for k in gate.CRITERIA}} for p in rendered['slides']]}


def test_review_complete_and_fail_closed_stale(fixture,monkeypatch):
    _,work,init=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    rt.export(work);_mock_render(monkeypatch)
    rendered=gate.render(work);doc=_review_doc(work,rendered)
    assert gate.review(work,doc)['ok']
    assert gate.verify(work)['status']=='final_delivery_ready'
    original=(work/'review.json').read_bytes()
    import copy
    for bad in ['missing_page','missing_observation','failed_observation','wrong_reference','unknown_field']:
        d=copy.deepcopy(doc)
        if bad=='missing_page':d['pages']=[]
        if bad=='missing_observation': del d['pages'][0]['hierarchy']
        if bad=='failed_observation':d['pages'][0]['readability']['status']='fail'
        if bad=='wrong_reference':d['pages'][0]['reference_sha256']='sha256:wrong'
        if bad=='unknown_field':d['script']='evil.py'
        with pytest.raises(ValueError):gate.review(work,d)
        assert (work/'review.json').read_bytes()==original
    image=Path(rendered['slides'][0]['image']);image.write_bytes(image.read_bytes()+b'x')
    with pytest.raises(ValueError,match='STALE_PREVIEW'):gate.verify(work)


def test_candidate_and_revision_invalidate_review(fixture,monkeypatch):
    _,work,init=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    rt.export(work);_mock_render(monkeypatch)
    rendered=gate.render(work);gate.review(work,_review_doc(work,rendered))
    candidate=work/'candidate.pptx';old=candidate.read_bytes();candidate.write_bytes(old+b'x')
    with pytest.raises(ValueError,match='CANDIDATE_HASH'):gate.verify(work)
    candidate.write_bytes(old)
    apply(work,1,{'op':'set_page','pageId':'a','notes':'A change'})
    with pytest.raises(ValueError,match='STALE_CANDIDATE'):gate.verify(work)
    rt.export(work)
    with pytest.raises(ValueError,match='STALE_REVIEW'):gate.verify(work)


def test_render_preview_symlink_refused(fixture,monkeypatch,tmp_path):
    _,work,init=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    rt.export(work);_mock_render(monkeypatch)
    rendered=gate.render(work);gate.review(work,_review_doc(work,rendered))
    p=Path(rendered['slides'][0]['referenceImage']);copy=tmp_path/'copy.png';copy.write_bytes(p.read_bytes());p.unlink();p.symlink_to(copy)
    with pytest.raises(ValueError,match='SYMLINK'):gate.verify(work)


def test_constraints_exact_maximum_target(tmp_path):
    template=make_template(tmp_path/'t.pptx')
    for mode,count in [('exact',2),('maximum',1),('target',3),('auto',None)]:
        w=tmp_path/mode;c={'mode':mode}
        if count is not None:c['count']=count
        init=rt.init(template,w,c)
        apply(w,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
        if mode=='exact':
            with pytest.raises(ValueError,match='EXACT_PAGE'):rt.export(w)
        else:
            report=rt.export(w)
            if mode=='target':assert not report['targetCountMet']
        if mode=='maximum':
            with pytest.raises(ValueError,match='MAXIMUM_PAGE'):apply(w,1,{'op':'add_page','id':'b','sourceSlideId':init['sourceSlideIds'][0]})


def test_styles_and_deterministic_chart_export(fixture):
    import time
    _,work,init=fixture
    inventory=rt.inspect(work)
    theme1=next(o for o in inventory['sourceSlides'][0]['objects'] if o['name']=='Main title')
    theme2=next(o for o in inventory['sourceSlides'][1]['objects'] if o['name']=='Main title')
    assert theme1['effectiveTextStyles'][0]['fontSizePt']['value']==28
    assert theme1['effectiveTextStyles'][0]['color']['baseRgb']!=theme2['effectiveTextStyles'][0]['color']['baseRgb']
    chart=next(o for o in inventory['sourceSlides'][0]['objects'] if o['type']=='chart')
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]},
          {'op':'edit_native','pageId':'a','objectId':chart['objectId'],'chart':{'categories':['A'],'series':[{'name':'X','values':[1]}]}})
    first=rt.export(work)['candidateHash'];time.sleep(1.1)
    assert rt.export(work)['candidateHash']==first


def test_immutable_style_bytes_and_picture_addition(fixture):
    template,work,init=fixture
    picture=next(o for o in rt.inspect(work)['sourceSlides'][2]['objects'] if o['type']=='picture')
    apply(work,0,{'op':'add_page','id':'picture','sourceSlideId':init['sourceSlideIds'][2]},
          {'op':'add_object','pageId':'picture','objectId':'100','type':'picture','styleSourceObjectId':picture['objectId'],
           'options':{'x':5,'y':2.4,'width':4,'height':1.5}})
    rt.export(work)
    with zipfile.ZipFile(template) as source,zipfile.ZipFile(work/'candidate.pptx') as output:
        for n in output.namelist():
            if n.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/','ppt/notesMasters/')):
                assert output.read(n)==source.read(n),n
        assert not any(n.startswith('docProps/') for n in output.namelist())
    p=Presentation(work/'candidate.pptx')
    pictures=[s for s in p.slides[0].shapes if s.shape_type==13]
    assert len(pictures)==2 and pictures[0].image.blob==pictures[1].image.blob


def test_presentation_order_not_slide_filename(fixture):
    template,work,_=fixture
    # Construct a separate synthetic source with presentation order different to
    # physical slide file numbers. The source IDs remain the true sldId values.
    from engine.native_template.package import xml,dump,pack,P
    parts=package(template.read_bytes());pres=xml(parts['ppt/presentation.xml'])
    ids=pres.find('{%s}sldIdLst'%P);first=ids[0];ids.remove(first);ids.append(first)
    parts['ppt/presentation.xml']=dump(pres)
    reordered=template.parent/'reordered.pptx';reordered.write_bytes(pack(parts))
    new_work=template.parent/'reordered-work';rt.init(reordered,new_work)
    inspection=rt.inspect(new_work)
    assert [p['sourceSlideId'] for p in inspection['sourceSlides']]==['257','258','256']
    apply(new_work,0,{'op':'add_page','id':'a','sourceSlideId':'257'})
    rt.export(new_work)
    out=Presentation(new_work/'candidate.pptx')
    assert any(s.has_table for s in out.slides[0].shapes)


def test_concurrent_revision_lock_allows_one_writer(fixture,tmp_path):
    import concurrent.futures,subprocess,sys
    _,work,init=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    paths=[]
    for i in (1,2):
        p=tmp_path/('request%d.json'%i)
        p.write_text(json.dumps({'expectedRevision':1,'operations':[{'op':'set_page','pageId':'a','notes':'Writer %d'%i}]}))
        paths.append(p)
    def invoke(p):
        return subprocess.run([sys.executable,'-m','engine','template','apply','--workspace',str(work),'--request',str(p)],capture_output=True,text=True)
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool: results=list(pool.map(invoke,paths))
    assert sorted(r.returncode for r in results)==[0,1]
    assert any('REVISION_CONFLICT' in r.stdout for r in results)
    assert rt.status(work)['revision']==2


def test_zip_compression_and_macro_content_type_refused(fixture):
    template,_,_=fixture
    for mode in ('bomb','macro','external','case-duplicate'):
        entries=package(template.read_bytes())
        if mode=='macro': entries['[Content_Types].xml']=entries['[Content_Types].xml'].replace(b'presentationml.presentation.main+xml',b'presentation.macroEnabled.main+xml')
        elif mode=='external':entries['_rels/.rels']=entries['_rels/.rels'].replace(b'Target="ppt/presentation.xml"',b'Target="https://example.invalid/a" TargetMode="External"')
        elif mode=='bomb':entries['bomb.bin']=b'x'*2000000
        elif mode=='case-duplicate':entries['PPT/PRESENTATION.XML']=entries['ppt/presentation.xml']
        from engine.native_template.package import pack
        with pytest.raises(ValueError):package(pack(entries))


def test_removed_page_id_cannot_rebind_to_different_source(fixture):
    _,work,init=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    apply(work,1,{'op':'remove_page','pageId':'a'})
    before=(work/'state.json').read_bytes()
    with pytest.raises(ValueError,match='BINDING_IMMUTABLE'):
        apply(work,2,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][1]})
    assert (work/'state.json').read_bytes()==before
    apply(work,2,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})


@pytest.mark.parametrize('aspect', ['4:3', '16:9'])
def test_render_forwards_actual_canvas_aspect(tmp_path,monkeypatch,aspect):
    template=make_template(tmp_path/'source.pptx',aspect=aspect)
    work=tmp_path/'work';initialized=rt.init(template,work)
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][0]})
    exported=rt.export(work)
    assert exported['compatibility']=={'powerpoint':'not_tested','wps':'not_tested'}
    assert exported['automaticFactChecking'] is False
    _mock_render(monkeypatch)
    import engine.native_template.rendering as rendering
    original=rendering.run_render_report;forwarded=[]
    def capture(*args,**kwargs):
        forwarded.append(kwargs['expected_aspect'])
        return original(*args,**kwargs)
    monkeypatch.setattr(rendering,'run_render_report',capture)
    rendered=gate.render(work)
    expected=4/3 if aspect=='4:3' else 16/9
    assert forwarded==pytest.approx([expected,expected])
    gate.review(work,_review_doc(work,rendered))
    assert gate.verify(work)['compatibility']=={'powerpoint':'not_tested','wps':'not_tested'}


def test_source_decoration_removal_blocks_final(fixture,monkeypatch):
    _,work,init=fixture
    stripe=next(o for o in rt.inspect(work)['sourceSlides'][0]['objects'] if o['name']=='Synthetic brand decoration')
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]},
          {'op':'remove_object','pageId':'a','objectId':stripe['objectId']})
    exported=rt.export(work)
    assert exported['qualityFloor']=='blocked'
    assert any(i['code']=='SOURCE_DECORATION_CHANGED_OR_MISSING' for i in exported['preservationFloor']['issues'])
    _mock_render(monkeypatch);rendered=gate.render(work);gate.review(work,_review_doc(work,rendered))
    with pytest.raises(ValueError,match='PRESERVATION_FLOOR_BLOCKED'):gate.verify(work)


def test_task_specific_notes_missing_blocks_final(fixture,monkeypatch):
    _,work,init=fixture
    rt.apply(work,{'expectedRevision':0,'operations':[{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]}]})
    exported=rt.export(work)
    assert any(i['code']=='OUTPUT_PAGE_NOTES_REQUIRED' for i in exported['preservationFloor']['issues'])
    _mock_render(monkeypatch);rendered=gate.render(work);gate.review(work,_review_doc(work,rendered))
    with pytest.raises(ValueError,match='OUTPUT_PAGE_NOTES_REQUIRED'):gate.verify(work)


def test_unsupported_lock_platform_is_explicit(monkeypatch):
    from engine.native_template import safety
    from engine.native_template.cli import capabilities
    monkeypatch.setattr(safety,'fcntl',None)
    assert capabilities()['platform']['windows']=='unsupported'
    with pytest.raises(ValueError,match='PLATFORM_UNSUPPORTED'):
        with safety.locked('unused'):pass


def test_changed_native_fontsize_is_not_preservation_pass(fixture):
    from pptx.util import Pt
    from engine.native_template.floor import check
    _,work,init=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    rt.export(work)
    candidate=Presentation(work/'candidate.pptx')
    title=next(s for s in candidate.slides[0].shapes if s.name=='Main title')
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:run.font.size=Pt(8)
    out=io.BytesIO();candidate.save(out)
    state,template=rt.load(work)
    report=check(template,out.getvalue(),state)
    assert report['status']=='blocked'
    assert report['webVisualFloor']=='not_implemented'
    assert any(i['code']=='SOURCE_FONT_SIZE_CHANGED' for i in report['issues'])


def test_inspect_compact_defaults_and_opt_in_details(fixture):
    _,work,_=fixture
    compact=rt.inspect(work);detailed=rt.inspect(work,details=True)
    assert compact['details'] is False and detailed['details'] is True
    for page in compact['sourceSlides']:
        for obj in page['objects']:
            assert 'inheritance' not in obj['style']
            assert 'sourceChain' in obj['style']
            for style in obj.get('effectiveTextStyles',[]):assert 'text' not in style and 'runCount' in style
    assert any(o['style']['inheritance'] for p in detailed['sourceSlides'] for o in p['objects'])
    assert len(json.dumps(compact))<len(json.dumps(detailed))
    assert all(r.get('origin')!='masterTextStyles' for p in detailed['sourceSlides'] for o in p['objects'] for r in o['style']['inheritance'])


def test_placeholder_font_matches_role_and_paragraph_level(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement
    from engine.native_template.styles import q
    source=make_template(tmp_path/'source.pptx')
    prs=Presentation(source);slide=prs.slides.add_slide(prs.slide_layouts[1])
    body=slide.placeholders[1];body.text='Nested body';body.text_frame.paragraphs[0].level=1
    # Layout level-1 default must not override body level-2 font evidence.
    layout_body=slide.slide_layout.placeholders[1]
    p=layout_body.text_frame.paragraphs[0]._p
    prop=p.find(q('pPr'))
    if prop is None:prop=OxmlElement('a:pPr');p.insert(0,prop)
    prop.set('lvl','0');r=prop.find(q('defRPr'))
    if r is None:r=OxmlElement('a:defRPr');prop.append(r)
    r.set('sz','3300')
    master=slide.slide_layout.slide_master
    # Remove placeholder paragraph evidence to isolate txStyles role+level.
    for placeholder in master.placeholders:
        if placeholder.has_text_frame:
            for p in placeholder.text_frame._txBody.findall(q('p')):
                for prop in p.findall(q('pPr')):p.remove(prop)
    styles=master._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}txStyles')
    for tag,size in [('bodyStyle',1200),('titleStyle',4400),('otherStyle',1800)]:
        node=styles.find('{http://schemas.openxmlformats.org/presentationml/2006/main}'+tag)
        n=node.find(q('lvl2pPr'))
        if n is None:n=OxmlElement('a:lvl2pPr');node.append(n)
        r=n.find(q('defRPr'))
        if r is None:r=OxmlElement('a:defRPr');n.append(r)
        r.set('sz',str(size))
    prs.save(source)
    work=tmp_path/'work';rt.init(source,work)
    obj=next(o for o in rt.inspect(work)['sourceSlides'][-1]['objects'] if o.get('text')=='Nested body')
    assert obj['effectiveTextStyles'][0]['fontSizePt']=={'value':12.0,'origin':'master.bodyStyle'}


def test_active_payload_content_type_without_active_name_or_relation(fixture):
    from lxml import etree
    from engine.native_template.package import pack,dump
    template,_,_=fixture;base=package(template.read_bytes())
    for suffix in ('oleObject','activeX+xml','vbaProject','controlproperties+xml'):
        parts=dict(base);types=etree.fromstring(parts['[Content_Types].xml'])
        etree.SubElement(types,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                         PartName='/ppt/media/payload.bin',ContentType='application/vnd.ms-office.'+suffix)
        parts['[Content_Types].xml']=dump(types);parts['ppt/media/payload.bin']=b'not filename identifiable'
        with pytest.raises(ValueError,match='ACTIVE_CONTENT_TYPE_REFUSED'):package(pack(parts))


def test_active_relationship_without_active_name_or_content_type(fixture):
    from lxml import etree
    from engine.native_template.package import pack,dump
    template,_,_=fixture;base=package(template.read_bytes())
    for suffix in ('oleObject','control','vbaProject','activeXControl','ctrlProp','vbaData'):
        parts=dict(base);rels=etree.fromstring(parts['ppt/slides/_rels/slide1.xml.rels'])
        etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                         Id='rIdPayload',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/'+suffix,
                         Target='../media/payload.bin')
        parts['ppt/slides/_rels/slide1.xml.rels']=dump(rels);parts['ppt/media/payload.bin']=b'opaque data'
        with pytest.raises(ValueError,match='ACTIVE_RELATIONSHIP_REFUSED'):package(pack(parts))


def test_active_xml_markers_even_without_type_or_relationship(fixture):
    from lxml import etree
    from engine.native_template.package import pack,dump
    template,_,_=fixture;base=package(template.read_bytes())
    for local in ('oleObj','oleObject','controls','control','ocx','script','vbaProject'):
        parts=dict(base);slide=etree.fromstring(parts['ppt/slides/slide1.xml'])
        etree.SubElement(slide,'{http://schemas.openxmlformats.org/presentationml/2006/main}'+local)
        parts['ppt/slides/slide1.xml']=dump(slide)
        with pytest.raises(ValueError,match='ACTIVE_OBJECT_XML_REFUSED'):package(pack(parts))
    # Declared generic XML is parsed even when deliberately renamed to .bin.
    parts=dict(base);types=etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(types,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                     PartName='/ppt/media/payload.bin',ContentType='application/xml')
    parts['[Content_Types].xml']=dump(types)
    parts['ppt/media/payload.bin']=b'<ocx xmlns="http://schemas.microsoft.com/office/2006/activeX"/>'
    with pytest.raises(ValueError,match='ACTIVE_OBJECT_XML_REFUSED'):package(pack(parts))


def test_cross_source_jump_is_rejected_without_state_change(fixture):
    from lxml import etree
    from engine.native_template.package import pack,dump
    template,work,_=fixture;parts=package(template.read_bytes())
    rels=etree.fromstring(parts['ppt/slides/_rels/slide1.xml.rels'])
    etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                     Id='rIdOtherPage',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide',Target='slide3.xml')
    parts['ppt/slides/_rels/slide1.xml.rels']=dump(rels)
    linked=template.parent/'linked.pptx';linked.write_bytes(pack(parts))
    newwork=template.parent/'linked-work';init=rt.init(linked,newwork);old=(newwork/'state.json').read_bytes()
    with pytest.raises(ValueError,match='CROSS_SLIDE_LINK'):
        apply(newwork,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    assert (newwork/'state.json').read_bytes()==old
    assert not (newwork/'candidate.pptx').exists()


def test_frozen_source_survives_original_removal_and_remains_tamper_bound(fixture):
    template,work,init=fixture
    template.write_bytes(b'original source replaced');template.unlink()
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':init['sourceSlideIds'][0]})
    rt.export(work)
    assert rt.status(work)['revision']==1
    frozen=work/'template.pptx';frozen.write_bytes(frozen.read_bytes()+b'tampered')
    with pytest.raises(ValueError,match='FROZEN_TEMPLATE_HASH_MISMATCH'):rt.status(work)


def test_master_relation_cannot_smuggle_hidden_source_slide(fixture):
    from lxml import etree
    from engine.native_template.package import pack,dump
    template,_,_=fixture;parts=package(template.read_bytes())
    name='ppt/slideMasters/_rels/slideMaster1.xml.rels';rels=etree.fromstring(parts[name])
    etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
                     Id='rIdHiddenSource',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide',
                     Target='../slides/slide3.xml')
    parts[name]=dump(rels)
    source=template.parent/'master-link.pptx';source.write_bytes(pack(parts))
    work=template.parent/'master-link-work';initialized=rt.init(source,work)
    before=(work/'state.json').read_bytes()
    with pytest.raises(ValueError,match='UNSELECTED_SLIDE_RELATIONSHIP'):
        apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][0]})
    assert (work/'state.json').read_bytes()==before


def test_fidelity_normal_native_edits_and_stable_candidate_bytes(fixture,monkeypatch):
    _,work,initialized=fixture
    objects=rt.inspect(work)['sourceSlides'][0]['objects']
    title=next(o for o in objects if o['name']=='Main title')
    chart=next(o for o in objects if o['type']=='chart')
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][0]},
          {'op':'edit_native','pageId':'a','objectId':title['objectId'],'text':'改写标题'},
          {'op':'edit_native','pageId':'a','objectId':chart['objectId'],'chart':{'categories':['A'],'series':[{'name':'Test','values':[7]}]}})
    first=rt.export(work)
    assert first['templateFidelity']['status']=='passed'
    assert first['templateFidelity']['pages'][0]['retention']['ratio']==1
    _mock_render(monkeypatch);rendered=gate.render(work);gate.review(work,_review_doc(work,rendered))
    assert gate.verify(work)['templateFidelity']['status']=='passed'
    second=rt.export(work)
    assert second['candidateHash']==first['candidateHash']
    assert gate.verify(work)['status']=='final_delivery_ready'  # Exact same bytes and render evidence.


@pytest.mark.parametrize('operation',['delete_decorations','move_stripe','delete_heading'])
def test_fidelity_blocks_source_decoration_or_heading_loss(fixture,operation):
    _,work,initialized=fixture
    objects=rt.inspect(work)['sourceSlides'][0]['objects']
    stripe=next(o for o in objects if o['name']=='Synthetic brand decoration')
    title=next(o for o in objects if o['name']=='Main title')
    edit=({'op':'remove_object','pageId':'a','objectId':title['objectId']} if operation=='delete_heading' else
          {'op':'edit_native','pageId':'a','objectId':stripe['objectId'],'options':{'y':1}} if operation=='move_stripe' else
          {'op':'remove_object','pageId':'a','objectId':stripe['objectId']})
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][0]},edit)
    exported=rt.export(work)
    assert exported['templateFidelity']['status']=='blocked'
    expected='TEMPLATE_TYPE_SCALE_LOST' if operation=='delete_heading' else 'TEMPLATE_DECORATION_LOST'
    assert expected in {f['code'] for f in exported['templateFidelity']['findings']}
    # Falsifying the persisted metric cannot bypass the current-byte recomputation.
    forged=json.loads((work/'candidate-report.json').read_text())
    forged['templateFidelity']={'status':'passed','findings':[]};forged['qualityFloor']='passed'
    (work/'candidate-report.json').write_text(json.dumps(forged))
    state,_=rt.load(work)
    assert rt.current_candidate(work,state)['templateFidelity']['status']=='blocked'
    with pytest.raises(ValueError,match='FIDELITY_BLOCKED|FLOOR_BLOCKED'):gate.verify(work)


@pytest.mark.parametrize('maximum,minimum,status',[(21,9,'passed'),(20.9,9,'blocked'),(21,8.9,'blocked')])
def test_fidelity_exact_web_type_scale_thresholds(fixture,maximum,minimum,status):
    from pptx.util import Pt
    from engine.native_template.fidelity import check
    _,work,initialized=fixture
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][0]})
    rt.export(work);candidate=Presentation(work/'candidate.pptx')
    for shape in candidate.slides[0].shapes:
        target=maximum if shape.name=='Main title' else minimum if shape.name=='Source label' else None
        if target is not None:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:r.font.size=Pt(target)
    out=io.BytesIO();candidate.save(out);state,template=rt.load(work)
    report=check(template,out.getvalue(),state)
    assert report['status']==status
    assert report['thresholds']=={'retentionMinimum':0.8,'headingSizeMinimumRatio':0.75,'smallestSizeMinimumRatio':0.9,'geometryTolerancePt':2.0}


def test_fidelity_unknown_font_size_is_not_pass(tmp_path):
    from engine.native_template.package import xml,dump,pack
    source=make_template(tmp_path/'source.pptx');parts=package(source.read_bytes())
    for name,data in list(parts.items()):
        if name.startswith('ppt/') and name.endswith('.xml'):
            tree=xml(data)
            for node in tree.iter():node.attrib.pop('sz',None)
            parts[name]=dump(tree)
    source.write_bytes(pack(parts));work=tmp_path/'work';initialized=rt.init(source,work)
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][0]})
    report=rt.export(work)['templateFidelity']
    assert report['status']=='blocked'
    assert any(f['code']=='TEMPLATE_FIDELITY_UNVERIFIED' and f['reason']=='native_font_size_unresolved' for f in report['findings'])


def test_fidelity_picture_retention_requires_original_id(fixture):
    _,work,initialized=fixture
    picture=next(o for o in rt.inspect(work)['sourceSlides'][2]['objects'] if o['type']=='picture')
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][2]},
          {'op':'remove_object','pageId':'a','objectId':picture['objectId']},
          {'op':'add_object','pageId':'a','objectId':'999','type':'picture','styleSourceObjectId':picture['objectId']})
    report=rt.export(work)['templateFidelity'];retention=report['pages'][0]['retention']
    assert report['status']=='blocked'
    assert picture['objectId'] in retention['lostObjectIds']
    assert retention['ratio']==0.5  # Same picture count does not replace the original ID/box evidence.


@pytest.mark.parametrize('removed,expected',[(1,'passed'),(2,'blocked')])
def test_fidelity_retention_eighty_percent_boundary(tmp_path,removed,expected):
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    source=make_template(tmp_path/'source.pptx');prs=Presentation(source)
    added=[]
    for i in range(4):
        shape=prs.slides[0].shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,Inches(1+i),Inches(.2),Inches(.1),Inches(.1))
        added.append(str(shape.shape_id))
    prs.save(source);work=tmp_path/'work';initialized=rt.init(source,work)
    apply(work,0,{'op':'add_page','id':'a','sourceSlideId':initialized['sourceSlideIds'][0]},
          *[{'op':'remove_object','pageId':'a','objectId':sid} for sid in added[:removed]])
    report=rt.export(work)['templateFidelity']
    assert report['pages'][0]['retention']['sourceDecorations']==5
    assert report['pages'][0]['retention']['ratio']==(5-removed)/5
    assert report['status']==expected


@pytest.mark.parametrize('action',['ppaction://macro?name=DoNotExecute','ppaction://program?path=DoNotExecute'])
def test_macro_program_action_alone_is_rejected(fixture,action):
    from lxml import etree
    from engine.native_template.package import pack,dump
    from engine.native_template.safety import xml
    template,_,_=fixture;parts=package(template.read_bytes())
    slide=etree.fromstring(parts['ppt/slides/slide1.xml'])
    etree.SubElement(slide,'{http://schemas.openxmlformats.org/drawingml/2006/main}hlinkClick',action=action)
    # No active CT, relationship, object marker or payload name in this fixture.
    data=dump(slide)
    with pytest.raises(ValueError,match='ACTIVE_OBJECT_ACTION_REFUSED'):xml(data)
    parts['ppt/slides/slide1.xml']=data
    with pytest.raises(ValueError,match='ACTIVE_OBJECT_ACTION_REFUSED'):package(pack(parts))


@pytest.mark.parametrize('local',['oleObj','control'])
def test_xml_parser_rejects_active_object_marker_without_package(local):
    from engine.native_template.safety import xml
    data=('<p:%s xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'%local).encode()
    with pytest.raises(ValueError,match='ACTIVE_OBJECT_XML_REFUSED'):xml(data)
