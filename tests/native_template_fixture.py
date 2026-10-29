"""Synthetic, redistributable fixtures for native-template integration tests.

All values are labelled test data. No user material, web service or model call is
required. Run under the registered test workspace; callers own output paths.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import zipfile

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}


def _text(slide, text, x, y, w, h, *, size=20, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    return shape


def _chrome(slide, title, *, width=10):
    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(width), Inches(.12)
    )
    stripe.name = "Synthetic brand decoration"
    stripe.fill.solid()
    stripe.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    stripe.line.fill.background()
    _text(slide, title, .5, .35, width - 1, .6, size=28, name="Main title")
    _text(slide, "Synthetic fixture · not business evidence", .5, 5.1, width - 1, .25, size=10,
          name="Source label")
    slide.notes_slide.notes_text_frame.text = "Synthetic fixture. All values are test data."


def make_template(path, *, aspect="16:9", second_theme=False):
    """Create a three-page native deck with chart, table, group and image."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    width = 10
    prs.slide_width = Inches(width)
    prs.slide_height = Inches(5.625 if aspect == "16:9" else 7.5)
    blank = prs.slide_layouts[6]
    first = prs.slides.add_slide(blank)
    _chrome(first, "Native chart fixture")
    _text(first, "Source text & literal <markup>", .6, 1.12, 8, .55, name="Editable paragraph")
    chart_data = CategoryChartData()
    chart_data.categories = ["Test A", "Test B", "Test C"]
    chart_data.add_series("Observed test values", [10, 20, 30])
    chart = first.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.6), Inches(1.9), Inches(5.5), Inches(2.8), chart_data
    )
    chart.name = "Native data chart"
    group = first.shapes.add_group_shape()
    group.name = "Native group"
    _text(group, "Nested editable label", 6.4, 2, 2.8, .6, size=18, name="Group label")

    second = prs.slides.add_slide(blank)
    _chrome(second, "Native table fixture")
    table_frame = second.shapes.add_table(3, 3, Inches(.6), Inches(1.4), Inches(7.8), Inches(2.1))
    table_frame.name = "Native test table"
    for row, values in enumerate([["Metric", "Test A", "Test B"], ["Count", "11", "22"], ["Rate", "0.1", "0.2"]]):
        for col, value in enumerate(values):
            cell = table_frame.table.cell(row, col)
            cell.text = value
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(18)
    _text(second, "A second-theme label", .6, 3.9, 8, .6, name="Theme probe")

    third = prs.slides.add_slide(blank)
    _chrome(third, "Unselected source page must not leak")
    _text(third, "UNSELECTED_SOURCE_SENTINEL", .6, 1.4, 8, .8, name="Private-looking test sentinel")
    image = Image.new("RGB", (320, 120), "#e8eef5")
    draw = ImageDraw.Draw(image)
    draw.rectangle((8, 8, 312, 112), outline="#355875", width=4)
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    image_bytes.seek(0)
    third.shapes.add_picture(image_bytes, Inches(.6), Inches(2.4), width=Inches(4))
    prs.save(path)
    if second_theme:
        _install_second_theme(path)
    return path


def _install_second_theme(path):
    """Give page 2 its own layout/master/theme; accent1 is visibly different.

    Duplicate only deterministic package parts, updating both relationship
    directions. This fixture catches choosing the first theme in the ZIP.
    """
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    theme = etree.fromstring(parts["ppt/theme/theme1.xml"])
    accent = theme.find(".//a:clrScheme/a:accent1", NS)
    for child in list(accent):
        accent.remove(child)
    etree.SubElement(accent, "{%s}srgbClr" % NS["a"], val="A13572")
    parts["ppt/theme/theme2.xml"] = etree.tostring(theme, xml_declaration=True, encoding="UTF-8")
    master = etree.fromstring(parts["ppt/slideMasters/slideMaster1.xml"])
    layout_ids = master.find("p:sldLayoutIdLst", NS)
    kept_layout_rid = layout_ids[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    for item in list(layout_ids)[1:]:
        layout_ids.remove(item)
    parts["ppt/slideMasters/slideMaster2.xml"] = etree.tostring(master, xml_declaration=True, encoding="UTF-8")
    master_rels = etree.fromstring(parts["ppt/slideMasters/_rels/slideMaster1.xml.rels"])
    for rel in list(master_rels):
        if rel.get("Type", "").endswith("/theme"):
            rel.set("Target", "../theme/theme2.xml")
        elif rel.get("Type", "").endswith("/slideLayout"):
            if rel.get("Id") != kept_layout_rid:
                master_rels.remove(rel)
            else:
                rel.set("Target", "../slideLayouts/slideLayout12.xml")
    parts["ppt/slideMasters/_rels/slideMaster2.xml.rels"] = etree.tostring(master_rels, xml_declaration=True, encoding="UTF-8")
    parts["ppt/slideLayouts/slideLayout12.xml"] = parts["ppt/slideLayouts/slideLayout7.xml"]
    layout_rels = etree.fromstring(parts["ppt/slideLayouts/_rels/slideLayout7.xml.rels"])
    for rel in layout_rels:
        if rel.get("Type", "").endswith("/slideMaster"):
            rel.set("Target", "../slideMasters/slideMaster2.xml")
    parts["ppt/slideLayouts/_rels/slideLayout12.xml.rels"] = etree.tostring(layout_rels, xml_declaration=True, encoding="UTF-8")
    slide_rels = etree.fromstring(parts["ppt/slides/_rels/slide2.xml.rels"])
    for rel in slide_rels:
        if rel.get("Type", "").endswith("/slideLayout"):
            rel.set("Target", "../slideLayouts/slideLayout12.xml")
    parts["ppt/slides/_rels/slide2.xml.rels"] = etree.tostring(slide_rels, xml_declaration=True, encoding="UTF-8")
    presentation = etree.fromstring(parts["ppt/presentation.xml"])
    master_ids = presentation.find("p:sldMasterIdLst", NS)
    master_id = str(max(int(item.get("id")) for item in master_ids) + 1)
    relation_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    etree.SubElement(master_ids, "{%s}sldMasterId" % NS["p"], id=master_id,
                     attrib={"{%s}id" % relation_ns: "rIdSecondMaster"})
    parts["ppt/presentation.xml"] = etree.tostring(presentation, xml_declaration=True, encoding="UTF-8")
    presentation_rels = etree.fromstring(parts["ppt/_rels/presentation.xml.rels"])
    etree.SubElement(presentation_rels, "{%s}Relationship" % NS["r"],
                     Id="rIdSecondMaster", Type=relation_ns + "/slideMaster",
                     Target="slideMasters/slideMaster2.xml")
    parts["ppt/_rels/presentation.xml.rels"] = etree.tostring(presentation_rels, xml_declaration=True, encoding="UTF-8")
    types = etree.fromstring(parts["[Content_Types].xml"])
    additions = {
        "/ppt/theme/theme2.xml": "application/vnd.openxmlformats-officedocument.theme+xml",
        "/ppt/slideMasters/slideMaster2.xml": "application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml",
        "/ppt/slideLayouts/slideLayout12.xml": "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml",
    }
    for name, content_type in additions.items():
        etree.SubElement(types, "{%s}Override" % NS["ct"], PartName=name, ContentType=content_type)
    parts["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
