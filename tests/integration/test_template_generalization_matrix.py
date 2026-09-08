from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from engine.component_atlas import build_component_atlas
from engine.component_composer import build_component_plan
from engine.template_evidence import analyze_template


STYLES = (
    ("light-consulting", (255, 255, 255), (176, 33, 42), 13.333, 7.5, None),
    ("dark-technology", (19, 28, 43), (0, 196, 180), 13.333, 7.5, None),
    ("minimal-business-4x3", (248, 246, 240), (34, 91, 145), 10.0, 7.5, None),
    ("chart-dense", (255, 255, 255), (20, 92, 140), 13.333, 7.5, "chart"),
    ("image-led", (250, 248, 245), (209, 95, 52), 13.333, 7.5, "image"),
    ("cn-enterprise", (246, 249, 252), (189, 0, 22), 13.333, 7.5, None),
)


def _template(path: Path, *, background: tuple[int, int, int], accent: tuple[int, int, int], width: float, height: float, extra: str | None) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*background)
    for index in range(3):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8 + index * (width - 1.6) / 3), Inches(1.7),
            Inches((width - 2.2) / 3), Inches(2.5),
        )
        card.name = f"native-card-{index + 1}"
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(*accent)
        label = slide.shapes.add_textbox(
            card.left + Inches(0.2), card.top + Inches(0.4),
            card.width - Inches(0.4), Inches(1.4),
        )
        label.name = f"native-label-{index + 1}"
        label.text = f"Evidence {index + 1}"
    if extra == "chart":
        data = ChartData()
        data.categories = ["A", "B", "C"]
        data.add_series("Series", (1, 2, 3))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(4.7),
            Inches(3.2), Inches(1.8), data,
        )
        chart.name = "template-native-chart"
    elif extra == "image":
        image_path = path.with_suffix(".png")
        Image.new("RGB", (320, 180), accent).save(image_path)
        picture = slide.shapes.add_picture(
            str(image_path), Inches(width - 3.8), Inches(4.7), Inches(3.2), Inches(1.8),
        )
        picture.name = "template-native-image"
    presentation.save(path)


def _review(component_id: str) -> dict:
    return {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "model", "id": "matrix"}},
        "components": [{
            "component_id": component_id,
            "family": "card_grid",
            "slide_index": 1,
            "semantic_uses": ["three comparable findings"],
            "topologies": ["comparison"],
            "archetypes": ["content"],
            "composition_roles": ["primary"],
            "groups": [
                {"role": "segment", "shape_names": [f"native-card-{i}" for i in range(1, 4)]},
                {"role": "label", "shape_names": [f"native-label-{i}" for i in range(1, 4)]},
            ],
            "parameters": {"element_count": {"minimum": 3, "maximum": 3}},
            "semantic_contract": {"required_fields": ["text"]},
            "text_capacity": {"text": 120},
            "data_contract": {"kinds": ["none"]},
            "native_fidelity": "exact",
            "renderer": "card_grid",
            "page_guidance": {
                "can_stand_alone": False,
                "supported_page_roles": ["body"],
                "minimum_information_units": 3,
                "recommended_page_recipes": ["cards-with-takeaway"],
                "required_companion_roles": ["implication"],
                "recommended_companion_families": ["statement"],
                "annotation_requirements": [],
                "prohibited_scenarios": ["single sparse card"],
            },
        }],
    }


def test_same_template_pipeline_handles_multiple_styles_and_aspect_ratios(tmp_path: Path) -> None:
    for index, (name, background, accent, width, height, extra) in enumerate(STYLES, 1):
        path = tmp_path / f"{name}.pptx"
        _template(path, background=background, accent=accent, width=width, height=height, extra=extra)
        component_id = f"unseen.vendor-{index}.native-cards"
        evidence = analyze_template(path)
        atlas = build_component_atlas(path, _review(component_id))
        component = atlas["components"][0]

        assert evidence["source"]["slide_size_in"] == {"width": round(width, 3), "height": round(height, 3)}
        if extra:
            assert evidence["slides"][0]["geometry"][extra] == 1
        assert component["component_id"] == component_id
        assert component["adaptation_contract"]["responsive"]["source_aspect_ratio"] > 0
        assert "reflow_grid" in component["adaptation_contract"]["responsive"]["modes"]

        plan = build_component_plan(atlas, {
            "schema_version": "1.0.0",
            "pages": [{
                "destination_slide_index": 1,
                "components": [{
                    "component_id": component_id,
                    "semantic_use": "three comparable findings",
                    "topology": "comparison",
                    "placement": {"x": 0.05, "y": 0.18, "w": 0.9, "h": 0.62},
                    "elements": [
                        {"text": f"Complete evidence statement {item}",
                         "binding_name": f"bind:block:s1:findings:item:{item}"}
                        for item in range(3)
                    ],
                }],
            }],
        })
        assert plan["selections"][0]["component_id"] == component_id


def test_template_runtime_has_no_brand_component_ids() -> None:
    root = Path(__file__).resolve().parents[2] / "engine"
    template_modules = [
        "component_atlas.py", "component_composer.py", "component_intent.py",
        "model_template_composer.py", "template_component_adaptation.py",
        "template_page_composition.py", "template_series_planner.py",
        "template_visual_fingerprint.py", "template_visual_quality.py",
        "manuscript_component_planner.py",
    ]
    combined = "\n".join((root / name).read_text(encoding="utf-8") for name in template_modules)

    assert "mckinsey." not in combined.lower()
    assert "mgi" not in combined.lower()
