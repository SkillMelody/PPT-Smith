from __future__ import annotations

import json
from pathlib import Path

from jsonschema import Draft202012Validator
from pptx import Presentation

from engine.speaker_notes import (
    format_speaker_notes,
    inspect_speaker_notes,
    validate_speaker_notes,
    write_speaker_notes,
)


ROOT = Path(__file__).resolve().parents[2]
IR_SCHEMA = json.loads(
    (ROOT / "schemas/v4/presentation-ir.schema.json").read_text(encoding="utf-8")
)


def _notes() -> dict:
    return {
        "narrative": (
            "Explain that adoption is broad but enterprise value depends on workflow redesign, "
            "leadership ownership, data readiness, and disciplined scaling across functions."
        ),
        "talking_points": ["Separate experimentation from scaled deployment."],
        "evidence_ids": ["src:para_1"],
        "source_refs": [{"source_id": "src", "loc": "para_1"}],
        "caveats": ["Survey evidence does not prove causality."],
    }


def _ir(notes: dict | None = None) -> dict:
    slide = {
        "id": "s1",
        "title": "Workflow redesign separates adoption from value",
        "slide_role": "content",
        "blocks": [{
            "id": "finding",
            "role": "fact",
            "text": "Adoption is broad.",
            "source_ref": {"source_id": "src", "loc": "para_1"},
        }],
        "component_intent": {
            "mode": "template_composition",
            "semantic_use": "single KPI comparison",
            "element_count": 1,
            "topology": "comparison",
            "required_slots": ["title", "detail", "metric"],
            "candidate_component_ids": ["template.kpi-card"],
            "selected_component_ids": ["template.kpi-card"],
            "rejected_candidates": [],
        },
    }
    if notes is not None:
        slide["speaker_notes"] = notes
    return {
        "schema_version": "4.1.0",
        "deck": {"title": "Model-directed deck", "language": "en"},
        "sources": [{"source_id": "src", "type": "text"}],
        "slides": [slide],
    }


def test_schema_accepts_model_notes_and_component_intent() -> None:
    assert not list(Draft202012Validator(IR_SCHEMA).iter_errors(_ir(_notes())))


def test_notes_contract_fails_closed_for_missing_body_notes() -> None:
    report = validate_speaker_notes(_ir(), require_body_notes=True)

    assert report["status"] == "fail"
    assert report["issues"] == [{"code": "SPEAKER_NOTES_MISSING", "slide_id": "s1"}]


def test_notes_are_written_and_read_back(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)

    written = write_speaker_notes(deck, {1: _notes()})
    report = inspect_speaker_notes(deck, required_slide_indices=[1])

    assert written["count"] == 1
    assert report["status"] == "pass"
    assert report["text_lengths"]["1"] == len(format_speaker_notes(_notes()))
