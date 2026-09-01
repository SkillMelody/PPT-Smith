from __future__ import annotations

import base64
import zipfile
import shutil
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from engine.template_native import (
    bind_chart_data,
    bind_table_data,
    bind_template_component,
    bind_text_shape,
    clone_chart_shape,
    clone_native_shapes,
    place_native_shapes,
    clone_table_shape,
)
from ppt_qa.package_inspector import inspect_package
from ppt_qa.verifier import run_render_report


def _source_deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = ChartData()
    data.categories = ["Current", "Target"]
    data.add_series("Adoption", (34, 72))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1), Inches(1), Inches(8), Inches(4), data,
    )
    chart.name = "adoption-chart"
    presentation.save(path)


def _blank_deck(path: Path) -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def _table_deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(8), Inches(2))
    table.name = "regional-template"
    for row_index, values in enumerate((("Region", "Growth"), ("US", "20%"))):
        for column_index, value in enumerate(values):
            table.table.cell(row_index, column_index).text = value
    table.table.cell(0, 0).fill.solid()
    table.table.cell(0, 0).fill.fore_color.rgb = RGBColor(181, 29, 37)
    presentation.save(path)


def _component_deck(path: Path, *, count: int = 3) -> None:
    """A minimal native component shell with template-styled prototype nodes."""
    from pptx.enum.shapes import MSO_SHAPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index in range(count):
        segment = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(2.0 + index * 0.3), Inches(1.2 + index * 1.1),
            Inches(5.2 - index * 0.6), Inches(0.8),
        )
        segment.name = f"template-segment-{index + 1}"
        segment.fill.solid()
        segment.fill.fore_color.rgb = RGBColor(181, 29, 37)
        label = slide.shapes.add_textbox(
            Inches(2.1 + index * 0.3), Inches(1.35 + index * 1.1),
            Inches(4.8 - index * 0.6), Inches(0.35),
        )
        label.name = f"template-label-{index + 1}"
        label.text = f"Template {index + 1}"
    presentation.save(path)


def _compound_component_deck(path: Path, *, count: int = 3) -> None:
    """Native funnel levels made of a body plus a separately styled rim."""
    from pptx.enum.shapes import MSO_SHAPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index in range(count):
        x, y, width = 2.0 + index * 0.3, 1.3 + index * 1.1, 5.2 - index * 0.6
        rim = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(0.14))
        rim.name = f"template-rim-{index + 1}"
        rim.fill.solid()
        rim.fill.fore_color.rgb = RGBColor(51, 65, 85)
        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.14), Inches(width), Inches(0.78))
        body.name = f"template-body-{index + 1}"
        body.fill.solid()
        body.fill.fore_color.rgb = RGBColor(181, 29, 37)
        label = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.35), Inches(width - 0.4), Inches(0.3))
        label.name = f"template-label-{index + 1}"
        label.text = f"Template {index + 1}"
    presentation.save(path)


def _extended_component_deck(path: Path, *, include_destination: bool = False) -> None:
    """A pyramid shell with uneven companions, two label fields and one shared base."""
    from pptx.enum.shapes import MSO_SHAPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, width in enumerate((2.0, 3.0, 4.0)):
        left = 4.0 - width / 2
        top = 1.0 + index * 1.1
        segment = slide.shapes.add_shape(
            MSO_SHAPE.TRAPEZOID, Inches(left), Inches(top), Inches(width), Inches(0.8),
        )
        segment.name = f"extended-segment-{index + 1}"
        title = slide.shapes.add_textbox(
            Inches(left + 0.2), Inches(top + 0.14), Inches(width * 0.55), Inches(0.3),
        )
        title.name = f"extended-title-{index + 1}"
        title.text = f"Sample title {index + 1}"
        value = slide.shapes.add_textbox(
            Inches(left + width * 0.68), Inches(top + 0.14), Inches(width * 0.25), Inches(0.3),
        )
        value.name = f"extended-value-{index + 1}"
        value.text = f"{index + 1}0%"
    for name, left, top, width in (
        ("extended-shadow-1-a", 3.0, 1.72, 0.8),
        ("extended-shadow-1-b", 4.2, 1.72, 0.8),
        ("extended-shadow-2", 2.7, 2.82, 1.0),
        ("extended-shadow-3", 2.2, 3.92, 1.2),
    ):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.08),
        )
        shape.name = name
    baseline = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.8), Inches(4.55), Inches(4.4), Inches(0.12),
    )
    baseline.name = "extended-baseline"
    if include_destination:
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def _card_grid_deck(path: Path) -> None:
    """Six editable cards arranged as a native 3x2 prototype grid."""
    from pptx.enum.shapes import MSO_SHAPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for row in range(3):
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.7), Inches(0.8 + row * 1.8), Inches(7.7), Inches(1.6),
        )
        background.name = f"grid-row-bg-{row + 1}"
        background.fill.solid()
        background.fill.fore_color.rgb = RGBColor(240, 242, 245)
    for index in range(6):
        row, column = divmod(index, 3)
        left = 1.0 + column * 2.5
        top = 1.0 + row * 1.8
        body = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left), Inches(top), Inches(2.1), Inches(1.2),
        )
        body.name = f"card-body-{index + 1}"
        title = slide.shapes.add_textbox(
            Inches(left + 0.15), Inches(top + 0.14), Inches(1.8), Inches(0.3),
        )
        title.name = f"card-title-{index + 1}"
        title.text = f"Sample title {index + 1}"
        detail = slide.shapes.add_textbox(
            Inches(left + 0.15), Inches(top + 0.52), Inches(1.8), Inches(0.48),
        )
        detail.name = f"card-detail-{index + 1}"
        detail.text = f"Sample detail {index + 1}"
    presentation.save(path)


def _icon_card_grid_deck(path: Path) -> None:
    """Five icon-card prototypes stored in one source row like a real template."""
    from pptx.enum.shapes import MSO_SHAPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index in range(5):
        left = 0.7 + index * 1.7
        body = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left), Inches(1.5), Inches(1.45), Inches(0.9),
        )
        body.name = f"card-body-{index + 1}"
        title = slide.shapes.add_textbox(
            Inches(left + 0.12), Inches(1.66), Inches(1.21), Inches(0.25),
        )
        title.name = f"card-title-{index + 1}"
        title.text = f"Sample title {index + 1}"
        detail = slide.shapes.add_textbox(
            Inches(left + 0.12), Inches(1.95), Inches(1.21), Inches(0.32),
        )
        detail.name = f"card-detail-{index + 1}"
        detail.text = f"Sample detail {index + 1}"
    presentation.save(path)


def _shape_by_name(slide, name: str):
    return next(shape for shape in slide.shapes if shape.name == name)


def test_bind_template_funnel_reflows_native_shapes_for_data_count_and_value(tmp_path: Path) -> None:
    deck = tmp_path / "funnel.pptx"
    _component_deck(deck, count=3)

    result = bind_template_component(
        deck,
        slide_index=1,
        component_type="funnel",
        segment_names=["template-segment-1", "template-segment-2", "template-segment-3"],
        label_names=["template-label-1", "template-label-2", "template-label-3"],
        elements=[
            {"text": "尝试", "value": 100, "binding_name": "bind:block:agents:stages:item:0"},
            {"text": "试点", "value": 58, "binding_name": "bind:block:agents:stages:item:1"},
            {"text": "规模化", "value": 24, "binding_name": "bind:block:agents:stages:item:2"},
            {"text": "实现具有治理护栏的企业级价值规模化", "value": 9, "binding_name": "bind:block:agents:stages:item:3"},
        ],
    )

    slide = Presentation(deck).slides[0]
    segment_names = [f"decoration:component:funnel:segment:{index}" for index in range(4)]
    widths = [_shape_by_name(slide, name).width for name in segment_names]
    assert result["element_count"] == 4
    assert widths == sorted(widths, reverse=True)
    label = _shape_by_name(slide, "bind:block:agents:stages:item:3")
    assert label.text == "实现具有治理护栏的企业级价值规模化"
    assert label.text_frame.word_wrap is True
    assert label.text_frame.paragraphs[0].runs[0].font.size <= Pt(14)
    assert inspect_package(deck).status == "passed"


def test_bind_template_funnel_reflows_all_native_parts_of_each_level(tmp_path: Path) -> None:
    deck = tmp_path / "compound-funnel.pptx"
    _compound_component_deck(deck, count=3)

    result = bind_template_component(
        deck,
        slide_index=1,
        component_type="funnel",
        segment_names=None,
        segment_groups=[
            ["template-body-1", "template-rim-1"],
            ["template-body-2", "template-rim-2"],
            ["template-body-3", "template-rim-3"],
        ],
        label_names=["template-label-1", "template-label-2", "template-label-3"],
        elements=[
            {"text": "探索", "value": 100, "binding_name": "bind:block:funnel:item:0"},
            {"text": "试点", "value": 60, "binding_name": "bind:block:funnel:item:1"},
            {"text": "规模化", "value": 30, "binding_name": "bind:block:funnel:item:2"},
            {"text": "固化价值", "value": 10, "binding_name": "bind:block:funnel:item:3"},
        ],
    )

    slide = Presentation(deck).slides[0]
    bodies = [_shape_by_name(slide, f"decoration:component:funnel:segment:{index}") for index in range(4)]
    rims = [_shape_by_name(slide, f"decoration:component:funnel:segment:{index}:part:1") for index in range(4)]
    assert result["source_segment_count"] == 3
    assert [body.width for body in bodies] == sorted((body.width for body in bodies), reverse=True)
    assert all(abs(rim.width - body.width) <= 1 for rim, body in zip(rims, bodies))
    assert all(rim.top < body.top for rim, body in zip(rims, bodies))
    assert inspect_package(deck).status == "passed"


def test_bind_template_component_places_the_reflowed_component_inside_a_declared_slot(tmp_path: Path) -> None:
    deck = tmp_path / "placed-funnel.pptx"
    _component_deck(deck, count=3)

    bind_template_component(
        deck,
        slide_index=1,
        component_type="funnel",
        segment_names=["template-segment-1", "template-segment-2", "template-segment-3"],
        label_names=["template-label-1", "template-label-2", "template-label-3"],
        elements=[
            {"text": "探索", "value": 100, "binding_name": "bind:block:funnel:item:0"},
            {"text": "试点", "value": 50, "binding_name": "bind:block:funnel:item:1"},
            {"text": "规模化", "value": 20, "binding_name": "bind:block:funnel:item:2"},
        ],
        placement={"x": 0.55, "y": 0.2, "w": 0.4, "h": 0.6},
    )

    presentation = Presentation(deck)
    slide = presentation.slides[0]
    placed = [
        shape for shape in slide.shapes
        if shape.name.startswith("decoration:component:funnel:")
        or shape.name.startswith("bind:block:funnel:")
    ]
    left = min(shape.left for shape in placed) / presentation.slide_width
    top = min(shape.top for shape in placed) / presentation.slide_height
    right = max(shape.left + shape.width for shape in placed) / presentation.slide_width
    bottom = max(shape.top + shape.height for shape in placed) / presentation.slide_height
    assert left >= 0.548
    assert top >= 0.198
    assert right <= 0.952
    assert bottom <= 0.802


def test_bind_template_component_names_decorations_with_a_component_instance_id(tmp_path: Path) -> None:
    deck = tmp_path / "instance-funnel.pptx"
    _component_deck(deck, count=3)

    bind_template_component(
        deck,
        slide_index=1,
        component_type="funnel",
        segment_names=["template-segment-1", "template-segment-2", "template-segment-3"],
        label_names=["template-label-1", "template-label-2", "template-label-3"],
        elements=[
            {"text": "探索", "value": 100, "binding_name": "bind:block:funnel:item:0"},
            {"text": "规模化", "value": 20, "binding_name": "bind:block:funnel:item:1"},
        ],
        component_instance_id="slide-4-component-1",
    )

    names = {shape.name for shape in Presentation(deck).slides[0].shapes}
    assert "decoration:component:funnel:slide-4-component-1:segment:0" in names
    assert "decoration:component:funnel:slide-4-component-1:segment:1" in names


def test_bind_template_pyramid_reflows_native_layers_for_data_count_and_value(tmp_path: Path) -> None:
    deck = tmp_path / "pyramid.pptx"
    _component_deck(deck, count=3)

    result = bind_template_component(
        deck,
        slide_index=1,
        component_type="pyramid",
        segment_names=["template-segment-1", "template-segment-2", "template-segment-3"],
        label_names=["template-label-1", "template-label-2", "template-label-3"],
        elements=[
            {"text": "战略方向", "value": 20, "binding_name": "bind:block:strategy:levels:item:0"},
            {"text": "组织能力", "value": 45, "binding_name": "bind:block:strategy:levels:item:1"},
            {"text": "流程重塑", "value": 70, "binding_name": "bind:block:strategy:levels:item:2"},
            {"text": "日常执行", "value": 90, "binding_name": "bind:block:strategy:levels:item:3"},
        ],
    )

    slide = Presentation(deck).slides[0]
    widths = [
        _shape_by_name(slide, f"decoration:component:pyramid:segment:{index}").width
        for index in range(4)
    ]
    assert result["element_count"] == 4
    assert widths == sorted(widths)
    assert _shape_by_name(slide, "bind:block:strategy:levels:item:0").text == "战略方向"
    assert inspect_package(deck).status == "passed"


def test_bind_template_component_consumes_extended_item_and_shared_contract(tmp_path: Path) -> None:
    deck = tmp_path / "extended-pyramid.pptx"
    _extended_component_deck(deck)

    result = bind_template_component(
        deck,
        slide_index=1,
        component_type="pyramid",
        segment_names=None,
        label_names=None,
        elements=[
            {
                "value": value,
                "labels": {
                    "title": {
                        "text": title,
                        "binding_name": f"bind:block:pyramid:levels:item:{index}",
                    },
                    "value": {
                        "text": detail,
                        "binding_name": f"bind:block:pyramid:levels:detail:{index}",
                    },
                },
            }
            for index, (title, detail, value) in enumerate((
                ("战略方向", "20%", 20),
                ("组织能力", "45%", 45),
                ("流程重塑", "70%", 70),
                ("日常执行", "90%", 90),
            ))
        ],
        item_groups=[
            {
                "segment_names": [
                    "extended-segment-1", "extended-shadow-1-a", "extended-shadow-1-b",
                ],
                "label_fields": {
                    "title": ["extended-title-1"],
                    "value": ["extended-value-1"],
                },
            },
            {
                "segment_names": ["extended-segment-2", "extended-shadow-2"],
                "label_fields": {
                    "title": ["extended-title-2"],
                    "value": ["extended-value-2"],
                },
            },
            {
                "segment_names": ["extended-segment-3", "extended-shadow-3"],
                "label_fields": {
                    "title": ["extended-title-3"],
                    "value": ["extended-value-3"],
                },
            },
        ],
        shared_names=["extended-baseline"],
        component_instance_id="slide-2-component-1",
    )

    slide = Presentation(deck).slides[0]
    names = {shape.name for shape in slide.shapes}
    assert result["binding_schema_version"] == "2.0.0"
    assert result["element_count"] == 4
    assert len(result["binding_names"]) == 8
    assert "decoration:component:pyramid:slide-2-component-1:segment:3" in names
    assert "decoration:component:pyramid:slide-2-component-1:shared:0" in names
    assert _shape_by_name(slide, "bind:block:pyramid:levels:item:3").text == "日常执行"
    assert _shape_by_name(slide, "bind:block:pyramid:levels:detail:3").text == "90%"
    assert not any("Sample title" in getattr(shape, "text", "") for shape in slide.shapes)
    assert inspect_package(deck).status == "passed"


def test_two_sided_contrast_places_secondary_labels_inside_opposing_segments(tmp_path: Path) -> None:
    deck = tmp_path / "two-sided-contrast.pptx"
    _extended_component_deck(deck)

    bind_template_component(
        deck,
        slide_index=1,
        component_type="two_sided_contrast",
        layout_mode="two_sided_contrast",
        segment_names=None,
        label_names=None,
        elements=[
            {
                "value": None,
                "labels": {
                    "title": {"text": "规模化", "binding_name": "bind:block:contrast:item:0"},
                    "value": {"text": "已进入业务流程", "binding_name": "bind:block:contrast:detail:0"},
                },
            },
            {
                "value": None,
                "labels": {
                    "title": {"text": "试验阶段", "binding_name": "bind:block:contrast:item:1"},
                    "value": {"text": "仍停留在局部试点", "binding_name": "bind:block:contrast:detail:1"},
                },
            },
        ],
        item_groups=[
            {
                "segment_names": ["extended-segment-1"],
                "label_fields": {
                    "title": ["extended-title-1"],
                    "value": ["extended-value-1"],
                },
            },
            {
                "segment_names": ["extended-segment-2"],
                "label_fields": {
                    "title": ["extended-title-2"],
                    "value": ["extended-value-2"],
                },
            },
        ],
        shared_names=[],
    )

    slide = Presentation(deck).slides[0]
    first_segment = _shape_by_name(slide, "decoration:component:two_sided_contrast:segment:0")
    second_segment = _shape_by_name(slide, "decoration:component:two_sided_contrast:segment:1")
    first_detail = _shape_by_name(slide, "bind:block:contrast:detail:0")
    second_detail = _shape_by_name(slide, "bind:block:contrast:detail:1")
    assert first_detail.top >= first_segment.top + int(first_segment.height * 0.48)
    assert first_detail.top + first_detail.height <= first_segment.top + first_segment.height
    assert second_detail.top <= second_segment.top + int(second_segment.height * 0.18)
    assert second_detail.top + second_detail.height <= second_segment.top + second_segment.height


def test_bind_template_card_grid_reflows_five_cards_as_centered_three_plus_two(tmp_path: Path) -> None:
    deck = tmp_path / "card-grid.pptx"
    _card_grid_deck(deck)

    result = bind_template_component(
        deck,
        slide_index=1,
        component_type="card_grid",
        segment_names=None,
        label_names=None,
        elements=[
            {
                "value": None,
                "labels": {
                    "title": {
                        "text": f"行动 {index + 1}",
                        "binding_name": f"bind:block:initiatives:cards:item:{index}",
                    },
                    "detail": {
                        "text": f"说明 {index + 1}",
                        "binding_name": f"bind:block:initiatives:cards:detail:{index}",
                    },
                },
            }
            for index in range(5)
        ],
        item_groups=[
            {
                "segment_names": [f"card-body-{index + 1}"],
                "label_fields": {
                    "title": [f"card-title-{index + 1}"],
                    "detail": [f"card-detail-{index + 1}"],
                },
            }
            for index in range(6)
        ],
        shared_names=["grid-row-bg-1", "grid-row-bg-2", "grid-row-bg-3"],
        component_instance_id="preview-5",
    )

    slide = Presentation(deck).slides[0]
    cards = [
        _shape_by_name(slide, f"decoration:component:card_grid:preview-5:segment:{index}")
        for index in range(5)
    ]
    assert result["element_count"] == 5
    assert len({card.top for card in cards[:3]}) == 1
    assert len({card.top for card in cards[3:]}) == 1
    assert cards[3].top > cards[0].top
    first_row_center = (cards[0].left + cards[2].left + cards[2].width) / 2
    second_row_center = (cards[3].left + cards[4].left + cards[4].width) / 2
    assert abs(first_row_center - second_row_center) <= 2
    assert _shape_by_name(slide, "bind:block:initiatives:cards:item:4").text == "行动 5"
    assert _shape_by_name(slide, "bind:block:initiatives:cards:detail:4").text == "说明 5"
    backgrounds = [
        shape for shape in slide.shapes
        if shape.name.startswith("decoration:component:card_grid:preview-5:shared:")
    ]
    assert len(backgrounds) == 2
    for background, card in zip(backgrounds, (cards[0], cards[3])):
        assert background.top < card.top
        assert background.top + background.height > card.top + card.height
    z_order = {shape.name: index for index, shape in enumerate(slide.shapes)}
    assert max(z_order[background.name] for background in backgrounds) < min(
        z_order[card.name] for card in cards
    )
    assert not any("Sample" in getattr(shape, "text", "") for shape in slide.shapes)
    assert inspect_package(deck).status == "passed"


def test_bind_template_icon_card_grid_keeps_five_cards_tall_enough_for_detail_text(tmp_path: Path) -> None:
    deck = tmp_path / "icon-card-grid.pptx"
    _icon_card_grid_deck(deck)

    result = bind_template_component(
        deck,
        slide_index=1,
        component_type="icon_card_grid",
        segment_names=None,
        label_names=None,
        elements=[
            {
                "value": None,
                "labels": {
                    "title": {
                        "text": f"能力 {index + 1}",
                        "binding_name": f"bind:block:capabilities:cards:item:{index}",
                    },
                    "detail": {
                        "text": f"能力说明 {index + 1}",
                        "binding_name": f"bind:block:capabilities:cards:detail:{index}",
                    },
                },
            }
            for index in range(5)
        ],
        item_groups=[
            {
                "segment_names": [f"card-body-{index + 1}"],
                "label_fields": {
                    "title": [f"card-title-{index + 1}"],
                    "detail": [f"card-detail-{index + 1}"],
                },
            }
            for index in range(5)
        ],
        shared_names=[],
        component_instance_id="icon-preview-5",
    )

    slide = Presentation(deck).slides[0]
    cards = [
        _shape_by_name(
            slide,
            f"decoration:component:icon_card_grid:icon-preview-5:segment:{index}",
        )
        for index in range(5)
    ]
    assert result["element_count"] == 5
    assert cards[3].top > cards[0].top
    assert all(card.height / card.width >= 0.55 for card in cards)


def test_bind_template_timeline_reflows_native_nodes_for_data_count(tmp_path: Path) -> None:
    deck = tmp_path / "timeline.pptx"
    _component_deck(deck, count=3)

    result = bind_template_component(
        deck,
        slide_index=1,
        component_type="timeline",
        segment_names=["template-segment-1", "template-segment-2", "template-segment-3"],
        label_names=["template-label-1", "template-label-2", "template-label-3"],
        elements=[
            {"text": f"阶段 {index}", "binding_name": f"bind:block:roadmap:steps:item:{index}"}
            for index in range(5)
        ],
    )

    slide = Presentation(deck).slides[0]
    lefts = [
        _shape_by_name(slide, f"decoration:component:timeline:segment:{index}").left
        for index in range(5)
    ]
    assert result["element_count"] == 5
    assert lefts == sorted(lefts)
    assert _shape_by_name(slide, "bind:block:roadmap:steps:item:4").text == "阶段 4"
    assert inspect_package(deck).status == "passed"


def test_clone_native_shapes_moves_an_editable_component_and_icon_between_slides(tmp_path: Path) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    deck = tmp_path / "native-component.pptx"
    icon_path = tmp_path / "icon.png"
    icon_path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/"
        "lRBQZQAAAABJRU5ErkJggg=="
    ))
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    segment = source.shapes.add_shape(
        MSO_SHAPE.TRAPEZOID, Inches(1), Inches(1), Inches(4), Inches(0.8),
    )
    segment.name = "component-segment"
    label = source.shapes.add_textbox(Inches(1.4), Inches(1.2), Inches(3), Inches(0.3))
    label.name = "component-label"
    label.text = "Template sample"
    icon = source.shapes.add_picture(str(icon_path), Inches(5.5), Inches(1), Inches(0.4), Inches(0.4))
    icon.name = "component-icon"
    destination = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = destination.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.4))
    title.name = "destination-title"
    title.text = "Keep me"
    presentation.save(deck)

    copied = clone_native_shapes(
        deck,
        source_slide_index=1,
        destination_slide_index=2,
        shape_names=["component-segment", "component-label", "component-icon"],
    )

    result = Presentation(deck)
    destination = result.slides[1]
    names = {shape.name for shape in destination.shapes}
    assert copied["shape_names"] == {
        "component-segment": "component-segment",
        "component-label": "component-label",
        "component-icon": "component-icon",
    }
    assert names == {
        "destination-title", "component-segment", "component-label", "component-icon",
    }
    assert _shape_by_name(destination, "component-segment").auto_shape_type == MSO_SHAPE.TRAPEZOID
    assert _shape_by_name(destination, "component-label").text == "Template sample"
    assert _shape_by_name(destination, "component-icon").shape_type is not None
    assert _shape_by_name(destination, "destination-title").text == "Keep me"
    assert inspect_package(deck).status == "passed"


def test_clone_chart_shape_copies_chart_relationship_closure(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    destination = tmp_path / "destination.pptx"
    _source_deck(source)
    _blank_deck(destination)

    result = clone_chart_shape(source, destination, source_slide_index=1, destination_slide_index=1)

    assert result["chart_name"] == "adoption-chart"
    assert result["copied_parts"]
    assert Presentation(destination).slides[0].shapes[0].has_chart is True

    with zipfile.ZipFile(destination) as archive:
        names = set(archive.namelist())
        chart_parts = {name for name in names if name.startswith("ppt/charts/chart") and name.endswith(".xml")}
        workbook_parts = {name for name in names if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")}
        assert len(chart_parts) == 1
        assert len(workbook_parts) == 1
        chart_relationships = archive.read("ppt/charts/_rels/chart1.xml.rels")
        assert b"../embeddings/" in chart_relationships
        copied_workbook = archive.read(next(iter(workbook_parts)))
    with zipfile.ZipFile(source) as archive:
        source_workbook = archive.read("ppt/embeddings/Microsoft_Excel_Sheet1.xlsx")
    assert copied_workbook == source_workbook
    assert inspect_package(destination).status == "passed"


def test_clone_chart_shape_remaps_parts_when_destination_already_has_chart(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    destination = tmp_path / "destination.pptx"
    _source_deck(source)
    _source_deck(destination)

    clone_chart_shape(source, destination, source_slide_index=1, destination_slide_index=1)

    with zipfile.ZipFile(destination) as archive:
        names = set(archive.namelist())
        chart_parts = {name for name in names if name.startswith("ppt/charts/chart") and name.endswith(".xml")}
        workbook_parts = {name for name in names if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")}
        assert len(chart_parts) == 2
        assert len(workbook_parts) == 2
    assert sum(shape.has_chart for shape in Presentation(destination).slides[0].shapes) == 2
    assert inspect_package(destination).status == "passed"


def test_clone_table_shape_preserves_template_style_and_binds_source_rows(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    destination = tmp_path / "destination.pptx"
    _table_deck(source)
    _blank_deck(destination)

    copied = clone_table_shape(source, destination, source_slide_index=1, destination_slide_index=1)
    bound = bind_table_data(
        destination,
        slide_index=1,
        table_name=copied["table_name"],
        binding_name="bind:table:growth:regional",
        rows=[["Region", "Growth"], ["EU", "15%"]],
    )

    table = Presentation(destination).slides[0].shapes[0]
    assert copied["table_name"] == "regional-template"
    assert bound["binding_name"] == "bind:table:growth:regional"
    assert table.name == "bind:table:growth:regional"
    assert [[cell.text for cell in row.cells] for row in table.table.rows] == [["Region", "Growth"], ["EU", "15%"]]
    assert str(table.table.cell(0, 0).fill.fore_color.rgb) == "B51D25"
    assert inspect_package(destination).status == "passed"
    if Path("/Applications/LibreOffice.app").exists() or shutil.which("soffice"):
        render = run_render_report(
            destination, tmp_path / "table-render", engine="libreoffice", expected_slides=1,
            expected_aspect=4 / 3,
        )
        assert render["status"] == "passed", render


def test_bind_chart_data_updates_native_workbook_without_replacing_the_chart_shape(tmp_path: Path) -> None:
    deck = tmp_path / "chart.pptx"
    _source_deck(deck)

    result = bind_chart_data(
        deck, slide_index=1, chart_name="adoption-chart", binding_name="bind:chart:impact",
        categories=["Innovation", "Employees", "Customers", "Differentiation"],
        series=[{"name": "Improved", "values": [64, 45, 45, 45]}],
    )

    chart = Presentation(deck).slides[0].shapes[0]
    assert result["binding_name"] == "bind:chart:impact"
    assert chart.name == "bind:chart:impact"
    assert [str(item.label) for item in chart.chart.plots[0].categories] == [
        "Innovation", "Employees", "Customers", "Differentiation",
    ]
    assert list(chart.chart.series[0].values) == [64.0, 45.0, 45.0, 45.0]
    assert inspect_package(deck).status == "passed"


def test_bind_text_shape_retains_an_existing_template_text_frame(tmp_path: Path) -> None:
    deck = tmp_path / "text.pptx"
    presentation = Presentation()
    shape = presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1),
    )
    shape.name = "template-title"
    shape.text = "Old title"
    presentation.save(deck)

    result = bind_text_shape(
        deck, slide_index=1, shape_name="template-title",
        binding_name="bind:slide:impact:title", text="AI impact",
    )

    text = Presentation(deck).slides[0].shapes[0]
    assert result["binding_name"] == "bind:slide:impact:title"
    assert text.name == "bind:slide:impact:title"
    assert text.text == "AI impact"


def test_place_native_shapes_can_stretch_a_text_shell_to_the_declared_box(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.5))
    shape.name = "title-shell"
    shape.text = "Title"
    presentation.save(deck)

    result = place_native_shapes(
        deck,
        slide_index=1,
        shape_names=["title-shell"],
        placement={"x": 0.1, "y": 0.2, "w": 0.8, "h": 0.1, "fit_mode": "stretch"},
    )

    updated = Presentation(deck)
    placed = updated.slides[0].shapes[0]
    assert round(placed.left / updated.slide_width, 3) == 0.1
    assert round(placed.top / updated.slide_height, 3) == 0.2
    assert round(placed.width / updated.slide_width, 3) == 0.8
    assert round(placed.height / updated.slide_height, 3) == 0.1
    assert result["fit_mode"] == "stretch"


@pytest.mark.skipif(
    not Path("/Applications/LibreOffice.app").exists() and not shutil.which("soffice"),
    reason="real renderer required",
)
def test_cloned_chart_renders_with_its_embedded_workbook(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    destination = tmp_path / "destination.pptx"
    _source_deck(source)
    _blank_deck(destination)

    clone_chart_shape(source, destination, source_slide_index=1, destination_slide_index=1)
    render = run_render_report(
        destination, tmp_path / "render", engine="libreoffice", expected_slides=1,
        expected_aspect=4 / 3,
    )

    assert render["status"] == "passed", render
