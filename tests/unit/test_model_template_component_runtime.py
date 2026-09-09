from __future__ import annotations

import hashlib
from pathlib import Path
import shutil

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from engine.model_template_component_runtime import author_model_template_component


def _template(path: Path) -> Path:
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    style = source.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(4), Inches(1),
    )
    style.fill.solid()
    style.fill.fore_color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    style.line.color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    style.text = "Template style"
    run = style.text_frame.paragraphs[0].runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 0, 0)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    return path


def _ir() -> dict:
    return {
        "slides": [{
            "id": "s1",
            "title": "Automation adoption",
            "message": "Adoption rises from 34% to 72%.",
            "chart": {
                "type": "bar",
                "data": {
                    "categories": ["Current", "Target"],
                    "series": [{"name": "Adoption", "values": [34, 72]}],
                },
                "source_ref": {"source_id": "doc", "loc": "table_1"},
            },
            "blocks": [],
        }],
    }


def _sha(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _write_script(path: Path, *, unbound: bool = False, wrong_color: bool = False) -> Path:
    name = "TextBox 1" if unbound else "bind:slide:s1:title"
    color = "0x00, 0x66, 0xCC" if wrong_color else "0xC8, 0x10, 0x2E"
    path.write_text(
        "\n".join([
            "from pptx.chart.data import ChartData",
            "from pptx.dml.color import RGBColor",
            "from pptx.enum.chart import XL_CHART_TYPE",
            "from pptx.enum.shapes import MSO_SHAPE",
            "from pptx.util import Inches, Pt",
            "",
            "def build(slide, context):",
            "    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4.5), Inches(0.6))",
            f"    title.name = {name!r}",
            "    title.text = context['slide']['title']",
            "    run = title.text_frame.paragraphs[0].runs[0]",
            "    run.font.name = 'Arial'",
            "    run.font.size = Pt(24)",
            "    run.font.color.rgb = RGBColor(0, 0, 0)",
            "    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.4), Inches(4.5), Inches(3.2))",
            "    panel.name = 'decoration:metric-panel'",
            "    panel.fill.solid()",
            f"    panel.fill.fore_color.rgb = RGBColor({color})",
            f"    panel.line.color.rgb = RGBColor({color})",
            "    data = ChartData()",
            "    data.categories = context['slide']['chart']['data']['categories']",
            "    series = context['slide']['chart']['data']['series'][0]",
            "    data.add_series(series['name'], tuple(series['values']))",
            "    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.8), Inches(3.9), Inches(2.4), data)",
            "    chart.name = 'bind:chart:s1'",
        ]) + "\n",
        encoding="utf-8",
    )
    return path


def test_model_template_component_authors_native_bound_shapes_with_template_tokens(
    tmp_path: Path,
) -> None:
    template = _template(tmp_path / "template.pptx")
    candidate = tmp_path / "candidate.pptx"
    shutil.copyfile(template, candidate)
    script = _write_script(tmp_path / "author.py")

    result = author_model_template_component(
        candidate,
        template_pptx=template,
        script_path=script,
        slide_index=2,
        slide_id="s1",
        component_id="model.adoption-chart",
        placement={"x": 0, "y": 0, "w": 0.6, "h": 0.8},
        ir=_ir(),
        required_binding_names=["bind:slide:s1:title", "bind:chart:s1"],
    )

    assert result["status"] == "pass"
    assert result["new_shape_count"] == 3
    assert result["binding_names"] == ["bind:chart:s1", "bind:slide:s1:title"]
    output = Presentation(candidate)
    assert any(shape.has_chart for shape in output.slides[1].shapes)


@pytest.mark.parametrize(
    ("unbound", "wrong_color", "code"),
    [
        (True, False, "MODEL_TEMPLATE_COMPONENT_CONTENT_UNBOUND"),
        (False, True, "MODEL_TEMPLATE_COMPONENT_COLOR_OUTSIDE_TEMPLATE"),
    ],
)
def test_model_template_component_rejects_unbound_or_off_style_content(
    tmp_path: Path, unbound: bool, wrong_color: bool, code: str,
) -> None:
    template = _template(tmp_path / "template.pptx")
    candidate = tmp_path / "candidate.pptx"
    shutil.copyfile(template, candidate)
    script = _write_script(
        tmp_path / "author.py", unbound=unbound, wrong_color=wrong_color,
    )

    with pytest.raises(ValueError, match=code):
        author_model_template_component(
            candidate,
            template_pptx=template,
            script_path=script,
            slide_index=2,
            slide_id="s1",
            component_id="model.adoption-chart",
            placement={"x": 0, "y": 0, "w": 0.6, "h": 0.8},
            ir=_ir(),
        )


def test_model_template_component_allows_only_audited_standalone_source_image(
    tmp_path: Path,
) -> None:
    template = _template(tmp_path / "template.pptx")
    candidate = tmp_path / "candidate.pptx"
    shutil.copyfile(template, candidate)
    image_path = tmp_path / "illustration.png"
    Image.new("RGB", (80, 50), (200, 16, 46)).save(image_path)
    script = tmp_path / "image_author.py"
    script.write_text(
        "from pptx.util import Inches\n"
        "def build(slide, context):\n"
        "    asset = context['asset_bindings'][0]\n"
        "    picture = slide.shapes.add_picture(asset['asset_path'], Inches(0.5), Inches(0.5), Inches(3), Inches(2))\n"
        "    picture.name = asset['binding_name']\n",
        encoding="utf-8",
    )
    ir = _ir()
    ir["slides"][0]["blocks"] = [{
        "id": "robot",
        "role": "image",
        "asset_ref": "source.robot-illustration",
        "asset_kind": "illustration",
        "reconstructable": False,
        "crop_audit": {
            "includes_page_header": False,
            "includes_page_footer": False,
            "includes_navigation": False,
            "includes_body_prose": False,
            "text_area_ratio": 0.0,
        },
        "source_ref": {"source_id": "doc", "loc": "img_1"},
    }]
    asset_binding = {
        "binding_name": "bind:image:s1:robot",
        "asset_ref": "source.robot-illustration",
        "asset_path": str(image_path),
        "sha256": _sha(image_path),
    }

    result = author_model_template_component(
        candidate,
        template_pptx=template,
        script_path=script,
        slide_index=2,
        slide_id="s1",
        component_id="model.illustration",
        placement={"x": 0, "y": 0, "w": 0.6, "h": 0.8},
        ir=ir,
        required_binding_names=["bind:image:s1:robot"],
        asset_bindings=[asset_binding],
    )

    assert result["approved_source_images"] == ["bind:image:s1:robot"]

    ir["slides"][0]["blocks"][0]["asset_origin"] = "external"
    shutil.copyfile(template, candidate)
    with pytest.raises(
        ValueError,
        match="MODEL_TEMPLATE_COMPONENT_EXTERNAL_ASSET_RIGHTS_REQUIRED",
    ):
        author_model_template_component(
            candidate,
            template_pptx=template,
            script_path=script,
            slide_index=2,
            slide_id="s1",
            component_id="model.illustration",
            placement={"x": 0, "y": 0, "w": 0.6, "h": 0.8},
            ir=ir,
            asset_bindings=[asset_binding],
        )
    ir["slides"][0]["blocks"][0].pop("asset_origin")

    ir["slides"][0]["blocks"][0]["asset_kind"] = "page_screenshot"
    shutil.copyfile(template, candidate)
    with pytest.raises(ValueError, match="MODEL_TEMPLATE_COMPONENT_IMAGE_NOT_STANDALONE"):
        author_model_template_component(
            candidate,
            template_pptx=template,
            script_path=script,
            slide_index=2,
            slide_id="s1",
            component_id="model.illustration",
            placement={"x": 0, "y": 0, "w": 0.6, "h": 0.8},
            ir=ir,
            asset_bindings=[asset_binding],
        )
