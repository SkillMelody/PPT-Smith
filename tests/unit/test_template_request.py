from __future__ import annotations

import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from engine.production_request import prepare_template_context


def _reference(path: Path, *, slides: int = 3) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8), Inches(0.6))
        box.text = f"Reference page {index + 1}"
    prs.save(path)


def _sha(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _request(path: Path, *, mode: str = "style_transfer", indices: list[int] | None = None) -> dict:
    template = {"use_mode": mode, "source_sha256": _sha(path), "required": True}
    if indices is not None:
        template["preferred_slide_indices"] = indices
    return {
        "schema_version": "1.0.0",
        "request_id": "template-context-test",
        "route": "template",
        "page_contract": {
            "mode": "exact", "exact_pages": 12,
            "overflow_policy": "ask", "underflow_policy": "fewer_pages",
        },
        "template": template,
    }


def test_uploaded_template_context_is_source_bound_and_visual_review_gated(tmp_path: Path) -> None:
    template = tmp_path / "reference.pptx"
    _reference(template)

    result = prepare_template_context(_request(template), template)

    assert result["status"] == "ready_for_visual_interpretation"
    assert result["use_mode"] == "style_transfer"
    assert result["source"]["sha256"] == _sha(template)
    assert result["source"]["slide_count"] == 3
    assert result["review"] == {"required": True, "state": "unreviewed"}
    assert result["generation_policy"]["template_slide_count_limits_output"] is False
    assert "evidence" in result


def test_template_context_does_not_replace_user_page_contract(tmp_path: Path) -> None:
    template = tmp_path / "reference.pptx"
    _reference(template, slides=3)

    result = prepare_template_context(_request(template), template)

    assert result["source"]["slide_count"] == 3
    assert result["requested_output_pages"] == 12


def test_template_context_rejects_wrong_file_for_declared_hash(tmp_path: Path) -> None:
    expected = tmp_path / "expected.pptx"
    actual = tmp_path / "actual.pptx"
    _reference(expected, slides=2)
    _reference(actual, slides=4)
    request = _request(expected)

    result = prepare_template_context(request, actual)

    assert result["status"] == "rejected"
    assert result["code"] == "TEMPLATE_SOURCE_HASH_MISMATCH"


def test_template_context_rejects_unknown_preferred_slide(tmp_path: Path) -> None:
    template = tmp_path / "reference.pptx"
    _reference(template, slides=3)

    result = prepare_template_context(_request(template, indices=[1, 4]), template)

    assert result["status"] == "rejected"
    assert result["code"] == "TEMPLATE_SLIDE_UNKNOWN"
    assert result["unknown_slide_indices"] == [4]


def test_required_template_cannot_be_omitted(tmp_path: Path) -> None:
    template = tmp_path / "reference.pptx"
    _reference(template)

    result = prepare_template_context(_request(template), None)

    assert result["status"] == "rejected"
    assert result["code"] == "TEMPLATE_REQUIRED_MISSING"
