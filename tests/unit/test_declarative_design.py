"""Observable rendering, integrity, and adversarial-input regression cases."""
import copy
import io
import json
import os
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData

from engine.design_scene.backend import inspect_objects, inspect_package, render
from engine.design_scene.contract import validate, versions
from engine.design_scene.pipeline import (assess, build, build_prefix, deliver, image_request,
                                         ingest, new_task, read_assets, review_skeleton, verify_build)
from engine.design_scene.store import DesignError, Store, digest, parse_json, normalized_png


def png(color="white", size=(960, 540)):
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "PNG")
    return b.getvalue()


FONT = {"family": "Arial", "size": 20, "color": "182535"}


def node(kind, ident, *, box=None, **fields):
    roles = {"text": "native_text", "shape": "native_shape", "path": "native_shape", "group": "group",
             "table": "native_table", "chart": "native_chart", "image": "replaceable_image"}
    return {"id": ident, "page_id": "page-1", "type": kind, "role": roles[kind], "z": 1,
            "box": box or {"x": 30, "y": 30, "width": 600, "height": 50}, **fields}


def content(kind, ident, **fields):
    return {"id": ident, "page_id": "page-1", "type": kind, "required_edit": "native",
            "provenance": {"kind": "user", "source": "test data", "checked": True, "uncertainties": []}, **fields}


@pytest.fixture
def task_root(tmp_path):
    root = tmp_path / "task"
    task = new_task("fixture-task", mode="recreate", width=960, height=540, author="author")
    with Store(root, create=True) as store:
        store.put_json("task.json", task)
        ingest(store, task, asset_id="reference-a", data=png(), page_ids=["page-1"], role="reference", source="synthetic input")
        task = store.json("task.json")
        task["design"]["references"] = [{"asset_id": "reference-a", "page_ids": ["page-1"],
                                          "roles": ["reconstruction_target"], "priority": 1, "features": []}]
        task["design"]["pages"][0].update(target_asset_id="reference-a", target_kind="original", confirmed_by="user")
        task["content"]["items"] = [content("text", "title", text="A meaningful title\nEditable continuation")]
        task["scene"]["pages"][0]["nodes"] = [node("text", "title-object", content_id="title", font=FONT,
                                                        align="left", valign="top", wrap=True,
                                                        box={"x": 30, "y": 30, "width": 880, "height": 70})]
        store.put_json("task.json", task)
    return root


def load(root):
    with Store(root) as s:
        return s.json("task.json")


def test_native_text_retains_paragraph_editing(task_root):
    with Store(task_root) as s:
        t = s.json("task.json")
        data = render(t, read_assets(s, t))
    assert inspect_objects(data, t)["pages"][0]["objects"] == {"native_text": 1}
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides[0].shapes) == 1
    prs.slides[0].shapes[0].text = "Changed body remains editable"
    out = io.BytesIO()
    prs.save(out)
    assert Presentation(io.BytesIO(out.getvalue())).slides[0].shapes[0].text == "Changed body remains editable"


def test_chart_and_table_data_are_native_and_editable(task_root):
    with Store(task_root) as s:
        t = s.json("task.json")
        t["content"]["items"] += [content("chart", "data", categories=["A", "B"], series=[{"name": "Revenue", "values": [18, 43]}], data_basis="user_data"),
                                    content("table", "rows", cells=[["Name", "Count"], ["A", "18"], ["B", "43"]])]
        t["scene"]["pages"][0]["nodes"] += [node("chart", "chart", content_id="data", chart_type="column", font=FONT,
                                                     colors=["22AA99"], legend=False, labels=True,
                                                     box={"x": 30, "y": 150, "width": 430, "height": 300}),
                                                node("table", "table", content_id="rows", font=FONT,
                                                     fill="FFFFFF", header_fill="182535", header_color="FFFFFF",
                                                     box={"x": 510, "y": 180, "width": 390, "height": 170})]
        data = render(t, read_assets(s, t))
    assert inspect_objects(data, t)["status"] == "passed"
    prs = Presentation(io.BytesIO(data))
    chart = prs.slides[0].shapes[1].chart
    replacement = CategoryChartData()
    replacement.categories = ["A", "B"]
    replacement.add_series("Revenue", [50, 70])
    chart.replace_data(replacement)
    prs.slides[0].shapes[2].table.cell(1, 1).text = "50"
    out = io.BytesIO()
    prs.save(out)
    edited = Presentation(io.BytesIO(out.getvalue()))
    assert list(edited.slides[0].shapes[1].chart.series[0].values) == [50, 70]
    assert edited.slides[0].shapes[2].table.cell(1, 1).text == "50"
    with ZipFile(io.BytesIO(out.getvalue())) as z:
        assert any(n.endswith(".xlsx") for n in z.namelist())


def test_group_paths_and_explicit_z_order(task_root):
    with Store(task_root) as s:
        t = s.json("task.json")
        title = t["scene"]["pages"][0]["nodes"][0]
        title["box"] = {"x": 10, "y": 10, "width": 700, "height": 70}
        title["z"] = 4
        path = node("path", "curve", box={"x": 20, "y": 130, "width": 240, "height": 180},
                    commands=[{"op": "M", "points": [[0, 100]]}, {"op": "C", "points": [[80, 0], [160, 180], [240, 100]]}],
                    style={"fill": None, "stroke": "118877", "stroke_width": 3})
        group = node("group", "group", box={"x": 40, "y": 50, "width": 800, "height": 360}, children=[title, path])
        t["scene"]["pages"][0]["nodes"] = [group]
        data = render(t, read_assets(s, t))
    assert inspect_objects(data, t)["status"] == "passed"
    prs = Presentation(io.BytesIO(data))
    shape = prs.slides[0].shapes[0]
    assert (shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt) == (40, 50, 800, 360)
    assert [s.name for s in shape.shapes] == ["curve", "title-object"]
    with ZipFile(io.BytesIO(data)) as z:
        assert b"cubicBezTo" in z.read("ppt/slides/slide1.xml")


@pytest.mark.parametrize("mutation,match", [
    (lambda t: t.update(script="print('oops')"), "SCHEMA_INVALID"),
    (lambda t: t["scene"]["pages"][0]["nodes"][0].update(expression="eval(x)"), "SCHEMA_INVALID"),
    (lambda t: t["scene"]["pages"][0]["nodes"][0]["box"].update(width=float("inf")), "INVALID_TASK"),
    (lambda t: t["scene"]["pages"][0]["nodes"][0]["box"].update(x=900), "NODE_OUT_OF_BOUNDS"),
    (lambda t: t["content"]["items"][0]["provenance"].update(checked=False), "CONTENT_NOT_CHECKED"),
    (lambda t: t["scene"]["pages"][0]["nodes"][0].update(page_id="page-2"), "NODE_ID_OR_PAGE"),
    (lambda t: t["content"]["items"].append(content("text", "missing", text="Not shown")), "CONTENT_COVERAGE"),
    (lambda t: t["canvas"].update(width=720), "ASPECT_RATIO"),
    (lambda t: t["design"]["pages"][0].update(target_kind="generated"), "ACTUAL_IMAGE_TOOL"),
    (lambda t: t["design"]["references"].clear(), "ORIGINAL_REFERENCE"),
])
def test_invalid_contract_rejected(task_root, mutation, match):
    t = load(task_root)
    mutation(t)
    with pytest.raises(DesignError, match=match):
        validate(t, complete=True)


def test_style_source_is_not_a_pixel_target(task_root):
    t = load(task_root)
    t["mode"] = "style_transfer"
    t["design"]["references"][0]["roles"] = ["style"]
    t["design"]["pages"][0]["target_kind"] = "provided_design"
    validate(t, complete=True)
    t["design"]["references"][0]["roles"] = ["content"]
    with pytest.raises(DesignError, match="STYLE_REFERENCE"):
        validate(t, complete=True)


def test_cross_page_asset_binding_rejected(task_root):
    t = load(task_root)
    t["assets"]["items"][0]["page_ids"] = ["page-2"]
    with pytest.raises(DesignError, match="ASSET_PAGE"):
        validate(t)


def test_whole_reference_cannot_masquerade_as_native_slide(task_root):
    t = load(task_root)
    t["scene"]["pages"][0]["nodes"].append(node("image", "screenshot", asset_id="reference-a", fit="contain"))
    with pytest.raises(DesignError, match="WHOLE_TARGET_IMAGE"):
        validate(t, complete=True)


def test_image_request_excludes_notes_ids_and_source_paths(task_root):
    t = load(task_root)
    t["content"]["notes"] = [{"page_id": "page-1", "text": "PRIVATE_MARKER_CREDENTIAL"}]
    t["content"]["items"][0]["provenance"]["source"] = "PRIVATE_SOURCE_PATH"
    request = json.dumps(image_request(t))
    assert "PRIVATE_" not in request
    assert "title-object" not in request
    assert "meaningful title" in request
    assert image_request(t)["generated"] is False


def test_cropping_embeds_real_asset_region(task_root):
    source = Image.new("RGB", (200, 100), "red")
    source.paste("blue", (100, 0, 200, 100))
    b = io.BytesIO()
    source.save(b, "PNG")
    data, size = normalized_png(b.getvalue(), [100, 0, 100, 100])
    assert size == (100, 100)
    assert Image.open(io.BytesIO(data)).getpixel((50, 50))[:3] == (0, 0, 255)
    with pytest.raises(DesignError, match="INVALID_SOURCE_CROP"):
        normalized_png(b.getvalue(), [199, 0, 100, 100])


def test_payloads_cannot_be_svg_or_scripts():
    for data in [b"<svg><script>alert(1)</script></svg>", b"#!/bin/sh\necho hi"]:
        with pytest.raises(DesignError, match="INVALID_IMAGE"):
            normalized_png(data)
    with pytest.raises(DesignError, match="DUPLICATE_JSON_KEY"):
        parse_json(b'{"x": 1, "x": 2}')
    with pytest.raises(DesignError, match="NONFINITE"):
        parse_json(b'{"x": NaN}')


def test_symlink_and_traversal_do_not_read_or_overwrite_outside(task_root, tmp_path):
    outside = tmp_path / "outside"
    outside.write_text("preserve me")
    (task_root / "linked").symlink_to(outside)
    (task_root / "linked-dir").symlink_to(tmp_path, target_is_directory=True)
    with Store(task_root) as s:
        for action in [lambda: s.read("linked"), lambda: s.write("linked", b"oops"),
                       lambda: s.write("linked-dir/outside", b"oops"), lambda: s.read("../outside")]:
            with pytest.raises(DesignError):
                action()
    assert outside.read_text() == "preserve me"
    with pytest.raises(DesignError):
        with Store(task_root / "linked-dir"):
            pass


def test_output_package_external_relationship_rejected(task_root):
    with Store(task_root) as s:
        t = s.json("task.json")
        raw = render(t, read_assets(s, t))
    with ZipFile(io.BytesIO(raw)) as src:
        out = io.BytesIO()
        with ZipFile(out, "w", ZIP_DEFLATED) as dst:
            for n in src.namelist():
                data = src.read(n)
                if n == "ppt/slides/_rels/slide1.xml.rels":
                    data = data.replace(b"</Relationships>", b'<Relationship Id="evil" Type="hyperlink" Target="https://example.com/secret" TargetMode="External"/></Relationships>')
                dst.writestr(n, data)
    with pytest.raises(DesignError, match="EXTERNAL_RELATIONSHIP"):
        inspect_package(out.getvalue())


def fake_worker(stage, **kwargs):
    # Unit-only deterministic stub; never used as real-render acceptance evidence.
    with Store(stage) as s:
        t = s.json("task.json")
        assets = {a["id"]: s.read("assets/" + a["sha256"] + ".png") for a in t["assets"]["items"]}
        data = render(t, assets)
        s.write("deck.pptx", data)
        s.put_json("objects.json", inspect_objects(data, t))
        image = png()
        s.write("preview-1.png", image)
        s.write("deck.pdf", b"UNIT TEST STUB; NOT A RENDER")
        s.put_json("render.json", {"status": "passed", "engine": "unit-stub", "pages": [{"page_id": "page-1", "file": "preview-1.png", "sha256": digest(image)}]})


@pytest.fixture
def candidate(task_root, monkeypatch):
    monkeypatch.setattr("engine.design_scene.pipeline.launch_worker", fake_worker)
    monkeypatch.setattr("engine.design_scene.pipeline.fonts_for", lambda t: [])
    report = build(task_root, isolation="host")
    assert report["status"] == "candidate_unreviewed", report
    return task_root, report


def complete_review(s, bid):
    r = review_skeleton(s, bid)
    r["reviewer"]["id"] = "independent-test-reviewer"
    for p in r["pages"]:
        for k in ("design", "reconstruction", "content", "editability", "fonts"):
            p[k] = "pass"
        p["observations"] = ["UNIT TEST: title-object retains both specified paragraphs"]
        p["observed_node_ids"] = ["title-object"]
    r["target_software"].update(version="UNIT TEST", status="pass", actions=["unit fixture text edit and save"])
    return r


@pytest.mark.parametrize("part", ["content", "design", "scene", "assets"])
def test_every_input_version_invalidates_old_review(candidate, part):
    root, m = candidate
    with Store(root) as s:
        t = s.json("task.json")
        if part == "content": t[part]["items"][0]["text"] += "changed"
        if part == "design": t[part]["brief"] = "new task-specific visual language"
        if part == "scene": t[part]["pages"][0]["background"] = "112233"
        if part == "assets": t[part]["items"][0]["source"] = "new source"
        s.put_json("task.json", t)
        with pytest.raises(DesignError, match="REVIEW_STALE_INPUTS"):
            verify_build(s, m["build_id"])


@pytest.mark.parametrize("artifact", ["deck.pptx", "preview-1.png", "objects.json", "render.json"])
def test_same_name_different_bytes_invalidates_review(candidate, artifact):
    root, m = candidate
    with Store(root) as s:
        prefix = build_prefix(m["build_id"])
        s.write(prefix + artifact, s.read(prefix + artifact) + b"tamper")
        with pytest.raises(DesignError, match="REVIEW_STALE_ARTIFACT"):
            verify_build(s, m["build_id"])


def test_no_self_approval_and_no_missing_target_software(candidate):
    root, m = candidate
    with Store(root) as s:
        r = complete_review(s, m["build_id"])
        r["reviewer"]["id"] = "author"
        with pytest.raises(DesignError, match="INDEPENDENT_REVIEW"):
            assess(s, m["build_id"], r)
        r["reviewer"]["id"] = "reviewer"
        r["target_software"]["status"] = "not_tested"
        assert "TARGET_SOFTWARE_EDIT_TEST_REQUIRED" in assess(s, m["build_id"], r)["blockers"]


def test_delivery_pairs_exact_task_target_deck_and_preview(candidate):
    root, m = candidate
    with Store(root) as s:
        r = complete_review(s, m["build_id"])
        result = deliver(s, m["build_id"], r)
        assert result["status"] == "final_delivery_ready"
        with ZipFile(io.BytesIO(s.read(result["delivery"]))) as z:
            assert {"task.json", "deck.pptx", "preview-1.png", "manifest.json", "review.json", "acceptance.json"} <= set(z.namelist())
            assert any(n.startswith("originals/") for n in z.namelist())
            assert digest(z.read("deck.pptx")) == m["artifacts"]["deck.pptx"]
        with pytest.raises(DesignError, match="OUTPUT_EXISTS"):
            deliver(s, m["build_id"], r)


def test_revision_exhaustion_does_not_approve(candidate):
    root, m = candidate
    with Store(root) as s:
        prefix = build_prefix(m["build_id"])
        m["revision"] = 3
        s.put_json(prefix + "manifest.json", m)
    with pytest.raises(DesignError, match="REVISION_BUDGET_EXHAUSTED"):
        build(root, parent_build=m["build_id"])


def test_missing_tool_keeps_failed_state(task_root, monkeypatch):
    def fail(*a, **k):
        raise DesignError("OS_SANDBOX_UNAVAILABLE")
    monkeypatch.setattr("engine.design_scene.pipeline.launch_worker", fail)
    m = build(task_root)
    assert m["status"] == "build_failed"
    assert m["artifacts"] == {}
    assert m["isolation"]["verified"] is False


def test_unknown_data_requires_explicit_image_acceptance(task_root):
    with Store(task_root) as s:
        t = s.json("task.json")
        ingest(s, t, asset_id="unreadable-plot", data=png("navy", (100, 100)), page_ids=["page-1"],
               role="raster_data", source="a supplied chart whose data cannot be read")
        t = s.json("task.json")
        item = content("unknown_data", "unknown-chart", description="underlying numeric values unknown")
        item["provenance"].update(kind="unknown", checked=False)
        t["content"]["items"].append(item)
        t["scene"]["pages"][0]["nodes"].append(node("image", "unreadable-image", asset_id="unreadable-plot", fit="contain", content_id="unknown-chart",
                                                         box={"x": 30, "y": 180, "width": 200, "height": 200}))
        with pytest.raises(DesignError, match="NATIVE_CONTENT_REPLACED"):
            validate(t, complete=True)
        item["required_edit"] = "raster_accepted"
        t["design"]["pages"][0]["raster_acceptances"] = [{"content_id": item["id"], "reason": "numeric data unavailable; replaceable image only", "accepted_by": "user"}]
        validate(t, complete=True)
        data = render(t, read_assets(s, t))
        assert inspect_objects(data, t)["pages"][0]["objects"]["replaceable_image"] == 1
        assert not any(x.has_chart for x in Presentation(io.BytesIO(data)).slides[0].shapes)


@pytest.mark.parametrize("ratio,bg,text_color", [((720, 540), "102437", "FFFFFF"), ((540, 720), "FFFFFF", "005577"), ((960, 540), "FFF4E0", "552244")])
def test_new_task_geometry_and_palette_do_not_leak(task_root, ratio, bg, text_color):
    t = load(task_root)
    t["canvas"].update(width=ratio[0], height=ratio[1])
    t["design"]["pages"][0]["aspect_policy"] = "contain"
    t["scene"]["pages"][0]["background"] = bg
    title = t["scene"]["pages"][0]["nodes"][0]
    title["box"].update(width=ratio[0]-60)
    title["font"]["color"] = text_color
    with Store(task_root) as s:
        raw = render(t, read_assets(s, t))
    p = Presentation(io.BytesIO(raw))
    assert (p.slide_width.pt, p.slide_height.pt) == ratio
    assert str(p.slides[0].background.fill.fore_color.rgb) == bg
    assert len(p.slides[0].shapes) == 1
    assert not p.slides[0].shapes[0].has_chart


def test_preplanted_temporary_symlink_is_never_followed(task_root, tmp_path, monkeypatch):
    class FixedId:
        hex = "unit-test-fixed"
    outside = tmp_path / "sentinel"
    outside.write_bytes(b"unchanged")
    (task_root / ".write-unit-test-fixed").symlink_to(outside)
    monkeypatch.setattr("engine.design_scene.store.uuid.uuid4", lambda: FixedId())
    with Store(task_root) as s, pytest.raises(DesignError):
        s.write("new.json", b"payload")
    assert outside.read_bytes() == b"unchanged"


def test_image_frame_and_pixel_limits():
    data = io.BytesIO()
    frames = [Image.new("RGB", (10, 10), c) for c in ("red", "blue")]
    frames[0].save(data, "PNG", save_all=True, append_images=frames[1:])
    with pytest.raises(DesignError, match="IMAGE_RESOURCE_LIMIT"):
        normalized_png(data.getvalue())
    with pytest.raises(DesignError, match="IMAGE_RESOURCE_LIMIT"):
        normalized_png(png(size=(8001, 4000)))


def test_same_asset_filename_changed_is_detected(task_root):
    with Store(task_root) as s:
        t = s.json("task.json")
        a = t["assets"]["items"][0]
        s.write("assets/" + a["sha256"] + ".png", png("red"))
        with pytest.raises(DesignError, match="ASSET_CONTENT_CHANGED"):
            read_assets(s, t)


def test_final_delivery_cannot_invent_accepted_differences(candidate):
    root, m = candidate
    with Store(root) as s:
        r = complete_review(s, m["build_id"])
        r["pages"][0]["accepted_differences"] = ["silently redesigned layout"]
        assert "DIFFERENCES_NOT_APPROVED_IN_DESIGN:page-1" in assess(s, m["build_id"], r)["blockers"]


@pytest.mark.parametrize("system", ["Linux", "Windows", "Darwin"])
def test_auto_host_can_deliver_without_claiming_os_isolation(task_root, monkeypatch, system):
    monkeypatch.setattr("engine.design_scene.runtime.platform.system", lambda: system)
    monkeypatch.setattr("engine.design_scene.runtime.shutil.which", lambda name: None)
    monkeypatch.setattr("engine.design_scene.pipeline.launch_worker", fake_worker)
    monkeypatch.setattr("engine.design_scene.pipeline.fonts_for", lambda task: [])
    m = build(task_root)  # The API and CLI now share the portable auto default.
    assert m["status"] == "candidate_unreviewed"
    assert m["isolation"] == {"requested": "auto", "mode": "host", "verified": False}
    with Store(task_root) as s:
        result = deliver(s, m["build_id"], complete_review(s, m["build_id"]))
        assert result["status"] == "final_delivery_ready"
        assert result["blockers"] == []
        assert result["warnings"] == ["OS_ISOLATION_NOT_VERIFIED"]
        with ZipFile(io.BytesIO(s.read(result["delivery"]))) as z:
            acceptance = json.loads(z.read("acceptance.json"))
            assert acceptance["execution"]["os_isolation_verified"] is False
            assert json.loads(z.read("manifest.json"))["isolation"]["verified"] is False
            assert "未验证" in z.read("交付说明.md").decode()
            assert Presentation(io.BytesIO(z.read("deck.pptx"))).slides[0].shapes[0].has_text_frame


def test_auto_keeps_available_macos_isolation_and_explicit_choice_fails_closed(monkeypatch):
    from engine.design_scene.runtime import select_isolation
    monkeypatch.setattr("engine.design_scene.runtime.platform.system", lambda: "Darwin")
    monkeypatch.setattr("engine.design_scene.runtime.shutil.which", lambda name: "/usr/bin/sandbox-exec")
    assert select_isolation() == "macos"
    assert select_isolation("host") == "host"
    monkeypatch.setattr("engine.design_scene.runtime.platform.system", lambda: "Linux")
    with pytest.raises(DesignError, match="OS_SANDBOX_UNAVAILABLE"):
        select_isolation("macos")
    with pytest.raises(DesignError, match="UNKNOWN_ISOLATION_MODE"):
        select_isolation("unverified-backend")


def test_strict_build_never_starts_host_worker(task_root, monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Strict builds must not execute an unisolated worker")
    monkeypatch.setattr("engine.design_scene.pipeline.launch_worker", forbidden)
    m = build(task_root, isolation="host", require_os_isolation=True)
    assert m["status"] == "build_failed"
    assert "OS_ISOLATION_REQUIRED" in m["error"]
    assert m["artifacts"] == {}


def test_strict_delivery_blocks_host_without_creating_archive(candidate):
    root, m = candidate
    with Store(root) as s:
        review = complete_review(s, m["build_id"])
        result = deliver(s, m["build_id"], review, require_os_isolation=True)
        assert result["status"] == "revision_required"
        assert result["blockers"] == ["OS_ISOLATION_NOT_VERIFIED"]
    assert not (root / "deliveries").exists()


@pytest.mark.parametrize("policy", [None, {"require_os_isolation": True}])
def test_manifest_strict_policy_cannot_be_dropped_by_omitting_cli_flag(candidate, policy):
    root, m = candidate
    with Store(root) as s:
        if policy is None:
            m.pop("delivery_policy")  # Legacy manifests retain their strict policy.
        else:
            m["delivery_policy"] = policy
        s.put_json(build_prefix(m["build_id"]) + "manifest.json", m)
        result = deliver(s, m["build_id"], complete_review(s, m["build_id"]))
        assert result["blockers"] == ["OS_ISOLATION_NOT_VERIFIED"]


@pytest.mark.parametrize("check", ["design", "reconstruction", "content", "editability", "fonts"])
def test_host_still_requires_every_visual_and_content_review(candidate, check):
    root, m = candidate
    with Store(root) as s:
        review = complete_review(s, m["build_id"])
        review["pages"][0][check] = "revise"
        result = deliver(s, m["build_id"], review)
        assert check.upper() + "_REVIEW_REQUIRED:page-1" in result["blockers"]
        assert "delivery" not in result


@pytest.mark.parametrize("missing,blocker", [("render", "ACTUAL_RENDER_REQUIRED"),
                                           ("font", "FONT_ENVIRONMENT_UNRESOLVED"),
                                           ("draft", "DRAFT_NOT_FINAL")])
def test_host_still_blocks_incomplete_artifacts(candidate, monkeypatch, missing, blocker):
    root, m = candidate
    if missing == "render":
        m["render"]["status"] = "not_requested"
    elif missing == "draft":
        m["status"] = "draft"
    else:
        m["fonts"] = [{"requested": "Missing Font", "status": "unverified"}]
        monkeypatch.setattr("engine.design_scene.pipeline.fonts_for", lambda task: m["fonts"])
    with Store(root) as s:
        s.put_json(build_prefix(m["build_id"]) + "manifest.json", m)
        result = deliver(s, m["build_id"], complete_review(s, m["build_id"]))
        assert blocker in result["blockers"]
        assert "delivery" not in result


def test_policy_change_invalidates_existing_review(candidate):
    root, m = candidate
    with Store(root) as s:
        review = complete_review(s, m["build_id"])
        m["delivery_policy"]["require_os_isolation"] = True
        s.put_json(build_prefix(m["build_id"]) + "manifest.json", m)
        with pytest.raises(DesignError, match="REVIEW_STALE_MANIFEST"):
            deliver(s, m["build_id"], review)


def test_cli_deliver_honors_strict_option_and_normal_host_delivery(candidate, capsys):
    from engine.design_scene.cli import main
    root, m = candidate
    with Store(root) as s:
        s.put_json("review.json", complete_review(s, m["build_id"]))
    args = ["deliver", "--task-dir", str(root), "--build-id", m["build_id"], "--review-file", "review.json"]
    assert main([*args, "--require-os-isolation"]) == 1
    result = json.loads(capsys.readouterr().out)
    assert result["blockers"] == ["OS_ISOLATION_NOT_VERIFIED"]
    assert not (root / "deliveries").exists()
    assert main(args) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["execution"]["os_isolation_verified"] is False
    assert (root / result["delivery"]).is_file()


def test_verified_macos_build_still_satisfies_strict_delivery(task_root, monkeypatch):
    monkeypatch.setattr("engine.design_scene.runtime.platform.system", lambda: "Darwin")
    monkeypatch.setattr("engine.design_scene.runtime.shutil.which", lambda name: "/usr/bin/sandbox-exec" if name == "sandbox-exec" else None)
    monkeypatch.setattr("engine.design_scene.pipeline.launch_worker", fake_worker)
    monkeypatch.setattr("engine.design_scene.pipeline.fonts_for", lambda task: [])
    m = build(task_root, require_os_isolation=True)
    assert m["isolation"]["mode"] == "macos"
    with Store(task_root) as s:
        result = deliver(s, m["build_id"], complete_review(s, m["build_id"]))
        assert result["status"] == "final_delivery_ready"
        assert result["delivery_policy"]["require_os_isolation"] is True
        assert result["warnings"] == []


def test_actual_host_worker_produces_editable_pptx_without_posix_assumptions(task_root):
    # Real subprocess and real native PPTX generation; no simulated render or review.
    # This test runs on the native Windows/Linux/macOS CI matrix.
    m = build(task_root, isolation="host", preview=False)
    assert m["status"] == "render_incomplete", m
    assert m["isolation"]["verified"] is False
    with Store(task_root) as s:
        raw = s.read(build_prefix(m["build_id"]) + "deck.pptx")
        prs = Presentation(io.BytesIO(raw))
        prs.slides[0].shapes[0].text = "跨平台可编辑内容"
        edited = io.BytesIO()
        prs.save(edited)
        assert Presentation(io.BytesIO(edited.getvalue())).slides[0].shapes[0].text == "跨平台可编辑内容"
        result = deliver(s, m["build_id"], complete_review(s, m["build_id"]))
        assert "ACTUAL_RENDER_REQUIRED" in result["blockers"]


def test_bounded_process_timeout_is_actionable_on_each_os(tmp_path):
    from engine.design_scene.runtime import run_process
    with pytest.raises(DesignError, match="PROCESS_TIMEOUT"):
        run_process([sys.executable, "-c", "import time; time.sleep(30)"], cwd=tmp_path, timeout=.1)


def test_clean_worker_environment_keeps_required_tools_but_not_credentials(tmp_path, monkeypatch):
    from engine.design_scene.runtime import clean_environment, find_tool
    monkeypatch.setenv("OPENAI_API_KEY", "unit-secret")
    monkeypatch.setenv("PYTHONSTARTUP", "/untrusted/hook.py")
    monkeypatch.setenv("HTTPS_PROXY", "http://secret.invalid")
    env = clean_environment(tmp_path)
    assert not {"OPENAI_API_KEY", "PYTHONSTARTUP", "HTTPS_PROXY"} & env.keys()
    assert env["TEMP"] == str(tmp_path / "tmp")
    for tool in ("soffice", "pdftoppm", "fc-match"):
        if path := find_tool(tool):
            assert str(Path(path).parent) in env["PATH"].split(os.pathsep)
    if os.name == "nt":
        assert env["SystemRoot"] == os.environ["SystemRoot"]


@pytest.mark.skipif(os.name != "nt", reason="requires actual Windows handle and junction semantics")
def test_windows_junction_cannot_escape_task_directory(task_root, tmp_path):
    outside = tmp_path / "outside-dir"
    outside.mkdir()
    (outside / "sentinel").write_bytes(b"unchanged")
    junction = task_root / "junction"
    subprocess.run(["cmd", "/c", "mklink", "/J", str(junction), str(outside)], check=True, capture_output=True)
    with Store(task_root) as s:
        for action in (lambda: s.read("junction/sentinel"), lambda: s.write("junction/sentinel", b"bad")):
            with pytest.raises(DesignError):
                action()
        for name in ("CON", "NUL.txt", "trailing. ", "stream:payload"):
            with pytest.raises(DesignError):
                s.write(name, b"bad")
    assert (outside / "sentinel").read_bytes() == b"unchanged"


@pytest.mark.skipif(os.name != "nt", reason="requires actual Windows delete sharing semantics")
def test_windows_task_ancestors_are_locked_during_io(task_root):
    with Store(task_root) as s:
        with pytest.raises(OSError):
            task_root.rename(task_root.with_name("moved"))
        s.write("nested/中文.json", b"{}", exclusive=True)
        assert s.json("nested/中文.json") == {}
    task_root.rename(task_root.with_name("moved"))  # Handles really close.
