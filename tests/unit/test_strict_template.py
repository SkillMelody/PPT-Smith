from __future__ import annotations

import hashlib
import json
from pathlib import Path
import subprocess
import sys
import shutil
import zipfile

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from engine.component_atlas import build_component_atlas
from engine.component_composer import build_component_plan
from engine.strict_template import (
    _asset_preservation,
    _inspection_delta,
    execute_strict_template as _execute_strict_template,
)
from engine.structural_parser import parse_source


ROOT = Path(__file__).resolve().parents[2]


def _sha(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _reference_template(path: Path) -> None:
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = ChartData()
    data.categories = ["Current", "Target"]
    data.add_series("Adoption", (34, 72))
    chart = source.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(0.8), Inches(5), Inches(3), data,
    )
    chart.name = "adoption-template"
    table = source.shapes.add_table(2, 2, Inches(6.2), Inches(0.8), Inches(3), Inches(2))
    table.name = "regional-template"
    for row_index, values in enumerate((("Region", "Growth"), ("US", "20%"))):
        for column_index, value in enumerate(values):
            table.table.cell(row_index, column_index).text = value
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def _request(template: Path) -> dict:
    return {
        "schema_version": "1.0.0",
        "request_id": "strict-template-test",
        "route": "template",
        "page_contract": {
            "mode": "exact", "exact_pages": 1,
            "overflow_policy": "ask", "underflow_policy": "fewer_pages",
        },
        "template": {
            "use_mode": "strict", "source_sha256": _sha(template), "required": True,
        },
    }


def _ir() -> dict:
    return {
        "slides": [{
            "id": "growth",
            "chart": {
                "data": {
                    "categories": ["Current", "Target"],
                    "series": [{"name": "Adoption", "values": [34, 72]}],
                },
            },
            "blocks": [{
                "id": "regional", "role": "table",
                "source_ref": {"source_id": "doc", "loc": "table_1"},
            }],
        }],
    }


def _plan() -> dict:
    return {
        "operations": [
            {
                "kind": "chart_clone", "source_slide_index": 1, "destination_slide_index": 2,
                "chart_name": "adoption-template", "binding_name": "bind:chart:growth",
            },
            {
                "kind": "table_clone", "source_slide_index": 1, "destination_slide_index": 2,
                "table_name": "regional-template", "binding_name": "bind:table:growth:regional",
                "rows": [["Region", "Growth"], ["EU", "15%"]],
            },
        ],
    }


def _content_integrity_inputs() -> tuple[dict, dict]:
    evidence_ledger = {
        "schema_version": "1.0.0",
        "sources": [{"source_id": "doc", "sha256": "sha256:test"}],
        "evidence_units": [{
            "evidence_id": "doc:para_1",
            "source_ref": {"source_id": "doc", "loc": "para_1"},
            "kind": "claim",
            "text": "The delivery assertion is source-backed.",
            "entities": [],
            "numbers": [],
            "relations": [],
            "salience": 1.0,
            "must_keep": False,
            "exhibit_id": None,
        }],
    }
    content_bindings = {"slides": [{
        "id": "growth",
        "archetype": "body",
        "assertion": "The delivery assertion is source-backed.",
        "evidence_ids": ["doc:para_1"],
        "required_evidence_ids": [],
        "omissions": [],
    }]}
    return evidence_ledger, content_bindings


def execute_strict_template(**kwargs) -> dict:
    evidence_ledger, content_bindings = _content_integrity_inputs()
    kwargs.setdefault("evidence_ledger", evidence_ledger)
    kwargs.setdefault("content_bindings", content_bindings)
    return _execute_strict_template(**kwargs)


def _asset_zip(
    path: Path,
    *,
    media: dict[str, bytes],
    relationships: dict[str, str] | None = None,
) -> Path:
    relationship_xml = "".join(
        f'<Relationship Id="{relationship_id}" '
        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        f'Target="{target}"/>'
        for relationship_id, target in (relationships or {}).items()
    )
    with zipfile.ZipFile(path, "w") as archive:
        for name, payload in media.items():
            archive.writestr(f"ppt/media/{name}", payload)
        if relationships is not None:
            archive.writestr(
                "ppt/slides/_rels/slide1.xml.rels",
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                f"{relationship_xml}</Relationships>",
            )
    return path


def test_asset_preservation_ignores_unreachable_source_media(tmp_path: Path) -> None:
    source = _asset_zip(
        tmp_path / "source.pptx",
        media={"image1.png": b"used", "image2.png": b"unused"},
    )
    output = _asset_zip(
        tmp_path / "output.pptx",
        media={"image1.png": b"used"},
    )

    assert _asset_preservation(source, output) == {
        "status": "pass",
        "verified_parts": 1,
        "changed_parts": [],
        "broken_relationships": [],
    }


def test_asset_preservation_rejects_changed_output_media(tmp_path: Path) -> None:
    source = _asset_zip(
        tmp_path / "source.pptx", media={"image1.png": b"original"},
    )
    output = _asset_zip(
        tmp_path / "output.pptx", media={"image1.png": b"changed"},
    )

    report = _asset_preservation(source, output)

    assert report["status"] == "fail"
    assert report["changed_parts"] == ["ppt/media/image1.png"]


def test_asset_preservation_rejects_broken_output_asset_relationship(tmp_path: Path) -> None:
    source = _asset_zip(
        tmp_path / "source.pptx", media={"image1.png": b"original"},
    )
    output = _asset_zip(
        tmp_path / "output.pptx",
        media={},
        relationships={"rId1": "../media/missing.png"},
    )

    report = _asset_preservation(source, output)

    assert report["status"] == "fail"
    assert report["broken_relationships"] == [{
        "relationship_part": "ppt/slides/_rels/slide1.xml.rels",
        "relationship_id": "rId1",
        "target": "ppt/media/missing.png",
    }]


def test_strict_delivery_rejects_missing_content_integrity_inputs(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _reference_template(template)

    result = _execute_strict_template(
        request=_request(template),
        template_pptx=template,
        strict_plan=_plan(),
        ir=_ir(),
        source_docs={},
        output_pptx=tmp_path / "candidate.pptx",
    )

    assert result == {"ok": False, "code": "CONTENT_INTEGRITY_REQUIRED"}


def test_strict_delivery_rejects_failed_content_integrity_before_writing(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "candidate.pptx"
    _reference_template(template)
    evidence_ledger, _ = _content_integrity_inputs()

    result = _execute_strict_template(
        request=_request(template),
        template_pptx=template,
        strict_plan=_plan(),
        ir=_ir(),
        source_docs={},
        output_pptx=output,
        evidence_ledger=evidence_ledger,
        content_bindings={"slides": [{
            "id": "growth",
            "archetype": "body",
            "assertion": "Unsupported assertion",
            "evidence_ids": [],
        }]},
    )

    assert result["ok"] is False
    assert result["code"] == "CONTENT_INTEGRITY_FAILED"
    assert result["content_integrity"]["status"] == "fail"
    assert not output.exists()


def test_strict_template_preserves_template_assets_and_only_applies_bound_operations(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    _reference_template(template)
    source_docs = {
        "doc": parse_source(
            "| Region | Growth |\n| --- | --- |\n| EU | 15% |", "markdown", "doc",
        ),
    }

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=_plan(),
        ir=_ir(), source_docs=source_docs, output_pptx=output, render_engine="libreoffice",
    )

    assert result["ok"] is True, result
    assert result["template"]["use_mode"] == "strict"
    assert result["binding"]["status"] == "pass"
    assert result["inspection"]["status"] == "passed"
    assert result["render"]["status"] == "passed"
    assert output.is_file()
    assert result["delivery"] == {
        "source_slide_count": 2,
        "output_slide_count": 1,
        "slide_numbers": [2],
    }
    with zipfile.ZipFile(template) as source, zipfile.ZipFile(output) as produced:
        assert source.read("ppt/slideMasters/slideMaster1.xml") == produced.read("ppt/slideMasters/slideMaster1.xml")
    assert len(Presentation(output).slides) == 1
    copied = Presentation(output).slides[0].shapes
    assert {shape.name for shape in copied} == {"bind:chart:growth", "bind:table:growth:regional"}
    assert (tmp_path / "qa" / "candidate" / "slides").is_dir()
    assert not (tmp_path / "qa" / "template-reference").exists()


def test_strict_template_rejects_a_template_that_does_not_match_its_declared_hash(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    _reference_template(template)
    request = _request(template)
    request["template"]["source_sha256"] = "sha256:wrong"

    result = execute_strict_template(
        request=request, template_pptx=template, strict_plan=_plan(),
        ir=_ir(), source_docs={}, output_pptx=output,
    )

    assert result["ok"] is False
    assert result["code"] == "TEMPLATE_SOURCE_HASH_MISMATCH"
    assert not output.exists()


def test_strict_template_rejects_a_bespoke_route_request(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    _reference_template(template)
    request = _request(template)
    request["route"] = "bespoke"

    result = execute_strict_template(
        request=request, template_pptx=template, strict_plan=_plan(),
        ir=_ir(), source_docs={}, output_pptx=output,
    )

    assert result == {"ok": False, "code": "TEMPLATE_ROUTE_REQUIRED"}
    assert not output.exists()


def test_strict_template_allows_inherited_template_findings_but_rejects_new_ones() -> None:
    baseline = {
        "status": "failed",
        "issues": [{"issue_code": "PPTX_CONNECTOR_CROSSING", "severity": "error", "slide_id": "S01", "object_id": "a", "evidence": {}}],
    }
    unchanged = _inspection_delta(baseline, baseline)
    changed = _inspection_delta(baseline, {
        "status": "failed",
        "issues": baseline["issues"] + [{"issue_code": "PPTX_BLANK_SLIDE", "severity": "error", "slide_id": "S02", "object_id": None, "evidence": {}}],
    })

    assert unchanged["status"] == "passed"
    assert unchanged["new_issue_count"] == 0
    assert changed["status"] == "failed"
    assert changed["new_issue_count"] == 1


def test_strict_template_can_bind_data_into_an_existing_template_chart(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    shutil.rmtree(tmp_path / "strict-output-render", ignore_errors=True)
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = ChartData()
    data.categories = ["Current", "Target"]
    data.add_series("Adoption", (34, 72))
    chart = source.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(0.8), Inches(5), Inches(3), data,
    )
    chart.name = "adoption-template"
    title = source.shapes.add_textbox(Inches(0.8), Inches(4), Inches(5), Inches(0.5))
    title.name = "template-title"
    title.text = "Old title"
    presentation.save(template)
    plan = {
        "operations": [{
            "kind": "chart_bind_existing", "source_slide_index": 1, "destination_slide_index": 1,
            "chart_name": "adoption-template", "binding_name": "bind:chart:growth",
            "data": {"categories": ["Current", "Target"], "series": [{"name": "Adoption", "values": [34, 72]}]},
        }, {
            "kind": "text_bind_existing", "source_slide_index": 1, "destination_slide_index": 1,
            "shape_name": "template-title", "binding_name": "bind:slide:growth:title", "text": "Growth",
        }],
    }

    ir = _ir()
    ir["slides"][0]["title"] = "Growth"
    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
    )

    assert result["ok"] is True, result
    assert result["binding"]["status"] == "pass"
    assert Presentation(output).slides[0].shapes[0].name == "bind:chart:growth"
    assert Presentation(output).slides[0].shapes[1].name == "bind:slide:growth:title"


def test_strict_template_binds_a_reviewed_multi_chart_dashboard_without_sample_leakage(tmp_path: Path) -> None:
    template = tmp_path / "dashboard-template.pptx"
    output = tmp_path / "dashboard-output.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, name in enumerate(("innovation-chart", "ebit-chart")):
        data = ChartData()
        data.categories = ["Sample yes", "Sample no"]
        data.add_series("Sample", (50, 50))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            Inches(0.8 + index * 4), Inches(1.2), Inches(3.2), Inches(3.2), data,
        )
        chart.name = name
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(7), Inches(0.5))
    title.name = "dashboard-title"
    title.text = "Sample dashboard"
    note = slide.shapes.add_textbox(Inches(8.4), Inches(1.2), Inches(3.5), Inches(1.2))
    note.name = "dashboard-note"
    note.text = "Sample conclusion must disappear"
    presentation.save(template)
    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "dashboard-review"}},
        "components": [{
            "component_id": "dashboard.research-impact",
            "family": "chart_dashboard",
            "slide_index": 1,
            "semantic_uses": ["research impact evidence"],
            "groups": [{
                "role": "chart", "scope": "item", "item_indices": [0, 1],
                "shape_names": ["innovation-chart", "ebit-chart"],
            }, {
                "role": "text", "scope": "shared",
                "shape_names": ["dashboard-title", "dashboard-note"],
            }],
            "parameters": {"element_count": {"minimum": 2, "maximum": 2}},
        }],
    }
    atlas = build_component_atlas(template, review)
    ir = {"slides": [{
        "id": "impact",
        "title": "AI impact is strongest in innovation",
        "charts": [{
            "id": "innovation",
            "data": {
                "categories": ["Reported impact", "No reported impact"],
                "series": [{"name": "Share", "values": [64, 36]}],
            },
        }, {
            "id": "ebit",
            "data": {
                "categories": ["Reported impact", "No reported impact"],
                "series": [{"name": "Share", "values": [39, 61]}],
            },
        }],
    }]}
    plan = {"operations": [{
        "kind": "chart_dashboard_bind_existing",
        "source_slide_index": 1,
        "destination_slide_index": 1,
        "component_requirement": {
            "semantic_use": "research impact evidence",
            "family": "chart_dashboard",
            "element_count": 2,
        },
        "charts": [{
            "id": "innovation",
            "binding_name": "bind:chart:impact:innovation",
            "data": ir["slides"][0]["charts"][0]["data"],
        }, {
            "id": "ebit",
            "binding_name": "bind:chart:impact:ebit",
            "data": ir["slides"][0]["charts"][1]["data"],
        }],
        "text_bindings": [{
            "shape_name": "dashboard-title",
            "binding_name": "bind:slide:impact:title",
            "text": "AI impact is strongest in innovation",
        }],
        "clear_text_names": ["dashboard-note"],
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, result
    assert result["binding"]["status"] == "pass"
    output_slide = Presentation(output).slides[0]
    assert {shape.name for shape in output_slide.shapes if getattr(shape, "has_chart", False)} == {
        "bind:chart:impact:innovation", "bind:chart:impact:ebit",
    }
    cleared = next(shape for shape in output_slide.shapes if shape.name.startswith("decoration:cleared-text:"))
    assert cleared.text == ""
    assert "Sample conclusion must disappear" not in output.read_bytes().decode("latin1", errors="ignore")


def test_strict_template_clones_a_reviewed_multi_chart_dashboard_to_a_new_slide(tmp_path: Path) -> None:
    template = tmp_path / "dashboard-clone-template.pptx"
    output = tmp_path / "dashboard-clone-output.pptx"
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, name in enumerate(("innovation-chart", "ebit-chart")):
        data = ChartData()
        data.categories = ["Sample yes", "Sample no"]
        data.add_series("Sample", (50, 50))
        chart = source.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            Inches(0.8 + index * 4), Inches(1.2), Inches(3.2), Inches(3.2), data,
        )
        chart.name = name
    title = source.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(7), Inches(0.5))
    title.name = "dashboard-title"
    title.text = "Sample dashboard"
    note = source.shapes.add_textbox(Inches(8.4), Inches(1.2), Inches(3.5), Inches(1.2))
    note.name = "dashboard-note"
    note.text = "Sample conclusion must disappear"
    rule = source.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.8), Inches(5.8), Inches(11.5), Inches(5.8),
    )
    rule.name = "dashboard-rule"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "dashboard-review"}},
        "components": [{
            "component_id": "dashboard.research-impact",
            "family": "chart_dashboard",
            "slide_index": 1,
            "semantic_uses": ["research impact evidence"],
            "groups": [{
                "role": "chart", "scope": "item", "item_indices": [0, 1],
                "shape_names": ["innovation-chart", "ebit-chart"],
            }, {
                "role": "text", "scope": "shared",
                "shape_names": ["dashboard-title", "dashboard-note"],
            }, {
                "role": "decoration", "scope": "shared",
                "shape_names": ["dashboard-rule"],
            }],
            "parameters": {"element_count": {"minimum": 2, "maximum": 2}},
        }],
    }
    atlas = build_component_atlas(template, review)
    charts = [{
        "id": chart_id,
        "binding_name": f"bind:chart:impact:{chart_id}",
        "data": {
            "categories": ["Reported impact", "No reported impact"],
            "series": [{"name": "Share", "values": values}],
        },
    } for chart_id, values in (("innovation", [64, 36]), ("ebit", [39, 61]))]
    ir = {"slides": [{
        "id": "impact",
        "title": "AI impact is strongest in innovation",
        "charts": [{"id": chart["id"], "data": chart["data"]} for chart in charts],
    }]}
    plan = {"operations": [{
        "kind": "chart_dashboard_clone",
        "source_slide_index": 1,
        "destination_slide_index": 2,
        "component_requirement": {
            "semantic_use": "research impact evidence",
            "family": "chart_dashboard",
            "element_count": 2,
        },
        "charts": charts,
        "text_bindings": [{
            "shape_name": "dashboard-title",
            "binding_name": "bind:slide:impact:title",
            "text": "AI impact is strongest in innovation",
        }],
        "clear_text_names": ["dashboard-note"],
        "placement": {"x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0},
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, result
    source_slide = Presentation(template).slides[0]
    output_slide = Presentation(output).slides[0]
    assert {shape.name for shape in source_slide.shapes if getattr(shape, "has_chart", False)} == {
        "innovation-chart", "ebit-chart",
    }
    assert {shape.name for shape in output_slide.shapes if getattr(shape, "has_chart", False)} == {
        "bind:chart:impact:innovation", "bind:chart:impact:ebit",
    }
    assert next(
        shape for shape in output_slide.shapes if shape.name == "bind:slide:impact:title"
    ).text == "AI impact is strongest in innovation"
    assert next(shape for shape in output_slide.shapes if shape.name.startswith("decoration:cleared-text:")).text == ""
    assert any(
        shape.name.startswith("decoration:component:chart_dashboard:dashboard.research-impact:shared:0")
        for shape in output_slide.shapes
    )


def test_strict_template_clones_nested_dashboard_groups_with_independent_chart_data(tmp_path: Path) -> None:
    template = tmp_path / "nested-dashboard-template.pptx"
    output = tmp_path / "nested-dashboard-output.pptx"
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, name in enumerate(("innovation-chart", "ebit-chart")):
        group = source.shapes.add_group_shape()
        group.name = f"metric-card-{index + 1}"
        data = ChartData()
        data.categories = ["Sample yes", "Sample no"]
        data.add_series("Sample", (50, 50))
        chart = group.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            Inches(0), Inches(0.45), Inches(3.0), Inches(2.4), data,
        )
        chart.name = name
        label = group.shapes.add_textbox(Inches(0.2), Inches(0), Inches(2.6), Inches(0.4))
        label.name = f"metric-label-{index + 1}"
        label.text = f"Sample metric {index + 1}"
        group.left = Inches(0.8 + index * 4.2)
        group.top = Inches(1.2)
        group.width = Inches(3.2)
        group.height = Inches(3.0)
    title = source.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(7), Inches(0.5))
    title.name = "dashboard-title"
    title.text = "Sample dashboard"
    note = source.shapes.add_textbox(Inches(8.4), Inches(1.2), Inches(3.5), Inches(1.2))
    note.name = "dashboard-note"
    note.text = "Sample conclusion must disappear"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "nested-dashboard-review"}},
        "components": [{
            "component_id": "dashboard.research-impact",
            "family": "chart_dashboard",
            "slide_index": 1,
            "semantic_uses": ["research impact evidence"],
            "groups": [{
                "role": "chart", "scope": "item", "item_indices": [0, 1],
                "shape_names": ["innovation-chart", "ebit-chart"],
            }, {
                "role": "text", "scope": "shared",
                "shape_names": [
                    "dashboard-title", "metric-label-1", "metric-label-2", "dashboard-note",
                ],
            }],
            "parameters": {"element_count": {"minimum": 2, "maximum": 2}},
        }],
    }
    atlas = build_component_atlas(template, review)
    charts = [{
        "id": chart_id,
        "binding_name": f"bind:chart:impact:{chart_id}",
        "data": {
            "categories": ["Reported impact", "No reported impact"],
            "series": [{"name": "Share", "values": values}],
        },
    } for chart_id, values in (("innovation", [64, 36]), ("ebit", [39, 61]))]
    ir = {"slides": [{
        "id": "impact",
        "title": "AI impact is strongest in innovation",
        "charts": [{"id": chart["id"], "data": chart["data"]} for chart in charts],
        "blocks": [{
            "id": "metrics", "role": "list",
            "items": [{"text": "Innovation 64%"}, {"text": "EBIT 39%"}],
        }],
    }]}
    plan = {"operations": [{
        "kind": "chart_dashboard_clone",
        "source_slide_index": 1,
        "destination_slide_index": 2,
        "component_requirement": {
            "semantic_use": "research impact evidence",
            "family": "chart_dashboard",
            "element_count": 2,
        },
        "charts": charts,
        "text_bindings": [{
            "shape_name": "dashboard-title",
            "binding_name": "bind:slide:impact:title",
            "text": "AI impact is strongest in innovation",
        }, {
            "shape_name": "metric-label-1",
            "binding_name": "bind:block:impact:metrics:item:0",
            "text": "Innovation 64%",
        }, {
            "shape_name": "metric-label-2",
            "binding_name": "bind:block:impact:metrics:item:1",
            "text": "EBIT 39%",
        }],
        "clear_text_names": ["dashboard-note"],
        "placement": {"x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0},
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, result

    def descendants(shapes):
        for shape in shapes:
            yield shape
            if hasattr(shape, "shapes"):
                yield from descendants(shape.shapes)

    source_slide = Presentation(template).slides[0]
    output_slide = Presentation(output).slides[0]
    source_shapes = list(descendants(source_slide.shapes))
    output_shapes = list(descendants(output_slide.shapes))
    source_charts = [shape for shape in source_shapes if getattr(shape, "has_chart", False)]
    assert [list(chart.chart.series[0].values) for chart in source_charts] == [[50.0, 50.0], [50.0, 50.0]]
    assert {shape.name for shape in output_shapes if getattr(shape, "has_chart", False)} == {
        "bind:chart:impact:innovation", "bind:chart:impact:ebit",
    }
    assert {shape.name for shape in output_shapes if getattr(shape, "has_text_frame", False)} >= {
        "bind:slide:impact:title",
        "bind:block:impact:metrics:item:0",
        "bind:block:impact:metrics:item:1",
    }
    assert sum(hasattr(shape, "shapes") for shape in output_slide.shapes) == 2


def test_strict_template_clones_a_nested_kpi_chart_card_as_one_editable_group(tmp_path: Path) -> None:
    template = tmp_path / "nested-kpi-card-template.pptx"
    output = tmp_path / "nested-kpi-card-output.pptx"
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    card = source.shapes.add_group_shape()
    card.name = "kpi-card"
    background = card.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0), Inches(0), Inches(3.2), Inches(3.0),
    )
    background.name = "kpi-card-background"
    data = ChartData()
    data.categories = ["Sample", "Other"]
    data.add_series("Share", (50, 50))
    chart = card.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.2), Inches(0.7), Inches(2.8), Inches(2.0), data,
    )
    chart.name = "kpi-card-chart"
    title = card.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(1.8), Inches(0.4))
    title.name = "kpi-card-title"
    title.text = "Sample KPI"
    metric = card.shapes.add_textbox(Inches(2.3), Inches(0.15), Inches(0.6), Inches(0.4))
    metric.name = "kpi-card-metric"
    metric.text = "50%"
    card.left = Inches(1.0)
    card.top = Inches(1.0)
    card.width = Inches(3.2)
    card.height = Inches(3.0)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)

    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "kpi-card-review"}},
        "components": [{
            "component_id": "kpi.chart-card",
            "family": "kpi_chart_card",
            "granularity": "micro",
            "slide_index": 1,
            "semantic_uses": ["single KPI comparison"],
            "groups": [{
                "role": "chart", "scope": "item", "item_indices": [0],
                "shape_names": ["kpi-card-chart"],
            }, {
                "role": "label", "scope": "item", "bind_field": "title",
                "item_indices": [0], "shape_names": ["kpi-card-title"],
            }, {
                "role": "label", "scope": "item", "bind_field": "metric",
                "item_indices": [0], "shape_names": ["kpi-card-metric"],
            }, {
                "role": "card_background", "scope": "shared",
                "shape_names": ["kpi-card-background"],
            }],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }],
    }
    atlas = build_component_atlas(template, review)
    ir = {"slides": [{
        "id": "impact",
        "title": "Innovation impact",
        "charts": [{
            "id": "innovation",
            "data": {
                "categories": ["Promoted", "Other"],
                "series": [{"name": "Share", "values": [64, 36]}],
            },
        }],
        "blocks": [{
            "id": "metric", "role": "list",
            "items": [{"text": "Innovation"}, {"text": "64%"}],
        }],
    }]}
    plan = {"operations": [{
        "kind": "chart_component_clone",
        "source_slide_index": 1,
        "destination_slide_index": 2,
        "component_requirement": {
            "component_id": "kpi.chart-card",
            "semantic_use": "single KPI comparison",
            "family": "kpi_chart_card",
            "element_count": 1,
        },
        "chart": {
            "binding_name": "bind:chart:impact:innovation",
            "data": {
                "categories": ["Promoted", "Other"],
                "series": [{"name": "Share", "values": [64, 36]}],
            },
        },
        "text_bindings": [{
            "field": "title",
            "binding_name": "bind:block:impact:metric:item:0",
            "text": "Innovation",
        }, {
            "field": "metric",
            "binding_name": "bind:block:impact:metric:item:1",
            "text": "64%",
        }],
        "placement": {"x": 0.1, "y": 0.15, "w": 0.3, "h": 0.5},
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, result

    def descendants(shapes):
        for shape in shapes:
            yield shape
            if hasattr(shape, "shapes"):
                yield from descendants(shape.shapes)

    source_slide = Presentation(template).slides[0]
    output_slide = Presentation(output).slides[0]
    source_chart = next(shape for shape in descendants(source_slide.shapes) if getattr(shape, "has_chart", False))
    assert list(source_chart.chart.series[0].values) == [50.0, 50.0]
    assert sum(hasattr(shape, "shapes") for shape in output_slide.shapes) == 1
    names = {shape.name for shape in descendants(output_slide.shapes)}
    assert "bind:chart:impact:innovation" in names
    assert "bind:block:impact:metric:item:0" in names
    assert "bind:block:impact:metric:item:1" in names
    assert any(name.startswith("decoration:component:kpi_chart_card:kpi.chart-card:shared:0") for name in names)


def test_strict_template_clones_a_reviewed_fixed_native_group_from_semantic_fields(tmp_path: Path) -> None:
    template = tmp_path / "native-group-template.pptx"
    output = tmp_path / "native-group-output.pptx"
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = source.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(5.0), Inches(3.0),
    )
    background.name = "group-background"
    title = source.shapes.add_textbox(Inches(1.1), Inches(1.0), Inches(4.2), Inches(0.5))
    title.name = "group-title"
    title.text = "Sample title"
    for index in range(2):
        detail = source.shapes.add_textbox(
            Inches(1.3), Inches(1.7 + index * 0.8), Inches(3.8), Inches(0.5),
        )
        detail.name = f"group-detail-{index + 1}"
        detail.text = f"Sample detail {index + 1}"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "native-group"}},
        "components": [{
            "component_id": "insight.list",
            "family": "fixed_native_group",
            "granularity": "section",
            "renderer": "native_group",
            "slide_index": 1,
            "semantic_uses": ["evidence insight list"],
            "groups": [{
                "role": "label", "scope": "shared", "bind_field": "title",
                "shape_names": ["group-title"],
            }, {
                "role": "label", "scope": "item", "bind_field": "detail",
                "item_indices": [0, 1],
                "shape_names": ["group-detail-1", "group-detail-2"],
            }, {
                "role": "decoration", "scope": "shared",
                "shape_names": ["group-background"],
            }],
            "parameters": {"element_count": {"minimum": 2, "maximum": 2}},
        }],
    }
    atlas = build_component_atlas(template, review)
    ir = {"slides": [{
        "id": "impact",
        "title": "Impact evidence",
        "blocks": [{
            "id": "insights", "role": "list",
            "items": [{"text": text} for text in [
                "Management implications",
                "Innovation impact is broader than EBIT impact.",
                "Workflow redesign remains the scaling constraint.",
            ]],
        }],
    }]}
    plan = {"operations": [{
        "kind": "native_group_component_clone",
        "destination_slide_index": 2,
        "component_requirement": {
            "component_id": "insight.list",
            "semantic_use": "evidence insight list",
            "family": "fixed_native_group",
            "element_count": 2,
        },
        "text_bindings": [{
            "field": "title",
            "binding_name": "bind:block:impact:insights:item:0",
            "text": "Management implications",
        }, {
            "field": "detail",
            "binding_name": "bind:block:impact:insights:item:1",
            "text": "Innovation impact is broader than EBIT impact.",
        }, {
            "field": "detail",
            "binding_name": "bind:block:impact:insights:item:2",
            "text": "Workflow redesign remains the scaling constraint.",
        }],
        "placement": {"x": 0.2, "y": 0.2, "w": 0.6, "h": 0.5},
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, result
    source_slide = Presentation(template).slides[0]
    output_slide = Presentation(output).slides[0]
    assert {shape.text for shape in source_slide.shapes if getattr(shape, "has_text_frame", False)} >= {
        "Sample title", "Sample detail 1", "Sample detail 2",
    }
    output_names = {shape.name for shape in output_slide.shapes}
    assert output_names >= {
        "bind:block:impact:insights:item:0",
        "bind:block:impact:insights:item:1",
        "bind:block:impact:insights:item:2",
    }
    assert any(
        name.startswith("decoration:component:fixed_native_group:insight.list:shared:0")
        for name in output_names
    )


def test_strict_template_selects_and_clones_a_native_funnel_without_filling_the_whole_slide(tmp_path: Path) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.4))
    title.name = "template-title"
    title.text = "Unchanged page heading"
    for index in range(3):
        segment = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(2.0 + index * 0.3), Inches(1.2 + index * 1.1),
            Inches(5.2 - index * 0.6), Inches(0.8),
        )
        segment.name = f"template-segment-{index + 1}"
        label = slide.shapes.add_textbox(
            Inches(2.1 + index * 0.3), Inches(1.35 + index * 1.1),
            Inches(4.8 - index * 0.6), Inches(0.35),
        )
        label.name = f"template-label-{index + 1}"
        label.text = f"Example {index + 1}"
    destination = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = destination.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.4))
    title.name = "destination-title"
    title.text = "Unchanged page heading"
    presentation.save(template)
    ir = {
        "slides": [{
            "id": "agents",
            "blocks": [{
                "id": "stages", "role": "list",
                "items": [{"text": "尝试"}, {"text": "试点"}, {"text": "规模化"}, {"text": "价值实现"}],
            }],
        }],
    }
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {
            "state": "reviewed",
            "reviewer": {"method": "human", "id": "design-review"},
        },
        "components": [{
            "component_id": "funnel.primary",
            "family": "funnel",
            "slide_index": 1,
            "semantic_uses": ["conversion"],
            "groups": [
                {"role": "segment", "shape_names": [
                    "template-segment-1", "template-segment-2", "template-segment-3",
                ]},
                {"role": "label", "shape_names": [
                    "template-label-1", "template-label-2", "template-label-3",
                ]},
            ],
            "parameters": {
                "element_count": {"minimum": 1, "maximum": 8},
                "size_driver": "value",
                "direction": "descending",
            },
        }],
    })
    plan = {
        "operations": [{
            "kind": "component_clone", "destination_slide_index": 2,
            "component_requirement": {"semantic_use": "conversion", "element_count": 4},
            "placement": {"x": 0.55, "y": 0.2, "w": 0.4, "h": 0.6},
            "elements": [
                {"text": "尝试", "value": 100, "binding_name": "bind:block:agents:stages:item:0"},
                {"text": "试点", "value": 56, "binding_name": "bind:block:agents:stages:item:1"},
                {"text": "规模化", "value": 23, "binding_name": "bind:block:agents:stages:item:2"},
                {"text": "价值实现", "value": 8, "binding_name": "bind:block:agents:stages:item:3"},
            ],
        }],
    }

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, result
    assert result["binding"]["status"] == "pass"
    output_slide = Presentation(output).slides[0]
    widths = [
        next(shape for shape in output_slide.shapes if shape.name == f"decoration:component:funnel:segment:{index}").width
        for index in range(4)
    ]
    assert widths == sorted(widths, reverse=True)
    placed_shapes = [
        shape for shape in output_slide.shapes
        if shape.name.startswith("decoration:component:funnel:")
        or shape.name.startswith("bind:block:agents:")
    ]
    assert min(shape.left for shape in placed_shapes) / Presentation(output).slide_width >= 0.548
    assert next(shape for shape in output_slide.shapes if shape.name == "destination-title").text == "Unchanged page heading"
    assert result["operations"][0]["copied"]["component_id"] == "funnel.primary"
    assert result["operations"][0]["copied"]["source_slide_index"] == 1
    assert result["operations"][0]["copied"]["destination_slide_index"] == 2


def test_strict_template_clones_places_and_binds_a_native_shape_group(tmp_path: Path) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = source.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(5), Inches(2),
    )
    background.name = "native-background"
    title = source.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(4.2), Inches(0.4))
    title.name = "native-title"
    title.text = "Sample title"
    detail = source.shapes.add_textbox(Inches(1.4), Inches(1.9), Inches(4.2), Inches(0.7))
    detail.name = "native-detail"
    detail.text = "Sample detail"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)

    ir = {"slides": [{
        "id": "unsupported",
        "blocks": [{
            "id": "status",
            "role": "list",
            "items": [{
                "text": "组件待补充",
                "detail": "此页所需组件尚未进入 reviewed Atlas。",
            }],
        }],
    }]}
    plan = {"operations": [{
        "kind": "native_group_clone",
        "source_slide_index": 1,
        "destination_slide_index": 2,
        "shape_names": ["native-background", "native-title", "native-detail"],
        "placement": {"x": 0.15, "y": 0.25, "w": 0.7, "h": 0.45},
        "text_bindings": [{
            "shape_name": "native-title",
            "binding_name": "bind:block:unsupported:status:item:0",
            "text": "组件待补充",
        }, {
            "shape_name": "native-detail",
            "binding_name": "bind:block:unsupported:status:detail:0",
            "text": "此页所需组件尚未进入 reviewed Atlas。",
        }],
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
    )

    assert result["ok"] is True, result
    copied = Presentation(output).slides[0].shapes
    assert {shape.name for shape in copied if shape.name.startswith("bind:")} == {
        "bind:block:unsupported:status:item:0",
        "bind:block:unsupported:status:detail:0",
    }
    assert "Sample title" not in {shape.text for shape in copied if getattr(shape, "has_text_frame", False)}
    assert result["operations"][0]["copied"]["placement"]["w"] == 0.7


def test_strict_template_clones_and_binds_an_extended_atlas_component(tmp_path: Path) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, width in enumerate((2.0, 3.0, 4.0)):
        left = 4.0 - width / 2
        top = 1.0 + index * 1.1
        segment = source.shapes.add_shape(
            MSO_SHAPE.TRAPEZOID, Inches(left), Inches(top), Inches(width), Inches(0.8),
        )
        segment.name = f"pyramid-segment-{index + 1}"
        title = source.shapes.add_textbox(
            Inches(left + 0.2), Inches(top + 0.14), Inches(width * 0.55), Inches(0.3),
        )
        title.name = f"pyramid-title-{index + 1}"
        title.text = f"Sample {index + 1}"
        value = source.shapes.add_textbox(
            Inches(left + width * 0.68), Inches(top + 0.14), Inches(width * 0.25), Inches(0.3),
        )
        value.name = f"pyramid-value-{index + 1}"
        value.text = f"{index + 1}0%"
        shadow = source.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top + 0.72), Inches(width * 0.5), Inches(0.08),
        )
        shadow.name = f"pyramid-shadow-{index + 1}"
    baseline = source.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.8), Inches(4.55), Inches(4.4), Inches(0.12),
    )
    baseline.name = "pyramid-baseline"
    destination = presentation.slides.add_slide(presentation.slide_layouts[6])
    destination.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.4)).text = "Keep"
    presentation.save(template)

    review_groups = [{
        "role": "segment",
        "scope": "item",
        "item_indices": [0, 1, 2],
        "shape_names": [f"pyramid-segment-{index}" for index in range(1, 4)],
    }, {
        "role": "shadow",
        "scope": "item",
        "item_indices": [0, 1, 2],
        "shape_names": [f"pyramid-shadow-{index}" for index in range(1, 4)],
    }, {
        "role": "label",
        "scope": "item",
        "bind_field": "title",
        "item_indices": [0, 1, 2],
        "shape_names": [f"pyramid-title-{index}" for index in range(1, 4)],
    }, {
        "role": "label",
        "scope": "item",
        "bind_field": "value",
        "item_indices": [0, 1, 2],
        "shape_names": [f"pyramid-value-{index}" for index in range(1, 4)],
    }, {
        "role": "attachment",
        "scope": "shared",
        "shape_names": ["pyramid-baseline"],
    }]
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "review"}},
        "components": [{
            "component_id": "pyramid.extended",
            "family": "pyramid",
            "slide_index": 1,
            "semantic_uses": ["hierarchy"],
            "groups": review_groups,
            "parameters": {"element_count": {"minimum": 1, "maximum": 6}},
        }],
    })
    data = (("战略方向", "20%", 20), ("组织能力", "45%", 45), ("日常执行", "90%", 90))
    ir = {"slides": [{"id": "pyramid", "blocks": [{
        "id": "levels",
        "role": "list",
        "items": [{"text": title, "detail": detail} for title, detail, _value in data],
    }]}]}
    plan = {"operations": [{
        "kind": "component_clone",
        "destination_slide_index": 2,
        "component_requirement": {"semantic_use": "hierarchy", "element_count": 3},
        "placement": {"x": 0.12, "y": 0.18, "w": 0.76, "h": 0.7},
        "elements": [{
            "value": numeric_value,
            "labels": {
                "title": {
                    "text": title,
                    "binding_name": f"bind:block:pyramid:levels:item:{index}",
                },
                "value": {
                    "text": detail,
                    "binding_name": f"bind:block:pyramid:levels:detail:{index}",
                },
            },
        } for index, (title, detail, numeric_value) in enumerate(data)],
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, json.dumps(result, ensure_ascii=False, indent=2)
    assert result["binding"]["status"] == "pass"
    slide = Presentation(output).slides[0]
    names = {shape.name for shape in slide.shapes}
    assert "decoration:component:pyramid:shared:0" in names
    assert "bind:block:pyramid:levels:item:2" in names
    assert "bind:block:pyramid:levels:detail:2" in names
    assert not any("Sample" in getattr(shape, "text", "") for shape in slide.shapes)


def test_strict_template_rejects_an_atlas_from_a_different_template(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    _reference_template(template)
    atlas = {
        "schema_version": "1.0.0",
        "status": "reviewed",
        "source": {"sha256": "sha256:not-this-template"},
        "components": [],
    }

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=_plan(),
        ir=_ir(), source_docs={}, output_pptx=output, component_atlas=atlas,
    )

    assert result["ok"] is False
    assert result["code"] == "COMPONENT_ATLAS_SOURCE_HASH_MISMATCH"
    assert not output.exists()


def test_strict_template_clones_and_binds_one_native_ring_as_a_micro_component(tmp_path: Path) -> None:
    template = tmp_path / "ring-template.pptx"
    output = tmp_path / "ring-output.pptx"
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = ChartData()
    data.categories = ["Done", "Remaining"]
    data.add_series("Progress", (45, 55))
    chart = source.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(1), Inches(1), Inches(2), Inches(2), data,
    )
    chart.name = "ring-chart"
    value = source.shapes.add_textbox(Inches(1.65), Inches(1.75), Inches(0.8), Inches(0.35))
    value.name = "ring-value"
    value.text = "45%"
    label = source.shapes.add_textbox(Inches(1.4), Inches(3.1), Inches(1.2), Inches(0.35))
    label.name = "ring-label"
    label.text = "Sample"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "ring"}},
        "components": [{
            "component_id": "ring.metric", "family": "metric_ring", "granularity": "micro",
            "slide_index": 1, "semantic_uses": ["progress metric"],
            "groups": [{"role": "chart", "scope": "item", "item_indices": [0], "shape_names": ["ring-chart"]},
                       {"role": "label", "scope": "item", "bind_field": "value", "item_indices": [0], "shape_names": ["ring-value"]},
                       {"role": "label", "scope": "item", "bind_field": "label", "item_indices": [0], "shape_names": ["ring-label"]}],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }],
    })
    ir = {"slides": [{
        "id": "progress",
        "charts": [{"id": "ring", "data": {"categories": ["Done", "Remaining"], "series": [{"name": "Progress", "values": [72, 28]}]}}],
        "blocks": [{"id": "summary", "role": "list", "items": [{"text": "72%", "detail": "完成率"}]}],
    }]}
    plan = {"operations": [{
        "kind": "chart_component_clone", "destination_slide_index": 2,
        "component_instance_id": "slide-2-ring-1",
        "component_requirement": {"component_id": "ring.metric", "semantic_use": "progress metric", "family": "metric_ring", "element_count": 1},
        "placement": {"x": 0.2, "y": 0.2, "w": 0.3, "h": 0.45},
        "chart": {"binding_name": "bind:chart:progress:ring", "data": ir["slides"][0]["charts"][0]["data"]},
        "text_bindings": [{"field": "value", "binding_name": "bind:block:progress:summary:item:0", "text": "72%"},
                          {"field": "label", "binding_name": "bind:block:progress:summary:detail:0", "text": "完成率"}],
    }]}

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, json.dumps(result, ensure_ascii=False, indent=2)
    destination = Presentation(output).slides[0]
    names = {shape.name for shape in destination.shapes}
    assert "bind:chart:progress:ring" in names
    assert "bind:block:progress:summary:item:0" in names
    assert "bind:block:progress:summary:detail:0" in names
    assert not any("Sample" in getattr(shape, "text", "") for shape in destination.shapes)


def test_strict_template_executes_two_atlas_components_in_independent_page_slots(tmp_path: Path) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    template = tmp_path / "template.pptx"
    output = tmp_path / "strict-output.pptx"
    presentation = Presentation()
    funnel_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    timeline_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    destination = presentation.slides.add_slide(presentation.slide_layouts[6])
    destination.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8), Inches(0.5)).text = "Multi component"
    for index in range(3):
        funnel_segment = funnel_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1 + index), Inches(4 - index * 0.5), Inches(0.7),
        )
        funnel_segment.name = f"funnel-segment-{index + 1}"
        funnel_label = funnel_slide.shapes.add_textbox(
            Inches(1.2), Inches(1.15 + index), Inches(3), Inches(0.3),
        )
        funnel_label.name = f"funnel-label-{index + 1}"
        funnel_label.text = f"Funnel {index + 1}"
        timeline_node = timeline_slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1 + index * 2), Inches(2), Inches(0.5), Inches(0.5),
        )
        timeline_node.name = f"timeline-node-{index + 1}"
        timeline_label = timeline_slide.shapes.add_textbox(
            Inches(0.6 + index * 2), Inches(2.7), Inches(1.4), Inches(0.4),
        )
        timeline_label.name = f"timeline-label-{index + 1}"
        timeline_label.text = f"Step {index + 1}"
    presentation.save(template)
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "design-review"}},
        "components": [{
            "component_id": "funnel.primary", "family": "funnel", "slide_index": 1,
            "semantic_uses": ["conversion"],
            "groups": [
                {"role": "segment", "shape_names": [f"funnel-segment-{index}" for index in range(1, 4)]},
                {"role": "label", "shape_names": [f"funnel-label-{index}" for index in range(1, 4)]},
            ],
            "parameters": {"element_count": {"minimum": 1, "maximum": 6}},
        }, {
            "component_id": "timeline.primary", "family": "timeline", "slide_index": 2,
            "semantic_uses": ["sequence"],
            "groups": [
                {"role": "segment", "shape_names": [f"timeline-node-{index}" for index in range(1, 4)]},
                {"role": "label", "shape_names": [f"timeline-label-{index}" for index in range(1, 4)]},
            ],
            "parameters": {"element_count": {"minimum": 2, "maximum": 6}},
        }],
    })
    funnel_items = [{"text": value} for value in ("探索", "试点", "规模化")]
    timeline_items = [{"text": value} for value in ("发现", "验证", "部署", "优化")]
    ir = {"slides": [{"id": "multi", "blocks": [
        {"id": "conversion", "role": "list", "items": funnel_items},
        {"id": "sequence", "role": "list", "items": timeline_items},
    ]}]}
    plan = build_component_plan(atlas, {
        "schema_version": "1.0.0",
        "pages": [{"destination_slide_index": 3, "components": [{
            "semantic_use": "conversion",
            "placement": {"x": 0.05, "y": 0.18, "w": 0.42, "h": 0.68},
            "elements": [
                {"text": item["text"], "value": 100 - index * 35,
                 "binding_name": f"bind:block:multi:conversion:item:{index}"}
                for index, item in enumerate(funnel_items)
            ],
        }, {
            "semantic_use": "sequence",
            "placement": {"x": 0.53, "y": 0.18, "w": 0.42, "h": 0.68},
            "elements": [
                {"text": item["text"],
                 "binding_name": f"bind:block:multi:sequence:item:{index}"}
                for index, item in enumerate(timeline_items)
            ],
        }]}],
    })

    result = execute_strict_template(
        request=_request(template), template_pptx=template, strict_plan=plan,
        ir=ir, source_docs={}, output_pptx=output, render_engine="libreoffice",
        component_atlas=atlas,
    )

    assert result["ok"] is True, json.dumps(result.get("inspection", result), ensure_ascii=False, indent=2)
    slide = Presentation(output).slides[0]
    names = [shape.name for shape in slide.shapes]
    assert len(names) == len(set(names))
    assert any("funnel:slide-3-component-1" in name for name in names)
    assert any("timeline:slide-3-component-2" in name for name in names)
    assert result["binding"]["status"] == "pass"


def test_strict_template_cli_runs_only_the_declared_native_operations(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    request = tmp_path / "request.json"
    plan = tmp_path / "strict-plan.json"
    ir = tmp_path / "ir.json"
    source = tmp_path / "source.md"
    output = tmp_path / "strict-output.pptx"
    report = tmp_path / "report.json"
    atlas_path = tmp_path / "component-atlas.json"
    evidence_path = tmp_path / "evidence-ledger.json"
    bindings_path = tmp_path / "content-bindings.json"
    _reference_template(template)
    request.write_text(json.dumps(_request(template)), encoding="utf-8")
    plan.write_text(json.dumps(_plan()), encoding="utf-8")
    ir.write_text(json.dumps(_ir()), encoding="utf-8")
    source.write_text("| Region | Growth |\n| --- | --- |\n| EU | 15% |", encoding="utf-8")
    atlas_path.write_text(json.dumps({
        "schema_version": "1.0.0",
        "status": "reviewed",
        "source": {"sha256": _sha(template)},
        "components": [],
    }), encoding="utf-8")
    evidence_ledger, content_bindings = _content_integrity_inputs()
    evidence_path.write_text(json.dumps(evidence_ledger), encoding="utf-8")
    bindings_path.write_text(json.dumps(content_bindings), encoding="utf-8")

    completed = subprocess.run(
        [
            sys.executable, "-m", "engine", "strict-template",
            "--request", str(request), "--template-pptx", str(template),
            "--plan", str(plan), "--ir", str(ir),
            "--evidence-ledger", str(evidence_path),
            "--content-bindings", str(bindings_path),
            "--component-atlas", str(atlas_path),
            "--source", f"doc:markdown:{source}", "--output-pptx", str(output),
            "--render-engine", "libreoffice", "--json-out", str(report),
        ],
        cwd=ROOT, text=True, capture_output=True, check=False,
    )

    assert completed.returncode == 0, completed.stderr or completed.stdout
    assert json.loads(report.read_text(encoding="utf-8"))["ok"] is True
