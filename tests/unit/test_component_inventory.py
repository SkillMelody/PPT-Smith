from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from engine.component_atlas import build_component_atlas
from engine.component_inventory import build_template_component_inventory


ROOT = Path(__file__).resolve().parents[2]


def _template_and_atlas(tmp_path: Path) -> tuple[Path, dict]:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    segment = first.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1),
    )
    segment.name = "metric-surface"
    label = first.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(2), Inches(0.4))
    label.name = "metric-label"
    label.text = "Revenue"
    note = first.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(0.5))
    note.name = "unreviewed-note"
    note.text = "Still needs component review"
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.5))
    title.name = "second-title"
    title.text = "Uncovered slide"
    presentation.save(template)
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "inventory"}},
        "components": [{
            "component_id": "metric.card", "family": "metric_card", "granularity": "micro",
            "slide_index": 1, "semantic_uses": ["metric"],
            "groups": [{"role": "segment", "shape_names": ["metric-surface"]},
                       {"role": "label", "shape_names": ["metric-label"]}],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }],
    })
    return template, atlas


def test_template_inventory_reports_covered_and_uncovered_native_content(tmp_path: Path) -> None:
    template, atlas = _template_and_atlas(tmp_path)

    inventory = build_template_component_inventory(template, atlas)

    assert inventory["summary"] == {
        "slide_count": 2,
        "native_content_count": 4,
        "covered_content_count": 2,
        "uncovered_content_count": 2,
        "page_recipe_only_content_count": 0,
        "coverage_ratio": 0.5,
        "fully_covered_slides": 0,
        "partially_covered_slides": 1,
        "uncovered_slides": 1,
    }
    assert inventory["slides"][0]["status"] == "partial"
    assert inventory["slides"][0]["uncovered"][0]["shape_name"] == "unreviewed-note"
    assert inventory["slides"][1]["status"] == "uncovered"
    assert inventory["slides"][1]["uncovered"][0]["text"] == "Uncovered slide"


def test_template_inventory_cli_writes_the_full_template_coverage_matrix(tmp_path: Path) -> None:
    template, atlas = _template_and_atlas(tmp_path)
    atlas_path = tmp_path / "atlas.json"
    output_path = tmp_path / "inventory.json"
    atlas_path.write_text(json.dumps(atlas), encoding="utf-8")

    completed = subprocess.run(
        [sys.executable, "-m", "engine", "component-inventory",
         "--template-pptx", str(template), "--component-atlas", str(atlas_path),
         "--json-out", str(output_path)],
        cwd=ROOT, text=True, capture_output=True, check=False,
    )

    assert completed.returncode == 0, completed.stderr
    report = json.loads(output_path.read_text(encoding="utf-8"))
    assert report["summary"]["uncovered_content_count"] == 2
    assert [slide["slide_index"] for slide in report["slides"]] == [1, 2]


def test_template_inventory_does_not_count_a_page_recipe_as_reusable_component_coverage(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.5))
    title.name = "page-title"
    title.text = "Whole page preset"
    presentation.save(template)
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "preset"}},
        "components": [{
            "component_id": "page.preset", "family": "dashboard", "granularity": "page_recipe",
            "slide_index": 1, "semantic_uses": ["dashboard"],
            "groups": [{"role": "text", "shape_names": ["page-title"]}],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }],
    })

    inventory = build_template_component_inventory(template, atlas)

    assert inventory["summary"]["covered_content_count"] == 0
    assert inventory["summary"]["uncovered_content_count"] == 1
    assert inventory["summary"]["page_recipe_only_content_count"] == 1
    assert inventory["slides"][0]["uncovered"][0]["page_recipe_covered_by"] == ["page.preset"]


def test_template_inventory_counts_reviewed_equivalent_instances_without_new_components(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for prefix, left in (("primary", 1), ("equivalent", 5)):
        surface = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(1), Inches(3), Inches(1),
        )
        surface.name = f"{prefix}-surface"
        label = slide.shapes.add_textbox(Inches(left + 0.2), Inches(1.2), Inches(2), Inches(0.4))
        label.name = f"{prefix}-label"
        label.text = prefix
    presentation.save(template)
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "aliases"}},
        "components": [{
            "component_id": "metric.card",
            "family": "metric_card",
            "granularity": "micro",
            "slide_index": 1,
            "semantic_uses": ["metric"],
            "groups": [
                {"role": "segment", "shape_names": ["primary-surface"]},
                {"role": "label", "shape_names": ["primary-label"]},
            ],
            "source_instances": [{
                "instance_id": "equivalent-card",
                "groups": [
                    {"role": "segment", "shape_names": ["equivalent-surface"]},
                    {"role": "label", "shape_names": ["equivalent-label"]},
                ],
            }],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }],
    })

    inventory = build_template_component_inventory(template, atlas)

    assert len(atlas["components"]) == 1
    assert inventory["summary"]["covered_content_count"] == 4
    assert inventory["summary"]["uncovered_content_count"] == 0
    assert inventory["slides"][0]["status"] == "full"
