from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from engine.template_evidence import analyze_template, validate_template_evidence


def _text(slide, value: str, x: float, y: float, w: float, h: float, size: float) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(68, 84, 106)


def _reference_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    first = prs.slides.add_slide(blank)
    _text(first, "Dense operating review", 0.5, 0.3, 8.0, 0.5, 26)
    for index in range(12):
        _text(first, f"Evidence {index + 1}", 0.7 + (index % 4) * 3.0,
              1.5 + (index // 4) * 1.2, 2.4, 0.35, 10)
        block = first.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.7 + (index % 4) * 3.0),
            Inches(1.9 + (index // 4) * 1.2), Inches(2.4), Inches(0.18),
        )
        block.fill.solid()
        block.fill.fore_color.rgb = RGBColor(181, 29, 37) if index == 0 else RGBColor(68, 84, 106)

    second = prs.slides.add_slide(blank)
    _text(second, "63%", 0.8, 2.0, 3.0, 1.0, 42)
    _text(second, "Skill gaps constrain growth", 4.5, 2.1, 5.5, 0.6, 18)
    prs.save(path)


def test_analyze_template_emits_private_safe_unreviewed_evidence(tmp_path: Path) -> None:
    pptx = tmp_path / "reference.pptx"
    _reference_pptx(pptx)

    report = analyze_template(pptx)
    encoded = json.dumps(report, ensure_ascii=False, sort_keys=True)

    assert report["status"] == "unreviewed"
    assert report["source"]["filename"] == "reference.pptx"
    assert report["source"]["sha256"].startswith("sha256:")
    assert report["source"]["slide_count"] == 2
    assert str(tmp_path) not in encoded
    assert "/Users/" not in encoded
    assert "routing_rules" not in encoded
    assert "layout_prototype" not in encoded
    assert "executable" not in encoded.lower()
    assert validate_template_evidence(report) == []


def test_analyze_template_records_geometry_without_claiming_page_family(tmp_path: Path) -> None:
    pptx = tmp_path / "reference.pptx"
    _reference_pptx(pptx)

    report = analyze_template(pptx)
    first = report["slides"][0]

    assert first["slide_index"] == 1
    assert first["object_count"] >= 25
    assert first["density"] == "high"
    assert first["title_candidates"][0]["text"] == "Dense operating review"
    assert first["geometry"]["text"] >= 13
    assert first["geometry"]["shape"] >= 12
    assert "family" not in first
    assert "prototype" not in first


def test_template_evidence_is_deterministic(tmp_path: Path) -> None:
    pptx = tmp_path / "reference.pptx"
    _reference_pptx(pptx)

    first = analyze_template(pptx)
    second = analyze_template(pptx)

    assert first == second


def test_template_evidence_requires_visual_review_before_use(tmp_path: Path) -> None:
    pptx = tmp_path / "reference.pptx"
    _reference_pptx(pptx)

    report = analyze_template(pptx)

    assert report["review"]["required"] is True
    assert report["review"]["state"] == "unreviewed"
    assert report["limitations"]
    assert any("generation" in item.lower() for item in report["limitations"])
