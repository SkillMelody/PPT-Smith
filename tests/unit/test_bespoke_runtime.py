from __future__ import annotations

import copy
import hashlib
import json
import subprocess
import sys
from pathlib import Path

from engine.bespoke import build_authoring_manifest
from engine.bespoke_runtime import run_bespoke_author
from engine.production_request import prepare_template_context
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]


def _ir() -> dict:
    return {
        "schema_version": "4.1.0",
        "deck": {"title": "Bespoke proof", "language": "en-US"},
        "sources": [{"source_id": "doc", "type": "markdown", "path": "source.md"}],
        "slides": [{
            "id": "proof", "title": "High performers redesign workflows",
            "message": "Workflow redesign separates high performers from peers.",
            "blocks": [{
                "id": "fact", "role": "fact",
                "text": "High performers redesign complete workflows.",
                "source_ref": {"source_id": "doc", "loc": "para_1"},
            }],
        }],
    }


def _manifest(ir: dict) -> dict:
    request = {
        "schema_version": "1.0.0", "request_id": "runtime-proof", "route": "bespoke",
        "page_contract": {
            "mode": "exact", "exact_pages": 1,
            "overflow_policy": "ask", "underflow_policy": "fewer_pages",
        },
    }
    assessment = {
        "minimum_viable_pages": 1, "recommended_pages": 1,
        "maximum_useful_pages": 1, "basis": ["one decision message"],
    }
    return build_authoring_manifest(request, ir, assessment)


def _complete_manifest(ir: dict) -> dict:
    request = {
        "schema_version": "1.0.0", "request_id": "runtime-complete-proof",
        "route": "bespoke", "delivery_scope": "complete_deck",
        "page_contract": {
            "mode": "exact", "exact_pages": 1,
            "overflow_policy": "ask", "underflow_policy": "fewer_pages",
        },
    }
    assessment = {
        "minimum_viable_pages": 1, "recommended_pages": 1,
        "maximum_useful_pages": 1, "basis": ["one decision message"],
    }
    raw = json.dumps(ir, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    contract = {
        "schema_version": "1.0.0",
        "contract_id": "runtime-professional-narrative",
        "source_ir_sha256": "sha256:" + hashlib.sha256(raw.encode("utf-8")).hexdigest(),
        "quality_principles": [
            "audience_decision_alignment", "assertion_evidence",
            "coherence_progression", "cognitive_load", "semantic_visual_encoding",
        ],
        "freedom": {
            "model_owns": [
                "storyline", "page_composition", "visual_encoding",
                "geometry", "cross_page_rhythm",
            ],
            "engine_verifies": [
                "source_fidelity", "page_contract", "native_editability",
                "render_integrity", "visual_review",
            ],
            "forbidden_prescriptions": [
                "composition_enum", "coordinate_grid",
                "template_slide_mapping", "archetype_selection",
            ],
        },
        "deck": {
            "audience_decision": "Decide whether to redesign the workflow.",
            "core_thesis": "Workflow redesign separates high performers.",
            "narrative_arc": "State the decision and prove it with source evidence.",
        },
        "chapters": [{
            "chapter_id": "proof", "title": "Proof", "job": "Support the decision.",
            "slide_ids": ["proof"],
        }],
        "pages": [{
            "slide_id": "proof",
            "audience_question": "What separates high performers?",
            "assertion_ref": "slide:proof:message",
            "evidence_refs": ["block:proof:fact"],
            "narrative_function": "answer_with_evidence",
            "transition_to_next": None,
        }],
    }
    return build_authoring_manifest(
        request, ir, assessment, narrative_contract=contract,
    )


def _reference_template(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for index in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8), Inches(0.5))
        box.text = f"Reference page {index + 1}"
    presentation.save(path)


def _manifest_with_template(ir: dict, template: Path) -> dict:
    template_hash = "sha256:" + hashlib.sha256(template.read_bytes()).hexdigest()
    request = {
        "schema_version": "1.0.0", "request_id": "runtime-template-proof", "route": "bespoke",
        "page_contract": {
            "mode": "exact", "exact_pages": 1,
            "overflow_policy": "ask", "underflow_policy": "fewer_pages",
        },
        "template": {
            "use_mode": "style_transfer", "source_sha256": template_hash,
            "required": True, "preferred_slide_indices": [1],
        },
    }
    assessment = {
        "minimum_viable_pages": 1, "recommended_pages": 1,
        "maximum_useful_pages": 1, "basis": ["one decision message"],
    }
    return build_authoring_manifest(
        request, ir, assessment, prepare_template_context(request, template),
    )


AUTHOR_SCRIPT = '''
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

def add_text(slide, name, text, x, y, w, h, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)

def build(prs, context):
    ir = context["ir"]
    source = ir["slides"][0]
    block = source["blocks"][0]
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "bind:slide:proof:title", source["title"], 0.6, 0.5, 11, 0.5, 24)
    add_text(slide, "bind:slide:proof:message", source["message"], 0.6, 1.1, 11, 0.4, 13)
    add_text(slide, "bind:block:proof:fact:text", block["text"], 0.8, 2.1, 7, 0.8, 18)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.2), Inches(5), Inches(0.05))
    rule.name = "decoration:accent_rule"
'''


AUTHOR_CHART_SCRIPT = '''
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

def add_text(slide, name, text, x, y, w, h, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)

def build(prs, context):
    source = context["ir"]["slides"][0]
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "bind:slide:proof:title", source["title"], 0.6, 0.5, 11, 0.5, 24)
    add_text(slide, "bind:slide:proof:message", source["message"], 0.6, 1.1, 11, 0.4, 13)
    add_text(slide, "bind:block:proof:fact:text", source["blocks"][0]["text"], 0.8, 1.7, 7, 0.5, 14)
    data = CategoryChartData()
    data.categories = source["chart"]["data"]["categories"]
    for series in source["chart"]["data"]["series"]:
        data.add_series(series["name"], series["values"])
    graphic = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(2.4), Inches(8), Inches(3.5), data)
    graphic.name = "bind:chart:proof"
'''


BAD_QUALITY_AUTHOR_SCRIPT = '''
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

def add_text(slide, name, text, x, y, w, h, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    box.text_frame.word_wrap = False
    box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)

def build(prs, context):
    source = context["ir"]["slides"][0]
    block = source["blocks"][0]
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, "bind:slide:proof:title", source["title"], 0.6, 0.5, 11, 0.5, 24)
    add_text(slide, "bind:slide:proof:message", source["message"], 0.6, 1.1, 11, 0.4, 9)
    add_text(slide, "bind:block:proof:fact:text", block["text"], 0.8, 2.1, 3, 0.35, 8)
'''


def test_bespoke_runtime_rejects_deterministic_visual_floor_failures(tmp_path: Path) -> None:
    ir = _ir()
    script = tmp_path / "author.py"
    script.write_text(BAD_QUALITY_AUTHOR_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(
        script, _manifest(ir), ir, tmp_path / "output",
        render_engine="libreoffice",
    )

    assert result["ok"] is False
    assert result["status"] == "rejected"
    assert result["code"] == "BESPOKE_VISUAL_FLOOR_FAILED"
    codes = {item["code"] for item in result["visual_quality"]["findings"]}
    assert "BESPOKE_BODY_FONT_TOO_SMALL" in codes
    assert "BESPOKE_SHRINK_TO_FIT_FORBIDDEN" in codes
    assert "BESPOKE_WORD_WRAP_REQUIRED" in codes
    assert result["render"]["status"] == "passed"


def test_bespoke_runtime_authors_from_blank_and_stops_at_visual_unreviewed(tmp_path: Path) -> None:
    ir = _ir()
    script = tmp_path / "author.py"
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(
        script, _manifest(ir), ir, tmp_path / "output",
        render_engine="libreoffice",
    )

    assert result["ok"] is True
    assert result["status"] == "visual_unreviewed"
    assert result["started_from_blank"] is True
    assert result["binding"]["status"] == "pass"
    assert result["structural_blockers"] == []
    assert result["render"]["status"] == "passed"
    assert result["native_editability"]["bound_native_text_shapes"] == 3
    assert Path(result["deck"]).is_file()
    assert Path(result["report"]).is_file()


def test_bespoke_runtime_rejects_manifest_for_different_ir_before_authoring(tmp_path: Path) -> None:
    ir = _ir()
    manifest = _manifest(ir)
    changed_ir = copy.deepcopy(ir)
    changed_ir["slides"][0]["blocks"][0]["text"] = "Changed after approval."
    script = tmp_path / "author.py"
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(script, manifest, changed_ir, tmp_path / "output")

    assert result["ok"] is False
    assert result["status"] == "rejected"
    assert result["code"] == "AUTHORING_MANIFEST_IR_HASH_MISMATCH"
    assert not (tmp_path / "output" / "deck.pptx").exists()


def test_bespoke_runtime_rejects_tampered_narrative_contract_before_authoring(tmp_path: Path) -> None:
    ir = _ir()
    manifest = _complete_manifest(ir)
    manifest["narrative_contract"]["deck"]["core_thesis"] = "Tampered after approval."

    result = run_bespoke_author(
        tmp_path / "missing.py", manifest, ir, tmp_path / "output",
    )

    assert result["ok"] is False
    assert result["code"] == "AUTHORING_NARRATIVE_CONTRACT_HASH_MISMATCH"
    assert not (tmp_path / "output" / "deck.pptx").exists()


def test_bespoke_runtime_requires_locked_narrative_for_complete_deck(tmp_path: Path) -> None:
    ir = _ir()
    manifest = _manifest(ir)
    manifest["delivery_scope"] = "complete_deck"

    result = run_bespoke_author(
        tmp_path / "missing.py", manifest, ir, tmp_path / "output",
    )

    assert result["ok"] is False
    assert result["code"] == "AUTHORING_NARRATIVE_CONTRACT_REQUIRED"


def test_bespoke_runtime_verifies_native_chart_data_and_editability(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["chart"] = {
        "type": "column",
        "data": {
            "categories": ["Pilot", "Scale"],
            "series": [{"name": "2025 share", "values": [30, 70]}],
        },
        "source_ref": {"source_id": "doc", "loc": "table_1"},
    }
    script = tmp_path / "author.py"
    script.write_text(AUTHOR_CHART_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(
        script, _manifest(ir), ir, tmp_path / "output", render_engine="libreoffice",
    )

    assert result["ok"] is True
    assert result["binding"]["valid_bindings"] == [
        "block:proof:fact:text", "chart:proof", "slide:proof:message", "slide:proof:title",
    ]
    assert result["native_editability"]["bound_native_charts"] == 1


def test_bespoke_runtime_requires_every_manifest_binding_not_just_one_per_block(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"][0]["label"] = "Workflow redesign"
    script = tmp_path / "author.py"
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(script, _manifest(ir), ir, tmp_path / "output")

    assert result["ok"] is False
    assert result["status"] == "rejected"
    assert result["code"] == "AUTHORING_BINDINGS_INCOMPLETE"
    assert result["missing_required_bindings"] == ["block:proof:fact:label"]


def test_bespoke_runtime_requires_source_documents_for_native_table_binding(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"].append({
        "id": "regional", "role": "table",
        "source_ref": {"source_id": "doc", "loc": "table_1"},
    })
    script = tmp_path / "author.py"
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(script, _manifest(ir), ir, tmp_path / "output")

    assert result["ok"] is False
    assert result["code"] == "TABLE_SOURCE_DOCUMENTS_REQUIRED"


def test_bespoke_runtime_provides_hash_bound_template_and_rendered_reference_pages(tmp_path: Path) -> None:
    ir = _ir()
    template = tmp_path / "reference.pptx"
    script = tmp_path / "author.py"
    _reference_template(template)
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(
        script, _manifest_with_template(ir, template), ir, tmp_path / "output",
        template_pptx=template, render_engine="libreoffice",
    )

    assert result["ok"] is True
    reference = result["template_reference"]
    assert reference["status"] == "rendered"
    assert reference["reference_pptx"] == str(template)
    assert reference["selected_slide_indices"] == [1]
    assert len(reference["rendered_pages"]) == 1
    assert Path(reference["rendered_pages"][0]["image"]).is_file()


def test_bespoke_runtime_rejects_template_manifest_without_local_reference(tmp_path: Path) -> None:
    ir = _ir()
    template = tmp_path / "reference.pptx"
    script = tmp_path / "author.py"
    _reference_template(template)
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")

    result = run_bespoke_author(
        script, _manifest_with_template(ir, template), ir, tmp_path / "output",
    )

    assert result["ok"] is False
    assert result["code"] == "TEMPLATE_REFERENCE_REQUIRED"


def test_bespoke_runtime_does_not_falsely_claim_strict_template_master_reuse(tmp_path: Path) -> None:
    ir = _ir()
    template = tmp_path / "reference.pptx"
    script = tmp_path / "author.py"
    _reference_template(template)
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")
    manifest = _manifest_with_template(ir, template)
    manifest["template"]["use_mode"] = "strict"

    result = run_bespoke_author(
        script, manifest, ir, tmp_path / "output", template_pptx=template,
    )

    assert result["ok"] is False
    assert result["code"] == "TEMPLATE_STRICT_REUSE_NOT_IMPLEMENTED"


def test_author_bespoke_cli_runs_independent_author_runtime(tmp_path: Path) -> None:
    ir = _ir()
    script = tmp_path / "author.py"
    manifest = tmp_path / "manifest.json"
    ir_path = tmp_path / "ir.json"
    report = tmp_path / "result.json"
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")
    manifest.write_text(json.dumps(_manifest(ir)), encoding="utf-8")
    ir_path.write_text(json.dumps(ir), encoding="utf-8")

    completed = subprocess.run(
        [
            sys.executable, "-m", "engine", "author-bespoke",
            "--script", str(script), "--manifest", str(manifest), "--ir", str(ir_path),
            "--output-dir", str(tmp_path / "output"), "--render-engine", "libreoffice",
            "--json-out", str(report),
        ],
        cwd=ROOT, text=True, capture_output=True, check=False,
    )

    assert completed.returncode == 0, completed.stderr or completed.stdout
    result = json.loads(report.read_text(encoding="utf-8"))
    assert result["status"] == "visual_unreviewed"
    assert result["started_from_blank"] is True


def test_author_bespoke_cli_forwards_template_reference(tmp_path: Path) -> None:
    ir = _ir()
    template = tmp_path / "reference.pptx"
    script = tmp_path / "author.py"
    manifest = tmp_path / "manifest.json"
    ir_path = tmp_path / "ir.json"
    report = tmp_path / "result.json"
    _reference_template(template)
    script.write_text(AUTHOR_SCRIPT, encoding="utf-8")
    manifest.write_text(json.dumps(_manifest_with_template(ir, template)), encoding="utf-8")
    ir_path.write_text(json.dumps(ir), encoding="utf-8")

    completed = subprocess.run(
        [
            sys.executable, "-m", "engine", "author-bespoke",
            "--script", str(script), "--manifest", str(manifest), "--ir", str(ir_path),
            "--output-dir", str(tmp_path / "output"), "--template-pptx", str(template),
            "--render-engine", "libreoffice", "--json-out", str(report),
        ],
        cwd=ROOT, text=True, capture_output=True, check=False,
    )

    assert completed.returncode == 0, completed.stderr or completed.stdout
    result = json.loads(report.read_text(encoding="utf-8"))
    assert result["template_reference"]["reference_pptx"] == str(template)
