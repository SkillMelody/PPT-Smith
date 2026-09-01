"""P12-refine tests: Level 1 page-level composition intent + Level 2 code refine.

Level 1: IR slide.refine (split/wheel/contrast/spotlight) -> engine renderers
         -> QA-gated deck. Refine wins over chart inference; invalid types
         degrade honestly.
Level 2: scripts/refine_code.py runs a python-pptx refine script over the
         engine deck and QA-rejects only NEW blockers vs the base deck.
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from engine.compile import compile_deck  # noqa: E402
from engine.policy import decide_deck  # noqa: E402

MD = """# 2025 年度经营总结

## 增长飞轮

- 内容生产提升 40%
- 渠道分发扩大 3 倍
- 用户留存提高 25%
- 收入增长 36%

## 对比

海外业务贡献 52%，国内贡献 8%。
"""


def _ir():
    return {
        "schema_version": "4.1.0",
        "deck": {"title": "2025 年度经营总结", "language": "zh"},
        "sources": [{"source_id": "src", "type": "markdown"}],
        "slides": [
            {"id": "s1", "title": "增长飞轮",
             "refine": {"type": "wheel",
                        "hub": {"blocks": ["b1"]}, "spokes": ["b2", "b3", "b4"]},
             "blocks": [
                 {"id": "b1", "role": "fact", "text": "增长飞轮",
                  "source_ref": {"source_id": "src", "loc": "list_1"}},
                 {"id": "b2", "role": "fact", "text": "内容生产提升 40%",
                  "source_ref": {"source_id": "src", "loc": "list_1"}},
                 {"id": "b3", "role": "fact", "text": "渠道分发扩大 3 倍",
                  "source_ref": {"source_id": "src", "loc": "list_1"}},
                 {"id": "b4", "role": "fact", "text": "用户留存提高 25%",
                  "source_ref": {"source_id": "src", "loc": "list_1"}},
             ]},
            {"id": "s2", "title": "区域对比",
             "refine": {"type": "contrast",
                        "left": {"blocks": ["b1"]}, "right": {"blocks": ["b2"]}},
             "blocks": [
                 {"id": "b1", "role": "metric", "label": "海外", "value": "52%",
                  "source_ref": {"source_id": "src", "loc": "para_1"}},
                 {"id": "b2", "role": "metric", "label": "国内", "value": "8%",
                  "source_ref": {"source_id": "src", "loc": "para_1"}},
             ]},
        ],
    }


class TestLevel1Policy:
    def test_refine_drives_archetype(self):
        ir = _ir()
        decisions = decide_deck(ir)
        chosen = {d.slide_id: d.chosen for d in decisions}
        assert chosen["s1"] == "refine_wheel"
        assert chosen["s2"] == "refine_contrast"

    def test_unknown_refine_type_falls_back(self):
        ir = _ir()
        ir["slides"][0]["refine"] = {"type": "magic"}
        decisions = decide_deck(ir)
        d = decisions[0]
        assert not d.chosen.startswith("refine_magic")
        assert d.chosen_by in ("rule", "fallback")

    def test_refine_beats_chart_inference(self):
        """A slide with 2 metrics + refine must NOT get chart_column."""
        ir = _ir()
        # s2 has 2 metrics and refine contrast
        from engine.chart_inference import infer_charts
        from engine.structural_parser import parse_source
        doc = parse_source(MD, "markdown", "src")
        infer_charts(ir, {"src": doc})
        # refine slides must not have chart injected
        assert ir["slides"][1].get("chart") is None
        assert ir["slides"][1].get("refine") is not None


class TestLevel1Compile:
    @pytest.fixture()
    def deck(self, tmp_path):
        article = tmp_path / "a.md"
        article.write_text(MD, encoding="utf-8")
        irp = tmp_path / "ir.json"
        irp.write_text(json.dumps(_ir(), ensure_ascii=False), encoding="utf-8")
        report = compile_deck(
            sources=[("src", "markdown", str(article))],
            ir_path=str(irp), output_dir=str(tmp_path / "out"))
        assert report["ok"], report
        assert report["qa"]["error_count"] == 0
        return tmp_path

    def test_wheel_renders_hub_oval(self, deck):
        plan = json.loads((deck / "out" / "render-plan.json").read_text("utf-8"))
        s1 = plan["slides"][1]  # first content slide
        ovals = [e for e in s1["elements"] if e.get("shape") == "oval"]
        conns = [e for e in s1["elements"] if e.get("type") == "connector"]
        assert len(ovals) == 1, "wheel hub should be an oval"
        assert len(conns) == 3, "wheel should have 3 spokes connectors"

    def test_contrast_renders_divider(self, deck):
        plan = json.loads((deck / "out" / "render-plan.json").read_text("utf-8"))
        s2 = plan["slides"][2]
        dividers = [e for e in s2["elements"]
                    if e.get("semantic_role") == "divider"]
        assert len(dividers) >= 1


GOOD_REFINE = """#!/usr/bin/env python3
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
def build(prs, context):
    from pptx.enum.shapes import MSO_SHAPE
    for i, slide in enumerate(prs.slides):
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(12.2), Inches(0.18), Inches(0.7), Inches(0.28))
        badge.name = f"decoration:brand_badge_{i+1}"
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor.from_string("0A2233")
        badge.line.fill.background()
"""

BAD_REFINE = GOOD_REFINE + """
        box = slide.shapes.add_textbox(Inches(1), Inches(5.8), Inches(3), Inches(0.3))
        box.text = 'Unverified growth claim'
"""


class TestLevel2RefineCode:
    def _run(self, tmp_path, script_src, *, high_fidelity=False):
        src = ROOT / "tests" / "fixtures" / "v4-bench" / "product-launch"
        script = tmp_path / "refine.py"
        script.write_text(script_src, encoding="utf-8")
        out = tmp_path / "out"
        proc = subprocess.run(
            [sys.executable, "-m", "engine", "refine-code",
             "--script", str(script),
             "--source", f"doc:markdown:{src / 'source.md'}",
             "--ir", str(src / "irs" / "sim-strong.json"),
             "--output-dir", str(out)]
            + (["--high-fidelity", "--render-engine", "libreoffice"] if high_fidelity else []),
            capture_output=True, text=True, cwd=str(ROOT), timeout=180)
        return proc

    def test_good_refine_accepted(self, tmp_path):
        proc = self._run(tmp_path, GOOD_REFINE)
        assert proc.returncode == 0, proc.stderr[-2000:]
        result = json.loads(proc.stdout)
        assert result["ok"] is True
        assert result["mode"] == "refine"
        assert result["status"] == "accepted"
        assert result["content_lock"] is None
        assert result["render"] is None
        assert (tmp_path / "out" / "deck-refined.pptx").exists()

    def test_bad_refine_rejected(self, tmp_path):
        proc = self._run(tmp_path, BAD_REFINE, high_fidelity=True)
        assert proc.returncode == 1, proc.stdout
        result = json.loads(proc.stdout)
        assert result["ok"] is False
        assert result["content_lock"]["status"] == "fail"
        assert result["content_lock"]["added_text"]
        assert set(result["content_lock"]["added_text"]) == {"Unverified growth claim"}

    def test_high_fidelity_requires_real_render_success(self, tmp_path, monkeypatch):
        monkeypatch.setenv("PPTSMITH_TEST_RENDERERS", "none")
        proc = self._run(tmp_path, GOOD_REFINE, high_fidelity=True)
        assert proc.returncode == 1, proc.stdout
        result = json.loads(proc.stdout)
        assert result["mode"] == "high_fidelity"
        assert result["content_lock"]["status"] == "pass"
        assert result["render"]["status"] != "passed"
        assert result["status"] == "rejected"
        assert result["visual_review"] == {"state": "unreviewed", "required": True}
