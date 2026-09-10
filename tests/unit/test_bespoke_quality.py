from __future__ import annotations

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from engine.bespoke_quality import inspect_bespoke_visual_quality
from engine.template_readability import enforce_bound_text_floor


def _save(tmp_path, build):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(slide)
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def _bound_text(slide, *, size=14, auto_size=None, word_wrap=True):
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.name = "bind:block:s1:b1:text"
    shape.text_frame.word_wrap = word_wrap
    shape.text_frame.auto_size = auto_size
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "A professional body statement with enough words to require wrapping."
    run.font.size = Pt(size)
    return shape


def test_rejects_small_bound_body_text(tmp_path):
    path = _save(tmp_path, lambda slide: _bound_text(slide, size=9))

    report = inspect_bespoke_visual_quality(path)

    assert report["status"] == "fail"
    assert any(item["code"] == "BESPOKE_BODY_FONT_TOO_SMALL" for item in report["findings"])


def test_rejects_shrink_to_fit_on_business_text(tmp_path):
    path = _save(
        tmp_path,
        lambda slide: _bound_text(slide, size=14, auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE),
    )

    report = inspect_bespoke_visual_quality(path)

    assert report["status"] == "fail"
    assert any(item["code"] == "BESPOKE_SHRINK_TO_FIT_FORBIDDEN" for item in report["findings"])


def test_rejects_long_bound_text_without_word_wrap(tmp_path):
    path = _save(tmp_path, lambda slide: _bound_text(slide, size=14, word_wrap=False))

    report = inspect_bespoke_visual_quality(path)

    assert report["status"] == "fail"
    assert any(item["code"] == "BESPOKE_WORD_WRAP_REQUIRED" for item in report["findings"])


def test_rejects_excessive_outlined_containers(tmp_path):
    def build(slide):
        _bound_text(slide)
        for index in range(9):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.4 + index * 0.7),
                Inches(2),
                Inches(0.6),
                Inches(0.6),
            )
            shape.name = f"decoration:card:{index}"
            shape.fill.background()
            shape.line.color.rgb = RGBColor(100, 100, 100)
            shape.line.width = Pt(1)

    path = _save(tmp_path, build)
    report = inspect_bespoke_visual_quality(path)

    assert report["status"] == "fail"
    finding = next(item for item in report["findings"] if item["code"] == "BESPOKE_OUTLINE_DENSITY_HIGH")
    assert finding["slide_index"] == 1
    assert finding["evidence"]["outlined_containers"] == 9


def test_rejects_pie_without_explicit_distinct_point_colors(tmp_path):
    def build(slide):
        _bound_text(slide)
        data = ChartData()
        data.categories = ["A", "B", "C"]
        data.add_series("Share", (40, 35, 25))
        shape = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE,
            Inches(1),
            Inches(2),
            Inches(5),
            Inches(3),
            data,
        )
        shape.name = "bind:chart:s1"
        series = shape.chart.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor(201, 52, 61)

    path = _save(tmp_path, build)
    report = inspect_bespoke_visual_quality(path)

    assert report["status"] == "fail"
    assert any(item["code"] == "BESPOKE_PIE_POINT_COLORS_REQUIRED" for item in report["findings"])


def test_accepts_readable_wrapped_text_and_portable_pie(tmp_path):
    def build(slide):
        _bound_text(slide, size=14, word_wrap=True)
        data = ChartData()
        data.categories = ["A", "B", "C"]
        data.add_series("Share", (40, 35, 25))
        shape = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE,
            Inches(1),
            Inches(2),
            Inches(5),
            Inches(3),
            data,
        )
        shape.name = "bind:chart:s1"
        for point, color in zip(
            shape.chart.series[0].points,
            (RGBColor(201, 52, 61), RGBColor(23, 42, 58), RGBColor(100, 114, 126)),
        ):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = color

    path = _save(tmp_path, build)
    report = inspect_bespoke_visual_quality(path)

    assert report == {"status": "pass", "findings": []}


def test_template_delivery_can_raise_bound_body_text_without_enlarging_page_numbers(tmp_path):
    def build(slide):
        body = _bound_text(slide, size=9)
        page = slide.shapes.add_textbox(Inches(8), Inches(7), Inches(0.5), Inches(0.2))
        page.name = "bind:block:s1:page_no:value"
        run = page.text_frame.paragraphs[0].add_run()
        run.text = "1"
        run.font.size = Pt(9)

    path = _save(tmp_path, build)

    normalization = enforce_bound_text_floor(path)
    report = inspect_bespoke_visual_quality(path)

    assert normalization["adjusted_run_count"] == 1
    assert report["status"] == "pass"
