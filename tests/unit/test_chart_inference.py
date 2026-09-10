"""P10 tests: chart auto-inference from metric blocks and table data.

Verifies the engine can generate chart payloads from slide content alone,
so the model never needs to write chart/diagram_ir structures.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from engine.chart_inference import (
    _infer_chart_data,
    _infer_chart_type,
    _parse_metric_value,
    infer_charts,
)
from engine.structural_parser import parse_source


class TestParseMetricValue:
    def test_percentage(self):
        assert _parse_metric_value("36%") == 36.0
        assert _parse_metric_value("+36%") == 36.0
        assert _parse_metric_value("-5%") == -5.0

    def test_currency(self):
        assert _parse_metric_value("$49") == 49.0
        assert _parse_metric_value("¥2000") == 2000.0

    def test_chinese_units(self):
        assert _parse_metric_value("720万元") == 720.0
        assert _parse_metric_value("1.2亿") == 1.2

    def test_plain_numbers(self):
        assert _parse_metric_value("2000") == 2000.0
        assert _parse_metric_value("1,200") == 1200.0

    def test_non_numeric(self):
        assert _parse_metric_value("N/A") is None
        assert _parse_metric_value("") is None
        assert _parse_metric_value(None) is None

    def test_already_number(self):
        assert _parse_metric_value(42) == 42.0
        assert _parse_metric_value(3.14) == 3.14


class TestInferChartType:
    def test_comparison(self):
        slide = {"relations": [{"relation": "comparison", "between": ["a", "b"]}]}
        assert _infer_chart_type(slide) == "combo"

    def test_sequence(self):
        slide = {"relations": [{"relation": "sequence", "between": ["a", "b"]}]}
        assert _infer_chart_type(slide) == "line"

    def test_default(self):
        assert _infer_chart_type({}) == "column"
        assert _infer_chart_type({"relations": []}) == "column"


class TestInferChartData:
    def test_two_metrics(self):
        metrics = [
            {"label": "订阅收入", "value": "+31%"},
            {"label": "服务收入", "value": "+9%"},
        ]
        data = _infer_chart_data(metrics)
        assert data is not None
        assert data["categories"] == ["订阅收入", "服务收入"]
        assert data["series"][0]["values"] == [31.0, 9.0]

    def test_single_metric_rejected(self):
        data = _infer_chart_data([{"label": "收入", "value": "+36%"}])
        assert data is None

    def test_mixed_parseable_rejected(self):
        """If ANY metric can't be parsed, the whole chart is skipped."""
        data = _infer_chart_data([
            {"label": "收入", "value": "+36%"},
            {"label": "利润", "value": "N/A"},
        ])
        assert data is None

    def test_three_metrics(self):
        metrics = [
            {"label": "Q1", "value": "100"},
            {"label": "Q2", "value": "150"},
            {"label": "Q3", "value": "120"},
        ]
        data = _infer_chart_data(metrics)
        assert data is not None
        assert len(data["categories"]) == 3
        assert data["series"][0]["values"] == [100.0, 150.0, 120.0]


class TestInferChartsIntegration:
    def test_metric_blocks_trigger_inference(self):
        """Model-provided IR with metric blocks → chart auto-injected."""
        ir = {
            "deck": {"title": "Test"},
            "slides": [
                {"id": "s1", "title": "Revenue",
                 "blocks": [
                     {"role": "metric", "label": "Q1", "value": "100"},
                     {"role": "metric", "label": "Q2", "value": "200"},
                 ]}
            ]
        }
        infer_charts(ir, {})
        assert ir["slides"][0].get("chart") is not None
        assert ir["slides"][0]["chart"]["type"] == "column"
        assert ir["slides"][0]["chart"]["inferred"] is True

    def test_explicit_chart_not_overridden(self):
        """Model-provided chart is never overridden by inference."""
        ir = {
            "deck": {"title": "Test"},
            "slides": [
                {"id": "s1", "title": "Revenue",
                 "chart": {"type": "pie", "data": {"categories": [], "series": []}},
                 "blocks": [
                     {"role": "metric", "label": "Q1", "value": "100"},
                     {"role": "metric", "label": "Q2", "value": "200"},
                 ]}
            ]
        }
        infer_charts(ir, {})
        assert ir["slides"][0]["chart"]["type"] == "pie"  # model's choice kept
        assert not ir["slides"][0]["chart"].get("inferred")

    def test_single_metric_no_inference(self):
        ir = {
            "deck": {"title": "Test"},
            "slides": [
                {"id": "s1", "title": "Revenue",
                 "blocks": [{"role": "metric", "label": "Q1", "value": "100"}]}
            ]
        }
        infer_charts(ir, {})
        assert ir["slides"][0].get("chart") is None

    def test_table_block_triggers_inference(self):
        """Table block with source data → chart from table rows."""
        doc = parse_source(
            "| 版本 | 月费 | 目标 |\n| --- | --- | --- |\n"
            "| 基础版 | 49 | 2000 |\n| 专业版 | 199 | 600 |\n"
            "| 企业版 | 899 | 80 |\n",
            "markdown", "doc")
        ir = {
            "deck": {"title": "Test"},
            "slides": [
                {"id": "s1", "title": "Pricing",
                 "blocks": [
                     {"role": "table", "source_ref": {"source_id": "doc", "loc": "table_1"}}
                 ]}
            ]
        }
        infer_charts(ir, {"doc": doc})
        chart = ir["slides"][0].get("chart")
        assert chart is not None
        assert chart["type"] == "column"
        assert chart["inferred"] is True
        assert len(chart["data"]["categories"]) == 3

    def test_compiles_with_inferred_chart(self):
        """End-to-end: compile with inferred chart produces a deck with a chart."""
        from engine.compile import compile_deck
        import tempfile, os
        src = ROOT / "tests" / "fixtures" / "v4-bench" / "product-launch" / "source.md"
        if not src.exists():
            pytest.skip("product-launch fixture missing")
        with tempfile.TemporaryDirectory() as td:
            result = compile_deck(
                sources=[("doc", "markdown", str(src))],
                ir_path=None, degrade_on_error=False, output_dir=td)
            assert result["ok"], result.get("errors")
            # verify a chart was rendered
            from pptx import Presentation
            prs = Presentation(os.path.join(td, "deck.pptx"))
            chart_count = 0
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "chart"):
                        chart_count += 1
            assert chart_count >= 1, "Expected at least one chart in the deck"