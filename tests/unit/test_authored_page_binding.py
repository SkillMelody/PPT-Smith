from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from engine.authored_page import verify_authored_page
from engine.structural_parser import parse_markdown


def _ir() -> dict:
    return {
        "schema_version": "4.1.0",
        "deck": {"title": "Q3 review", "language": "en-US"},
        "sources": [{"source_id": "doc", "type": "markdown", "path": "source.md"}],
        "slides": [{
            "id": "growth", "title": "Growth accelerated",
            "message": "Growth accelerated in Q3.",
            "blocks": [
                {"id": "m_growth", "role": "metric", "label": "Q3 growth", "value": "24%",
                 "source_ref": {"source_id": "doc", "loc": "para_1"}},
                {"id": "f_driver", "role": "fact", "text": "New products drove most of the growth.",
                 "source_ref": {"source_id": "doc", "loc": "para_1"}},
            ],
        }],
    }


def _authored(path: Path, *, texts: list[tuple[str, str | None]]) -> None:
    """Build a model-authored page. Each text carries an optional binding name."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, (text, binding) in enumerate(texts):
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.6 + index * 0.9), Inches(9), Inches(0.6))
        box.text = text
        if binding:
            box.name = binding
    prs.save(path)


def _chart_authored(path: Path, *, values: list[float]) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Pilot", "Scale"]
    data.add_series("2025 share", values)
    graphic = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.2), Inches(8), Inches(4), data,
    )
    graphic.name = "bind:chart:growth"
    prs.save(path)


def _multi_chart_authored(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, (chart_id, values) in enumerate((("innovation", (64, 36)), ("ebit", (39, 61)))):
        data = CategoryChartData()
        data.categories = ["Reported impact", "No reported impact"]
        data.add_series("Share", values)
        graphic = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            Inches(0.8 + index * 4.2), Inches(1.2), Inches(3.4), Inches(3.4), data,
        )
        graphic.name = f"bind:chart:impact:{chart_id}"
    prs.save(path)


def test_authored_page_accepts_declared_ir_bindings(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Growth accelerated in Q3.", "bind:slide:growth:message"),
        ("24%", "bind:block:growth:m_growth:value"),
        ("Q3 growth", "bind:block:growth:m_growth:label"),
        ("New products drove most of the growth.", "bind:block:growth:f_driver:text"),
        ("Source: doc para_1", "bind:source:growth:f_driver"),
    ])
    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])
    assert report["status"] == "pass"
    assert report["unbound_text"] == []
    assert report["invalid_bindings"] == []
    assert report["mismatched_text"] == []


def test_authored_page_rejects_unbound_business_text(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Growth accelerated in Q3.", "bind:slide:growth:message"),
        ("High performers always win", None),
    ])
    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])
    assert report["status"] == "fail"
    assert report["unbound_text"] == ["High performers always win"]


def test_authored_page_rejects_binding_to_unknown_block(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Growth accelerated in Q3.", "bind:slide:growth:message"),
        ("48%", "bind:block:growth:m_invented:value"),
    ])
    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])
    assert report["status"] == "fail"
    assert any(item["binding"] == "bind:block:growth:m_invented:value" for item in report["invalid_bindings"])


def test_authored_page_rejects_text_not_matching_bound_value(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Growth accelerated in Q3.", "bind:slide:growth:message"),
        ("48%", "bind:block:growth:m_growth:value"),
    ])
    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])
    assert report["status"] == "fail"
    assert any(item["binding"] == "bind:block:growth:m_growth:value" for item in report["mismatched_text"])


def test_authored_page_reports_missing_required_ir_content(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[("Growth accelerated in Q3.", "bind:slide:growth:message")])
    report = verify_authored_page(deck, _ir(), slide_ids=["growth"], require_full_coverage=True)
    assert report["status"] == "fail"
    assert "slide:growth:title" in report["uncovered_content"]
    assert "slide:growth:message" not in report["uncovered_content"]
    assert "growth:m_growth" in report["uncovered_content"]
    assert "growth:f_driver" in report["uncovered_content"]


def test_authored_page_allows_pure_decoration_without_binding(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[("Growth accelerated in Q3.", "bind:slide:growth:message")])
    prs = Presentation(str(deck))
    slide = prs.slides[0]
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(6.6), Inches(3), Inches(0.05))
    shape.name = "decoration:accent_rule"
    prs.save(deck)
    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])
    assert report["status"] == "pass"


def test_authored_page_rejects_business_text_disguised_as_decoration(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Growth accelerated in Q3.", "bind:slide:growth:message"),
        ("Invented market leadership claim", "decoration:fake_badge"),
    ])

    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])

    assert report["status"] == "fail"
    assert report["invalid_bindings"] == [{
        "binding": "decoration:fake_badge",
        "reason": "decoration shapes may not contain text",
    }]


def test_authored_page_rejects_unbound_text_nested_in_group_shape(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[])
    prs = Presentation(str(deck))
    group = prs.slides[0].shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(6), Inches(0.5))
    box.text = "Invented claim hidden in group"
    prs.save(deck)

    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])

    assert report["status"] == "fail"
    assert report["unbound_text"] == ["Invented claim hidden in group"]


def test_authored_page_verifies_visible_source_text(tmp_path: Path) -> None:
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Growth accelerated in Q3.", "bind:slide:growth:message"),
        ("Source: invented report page_99", "bind:source:growth:f_driver"),
    ])

    report = verify_authored_page(deck, _ir(), slide_ids=["growth"])

    assert report["status"] == "fail"
    assert report["mismatched_text"] == [{
        "binding": "bind:source:growth:f_driver",
        "expected": ["Source: doc para_1"],
        "actual": "Source: invented report page_99",
    }]


def test_authored_page_accepts_value_with_source_declared_unit(tmp_path: Path) -> None:
    ir = _ir()
    metric = ir["slides"][0]["blocks"][0]
    metric["value"] = "3.6"
    metric["unit"] = "x"
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("3.6x", "bind:block:growth:m_growth:display_value"),
    ])

    report = verify_authored_page(deck, ir, slide_ids=["growth"])

    assert report["status"] == "pass"
    assert report["mismatched_text"] == []


def test_authored_page_rejects_display_unit_not_declared_in_ir(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"][0]["value"] = "3.6"
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("3.6x", "bind:block:growth:m_growth:display_value"),
    ])

    report = verify_authored_page(deck, ir, slide_ids=["growth"])

    assert report["status"] == "fail"
    assert report["mismatched_text"][0]["expected"] == ["3.6"]


def test_authored_page_accepts_indexed_list_item_binding_from_manifest(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"].append({
        "id": "actions", "role": "list", "semantics": "steps",
        "items": [
            {"text": "Confirm scope", "detail": "Lock the decision question."},
            {"text": "Ship pilot", "detail": "Validate with a real audience."},
        ],
    })
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Ship pilot", "bind:block:growth:actions:item:1"),
        ("Validate with a real audience.", "bind:block:growth:actions:detail:1"),
    ])

    report = verify_authored_page(deck, ir, slide_ids=["growth"])

    assert report["status"] == "pass"


def test_authored_page_rejects_unknown_list_item_index(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"].append({
        "id": "actions", "role": "list",
        "items": [{"text": "Confirm scope"}],
    })
    deck = tmp_path / "page.pptx"
    _authored(deck, texts=[
        ("Confirm scope", "bind:block:growth:actions:item:4"),
    ])

    report = verify_authored_page(deck, ir, slide_ids=["growth"])

    assert report["status"] == "fail"
    assert report["invalid_bindings"] == [{
        "binding": "bind:block:growth:actions:item:4",
        "reason": "item index is out of range",
    }]


def test_authored_page_accepts_native_chart_with_source_locked_categories_and_series(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["chart"] = {
        "type": "column",
        "data": {
            "categories": ["Pilot", "Scale"],
            "series": [{"name": "2025 share", "values": [30, 70]}],
        },
        "source_ref": {"source_id": "doc", "loc": "table_1"},
    }
    deck = tmp_path / "chart.pptx"
    _chart_authored(deck, values=[30, 70])

    report = verify_authored_page(deck, ir, slide_ids=["growth"], require_full_coverage=True)

    assert report["status"] == "fail"  # title/message/blocks remain deliberately uncovered
    assert "chart:growth" not in report["uncovered_content"]
    assert "chart:growth" in report["valid_bindings"]


def test_authored_page_rejects_native_chart_with_changed_series_value(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["chart"] = {
        "type": "column",
        "data": {
            "categories": ["Pilot", "Scale"],
            "series": [{"name": "2025 share", "values": [30, 70]}],
        },
        "source_ref": {"source_id": "doc", "loc": "table_1"},
    }
    deck = tmp_path / "chart.pptx"
    _chart_authored(deck, values=[30, 71])

    report = verify_authored_page(deck, ir, slide_ids=["growth"])

    assert report["status"] == "fail"
    assert report["mismatched_text"][0]["binding"] == "bind:chart:growth"


def test_authored_page_verifies_multiple_named_charts_on_one_slide(tmp_path: Path) -> None:
    ir = {
        "slides": [{
            "id": "impact",
            "charts": [{
                "id": "innovation",
                "data": {
                    "categories": ["Reported impact", "No reported impact"],
                    "series": [{"name": "Share", "values": [64, 36]}],
                },
            }, {
                "id": "ebit",
                "data": {
                    "categories": ["Reported impact", "No reported impact"],
                    "series": [{"name": "Share", "values": [39, 61]}],
                },
            }],
        }],
    }
    deck = tmp_path / "multi-chart.pptx"
    _multi_chart_authored(deck)

    report = verify_authored_page(
        deck, ir, slide_ids=["impact"], require_full_coverage=True,
    )

    assert report["status"] == "pass"
    assert report["valid_bindings"] == [
        "chart:impact:ebit",
        "chart:impact:innovation",
    ]
    assert report["uncovered_content"] == []


def test_authored_page_accepts_native_table_only_when_cells_match_source_anchor(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"].append({
        "id": "regional", "role": "table",
        "source_ref": {"source_id": "doc", "loc": "table_1"},
        "column_filter": [0, 2], "row_limit": 1,
    })
    docs = {"doc": parse_markdown("| Region | Revenue | Growth |\n| --- | --- | --- |\n| US | 10 | 20% |\n| EU | 8 | 15% |", "doc")}
    deck = tmp_path / "table.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(0.8), Inches(1), Inches(8), Inches(2))
    table_shape.name = "bind:table:growth:regional"
    for row_index, row in enumerate([["Region", "Growth"], ["US", "20%"]]):
        for column_index, value in enumerate(row):
            table_shape.table.cell(row_index, column_index).text = value
    prs.save(deck)

    report = verify_authored_page(deck, ir, slide_ids=["growth"], source_docs=docs)

    assert report["status"] == "pass"
    assert "table:growth:regional" in report["valid_bindings"]


def test_authored_page_accepts_a_source_table_row_window(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"].append({
        "id": "regional", "role": "table",
        "source_ref": {"source_id": "doc", "loc": "table_1"},
        "row_offset": 1, "row_limit": 1,
    })
    docs = {"doc": parse_markdown(
        "| Region | Revenue |\n| --- | --- |\n| US | 10 |\n| EU | 8 |",
        "doc",
    )}
    deck = tmp_path / "table-window.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(0.8), Inches(1), Inches(8), Inches(2))
    table_shape.name = "bind:table:growth:regional"
    for row_index, row in enumerate([["Region", "Revenue"], ["EU", "8"]]):
        for column_index, value in enumerate(row):
            table_shape.table.cell(row_index, column_index).text = value
    prs.save(deck)

    report = verify_authored_page(deck, ir, slide_ids=["growth"], source_docs=docs)

    assert report["status"] == "pass"
    assert "table:growth:regional" in report["valid_bindings"]


def test_authored_page_rejects_native_table_with_changed_source_cell(tmp_path: Path) -> None:
    ir = _ir()
    ir["slides"][0]["blocks"].append({
        "id": "regional", "role": "table",
        "source_ref": {"source_id": "doc", "loc": "table_1"},
    })
    docs = {"doc": parse_markdown("| Region | Revenue |\n| --- | --- |\n| US | 10 |", "doc")}
    deck = tmp_path / "table.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(0.8), Inches(1), Inches(8), Inches(2))
    table_shape.name = "bind:table:growth:regional"
    for row_index, row in enumerate([["Region", "Revenue"], ["US", "11"]]):
        for column_index, value in enumerate(row):
            table_shape.table.cell(row_index, column_index).text = value
    prs.save(deck)

    report = verify_authored_page(deck, ir, slide_ids=["growth"], source_docs=docs)

    assert report["status"] == "fail"
    assert report["mismatched_text"][0]["binding"] == "bind:table:growth:regional"
