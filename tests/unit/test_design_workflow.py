"""User-visible create, compact authoring, patch and review-inheritance contracts."""
import copy
import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation

from engine.design_scene.authoring import apply_updates, compact, expand
from engine.design_scene.backend import inspect_objects, render
from engine.design_scene.cli import build_summary, main, route
from engine.design_scene.contract import validate, versions
from engine.design_scene.incremental import prepare_inheritance
from engine.design_scene.pipeline import assess, build, deliver, new_task, read_assets, review_skeleton
from engine.design_scene.store import DesignError, Store, canonical, digest


def original_task():
    t = new_task("original", mode="create", pages=2, width=960, height=540, author="actual-author")
    t["design"]["brief"] = "A task-specific narrative, high contrast and generous whitespace."
    for i, page in enumerate(t["scene"]["pages"], 1):
        pid = page["id"]
        t["content"]["items"].append({
            "id": "text-" + str(i), "page_id": pid, "type": "text", "text": "Evidence " + str(i),
            "required_edit": "native",
            "provenance": {"kind": "document", "source": "Source report p1", "checked": True, "uncertainties": []}})
        t["content"]["notes"].append({"page_id": pid, "text": "Speaker evidence and qualifications, source page 1."})
        page["nodes"] = [{"id": "title-" + str(i), "page_id": pid, "type": "text", "role": "native_text",
                          "z": 0, "box": {"x": 40, "y": 50, "width": 800, "height": 100},
                          "content_id": "text-" + str(i), "font": {"family": "Arial", "size": 28, "color": "162C46"},
                          "align": "left", "valign": "top", "wrap": True}]
    return t


def mock_worker(stage, **kwargs):
    # Contract-only fixture. Actual rendering is verified in separate real runs.
    with Store(stage) as s:
        t = s.json("task.json")
        assets = {a["id"]: s.read("assets/" + a["sha256"] + ".png") for a in t["assets"]["items"]}
        data = render(t, assets)
        s.write("deck.pptx", data)
        s.put_json("objects.json", inspect_objects(data, t))
        pages = []
        for i, page in enumerate(t["scene"]["pages"], 1):
            b = io.BytesIO()
            text = [c for c in t["content"]["items"] if c["page_id"] == page["id"]]
            colour = "#" + digest(canonical([page, text]))[:6]
            Image.new("RGB", (960, 540), colour).save(b, "PNG")
            name = "preview-" + str(i) + ".png"
            s.write(name, b.getvalue())
            pages.append({"page_id": page["id"], "file": name, "sha256": digest(b.getvalue())})
        s.write("deck.pdf", b"UNIT TEST STUB, NOT REAL RENDER EVIDENCE")
        s.put_json("render.json", {"status": "passed", "engine": "unit-test-stub", "pages": pages})


@pytest.fixture
def original(tmp_path, monkeypatch):
    root = tmp_path / "task"
    with Store(root, create=True) as s:
        s.put_json("task.json", original_task())
    monkeypatch.setattr("engine.design_scene.pipeline.launch_worker", mock_worker)
    monkeypatch.setattr("engine.design_scene.pipeline.fonts_for", lambda t: [])
    return root


def reviewed(s, bid):
    r = review_skeleton(s, bid)
    r["reviewer"] = {"id": "independent-fixture", "method": "human", "environment": "Unit fixture; not real review"}
    for page in r["pages"]:
        for key in ("design", "content", "editability", "fonts"):
            page[key] = "pass"
        if page["reconstruction"] != "not_applicable":
            page["reconstruction"] = "pass"
        page["observations"] = ["Unit fixture for acceptance state transitions only."]
        n = int(page["page_id"].split("-")[-1])
        page["observed_node_ids"] = ["title-" + str(n)]
    r["target_software"] = {"name": "LibreOffice", "version": "unit-fixture", "status": "pass",
                            "actions": ["Unit fixture; actual edit evidence belongs to real runs."]}
    return r


def test_create_has_native_objects_and_can_deliver_without_image_target(original):
    m = build(original, isolation="host")
    assert m["status"] == "candidate_unreviewed"
    assert m["comparisons"] == []
    assert {d["kind"] for d in m["design_origins"]} == {"programmatic_preview"}
    with Store(original) as s:
        review = reviewed(s, m["build_id"])
        assert all(p["reconstruction"] == "not_applicable" for p in review["pages"])
        result = deliver(s, m["build_id"], review)
        with ZipFile(io.BytesIO(s.read(result["delivery"]))) as archive:
            pptx = archive.read("deck.pptx")
            assert len(Presentation(io.BytesIO(pptx)).slides) == 2
            assert b"Evidence" in archive.read("task.json")
    assert result["status"] == "final_delivery_ready"


@pytest.mark.parametrize("mode", ["recreate", "new_design", "style_transfer"])
def test_saved_target_routes_keep_their_target_requirement(mode):
    t = original_task()
    t["mode"] = mode
    with pytest.raises(DesignError, match="DESIGN_TARGET_REQUIRED"):
        validate(t, complete=True)


def test_create_cannot_claim_self_reconstruction_or_skip_notes(original):
    t = original_task()
    t["content"]["notes"] = []
    with pytest.raises(DesignError, match="SPEAKER_NOTES"):
        validate(t, complete=True)
    m = build(original, isolation="host")
    with Store(original) as s:
        r = reviewed(s, m["build_id"])
        r["pages"][0]["reconstruction"] = "pass"
        assert "RECONSTRUCTION_REVIEW_REQUIRED:page-1" in assess(s, m["build_id"], r)["blockers"]


def test_three_intents_and_legacy_init_alias(tmp_path, capsys):
    assert route()["intent"] == "create"
    assert route(template=True)["intent"] == "template"
    assert route(design_target=True)["intent"] == "recreate"
    with pytest.raises(DesignError, match="ROUTE_CONFLICT"):
        route(template=True, design_target=True)
    root = tmp_path / "alias"
    assert main(["init", "--task-dir", str(root), "--task-id", "alias", "--mode", "new_design"]) == 0
    capsys.readouterr()
    with Store(root) as s:
        assert s.json("task.json")["mode"] == "create"
    assert main(["init", "--task-dir", str(root), "--task-id", "different"]) == 1
    with Store(root) as s:
        assert s.json("task.json")["task_id"] == "alias"


def test_compact_roundtrip_preserves_text_notes_data_and_objects():
    task = original_task()
    task["content"]["items"][0]["text"] = "$literal business text"
    result = compact(task)
    assert canonical(expand(result)) == canonical(task)
    assert len(canonical(result)) < len(canonical(task))
    assert "source_ref" in result["task"]["content"]["items"][0]


def test_symbol_instances_expand_as_separately_editable_native_shapes():
    document = compact(original_task())
    document["tokens"] = {"accent": "11AABB"}
    document["symbols"] = {"dot": {"width": 20, "height": 20, "nodes": [
        {"id": "circle", "type": "shape", "shape": "ellipse",
         "box": {"x": 0, "y": 0, "width": 20, "height": 20}, "style": {"fill": "$accent"}}]}}
    nodes = document["task"]["scene"]["pages"][0]["nodes"]
    nodes.extend([{"id": "mark-" + str(i), "symbol_ref": "dot",
                   "box": {"x": 100 + i * 50, "y": 220, "width": 40, "height": 40}} for i in range(2)])
    t = expand(document)
    p = Presentation(io.BytesIO(render(t, {})))
    assert p.slides[0].shapes[1].shapes[0].name == "mark-0_circle"
    assert tuple(p.slides[0].shapes[2].shapes[0].fill.fore_color.rgb) == (17, 170, 187)
    nodes[-1]["box"]["width"] = 60
    with pytest.raises(DesignError, match="SYMBOL_ASPECT"):
        expand(document)


@pytest.mark.parametrize("mutation", [
    lambda d: d["styles"].update(bad={"script": "arbitrary code"}),
    lambda d: d["task"]["scene"]["pages"][0]["nodes"][0].update(expression="eval(1)"),
    lambda d: d["task"]["content"]["items"][0].update(source_ref="nonexistent"),
    lambda d: d.update(tokens={"recursive": "$recursive"}),
])
def test_compact_input_cannot_bypass_strict_execution(mutation):
    d = compact(original_task())
    mutation(d)
    with pytest.raises(DesignError):
        expand(d)


def test_atomic_patch_guard_and_original_task_preservation(original, capsys):
    with Store(original) as s:
        old = s.json("task.json")
        patch = {"base_task_sha256": versions(old)["task"],
                 "updates": [{"target": "node", "id": "title-1", "changes": {"font": {"size": 25}}},
                             {"target": "node", "id": "missing", "changes": {"font": {"size": 24}}}]}
        s.put_json("patch.json", patch)
    assert main(["patch", "--task-dir", str(original), "--file", "patch.json"]) == 1
    capsys.readouterr()
    with Store(original) as s:
        assert s.json("task.json") == old
        patch["updates"].pop()
        s.put_json("patch.json", patch)
    assert main(["patch", "--task-dir", str(original), "--file", "patch.json"]) == 0
    with Store(original) as s:
        updated = s.json("task.json")
        assert updated["scene"]["pages"][0]["nodes"][0]["font"]["family"] == "Arial"
        assert updated["scene"]["pages"][0]["nodes"][0]["font"]["size"] == 25
        assert s.json("history/" + versions(old)["task"] + ".json") == old
    assert main(["patch", "--task-dir", str(original), "--file", "patch.json"]) == 1
    assert "PATCH_STALE_INPUT" in capsys.readouterr().out


def test_summary_does_not_dump_manifest_but_details_remain(original, capsys):
    assert main(["build", "--task-dir", str(original), "--isolation", "host"]) == 0
    output = capsys.readouterr().out
    result = json.loads(output)
    assert len(output.encode()) < 2048
    assert "object_inspection" not in result
    with Store(original) as s:
        full = s.json(result["manifest"])
        assert len(full["object_inspection"]["pages"]) == 2
        full["status"] = "render_incomplete"
        full["render"].update(status="failed", reason="PROCESS_FAILED: office could not render this deck")
        assert "office could not render" in build_summary(full)["error"]
    assert main(["inspect", "--task-dir", str(original), "--build-id", result["build_id"], "--page", "page-2"]) == 0
    detail = json.loads(capsys.readouterr().out)
    assert detail["objects"]["page_id"] == "page-2"


def test_unchanged_review_inherits_but_changed_notes_do_not(original):
    first = build(original, isolation="host")
    with Store(original) as s:
        previous = reviewed(s, first["build_id"])
        t = s.json("task.json")
        t["content"]["notes"][0]["text"] += "Changed qualification."
        s.put_json("task.json", t)
    second = build(original, isolation="host", parent_build=first["build_id"])
    # The pixels did not change; page 1 must still be freshly reviewed.
    assert second["changed_pages"] == ["page-1"]
    with Store(original) as s:
        assert first["artifacts"]["preview-1.png"] == second["artifacts"]["preview-1.png"]
        draft = review_skeleton(s, second["build_id"])
        assert prepare_inheritance(s, t, second, draft, previous) == ["page-2"]
        assert draft["pages"][0]["content"] == "revise"
        assert draft["pages"][1]["content"] == "pass"
        fresh = reviewed(s, second["build_id"])
        fresh["pages"][1] = draft["pages"][1]
        assert assess(s, second["build_id"], fresh)["status"] == "final_delivery_ready"
        result = deliver(s, second["build_id"], fresh)
        with ZipFile(io.BytesIO(s.read(result["delivery"]))) as archive:
            assert any(n.startswith("review-evidence/") for n in archive.namelist())
        fresh["pages"][1]["observations"] = ["Pretend to have a new observation"]
        with pytest.raises(DesignError, match="INHERITED_REVIEW_ROW_CHANGED"):
            assess(s, second["build_id"], fresh)


def test_review_and_delivery_keep_resumable_state_and_diagnose_report_input(original, capsys):
    assert main(["build", "--task-dir", str(original), "--isolation", "host"]) == 0
    bid = json.loads(capsys.readouterr().out)["build_id"]
    with Store(original) as s:
        s.put_json("unfinished.json", review_skeleton(s, bid))
        s.put_json("independent.json", reviewed(s, bid))
    assert main(["review", "--task-dir", str(original), "--build-id", bid,
                 "--review-file", "unfinished.json"]) == 1
    assert "pages/0/observations" in capsys.readouterr().out
    assert main(["review", "--task-dir", str(original), "--build-id", bid,
                 "--review-file", "independent.json"]) == 0
    assessment = json.loads(capsys.readouterr().out)
    with Store(original) as s:
        state = s.json("state.json")
        assert state["build_id"] == bid
        assert state["review_file"] == "independent.json"
        assert state["report"] == assessment["report"]
    assert main(["review-template", "--task-dir", str(original), "--build-id", bid,
                 "--inherit-from", assessment["report"], "--output", "should-not-exist.json"]) == 1
    assert "REVIEW_INPUT_MUST_BE_RAW" in capsys.readouterr().out
    assert not (original / "should-not-exist.json").exists()
    assert main(["deliver", "--task-dir", str(original), "--build-id", bid,
                 "--review-file", "independent.json"]) == 0
    delivery = json.loads(capsys.readouterr().out)
    with Store(original) as s:
        state = s.json("state.json")
        assert state["phase"] == "final_delivery_ready"
        assert state["build_id"] == bid
        assert state["delivery"] == delivery["delivery"]
        assert s.read(state["delivery"])


def test_failed_other_page_does_not_discard_valid_independent_page(original):
    first = build(original, isolation="host")
    with Store(original) as s:
        previous = reviewed(s, first["build_id"])
        previous["pages"][0]["design"] = "revise"
        t = s.json("task.json")
        t["scene"]["pages"][0]["nodes"][0]["font"]["size"] = 24
        s.put_json("task.json", t)
    second = build(original, isolation="host", parent_build=first["build_id"])
    with Store(original) as s:
        draft = review_skeleton(s, second["build_id"])
        assert prepare_inheritance(s, t, second, draft, previous) == ["page-2"]
        previous["target_software"]["status"] = "not_tested"
        draft = review_skeleton(s, second["build_id"])
        assert prepare_inheritance(s, t, second, draft, previous) == []


def test_global_design_change_and_inherited_source_tamper_are_rejected(original):
    first = build(original, isolation="host")
    with Store(original) as s:
        previous = reviewed(s, first["build_id"])
        t = s.json("task.json")
        t["design"]["brief"] += "New audience emphasis."
        s.put_json("task.json", t)
    second = build(original, isolation="host", parent_build=first["build_id"])
    assert set(second["changed_pages"]) == {"page-1", "page-2"}
    with Store(original) as s:
        draft = review_skeleton(s, second["build_id"])
        assert prepare_inheritance(s, t, second, draft, previous) == []
    third = build(original, isolation="host", parent_build=second["build_id"])
    with Store(original) as s:
        previous = reviewed(s, second["build_id"])
        draft = reviewed(s, third["build_id"])
        assert len(prepare_inheritance(s, t, third, draft, previous)) == 2
        path = "review-evidence/" + digest(canonical(previous)) + ".json"
        s.put_json(path, {**previous, "reviewer": {"id": "forged", "method": "human", "environment": "changed"}})
        with pytest.raises(DesignError, match="INHERITED_REVIEW_CHANGED"):
            assess(s, third["build_id"], draft)
