from __future__ import annotations

import importlib
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _four_slide_deck(path: Path) -> None:
    presentation = Presentation()
    for slide_number in range(1, 5):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.5))
        title.name = f"title-{slide_number}"
        title.text = f"Slide {slide_number}"
        if slide_number == 4:
            data = ChartData()
            data.categories = ["完成", "剩余"]
            data.add_series("进度", (63, 37))
            chart = slide.shapes.add_chart(
                XL_CHART_TYPE.DOUGHNUT,
                Inches(2), Inches(1.5), Inches(4), Inches(4), data,
            )
            chart.name = "native-progress-ring"
    presentation.save(path)


def test_write_slide_subset_keeps_only_requested_native_slides(tmp_path: Path) -> None:
    try:
        module = importlib.import_module("engine.presentation_subset")
    except ModuleNotFoundError:
        pytest.fail("standalone PPTX subset support is missing")

    source = tmp_path / "engineering-full.pptx"
    output = tmp_path / "user-preview.pptx"
    _four_slide_deck(source)

    result = module.write_slide_subset(source, output, slide_numbers=[2, 4])

    preview = Presentation(output)
    assert result == {"source_slide_count": 4, "output_slide_count": 2, "slide_numbers": [2, 4]}
    assert len(preview.slides) == 2
    assert [slide.shapes[0].text for slide in preview.slides] == ["Slide 2", "Slide 4"]
    assert [shape.name for shape in preview.slides[1].shapes if getattr(shape, "has_chart", False)] == [
        "native-progress-ring"
    ]
    assert len(Presentation(source).slides) == 4


def test_write_slide_subset_rejects_unsorted_or_duplicate_slide_numbers(tmp_path: Path) -> None:
    try:
        module = importlib.import_module("engine.presentation_subset")
    except ModuleNotFoundError:
        pytest.fail("standalone PPTX subset support is missing")

    source = tmp_path / "engineering-full.pptx"
    _four_slide_deck(source)

    with pytest.raises(ValueError, match="strictly increasing"):
        module.write_slide_subset(source, tmp_path / "invalid.pptx", slide_numbers=[4, 2, 2])
