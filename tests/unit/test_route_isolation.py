from __future__ import annotations

import ast
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from ppt_qa.verifier import run_structural_inspection


ROOT = Path(__file__).resolve().parents[2]


def _engine_imports(module_name: str) -> set[str]:
    tree = ast.parse((ROOT / "engine" / f"{module_name}.py").read_text(encoding="utf-8"))
    imported: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module.rsplit(".", 1)[-1])
        elif isinstance(node, ast.Import):
            imported.update(alias.name.rsplit(".", 1)[-1] for alias in node.names)
    return imported


def _source_bound_fragments(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(8), Inches(0.5))
    title.name = "bind:slide:multi:title"
    title.text = "Structured labels"
    for index, text in enumerate(tuple("ABCDEFGHIJKL")):
        box = slide.shapes.add_textbox(
            Inches(0.7 + (index % 4) * 2.5),
            Inches(1.4 + (index // 4) * 2.0),
            Inches(1.8),
            Inches(0.5),
        )
        box.name = f"bind:block:multi:stages:item:{index}"
        box.text = text
    presentation.save(path)


def _issue_codes(report: dict) -> set[str]:
    return {issue["issue_code"] for issue in report["issues"]}


def test_component_fragmentation_exemption_is_template_route_only(tmp_path: Path) -> None:
    deck = tmp_path / "bound-fragments.pptx"
    _source_bound_fragments(deck)

    bespoke = run_structural_inspection(deck, route="bespoke")
    template = run_structural_inspection(deck, route="template")

    assert "PPTX_TEXT_FRAGMENTATION" in _issue_codes(bespoke)
    assert "PPTX_TEXT_FRAGMENTATION" not in _issue_codes(template)
    assert bespoke["qa_route"] == "bespoke"
    assert template["qa_route"] == "template"


def test_unknown_route_cannot_silently_inherit_a_qa_profile(tmp_path: Path) -> None:
    deck = tmp_path / "empty.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(deck)

    try:
        run_structural_inspection(deck, route="bespoke-template")
    except ValueError as exc:
        assert str(exc) == "UNKNOWN_PRODUCTION_ROUTE: bespoke-template"
    else:
        raise AssertionError("unknown production route was accepted")


def test_bespoke_and_template_authoring_modules_do_not_cross_import() -> None:
    template_only = {
        "strict_template", "template_native", "component_atlas", "component_composer",
        "component_inventory", "manuscript_component_planner", "manuscript_strict_preview",
        "presentation_subset", "template_page_composition", "template_visual_quality",
        "template_component_adaptation", "template_series_planner",
        "template_visual_fingerprint", "template_delivery_quality",
        "template_page_recipes",
    }
    bespoke_only = {"bespoke", "bespoke_runtime", "bespoke_quality", "visual_review"}

    for module in bespoke_only:
        assert not (_engine_imports(module) & template_only), module
    for module in template_only:
        assert not (_engine_imports(module) & bespoke_only), module


def test_template_page_composition_gates_are_not_imported_by_standard_or_bespoke() -> None:
    template_page_gates = {
        "template_page_composition", "template_visual_quality",
        "component_atlas_report", "component_intent",
        "template_component_adaptation", "template_series_planner",
        "template_visual_fingerprint", "template_delivery_quality",
        "template_page_recipes",
    }
    other_route_modules = {
        "compile", "layout", "policy", "extractive_ir",
        "bespoke", "bespoke_runtime", "bespoke_quality", "visual_review",
    }

    for module in other_route_modules:
        assert not (_engine_imports(module) & template_page_gates), module
