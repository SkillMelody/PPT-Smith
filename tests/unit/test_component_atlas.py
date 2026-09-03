from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys

import pytest
import engine.component_atlas as component_atlas_module
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from engine.component_atlas import (
    build_component_atlas,
    resolve_chart_component_binding,
    resolve_chart_dashboard_binding,
    resolve_component_binding,
    select_component,
)


ROOT = Path(__file__).resolve().parents[2]


def _component_template(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, width in enumerate((5.0, 4.0, 3.0), 1):
        segment = slide.shapes.add_shape(
            MSO_SHAPE.TRAPEZOID,
            Inches(2 + (5 - width) / 2), Inches(1 + index),
            Inches(width), Inches(0.7),
        )
        segment.name = f"funnel-segment-{index}"
        label = slide.shapes.add_textbox(
            Inches(3), Inches(1.15 + index), Inches(3), Inches(0.3),
        )
        label.name = f"funnel-label-{index}"
        label.text = f"Example {index}"
    presentation.save(path)


def _review() -> dict:
    return {
        "schema_version": "1.0.0",
        "review": {
            "state": "reviewed",
            "reviewer": {"method": "human", "id": "design-review"},
        },
        "components": [{
            "component_id": "funnel.primary",
            "family": "funnel",
            "slide_index": 1,
            "semantic_uses": ["stage narrowing", "conversion"],
            "groups": [
                {"role": "segment", "shape_names": [
                    "funnel-segment-1", "funnel-segment-2", "funnel-segment-3",
                ]},
                {"role": "label", "shape_names": [
                    "funnel-label-1", "funnel-label-2", "funnel-label-3",
                ]},
            ],
            "parameters": {
                "element_count": {"minimum": 1, "maximum": 8},
                "size_driver": "value",
                "direction": "descending",
            },
        }],
    }


def _extended_component_template(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, width in enumerate((2.0, 3.0, 4.0), 1):
        segment = slide.shapes.add_shape(
            MSO_SHAPE.TRAPEZOID,
            Inches(4 - width / 2), Inches(0.7 + index),
            Inches(width), Inches(0.7),
        )
        segment.name = f"pyramid-segment-{index}"
        title = slide.shapes.add_textbox(
            Inches(3), Inches(0.8 + index), Inches(2.0), Inches(0.25),
        )
        title.name = f"pyramid-title-{index}"
        value = slide.shapes.add_textbox(
            Inches(5), Inches(0.8 + index), Inches(1.0), Inches(0.25),
        )
        value.name = f"pyramid-value-{index}"
    for name, left, top in (
        ("pyramid-shadow-1-a", 3.0, 1.7),
        ("pyramid-shadow-1-b", 4.0, 1.7),
        ("pyramid-shadow-2", 3.0, 2.7),
        ("pyramid-shadow-3", 3.0, 3.7),
        ("pyramid-baseline", 2.0, 4.8),
    ):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(1.0), Inches(0.1),
        )
        shape.name = name
    presentation.save(path)


def _extended_review() -> dict:
    return {
        "schema_version": "1.0.0",
        "review": {
            "state": "reviewed",
            "reviewer": {"method": "human", "id": "extended-design-review"},
        },
        "components": [{
            "component_id": "pyramid.extended",
            "family": "pyramid",
            "slide_index": 1,
            "semantic_uses": ["hierarchy"],
            "groups": [
                {
                    "role": "segment",
                    "shape_names": [
                        "pyramid-segment-1", "pyramid-segment-2", "pyramid-segment-3",
                    ],
                    "scope": "item",
                    "item_indices": [0, 1, 2],
                },
                {
                    "role": "shadow",
                    "shape_names": [
                        "pyramid-shadow-1-a", "pyramid-shadow-1-b",
                        "pyramid-shadow-2", "pyramid-shadow-3",
                    ],
                    "scope": "item",
                    "item_indices": [0, 0, 1, 2],
                },
                {
                    "role": "label",
                    "bind_field": "title",
                    "shape_names": [
                        "pyramid-title-1", "pyramid-title-2", "pyramid-title-3",
                    ],
                    "scope": "item",
                    "item_indices": [0, 1, 2],
                },
                {
                    "role": "label",
                    "bind_field": "value",
                    "shape_names": [
                        "pyramid-value-1", "pyramid-value-2", "pyramid-value-3",
                    ],
                    "scope": "item",
                    "item_indices": [0, 1, 2],
                },
                {
                    "role": "attachment",
                    "shape_names": ["pyramid-baseline"],
                    "scope": "shared",
                },
            ],
            "parameters": {
                "element_count": {"minimum": 1, "maximum": 6},
                "size_driver": "value",
            },
        }],
    }


def test_component_atlas_resolves_reviewed_native_shapes_and_parameters(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)

    atlas = build_component_atlas(template, _review())
    component = atlas["components"][0]

    assert atlas["status"] == "reviewed"
    assert atlas["source"]["sha256"].startswith("sha256:")
    assert component["component_id"] == "funnel.primary"
    assert component["family"] == "funnel"
    assert component["parameters"]["element_count"] == {"minimum": 1, "maximum": 8}
    assert [group["role"] for group in component["groups"]] == ["segment", "label"]
    assert all(member["shape_id"] > 0 for group in component["groups"] for member in group["members"])
    assert all(member["frame"]["w"] > 0 for group in component["groups"] for member in group["members"])
    assert "Example" not in str(atlas)


def test_component_atlas_rejects_a_review_shape_missing_from_the_template(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    review = _review()
    review["components"][0]["groups"][0]["shape_names"][0] = "missing-segment"

    with pytest.raises(ValueError, match="missing-segment"):
        build_component_atlas(template, review)


def test_component_atlas_records_reviewed_equivalent_native_source_instances(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    presentation = Presentation(template)
    slide = presentation.slides[0]
    for index in range(1, 4):
        segment = slide.shapes.add_shape(
            MSO_SHAPE.TRAPEZOID,
            Inches(8), Inches(1 + index), Inches(2), Inches(0.7),
        )
        segment.name = f"funnel-alias-segment-{index}"
        label = slide.shapes.add_textbox(
            Inches(8), Inches(1.2 + index), Inches(2), Inches(0.3),
        )
        label.name = f"funnel-alias-label-{index}"
        label.text = f"Equivalent example {index}"
    presentation.save(template)
    review = _review()
    review["components"][0]["source_instances"] = [{
        "instance_id": "secondary-funnel-instance",
        "slide_index": 1,
        "groups": [
            {"role": "segment", "shape_names": [
                "funnel-alias-segment-1", "funnel-alias-segment-2",
                "funnel-alias-segment-3",
            ]},
            {"role": "label", "shape_names": [
                "funnel-alias-label-1", "funnel-alias-label-2", "funnel-alias-label-3",
            ]},
        ],
    }]

    component = build_component_atlas(template, review)["components"][0]

    assert component["source_instances"][0]["instance_id"] == "secondary-funnel-instance"
    assert component["source_instances"][0]["slide_index"] == 1
    assert [
        member["shape_name"]
        for group in component["source_instances"][0]["groups"]
        for member in group["members"]
    ] == [
        "funnel-alias-segment-1", "funnel-alias-segment-2", "funnel-alias-segment-3",
        "funnel-alias-label-1", "funnel-alias-label-2", "funnel-alias-label-3",
    ]


def test_component_atlas_rejects_an_equivalent_instance_with_a_different_group_contract(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    review = _review()
    review["components"][0]["source_instances"] = [{
        "instance_id": "invalid-instance",
        "groups": [{
            "role": "label",
            "shape_names": ["funnel-label-1"],
        }],
    }]

    with pytest.raises(ValueError, match="group contract"):
        build_component_atlas(template, review)


def test_component_selector_matches_semantics_and_element_capacity(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    atlas = build_component_atlas(template, _review())

    selection = select_component(atlas, {
        "semantic_use": "conversion",
        "element_count": 4,
    })

    assert selection == {
        "status": "selected",
        "component_id": "funnel.primary",
        "family": "funnel",
        "reason": "semantic_use=conversion; element_count=4 within 1..8",
    }


def test_component_selector_refuses_to_force_an_over_capacity_component(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    atlas = build_component_atlas(template, _review())

    selection = select_component(atlas, {
        "semantic_use": "conversion",
        "element_count": 9,
    })

    assert selection["status"] == "no_match"
    assert selection["reason"] == "no reviewed component satisfies semantic use and element capacity"


def test_component_selector_requires_reviewed_topology_capability(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    review = _review()
    review["components"][0]["topologies"] = ["sequence"]
    atlas = build_component_atlas(template, review)

    selected = select_component(atlas, {
        "semantic_use": "conversion",
        "element_count": 4,
        "topology": "sequence",
    })
    rejected = select_component(atlas, {
        "semantic_use": "conversion",
        "element_count": 4,
        "topology": "comparison",
    })

    assert selected["status"] == "selected"
    assert rejected == {
        "status": "no_match",
        "reason": "no reviewed component satisfies topology, semantic use and element capacity",
    }


def test_component_selector_requires_reviewed_semantic_slots(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    review = _review()
    review["components"][0]["semantic_contract"] = {
        "required_fields": ["stage", "value"],
    }
    atlas = build_component_atlas(template, review)

    selected = select_component(atlas, {
        "semantic_use": "conversion",
        "element_count": 4,
        "required_slots": ["stage", "value"],
    })
    rejected = select_component(atlas, {
        "semantic_use": "conversion",
        "element_count": 4,
        "required_slots": ["stage", "source_footer"],
    })

    assert selected["status"] == "selected"
    assert rejected == {
        "status": "no_match",
        "reason": "no reviewed component satisfies required semantic slots, semantic use and element capacity",
    }


def test_component_binding_resolves_semantics_to_native_shape_groups(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    atlas = build_component_atlas(template, _review())

    binding = resolve_component_binding(atlas, {
        "semantic_use": "conversion",
        "element_count": 4,
    })

    assert binding == {
        "component_id": "funnel.primary",
        "component_type": "funnel",
        "source_slide_index": 1,
        "segment_groups": [
            ["funnel-segment-1"],
            ["funnel-segment-2"],
            ["funnel-segment-3"],
        ],
        "label_names": [
            "funnel-label-1", "funnel-label-2", "funnel-label-3",
        ],
    }


def test_component_binding_refuses_a_requirement_without_a_reviewed_match(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _component_template(template)
    atlas = build_component_atlas(template, _review())

    with pytest.raises(ValueError, match="no reviewed component"):
        resolve_component_binding(atlas, {
            "semantic_use": "conversion",
            "element_count": 9,
        })


def test_component_atlas_resolves_a_reviewed_multi_chart_dashboard_contract(tmp_path: Path) -> None:
    template = tmp_path / "dashboard-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, name in enumerate(("innovation-chart", "ebit-chart")):
        data = ChartData()
        data.categories = ["Yes", "No"]
        data.add_series("Share", (50, 50))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            Inches(0.8 + index * 4), Inches(1.2), Inches(3.2), Inches(3.2), data,
        )
        chart.name = name
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(7), Inches(0.5))
    title.name = "dashboard-title"
    title.text = "Sample dashboard"
    note = slide.shapes.add_textbox(Inches(8.4), Inches(1.2), Inches(3.5), Inches(1.2))
    note.name = "dashboard-note"
    note.text = "Sample conclusion"
    presentation.save(template)
    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "dashboard-review"}},
        "components": [{
            "component_id": "dashboard.research-impact",
            "family": "chart_dashboard",
            "slide_index": 1,
            "semantic_uses": ["research impact evidence"],
            "groups": [{
                "role": "chart", "scope": "item", "item_indices": [0, 1],
                "shape_names": ["innovation-chart", "ebit-chart"],
            }, {
                "role": "text", "scope": "shared",
                "shape_names": ["dashboard-title", "dashboard-note"],
            }],
            "parameters": {"element_count": {"minimum": 2, "maximum": 2}},
        }],
    }

    atlas = build_component_atlas(template, review)
    binding = resolve_chart_dashboard_binding(atlas, {
        "semantic_use": "research impact evidence",
        "family": "chart_dashboard",
        "element_count": 2,
    })

    assert binding == {
        "component_id": "dashboard.research-impact",
        "component_type": "chart_dashboard",
        "source_slide_index": 1,
        "chart_names": ["innovation-chart", "ebit-chart"],
        "text_names": ["dashboard-title", "dashboard-note"],
        "decoration_names": [],
    }


def test_component_atlas_preserves_extended_group_contract(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _extended_component_template(template)

    atlas = build_component_atlas(template, _extended_review())
    groups = atlas["components"][0]["groups"]

    assert groups[1]["scope"] == "item"
    assert groups[1]["item_indices"] == [0, 0, 1, 2]
    assert groups[2]["bind_field"] == "title"
    assert groups[4]["scope"] == "shared"
    assert "item_indices" not in groups[4]


def test_component_binding_resolves_extended_native_shape_contract(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _extended_component_template(template)
    atlas = build_component_atlas(template, _extended_review())

    binding = resolve_component_binding(atlas, {
        "semantic_use": "hierarchy",
        "element_count": 3,
    })

    assert binding == {
        "component_id": "pyramid.extended",
        "component_type": "pyramid",
        "source_slide_index": 1,
        "binding_schema_version": "2.0.0",
        "binding_mode": "extended",
        "item_groups": [
            {
                "segment_names": [
                    "pyramid-segment-1", "pyramid-shadow-1-a", "pyramid-shadow-1-b",
                ],
                "label_fields": {
                    "title": ["pyramid-title-1"],
                    "value": ["pyramid-value-1"],
                },
            },
            {
                "segment_names": ["pyramid-segment-2", "pyramid-shadow-2"],
                "label_fields": {
                    "title": ["pyramid-title-2"],
                    "value": ["pyramid-value-2"],
                },
            },
            {
                "segment_names": ["pyramid-segment-3", "pyramid-shadow-3"],
                "label_fields": {
                    "title": ["pyramid-title-3"],
                    "value": ["pyramid-value-3"],
                },
            },
        ],
        "shared_names": ["pyramid-baseline"],
    }


def test_component_atlas_rejects_misaligned_item_indices(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    _extended_component_template(template)
    review = _extended_review()
    review["components"][0]["groups"][1]["item_indices"] = [0, 1]

    with pytest.raises(ValueError, match="item_indices must align"):
        build_component_atlas(template, review)


def test_component_atlas_cli_writes_a_reviewed_atlas(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    review_path = tmp_path / "review.json"
    output_path = tmp_path / "atlas.json"
    _component_template(template)
    review_path.write_text(json.dumps(_review()), encoding="utf-8")

    completed = subprocess.run(
        [
            sys.executable, "-m", "engine", "component-atlas",
            "--template-pptx", str(template),
            "--review", str(review_path),
            "--json-out", str(output_path),
        ],
        cwd=ROOT, text=True, capture_output=True, check=False,
    )

    assert completed.returncode == 0, completed.stderr
    atlas = json.loads(output_path.read_text(encoding="utf-8"))
    assert atlas["status"] == "reviewed"
    assert atlas["components"][0]["component_id"] == "funnel.primary"


def test_component_atlas_preserves_recursive_granularity_and_child_slots(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = ChartData()
    data.categories = ["Done", "Remaining"]
    data.add_series("Progress", (45, 55))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(1), Inches(1), Inches(2), Inches(2), data,
    )
    chart.name = "ring-chart"
    value = slide.shapes.add_textbox(Inches(1.6), Inches(1.7), Inches(0.8), Inches(0.3))
    value.name = "ring-value"
    value.text = "45%"
    label = slide.shapes.add_textbox(Inches(1.5), Inches(3), Inches(1), Inches(0.3))
    label.name = "ring-label"
    label.text = "Platform A"
    presentation.save(template)

    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "hierarchy"}},
        "components": [{
            "component_id": "ring.metric",
            "family": "metric_ring",
            "granularity": "micro",
            "slide_index": 1,
            "semantic_uses": ["proportion metric", "progress metric"],
            "semantic_contract": {
                "required_fields": ["value", "label"],
                "value_domain": "ratio",
            },
            "groups": [{
                "role": "chart", "scope": "item", "item_indices": [0],
                "shape_names": ["ring-chart"],
            }, {
                "role": "label", "scope": "item", "bind_field": "value",
                "item_indices": [0], "shape_names": ["ring-value"],
            }, {
                "role": "label", "scope": "item", "bind_field": "label",
                "item_indices": [0], "shape_names": ["ring-label"],
            }],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }, {
            "component_id": "ring.comparison",
            "family": "ring_comparison",
            "granularity": "composite",
            "slide_index": 1,
            "semantic_uses": ["before after comparison"],
            "children": [{
                "slot_id": "before", "component_id": "ring.metric",
                "placement": {"x": 0.0, "y": 0.0, "w": 0.45, "h": 1.0},
            }, {
                "slot_id": "after", "component_id": "ring.metric",
                "placement": {"x": 0.55, "y": 0.0, "w": 0.45, "h": 1.0},
            }],
            "parameters": {"element_count": {"minimum": 2, "maximum": 2}},
        }],
    }

    atlas = build_component_atlas(template, review)
    ring, comparison = atlas["components"]

    assert ring["granularity"] == "micro"
    assert ring["semantic_contract"]["required_fields"] == ["value", "label"]
    assert comparison["granularity"] == "composite"
    assert [child["slot_id"] for child in comparison["children"]] == ["before", "after"]
    assert comparison["children"][1]["placement"]["x"] == 0.55


def test_component_atlas_rejects_recursive_component_cycles(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "cycle"}},
        "components": [{
            "component_id": "recipe.a", "family": "recipe", "granularity": "composite",
            "slide_index": 1, "semantic_uses": ["a"],
            "children": [{"slot_id": "b", "component_id": "recipe.b", "placement": {"x": 0, "y": 0, "w": 1, "h": 1}}],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }, {
            "component_id": "recipe.b", "family": "recipe", "granularity": "composite",
            "slide_index": 1, "semantic_uses": ["b"],
            "children": [{"slot_id": "a", "component_id": "recipe.a", "placement": {"x": 0, "y": 0, "w": 1, "h": 1}}],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }],
    }

    with pytest.raises(ValueError, match="cycle"):
        build_component_atlas(template, review)


def test_component_atlas_resolves_a_reviewed_fixed_native_group_contract(tmp_path: Path) -> None:
    template = tmp_path / "native-group-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(5.0), Inches(3.0),
    )
    background.name = "group-background"
    title = slide.shapes.add_textbox(Inches(1.1), Inches(1.0), Inches(4.2), Inches(0.5))
    title.name = "group-title"
    title.text = "Sample title"
    for index in range(2):
        detail = slide.shapes.add_textbox(
            Inches(1.3), Inches(1.7 + index * 0.8), Inches(3.8), Inches(0.5),
        )
        detail.name = f"group-detail-{index + 1}"
        detail.text = f"Sample detail {index + 1}"
    presentation.save(template)
    review = {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "native-group"}},
        "components": [{
            "component_id": "insight.list",
            "family": "fixed_native_group",
            "granularity": "section",
            "renderer": "native_group",
            "slide_index": 1,
            "semantic_uses": ["evidence insight list"],
            "groups": [{
                "role": "label", "scope": "shared", "bind_field": "title",
                "shape_names": ["group-title"],
            }, {
                "role": "label", "scope": "item", "bind_field": "detail",
                "item_indices": [0, 1],
                "shape_names": ["group-detail-1", "group-detail-2"],
            }, {
                "role": "decoration", "scope": "shared",
                "shape_names": ["group-background"],
            }],
            "parameters": {"element_count": {"minimum": 2, "maximum": 2}},
        }],
    }
    atlas = build_component_atlas(template, review)
    resolver = getattr(component_atlas_module, "resolve_native_group_binding", None)

    assert resolver is not None, "fixed native groups need an Atlas resolver"
    binding = resolver(atlas, {
        "component_id": "insight.list",
        "semantic_use": "evidence insight list",
        "family": "fixed_native_group",
        "element_count": 2,
    })
    assert binding == {
        "component_id": "insight.list",
        "component_type": "fixed_native_group",
        "source_slide_index": 1,
        "label_fields": {
            "title": ["group-title"],
            "detail": ["group-detail-1", "group-detail-2"],
        },
        "decoration_names": ["group-background"],
    }


def test_chart_component_binding_keeps_one_ring_independently_addressable(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = ChartData()
    data.categories = ["Done", "Remaining"]
    data.add_series("Progress", (45, 55))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(1), Inches(1), Inches(2), Inches(2), data,
    )
    chart.name = "ring-chart"
    value = slide.shapes.add_textbox(Inches(1.6), Inches(1.7), Inches(0.8), Inches(0.3))
    value.name = "ring-value"
    label = slide.shapes.add_textbox(Inches(1.5), Inches(3), Inches(1), Inches(0.3))
    label.name = "ring-label"
    presentation.save(template)
    atlas = build_component_atlas(template, {
        "schema_version": "1.0.0",
        "review": {"state": "reviewed", "reviewer": {"method": "human", "id": "ring"}},
        "components": [{
            "component_id": "ring.metric", "family": "metric_ring", "granularity": "micro",
            "slide_index": 1, "semantic_uses": ["progress metric"],
            "groups": [{"role": "chart", "scope": "item", "item_indices": [0], "shape_names": ["ring-chart"]},
                       {"role": "label", "scope": "item", "bind_field": "value", "item_indices": [0], "shape_names": ["ring-value"]},
                       {"role": "label", "scope": "item", "bind_field": "label", "item_indices": [0], "shape_names": ["ring-label"]}],
            "parameters": {"element_count": {"minimum": 1, "maximum": 1}},
        }],
    })

    binding = resolve_chart_component_binding(atlas, {
        "component_id": "ring.metric", "semantic_use": "progress metric",
        "family": "metric_ring", "element_count": 1,
    })

    assert binding == {
        "component_id": "ring.metric",
        "component_type": "metric_ring",
        "source_slide_index": 1,
        "chart_names": ["ring-chart"],
        "label_fields": {"value": ["ring-value"], "label": ["ring-label"]},
        "decoration_names": [],
    }
