from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from engine.high_fidelity_guard import content_lock, text_inventory


def _deck(path: Path, *, extra: str | None = None, drop_source: bool = False) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, text in enumerate(["Growth accelerated", "Q3 growth: 24%", "Source: report p.4"]):
        if drop_source and text.startswith("Source:"):
            continue
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8 + index), Inches(8), Inches(0.5))
        box.text = text
    if extra:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(8), Inches(0.5))
        box.text = extra
    prs.save(path)


def test_content_lock_accepts_layout_only_change(tmp_path: Path) -> None:
    base, candidate = tmp_path / "base.pptx", tmp_path / "candidate.pptx"
    _deck(base)
    _deck(candidate)
    result = content_lock(base, candidate)
    assert result["status"] == "pass"
    assert result["missing_text"] == []
    assert result["added_text"] == []


def test_content_lock_rejects_added_unbound_text(tmp_path: Path) -> None:
    base, candidate = tmp_path / "base.pptx", tmp_path / "candidate.pptx"
    _deck(base)
    _deck(candidate, extra="High performers always win")
    result = content_lock(base, candidate)
    assert result["status"] == "fail"
    assert result["added_text"] == ["High performers always win"]


def test_content_lock_rejects_missing_source_text(tmp_path: Path) -> None:
    base, candidate = tmp_path / "base.pptx", tmp_path / "candidate.pptx"
    _deck(base)
    _deck(candidate, drop_source=True)
    result = content_lock(base, candidate)
    assert result["status"] == "fail"
    assert result["missing_text"] == ["Source: report p.4"]


def test_text_inventory_ignores_empty_and_whitespace(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    _deck(path)
    prs = Presentation(str(path))
    slide = prs.slides[0]
    box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(1), Inches(0.2))
    box.text = "  \n "
    prs.save(path)
    assert text_inventory(path) == ["Growth accelerated", "Q3 growth: 24%", "Source: report p.4"]
