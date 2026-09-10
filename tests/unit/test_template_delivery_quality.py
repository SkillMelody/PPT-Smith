from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from engine.template_delivery_quality import inspect_template_delivery_quality


def _deck(path: Path, *, invalid: bool) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(7), Inches(1))
    body.name = "bind:block:s1:content:item:0"
    run = body.text_frame.paragraphs[0].add_run()
    run.text = "undefined" if invalid else "A complete evidence-backed conclusion"
    run.font.size = Pt(9 if invalid else 12)
    page = slide.shapes.add_textbox(Inches(12), Inches(7.0), Inches(0.5), Inches(0.2))
    page.name = "bind:block:s1:page_no:value"
    page.text = "1"
    page.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    if invalid:
        duplicate = slide.shapes.add_textbox(
            Inches(12.1), Inches(7.02), Inches(0.5), Inches(0.2),
        )
        duplicate.name = "Slide Number Placeholder 0"
        duplicate.text = "01"
        duplicate.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    prs.save(path)


def _ir() -> dict:
    return {"slides": [{"id": "s1", "slide_role": "content"}]}


def test_template_delivery_quality_accepts_clean_readable_slide(tmp_path: Path) -> None:
    deck = tmp_path / "valid.pptx"
    _deck(deck, invalid=False)

    report = inspect_template_delivery_quality(deck, ir=_ir())

    assert report["status"] == "pass"
    assert report["issues"] == []


def test_template_delivery_quality_rejects_placeholder_small_text_and_duplicate_page_number(
    tmp_path: Path,
) -> None:
    deck = tmp_path / "invalid.pptx"
    _deck(deck, invalid=True)

    report = inspect_template_delivery_quality(deck, ir=_ir())
    codes = {issue["code"] for issue in report["issues"]}

    assert report["status"] == "fail"
    assert {
        "TEMPLATE_RENDERED_PLACEHOLDER_TEXT",
        "TEMPLATE_BODY_FONT_BELOW_FLOOR",
        "TEMPLATE_DUPLICATE_PAGE_NUMBER",
    } <= codes
