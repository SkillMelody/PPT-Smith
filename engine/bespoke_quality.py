"""Deterministic visual-floor checks for Bespoke PPTX delivery.

These checks do not choose layouts or constrain the author's narrative. They
catch portable-delivery failures that a visual reviewer should never have to
infer: silently shrinking business text, unreadable explicit font sizes,
missing wrapping on dense text, excessive outlined-container noise, and pie
charts whose colours are not bound per data point.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE

_EMU_PER_INCH = 914400
_PIE_TYPES = {XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT}


def _iter_shapes(shapes: Any) -> list[Any]:
    flattened: list[Any] = []
    for shape in shapes:
        flattened.append(shape)
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            flattened.extend(_iter_shapes(shape.shapes))
    return flattened


def _text_runs(shape: Any):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            yield run


def _font_floor(name: str) -> tuple[float, str]:
    if name.startswith("bind:source:") or name.endswith(":caption"):
        return 9.0, "source_or_caption"
    if name.startswith("bind:slide:") and name.endswith(":title"):
        return 22.0, "title"
    if name.startswith("bind:slide:") and name.endswith(":message"):
        return 12.0, "message"
    return 12.0, "body"


def _color_key(fill: Any) -> str | None:
    try:
        if fill.type is None:
            return None
        color = fill.fore_color
        if color.type is None:
            return None
        try:
            rgb = color.rgb
        except (TypeError, ValueError):
            rgb = None
        if rgb is not None:
            return f"rgb:{rgb}"
        try:
            theme = color.theme_color
        except (TypeError, ValueError):
            theme = None
        return f"theme:{theme}" if theme is not None else None
    except (AttributeError, TypeError, ValueError):
        return None


def _finding(code: str, slide_index: int, shape_name: str, message: str, **evidence: Any) -> dict:
    return {
        "code": code,
        "slide_index": slide_index,
        "shape_name": shape_name,
        "message": message,
        "evidence": evidence,
    }


def inspect_bespoke_visual_quality(pptx_path: str | Path) -> dict:
    """Return deterministic visual-floor findings for a Bespoke PPTX.

    ``pass`` is deliberately narrow: it means the file avoids known delivery
    anti-patterns. It is not an aesthetic approval and never replaces rendered
    human/model review.
    """
    prs = Presentation(str(pptx_path))
    findings: list[dict] = []

    for slide_index, slide in enumerate(prs.slides, 1):
        outlined_containers: list[str] = []
        for shape in _iter_shapes(slide.shapes):
            name = str(getattr(shape, "name", "") or "")

            if getattr(shape, "has_text_frame", False) and name.startswith("bind:"):
                text = " ".join((shape.text or "").split())
                if text:
                    floor, role = _font_floor(name)
                    explicit_sizes = [
                        run.font.size.pt for run in _text_runs(shape)
                        if run.font.size is not None
                    ]
                    for size in explicit_sizes:
                        if size + 1e-6 < floor:
                            findings.append(_finding(
                                "BESPOKE_BODY_FONT_TOO_SMALL",
                                slide_index,
                                name,
                                "Bound business text is below the Bespoke readability floor.",
                                font_size_pt=round(size, 2),
                                minimum_pt=floor,
                                role=role,
                                text=text[:160],
                            ))
                    if shape.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE:
                        findings.append(_finding(
                            "BESPOKE_SHRINK_TO_FIT_FORBIDDEN",
                            slide_index,
                            name,
                            "Shrink-to-fit can silently reduce text below the declared font size.",
                            text=text[:160],
                        ))
                    width_in = float(shape.width) / _EMU_PER_INCH
                    if len(text) >= 35 and width_in < 10 and shape.text_frame.word_wrap is not True:
                        findings.append(_finding(
                            "BESPOKE_WORD_WRAP_REQUIRED",
                            slide_index,
                            name,
                            "Dense bound text must explicitly enable word wrapping.",
                            width_in=round(width_in, 2),
                            characters=len(text),
                            text=text[:160],
                        ))

            if getattr(shape, "has_table", False) and name.startswith("bind:table:"):
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if run.font.size is not None and run.font.size.pt < 12.0:
                                    findings.append(_finding(
                                        "BESPOKE_TABLE_FONT_TOO_SMALL",
                                        slide_index,
                                        name,
                                        "Native table text is below the Bespoke readability floor.",
                                        row=row_index,
                                        column=column_index,
                                        font_size_pt=round(run.font.size.pt, 2),
                                        minimum_pt=12.0,
                                        text=run.text[:120],
                                    ))

            if getattr(shape, "has_chart", False) and shape.chart.chart_type in _PIE_TYPES:
                for series_index, series in enumerate(shape.chart.series):
                    colors = [_color_key(point.format.fill) for point in series.points]
                    distinct = {color for color in colors if color is not None}
                    required = min(3, len(series.points))
                    if len(colors) > 1 and (None in colors or len(distinct) < required):
                        findings.append(_finding(
                            "BESPOKE_PIE_POINT_COLORS_REQUIRED",
                            slide_index,
                            name,
                            "Pie and doughnut charts require explicit distinct point colours for portable rendering.",
                            series_index=series_index,
                            point_count=len(colors),
                            explicit_colors=colors,
                            minimum_distinct_colors=required,
                        ))

            if name.startswith("decoration:") and getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.AUTO_SHAPE:
                try:
                    has_line = shape.line.fill.type not in {None, MSO_FILL_TYPE.BACKGROUND}
                    area = (float(shape.width) / _EMU_PER_INCH) * (float(shape.height) / _EMU_PER_INCH)
                except (AttributeError, TypeError, ValueError):
                    has_line, area = False, 0.0
                if has_line and area >= 0.2:
                    outlined_containers.append(name)

        if len(outlined_containers) > 6:
            findings.append(_finding(
                "BESPOKE_OUTLINE_DENSITY_HIGH",
                slide_index,
                "<slide>",
                "Too many outlined containers create card-grid and wireframe noise.",
                outlined_containers=len(outlined_containers),
                maximum=6,
                names=outlined_containers,
            ))

    return {"status": "fail" if findings else "pass", "findings": findings}
