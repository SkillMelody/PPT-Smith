"""Speaker-notes contracts, formatting, PPTX writing, and read-back QA."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from .evidence_ledger import evidence_by_id


def format_speaker_notes(notes: dict) -> str:
    """Render structured model notes into a stable human-readable script."""
    if not isinstance(notes, dict):
        raise ValueError("speaker notes must be an object")
    narrative = notes.get("narrative")
    if not isinstance(narrative, str) or not narrative.strip():
        raise ValueError("speaker notes require narrative")
    sections = [narrative.strip()]
    talking_points = notes.get("talking_points", [])
    if talking_points:
        sections.append("Talking points:\n" + "\n".join(
            f"- {item.strip()}" for item in talking_points
            if isinstance(item, str) and item.strip()
        ))
    caveats = notes.get("caveats", [])
    if caveats:
        sections.append("Caveats:\n" + "\n".join(
            f"- {item.strip()}" for item in caveats
            if isinstance(item, str) and item.strip()
        ))
    evidence_ids = notes.get("evidence_ids", [])
    if evidence_ids:
        sections.append("Evidence IDs: " + ", ".join(evidence_ids))
    refs = notes.get("source_refs", [])
    if refs:
        sections.append("Sources:\n" + "\n".join(
            f"- {ref.get('source_id')} {ref.get('loc')}"
            for ref in refs if isinstance(ref, dict)
        ))
    return "\n\n".join(section for section in sections if section.strip())


def validate_speaker_notes(
    ir: dict,
    *,
    evidence_ledger: dict | None = None,
    require_body_notes: bool = True,
) -> dict:
    """Validate content depth and source/evidence identity before authoring."""
    known_evidence = evidence_by_id(evidence_ledger or {})
    issues: list[dict] = []
    body_count = 0
    covered_count = 0
    for index, slide in enumerate(ir.get("slides", []), 1):
        if not isinstance(slide, dict):
            continue
        slide_id = str(slide.get("id") or index)
        role = slide.get("slide_role", "content")
        body = role not in {"cover", "section", "closing"}
        if body:
            body_count += 1
        notes = slide.get("speaker_notes")
        if notes is None:
            if body and require_body_notes:
                issues.append({"code": "SPEAKER_NOTES_MISSING", "slide_id": slide_id})
            continue
        if not isinstance(notes, dict):
            issues.append({"code": "SPEAKER_NOTES_INVALID", "slide_id": slide_id})
            continue
        narrative = notes.get("narrative")
        if not isinstance(narrative, str) or len(narrative.strip()) < 80:
            issues.append({"code": "SPEAKER_NOTES_NARRATIVE_TOO_SHORT", "slide_id": slide_id})
        evidence_ids = notes.get("evidence_ids")
        if not isinstance(evidence_ids, list) or not evidence_ids:
            issues.append({"code": "SPEAKER_NOTES_EVIDENCE_MISSING", "slide_id": slide_id})
        elif evidence_ledger is not None:
            for evidence_id in evidence_ids:
                if evidence_id not in known_evidence:
                    issues.append({
                        "code": "SPEAKER_NOTES_EVIDENCE_UNKNOWN",
                        "slide_id": slide_id,
                        "evidence_id": evidence_id,
                    })
        refs = notes.get("source_refs")
        if not isinstance(refs, list) or not refs:
            issues.append({"code": "SPEAKER_NOTES_SOURCE_REFS_MISSING", "slide_id": slide_id})
        if not any(issue.get("slide_id") == slide_id for issue in issues):
            covered_count += int(body)
    return {
        "status": "pass" if not issues else "fail",
        "body_slide_count": body_count,
        "covered_body_slide_count": covered_count,
        "coverage_ratio": round(covered_count / body_count, 4) if body_count else 1.0,
        "issues": issues,
    }


def write_speaker_notes(
    pptx_path: str | Path,
    notes_by_slide_index: dict[int, dict],
) -> dict:
    """Write notes to existing slides with one open/save cycle."""
    path = Path(pptx_path)
    presentation = Presentation(path)
    written: list[int] = []
    for slide_index, notes in sorted(notes_by_slide_index.items()):
        if slide_index < 1 or slide_index > len(presentation.slides):
            raise ValueError(f"speaker-notes slide index out of range: {slide_index}")
        presentation.slides[slide_index - 1].notes_slide.notes_text_frame.text = (
            format_speaker_notes(notes)
        )
        written.append(slide_index)
    presentation.save(path)
    return {"status": "written", "slide_indices": written, "count": len(written)}


def inspect_speaker_notes(
    pptx_path: str | Path,
    *,
    required_slide_indices: list[int],
) -> dict:
    """Read back notes so a missing relationship cannot pass silently."""
    presentation = Presentation(str(pptx_path))
    missing: list[int] = []
    lengths: dict[str, int] = {}
    for slide_index in required_slide_indices:
        if slide_index < 1 or slide_index > len(presentation.slides):
            missing.append(slide_index)
            continue
        text = presentation.slides[slide_index - 1].notes_slide.notes_text_frame.text.strip()
        lengths[str(slide_index)] = len(text)
        if not text:
            missing.append(slide_index)
    return {
        "status": "pass" if not missing else "fail",
        "required_slide_count": len(required_slide_indices),
        "present_slide_count": len(required_slide_indices) - len(missing),
        "missing_slide_indices": missing,
        "text_lengths": lengths,
    }

