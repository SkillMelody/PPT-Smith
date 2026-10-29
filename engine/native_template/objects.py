"""Real cNvPr IDs, style-preserving edits and same-source native additions."""
from __future__ import annotations
from copy import deepcopy
import io
import re
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.opc.packuri import PackURI
from .package import A, P, R, NS, clone_pages, package, pack, prune, canonical_workbooks, strip_metadata
from .safety import require, fields, sha, MAX_EXPANDED, encode
from .styles import resolved_text, theme_context

EMU = 914400


def walk(shapes, parent=None):
    for shape in shapes:
        yield shape, parent
        if shape.shape_type == 6:
            yield from walk(shape.shapes, str(shape.shape_id))


def find(slide, object_id):
    require(isinstance(object_id, str), 'OBJECT_ID_MUST_BE_STRING')
    found = [s for s, _ in walk(slide.shapes) if str(s.shape_id) == object_id]
    require(len(found) == 1, 'OBJECT_MISSING_OR_AMBIGUOUS: ' + object_id)
    return found[0]


def kind(shape):
    if shape.shape_type == 6: return 'group'
    if shape.has_chart: return 'chart'
    if shape.has_table: return 'table'
    if shape.shape_type == 13: return 'picture'
    return 'text' if shape.has_text_frame and shape.text else 'shape'


def _inheritance(shape, slide):
    chain = [('object', shape._element)]
    if shape.is_placeholder:
        idx = shape.placeholder_format.idx
        for label, owner in [('layout', slide.slide_layout), ('master', slide.slide_layout.slide_master)]:
            match = next((p for p in owner.placeholders if (p.placeholder_format.idx == idx if label == 'layout' else p.placeholder_format.type == shape.placeholder_format.type)), None)
            if match is not None: chain.append((label, match._element))
    return chain


def style_evidence(shape, slide):
    evidence = []
    for origin, element in _inheritance(shape, slide):
        for node in element.iter():
            local = etree.QName(node).localname
            if local in {'rPr', 'defRPr', 'endParaRPr'}:
                record = {'origin': origin, 'attributes': dict(node.attrib)}
                record['fonts'] = [dict(c.attrib) for c in node if etree.QName(c).localname in {'latin', 'ea', 'cs'}]
                record['colors'] = [dict(kind=etree.QName(c).localname, **c.attrib) for c in node.iter() if etree.QName(c).localname in {'srgbClr', 'schemeClr', 'sysClr'}]
                if record['attributes'] or record['fonts'] or record['colors']: evidence.append(record)
    themes = []
    master = slide.slide_layout.slide_master
    for rel in master.part.rels.values():
        if rel.reltype.endswith('/theme'):
            tree = etree.fromstring(rel.target_part.blob)
            themes.append({'sha256': sha(rel.target_part.blob), 'part': str(rel.target_part.partname)})
    return {'inheritance': evidence, 'themes': themes, 'resolution': 'OOXML inheritance evidence; no guessed default font/color',
            'inheritedObjectsEditable': False}


def inspect_slide(slide, source_id, page_id=None, details=False):
    objects = []
    seen = set()
    for s, parent in walk(slide.shapes):
        sid = str(s.shape_id)
        require(sid not in seen, 'DUPLICATE_OBJECT_ID'); seen.add(sid)
        item = {'objectId': sid, 'type': kind(s), 'name': s.name, 'parentObjectId': parent,
                'geometry': {k: getattr(s, k) / EMU for k in ('left', 'top', 'width', 'height')},
                'style': {'resolution': 'effective_values_only; unresolved is unknown', 'sourceChain': [], 'inheritedObjectsEditable': False}}
        if details: item['style'] = style_evidence(s, slide)
        if s.has_text_frame:
            item['text'] = s.text
            resolved = resolved_text(s, slide)
            if details:
                item['effectiveTextStyles'] = resolved
            else:
                unique = {}
                for record in resolved:
                    compact = {k: v for k, v in record.items() if k not in {'text', 'paragraph', 'run'}}
                    key = encode(compact)
                    if key not in unique: unique[key] = {**compact, 'runCount': 0}
                    unique[key]['runCount'] += 1
                item['effectiveTextStyles'] = list(unique.values())
                item['style']['sourceChain'] = list(dict.fromkeys(v['origin'] for r in resolved for v in r.values() if isinstance(v, dict) and 'origin' in v))
        if s.has_table:
            item['rows'] = [[c.text for c in row.cells] for row in s.table.rows]
            item['editBoundary'] = 'existing rectangular cells only; no merge/structure changes'
        if s.has_chart:
            item['editBoundary'] = 'single-plot category chart replace_data only; other chart types unsupported'
            item['chartType'] = str(s.chart.chart_type)
        objects.append(item)
    require(len(objects) <= 1000, 'OBJECT_LIMIT')
    _, palette, color_map = theme_context(slide)
    return {'sourceSlideId': source_id, 'pageId': page_id, 'objects': objects,
            'themeColors': palette, 'colorMap': color_map,
            'layout': str(slide.slide_layout.part.partname), 'master': str(slide.slide_layout.slide_master.part.partname)}


def replace_text(frame, text):
    require(isinstance(text, str) and len(text) <= 100000 and not any(ord(c) < 32 and c not in '\n\t' for c in text), 'INVALID_TEXT')
    if frame.text == text:
        return  # Preserve mixed-run formatting byte-for-byte for a no-op edit.
    paragraphs = list(frame.paragraphs)
    prototype = deepcopy(paragraphs[-1]._p)
    body = frame._txBody
    lines = text.split('\n')
    for idx, line in enumerate(lines):
        if idx < len(paragraphs):
            p = paragraphs[idx]._p
        else:
            p = deepcopy(prototype); body.append(p)
        runs = p.findall('{%s}r' % A)
        # Keep paragraph properties and the first real run's complete native rPr.
        seed = deepcopy(runs[0]) if runs else etree.Element('{%s}r' % A)
        for child in list(p):
            if etree.QName(child).localname not in {'pPr', 'endParaRPr'}: p.remove(child)
        for child in list(seed):
            if etree.QName(child).localname != 'rPr': seed.remove(child)
        etree.SubElement(seed, '{%s}t' % A).text = line
        end = p.find('{%s}endParaRPr' % A)
        if end is None: p.append(seed)
        else: p.insert(p.index(end), seed)
    for p in paragraphs[len(lines):]: body.remove(p._p)


def geometry(shape, options, presentation):
    fields(options, {'x', 'y', 'width', 'height', 'rotation'})
    if options:
        require(shape._element.getparent().tag != '{%s}grpSp' % P, 'GROUP_CHILD_TRANSFORM_UNSUPPORTED')
    current = {'x': shape.left / EMU, 'y': shape.top / EMU, 'width': shape.width / EMU, 'height': shape.height / EMU}
    current.update({k: v for k, v in options.items() if k != 'rotation'})
    for k, v in current.items():
        require(type(v) in (int, float) and v >= 0 and (v > 0 if k in {'width', 'height'} else True), 'INVALID_GEOMETRY')
    require(current['x'] + current['width'] <= presentation.slide_width / EMU + 1e-6 and current['y'] + current['height'] <= presentation.slide_height / EMU + 1e-6, 'GEOMETRY_OUT_OF_BOUNDS')
    for k, attr in [('x', 'left'), ('y', 'top'), ('width', 'width'), ('height', 'height')]:
        if k in options: setattr(shape, attr, int(options[k] * EMU))
    if 'rotation' in options:
        require(type(options['rotation']) in (int, float) and 0 <= options['rotation'] < 360, 'INVALID_ROTATION')
        shape.rotation = options['rotation']


def edit(shape, operation, presentation):
    require(sum(k in operation for k in ('text', 'rows', 'chart')) <= 1, 'ONE_CONTENT_EDIT_PER_OPERATION')
    if 'text' in operation:
        require(shape.has_text_frame, 'OBJECT_HAS_NO_TEXT_FRAME')
        replace_text(shape.text_frame, operation['text'])
    if 'rows' in operation:
        require(shape.has_table, 'OBJECT_NOT_TABLE')
        rows = operation['rows']
        require(isinstance(rows, list) and len(rows) == len(shape.table.rows) and all(isinstance(r, list) and len(r) == len(shape.table.columns) for r in rows), 'TABLE_DIMENSIONS_MUST_MATCH')
        require(all(not c.is_merge_origin and not c.is_spanned for row in shape.table.rows for c in row.cells), 'MERGED_TABLE_EDIT_UNSUPPORTED')
        for r, values in enumerate(rows):
            for c, value in enumerate(values): replace_text(shape.table.cell(r, c).text_frame, value)
    if 'chart' in operation:
        require(shape.has_chart, 'OBJECT_NOT_CHART')
        chart = operation['chart']
        fields(chart, {'categories', 'series'}, {'categories', 'series'})
        categories, series = chart['categories'], chart['series']
        require(isinstance(categories, list) and 0 < len(categories) <= 1000 and all(isinstance(v, str) and len(v) <= 1000 for v in categories), 'INVALID_CHART_CATEGORIES')
        require(isinstance(series, list) and 0 < len(series) <= 100, 'INVALID_CHART_SERIES')
        native = shape.chart
        require(len(native.plots) == 1 and not any(x in str(native.chart_type) for x in ('XY', 'BUBBLE', 'STOCK', 'SURFACE')), 'CHART_TYPE_UNSUPPORTED')
        data = CategoryChartData(); data.categories = categories
        for s in series:
            fields(s, {'name', 'values'}, {'name', 'values'})
            require(isinstance(s['name'], str) and isinstance(s['values'], list) and len(s['values']) == len(categories) and all(type(v) in (int, float) for v in s['values']), 'INVALID_CHART_VALUES')
            data.add_series(s['name'], s['values'])
        native.replace_data(data)
    if 'options' in operation: geometry(shape, operation['options'], presentation)


def _clone_part(part, package_, allocated, mapping, budget):
    if part in mapping: return mapping[part]
    stem, ext = str(part.partname).rsplit('.', 1)
    n = 1
    while stem + '-native%d.' % n + ext in allocated: n += 1
    name = stem + '-native%d.' % n + ext; allocated.add(name)
    budget[0] += len(part.blob)
    require(budget[0] <= MAX_EXPANDED, 'NATIVE_OBJECT_PACKAGE_SIZE_LIMIT')
    clone = type(part).load(PackURI(name), part.content_type, package_, part.blob)
    mapping[part] = clone
    for rel in part.rels.values():
        require(not rel.is_external, 'EXTERNAL_RELATIONSHIP_REFUSED')
        child = _clone_part(rel.target_part, package_, allocated, mapping, budget)
        rid = clone.relate_to(child, rel.reltype)
        if rid != rel.rId and hasattr(clone, '_element'):
            for node in clone._element.iter():
                for key, value in list(node.attrib.items()):
                    if key.startswith('{%s}' % R) and value == rel.rId: node.set(key, rid)
    return clone


def add_object(slide, source_slide, operation, presentation):
    source = find(source_slide, operation['styleSourceObjectId'])
    expected = operation['type']
    require(expected in {'text', 'shape', 'table', 'chart', 'picture'}, 'OBJECT_TYPE_UNSUPPORTED')
    require(kind(source) == expected or expected == 'text' and source.has_text_frame, 'STYLE_SOURCE_TYPE_MISMATCH')
    require(source.shape_type != 6, 'GROUP_ADDITION_UNSUPPORTED')
    oid = operation['objectId']
    require(isinstance(oid, str) and re.fullmatch(r'[1-9][0-9]{0,8}', oid), 'NEW_OBJECT_ID_MUST_BE_POSITIVE_NUMERIC_STRING')
    require(not any(str(s.shape_id) == oid for s, _ in walk(slide.shapes)), 'DUPLICATE_OBJECT_ID')
    budget = [sum(len(p.blob) for p in presentation.part.package.iter_parts()) + len(etree.tostring(source._element))]
    require(budget[0] <= MAX_EXPANDED, 'NATIVE_OBJECT_PACKAGE_SIZE_LIMIT')
    element = deepcopy(source._element)
    props = element.find('.//{%s}cNvPr' % P); props.set('id', oid); props.set('name', source.name + '-native-' + oid)
    allocated = {str(p.partname) for p in presentation.part.package.iter_parts()}
    mapping = {}
    for node in element.iter():
        for key, value in list(node.attrib.items()):
            if key.startswith('{%s}' % R):
                rel = source_slide.part.rels[value]
                require(not rel.is_external, 'EXTERNAL_RELATIONSHIP_REFUSED')
                copied = _clone_part(rel.target_part, presentation.part.package, allocated, mapping, budget)
                node.set(key, slide.part.relate_to(copied, rel.reltype))
    slide.shapes._spTree.insert_element_before(element, 'p:extLst')
    edit(find(slide, oid), operation, presentation)


def build(template, state):
    presentation = Presentation(io.BytesIO(clone_pages(template, state['pages'])))
    source = Presentation(io.BytesIO(template))
    from .package import slide_inventory
    source_map = {sid: source.slides[i] for i, (sid, _) in enumerate(slide_inventory(package(template)))}
    for slide, page in zip(presentation.slides, state['pages']):
        for operation in page['edits']:
            if operation['op'] == 'add_object': add_object(slide, source_map[page['sourceSlideId']], operation, presentation)
            else:
                shape = find(slide, operation['objectId'])
                if operation['op'] == 'remove_object': shape._element.getparent().remove(shape._element)
                else: edit(shape, operation, presentation)
            if 'chart' in operation:
                require(sum(len(p.blob) for p in presentation.part.package.iter_parts()) <= MAX_EXPANDED, 'NATIVE_OBJECT_PACKAGE_SIZE_LIMIT')
        if 'notes' in page:
            replace_text(slide.notes_slide.notes_text_frame, page['notes'])
        inspect_slide(slide, page['sourceSlideId'], page['id'])
    buffer = io.BytesIO(); presentation.save(buffer)
    # python-pptx retains unused media rels after object deletion: prune those too.
    parts = package(buffer.getvalue())
    original_parts = package(template)
    for name in list(parts):
        if name.startswith(('ppt/slideLayouts/', 'ppt/slideMasters/', 'ppt/theme/', 'ppt/notesMasters/')) and name in original_parts:
            parts[name] = original_parts[name]
    from .package import slide_inventory, relname, xml, dump
    for _, part in slide_inventory(parts):
        tree = xml(parts[part]); used = {v for node in tree.iter() for k, v in node.attrib.items() if k.startswith('{%s}' % R)}
        rn = relname(part)
        if rn in parts:
            rr = xml(parts[rn])
            for rel in list(rr):
                if not rel.get('Type', '').endswith(('/slideLayout', '/notesSlide')) and rel.get('Id') not in used: rr.remove(rel)
            parts[rn] = dump(rr)
    return pack(prune(strip_metadata(canonical_workbooks(parts))))
