"""Source-backed retention/type-scale floor, not a visual or factual verdict.

Thresholds ported from Web lib/ppt-agent/template-style.mjs. Geometry is measured
in points (2pt tolerance); pictures count once and must retain their actual ID
and box, rather than letting unrelated replacement images increase retention.
"""
from __future__ import annotations
import io
import math
from pptx import Presentation
from .objects import walk
from .styles import resolved_text
from .package import A, P, slide_inventory
from .safety import package

RETENTION_MIN = 0.8
HEADING_SIZE_MIN = 0.75
SMALLEST_SIZE_MIN = 0.9
BOX_TOLERANCE_PT = 2.0
POINT = 12700


def _decoration(shape):
    if shape.shape_type == 13: return 'picture'
    if shape.shape_type in {6, 14, 17} or shape.has_chart or shape.has_table: return None
    if shape.has_text_frame and shape.text.strip(): return None
    return 'native_shape'


def _decorations(slide):
    result = {}
    def visit(shapes, transform=(1.0, 1.0, 0.0, 0.0), known=True):
        sx,sy,tx,ty=transform
        for shape in shapes:
            category=_decoration(shape)
            if category:
                box=[(tx+sx*shape.left)/POINT,(ty+sy*shape.top)/POINT,
                     sx*shape.width/POINT,sy*shape.height/POINT] if known else None
                result[str(shape.shape_id)]={'type':category,'boxPt':box,'rotation':shape.rotation}
            if shape.shape_type==6:
                x=shape._element.find('{%s}grpSpPr/{%s}xfrm'%(P,A))
                child=None
                if x is not None and not int(x.get('rot','0')) and x.get('flipH') not in {'1','true'} and x.get('flipV') not in {'1','true'}:
                    nodes=[x.find('{%s}%s'%(A,n)) for n in ('off','ext','chOff','chExt')]
                    if all(n is not None for n in nodes):
                        off,ext,co,ce=nodes
                        cx,cy=int(ce.get('cx','0')),int(ce.get('cy','0'))
                        if cx and cy:
                            rx,ry=int(ext.get('cx'))/cx,int(ext.get('cy'))/cy
                            child=(sx*rx,sy*ry,tx+sx*(int(off.get('x'))-int(co.get('x'))*rx),ty+sy*(int(off.get('y'))-int(co.get('y'))*ry))
                visit(shape.shapes,child or transform,known and child is not None)
    visit(slide.shapes)
    return result


def _typography(slide):
    distribution={};unknown=[];objects=0;visible_runs=0
    for shape,_ in walk(slide.shapes):
        if not shape.has_text_frame or not shape.text.strip():continue
        objects+=1
        for run in resolved_text(shape,slide):
            if not run.get('text','').strip():continue
            visible_runs+=1
            size=run.get('fontSizePt',{}).get('value')
            if not isinstance(size,(int,float)) or isinstance(size,bool) or not math.isfinite(size) or size<=0:
                unknown.append({'objectId':str(shape.shape_id),'paragraph':run['paragraph'],'run':run['run']})
            else:distribution[size]=distribution.get(size,0)+1
    return {'nativeTextObjectCount':objects,'visibleRunCount':visible_runs,'unknownSizeCount':len(unknown),
            'unknownRuns':unknown,'minFontSize':min(distribution) if distribution else None,
            'maxFontSize':max(distribution) if distribution else None,
            'distribution':[{'fontSize':s,'runCount':n} for s,n in sorted(distribution.items())]}


def check(template,candidate,state):
    source=Presentation(io.BytesIO(template));actual=Presentation(io.BytesIO(candidate))
    sources={sid:source.slides[i] for i,(sid,_) in enumerate(slide_inventory(package(template)))}
    pages=[];findings=[]
    if len(state['pages'])!=len(actual.slides):
        findings.append({'code':'TEMPLATE_FIDELITY_UNVERIFIED','reason':'physical_page_count_mismatch'})
    for index,(page,slide) in enumerate(zip(state['pages'],actual.slides),1):
        at={'page':index,'pageId':page['id'],'sourceSlideId':page['sourceSlideId']};issues=[]
        original=sources.get(page['sourceSlideId'])
        if original is None:
            issues.append({**at,'code':'TEMPLATE_FIDELITY_UNVERIFIED','reason':'unknown_source_page'})
            pages.append({**at,'status':'blocked','findings':issues});findings.extend(issues);continue
        fixed=_decorations(original);objects=_decorations(slide);kept=[];lost=[];unknown=[]
        for sid,obj in fixed.items():
            new=objects.get(sid)
            if obj['boxPt'] is None or (new is not None and new['boxPt'] is None):unknown.append(sid)
            elif new and obj['type']==new['type'] and abs(obj['rotation']-new['rotation'])<=1e-8 and all(abs(a-b)<=BOX_TOLERANCE_PT+1e-8 for a,b in zip(obj['boxPt'],new['boxPt'])):kept.append(sid)
            else:lost.append(sid)
        ratio=len(kept)/len(fixed) if fixed else 1.0
        retention={'sourceDecorations':len(fixed),'retainedDecorations':len(kept),'sourcePictures':sum(o['type']=='picture' for o in fixed.values()),
                   'retainedObjectIds':kept,'lostObjectIds':lost,'unresolvedObjectIds':unknown,'ratio':ratio,'minimum':RETENTION_MIN}
        if unknown:issues.append({**at,'code':'TEMPLATE_FIDELITY_UNVERIFIED','reason':'decoration_geometry_unresolved','objectIds':unknown})
        if ratio+1e-12<RETENTION_MIN:issues.append({**at,'code':'TEMPLATE_DECORATION_LOST','retention':retention})
        src=_typography(original);out=_typography(slide)
        if src['unknownSizeCount'] or out['unknownSizeCount']:
            issues.append({**at,'code':'TEMPLATE_FIDELITY_UNVERIFIED','reason':'native_font_size_unresolved',
                           'sourceUnknownSizeCount':src['unknownSizeCount'],'candidateUnknownSizeCount':out['unknownSizeCount']})
        if src['visibleRunCount'] and not out['visibleRunCount']:
            issues.append({**at,'code':'TEMPLATE_TYPE_SCALE_LOST','reason':'all_visible_native_text_removed'})
        for key,factor in (('maxFontSize',HEADING_SIZE_MIN),('minFontSize',SMALLEST_SIZE_MIN)):
            before,after=src[key],out[key]
            if before is not None and after is not None and after+1e-8<before*factor:
                issues.append({**at,'code':'TEMPLATE_TYPE_SCALE_LOST','reason':key+'_below_source_floor',
                               'sourceFontSize':before,'candidateFontSize':after,'minimumRatio':factor,'minimumFontSize':before*factor})
        pages.append({**at,'status':'passed' if not issues else 'blocked','retention':retention,
                      'sourceTypography':src,'candidateTypography':out,
                      'typographyStatus':'not_applicable' if not src['visibleRunCount'] and not out['visibleRunCount'] else 'evaluated',
                      'findings':issues})
        findings.extend(issues)
    return {'schema_version':'1.0.0','status':'passed' if not findings else 'blocked',
            'scope':'native_page_decoration_retention_and_native_text_type_scale_only',
            'thresholds':{'retentionMinimum':RETENTION_MIN,'headingSizeMinimumRatio':HEADING_SIZE_MIN,
                          'smallestSizeMinimumRatio':SMALLEST_SIZE_MIN,'geometryTolerancePt':BOX_TOLERANCE_PT},
            'pages':pages,'findings':findings,
            'limitations':['Unknown evidence blocks verify; external observations do not automatically waive missing machine evidence',
                           'Native text frames only; chart labels and table text are not included in this typography metric',
                           'Rotated/flipped group geometry is unresolved; immutable layout/master styles use the separate preservation floor',
                           'Source small print is not a valid body-style choice merely because this threshold passes; external review must assess roles and readability',
                           'Not the complete Web visual floor, aesthetic judgment, fact checking or whole-Skill acceptance']}
