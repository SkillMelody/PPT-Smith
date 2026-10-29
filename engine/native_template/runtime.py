"""Hash-bound, revision-locked workspace lifecycle and fixed operations."""
from __future__ import annotations
from copy import deepcopy
import io
import re
from pathlib import Path
from pptx import Presentation
from .safety import (require, fields, strict_json, encode, sha, read_file, atomic,
                     locked, package, safe_path)
from .package import slide_inventory
from .objects import build, inspect_slide

OPERATIONS = {
    'add_page': {'op', 'id', 'sourceSlideId', 'topicIds', 'notes'},
    'remove_page': {'op', 'pageId'},
    'reorder_pages': {'op', 'ids'},
    'set_page': {'op', 'pageId', 'notes', 'topicIds'},
    'edit_native': {'op', 'pageId', 'objectId', 'text', 'rows', 'chart', 'options'},
    'update_object': {'op', 'pageId', 'objectId', 'text', 'rows', 'chart', 'options'},
    'remove_object': {'op', 'pageId', 'objectId'},
    'add_object': {'op', 'pageId', 'objectId', 'type', 'styleSourceObjectId', 'text', 'rows', 'chart', 'options'},
}


def identifier(value):
    require(isinstance(value, str) and re.fullmatch(r'[A-Za-z0-9_-]{1,80}', value), 'INVALID_ID')


def constraints(value):
    fields(value, {'mode', 'count'}, {'mode'})
    require(value['mode'] in {'auto', 'target', 'exact', 'maximum'}, 'INVALID_PAGE_CONSTRAINT_MODE')
    if value['mode'] == 'auto': require('count' not in value, 'AUTO_HAS_NO_COUNT')
    else: require(type(value.get('count')) is int and 1 <= value['count'] <= 100, 'INVALID_PAGE_CONSTRAINT_COUNT')


def page_count(state):
    n = len(state['pages']); c = state['pageConstraints']
    require(0 < n <= 100, 'EXPORT_REQUIRES_1_TO_100_PAGES')
    if c['mode'] == 'exact': require(n == c['count'], 'EXACT_PAGE_COUNT_NOT_MET')
    if c['mode'] == 'maximum': require(n <= c['count'], 'MAXIMUM_PAGE_COUNT_EXCEEDED')


def state_hash(state): return sha(encode(state))


def load(root):
    state = strict_json(read_file(root / 'state.json'))
    fields(state, {'version', 'revision', 'templateHash', 'templatePath', 'pageConstraints', 'pages', 'lastChange', 'pageBindings'},
           {'version', 'revision', 'templateHash', 'templatePath', 'pageConstraints', 'pages', 'lastChange'})
    require(state['version'] == 1 and type(state['revision']) is int and state['revision'] >= 0, 'INVALID_STATE')
    template = read_file(root / 'template.pptx')
    require(sha(template) == state['templateHash'], 'FROZEN_TEMPLATE_HASH_MISMATCH')
    # templatePath is provenance only; workspace/template.pptx is the frozen source.
    constraints(state['pageConstraints'])
    require(isinstance(state['pages'], list) and len(state['pages']) <= 100, 'PAGE_LIMIT')
    return state, template


def init(template_path, workspace, page_constraints=None):
    data = read_file(template_path)
    parts = package(data)
    inv = slide_inventory(parts)
    require(inv, 'TEMPLATE_HAS_NO_SLIDES')
    pres = Presentation(io.BytesIO(data))
    require(0 < pres.slide_width <= 60 * 914400 and 0 < pres.slide_height <= 60 * 914400, 'SLIDE_SIZE_LIMIT')
    for i, (sid, _) in enumerate(inv): inspect_slide(pres.slides[i], sid)
    c = {'mode': 'auto'} if page_constraints is None else page_constraints; constraints(c)
    with locked(workspace, create=True) as root:
        require(not (root / 'state.json').exists() and not (root / 'template.pptx').exists(), 'WORKSPACE_ALREADY_INITIALIZED')
        state = {'version': 1, 'revision': 0, 'templateHash': sha(data),
                 'templatePath': str(safe_path(template_path)), 'pageConstraints': c,
                 'pages': [], 'pageBindings': {}, 'lastChange': {'changed': [], 'unchanged': []}}
        atomic(root / 'template.pptx', data)
        atomic(root / 'state.json', encode(state))
        return {'ok': True, 'status': 'initialized', **summary(state), 'sourceSlideIds': [x[0] for x in inv]}


def summary(state):
    return {'revision': state['revision'], 'templateHash': state['templateHash'],
            'pages': [{k: v for k, v in p.items() if k != 'edits'} for p in state['pages']],
            'pageConstraints': state['pageConstraints'], 'lastChange': state['lastChange']}


def inspect(workspace, offset=0, limit=20, details=False):
    require(type(details) is bool, 'INVALID_DETAILS_FLAG')
    require(type(offset) is int and offset >= 0 and type(limit) is int and 1 <= limit <= 100, 'INVALID_PAGINATION')
    with locked(workspace) as root:
        state, template = load(root)
        inv = slide_inventory(package(template)); pres = Presentation(io.BytesIO(template))
        source = [inspect_slide(pres.slides[i], sid, details=details) for i, (sid, _) in enumerate(inv) if offset <= i < offset + limit]
        current = []
        if state['pages']:
            deck = Presentation(io.BytesIO(build(template, state)))
            current = [inspect_slide(deck.slides[i], p['sourceSlideId'], p['id'], details=details) for i, p in enumerate(state['pages']) if offset <= i < offset + limit]
        return {'ok': True, 'revision': state['revision'], 'sourceSlides': source, 'pages': current,
                'details': details, 'offset': offset, 'limit': limit, 'totalSourceSlides': len(inv), 'totalPages': len(state['pages']),
                'nextOffset': offset + limit if offset + limit < max(len(inv), len(state['pages'])) else None}


def validate_operation(op):
    require(isinstance(op, dict) and op.get('op') in OPERATIONS, 'UNSUPPORTED_OPERATION')
    name = op['op']
    required = {'op'} | ({'id', 'sourceSlideId'} if name == 'add_page' else {'ids'} if name == 'reorder_pages' else {'pageId'})
    if name in {'edit_native', 'update_object', 'remove_object', 'add_object'}: required.add('objectId')
    if name == 'add_object': required.update({'type', 'styleSourceObjectId'})
    fields(op, OPERATIONS[name], required)
    if 'pageId' in op: identifier(op['pageId'])
    if 'notes' in op: require(isinstance(op['notes'], str) and len(op['notes']) <= 100000, 'INVALID_NOTES')
    if 'topicIds' in op: require(isinstance(op['topicIds'], list) and all(isinstance(t, str) and 0 < len(t) <= 200 for t in op['topicIds']), 'INVALID_TOPIC_IDS')
    if name in {'edit_native', 'update_object'}: require(any(k in op for k in ('text', 'rows', 'chart', 'options')), 'EMPTY_OBJECT_EDIT')


def apply(workspace, request):
    fields(request, {'expectedRevision', 'operations'}, {'expectedRevision', 'operations'})
    require(type(request['expectedRevision']) is int, 'INVALID_EXPECTED_REVISION')
    require(isinstance(request['operations'], list) and 0 < len(request['operations']) <= 1000, 'OPERATION_LIMIT')
    # Also validate programmatic callers, including non-finite floats.
    strict_json(encode(request))
    for op in request['operations']: validate_operation(op)
    with locked(workspace) as root:
        old, template = load(root)
        require(old['revision'] == request['expectedRevision'], 'REVISION_CONFLICT')
        state = deepcopy(old); source_ids = dict(slide_inventory(package(template)))
        bindings = state.setdefault('pageBindings', {p['id']: p['sourceSlideId'] for p in state['pages']})
        for op in request['operations']:
            name = op['op']
            pages = state['pages']; by_id = {p['id']: p for p in pages}
            if name == 'add_page':
                identifier(op['id'])
                require(op['id'] not in by_id, 'DUPLICATE_PAGE_ID')
                require(isinstance(op['sourceSlideId'], str) and op['sourceSlideId'] in source_ids, 'UNKNOWN_SOURCE_SLIDE_ID')
                require(bindings.get(op['id'], op['sourceSlideId']) == op['sourceSlideId'], 'PAGE_SOURCE_BINDING_IMMUTABLE')
                bindings[op['id']] = op['sourceSlideId']
                page = {'id': op['id'], 'sourceSlideId': op['sourceSlideId'], 'edits': []}
                for k in ('topicIds', 'notes'):
                    if k in op: page[k] = deepcopy(op[k])
                pages.append(page)
            elif name == 'reorder_pages':
                ids = op['ids']
                require(isinstance(ids, list) and all(isinstance(i, str) for i in ids) and len(ids) == len(pages) and len(set(ids)) == len(ids) and set(ids) == set(by_id), 'REORDER_REQUIRES_ALL_PAGES_ONCE')
                state['pages'] = [by_id[i] for i in ids]
            else:
                require(op['pageId'] in by_id, 'UNKNOWN_PAGE_ID')
                page = by_id[op['pageId']]
                if name == 'remove_page': pages.remove(page)
                elif name == 'set_page':
                    require('notes' in op or 'topicIds' in op, 'EMPTY_PAGE_EDIT')
                    for k in ('notes', 'topicIds'):
                        if k in op: page[k] = deepcopy(op[k])
                else: page['edits'].append(deepcopy(op))
            require(len(state['pages']) <= 100, 'PAGE_LIMIT')
            if state['pageConstraints']['mode'] == 'maximum': require(len(state['pages']) <= state['pageConstraints']['count'], 'MAXIMUM_PAGE_COUNT_EXCEEDED')
        require(sum(len(p['edits']) for p in state['pages']) <= 10000, 'ACCUMULATED_EDIT_LIMIT')
        require(len(encode(state)) <= 8 * 1024 * 1024, 'STATE_SIZE_LIMIT')
        if state['pages']: build(template, state)  # No state writes before every edit has succeeded.
        state['revision'] += 1
        old_pages = {p['id']: p for p in old['pages']}; new_pages = {p['id']: p for p in state['pages']}
        changed = sorted(i for i in old_pages.keys() | new_pages.keys() if old_pages.get(i) != new_pages.get(i))
        # Position is an observable dependency, even if the page's native content is unchanged.
        old_positions = {p['id']: i for i, p in enumerate(old['pages'])}
        changed = sorted(set(changed) | {p['id'] for i, p in enumerate(state['pages']) if old_positions.get(p['id']) != i})
        state['lastChange'] = {'changed': changed, 'unchanged': sorted(set(new_pages) - set(changed)), 'reviewReuse': False}
        require(len(encode(state)) <= 8 * 1024 * 1024, 'STATE_SIZE_LIMIT')
        atomic(root / 'state.json', encode(state))
        return {'ok': True, 'status': 'edited', **summary(state)}


def export(workspace):
    with locked(workspace) as root:
        state, template = load(root); page_count(state)
        candidate = build(template, state)
        from .floor import check
        floor = check(template, candidate, state)
        from .fidelity import check as check_fidelity
        fidelity = check_fidelity(template, candidate, state)
        report = {'ok': True, 'status': 'candidate_unreviewed', 'revision': state['revision'],
                  'stateHash': state_hash(state), 'templateHash': state['templateHash'], 'candidateHash': sha(candidate),
                  'candidate': str(root / 'candidate.pptx'), 'pageIds': [p['id'] for p in state['pages']],
                  'sourceSlideIds': [p['sourceSlideId'] for p in state['pages']],
                  'qualityFloor': 'passed' if floor['status'] == fidelity['status'] == 'passed' else 'blocked', 'preservationFloor': floor, 'templateFidelity': fidelity, 'scope': 'native_template_toolchain_only', 'lastChange': state['lastChange'],
                  'compatibility': {'powerpoint': 'not_tested', 'wps': 'not_tested'},
                  'automaticFactChecking': False,
                  'targetCountMet': state['pageConstraints']['mode'] != 'target' or len(state['pages']) == state['pageConstraints']['count'],
                  'limitations': ['No automatic factual verification', 'Office/WPS compatibility not actually verified', 'External review identity is declared, not authenticated']}
        atomic(root / 'candidate.pptx', candidate)
        atomic(root / 'candidate-report.json', encode(report))
        return report


def current_candidate(root, state):
    report = strict_json(read_file(root / 'candidate-report.json'))
    require(report['revision'] == state['revision'] and report['stateHash'] == state_hash(state), 'STALE_CANDIDATE_REVISION')
    require(report['templateHash'] == state['templateHash'], 'STALE_CANDIDATE_TEMPLATE')
    candidate_bytes = read_file(root / 'candidate.pptx')
    require(report['candidateHash'] == sha(candidate_bytes), 'CANDIDATE_HASH_MISMATCH')
    page_count(state)
    # Recompute from frozen source and current bytes, not the persisted report.
    from .fidelity import check as check_fidelity
    template_bytes = read_file(root / 'template.pptx')
    require(sha(template_bytes) == state['templateHash'], 'FROZEN_TEMPLATE_HASH_MISMATCH')
    report['templateFidelity'] = check_fidelity(template_bytes, candidate_bytes, state)
    if report['templateFidelity']['status'] != 'passed': report['qualityFloor'] = 'blocked'
    return report


def status(workspace):
    with locked(workspace) as root:
        state, _ = load(root)
        candidate_status = 'not_exported'
        if (root / 'candidate-report.json').exists():
            try:
                current_candidate(root, state); candidate_status = 'candidate_unreviewed'
            except ValueError as e: candidate_status = str(e)
        return {'ok': True, 'status': candidate_status, **summary(state)}
