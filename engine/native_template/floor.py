"""Conservative source-backed preservation checks, not an artistic/factual verdict."""
from __future__ import annotations
import io
from pptx import Presentation
from .objects import walk, kind
from .styles import resolved_text
from .package import slide_inventory
from .safety import package


def _xml(element):
    # Compare element/attributes/text, independent of serialization namespaces.
    return (element.tag, tuple(sorted(element.attrib.items())), element.text or '',
            tuple(_xml(c) for c in element))


def check(template, candidate, state):
    source=Presentation(io.BytesIO(template));output=Presentation(io.BytesIO(candidate))
    indices={sid:i for i,(sid,_) in enumerate(slide_inventory(package(template)))}
    issues=[];pages=[]
    original_parts=package(template);output_parts=package(candidate)
    for name,data in output_parts.items():
        if name.startswith(('ppt/slideLayouts/','ppt/slideMasters/','ppt/theme/','ppt/notesMasters/')) and name in original_parts and data!=original_parts[name]:
            issues.append({'code':'SOURCE_IMMUTABLE_STYLE_PART_CHANGED','part':name})
    for page,slide in zip(state['pages'],output.slides):
        original=source.slides[indices[page['sourceSlideId']]]
        originals={str(s.shape_id):s for s,_ in walk(original.shapes)}
        actual={str(s.shape_id):s for s,_ in walk(slide.shapes)}
        styles={op['objectId']:op['styleSourceObjectId'] for op in page['edits'] if op['op']=='add_object'}
        page_issues=[];decorations=0;text_checks=0
        for sid,s in originals.items():
            if kind(s) in {'picture','shape'} and not (s.has_text_frame and s.text.strip()):
                decorations+=1;c=actual.get(sid)
                if c is None or _xml(s._element)!=_xml(c._element):
                    page_issues.append({'code':'SOURCE_DECORATION_CHANGED_OR_MISSING','objectId':sid})
                elif kind(s)=='picture' and s.image.blob!=c.image.blob:
                    page_issues.append({'code':'SOURCE_PICTURE_BYTES_CHANGED','objectId':sid})
            if kind(s)=='group':
                c=actual.get(sid)
                old=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}grpSpPr')
                new=c._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}grpSpPr') if c is not None else None
                if old is None or new is None or _xml(old)!=_xml(new):
                    page_issues.append({'code':'SOURCE_GROUP_GEOMETRY_CHANGED','objectId':sid})
        for sid,c in actual.items():
            s=originals.get(styles.get(sid,sid))
            if c.has_text_frame and c.text.strip():
                text_checks+=1
                expected=resolved_text(s,original) if s is not None and s.has_text_frame else []
                observed=resolved_text(c,slide)
                sizes={r.get('fontSizePt',{}).get('value') for r in expected}
                actual_sizes=[r.get('fontSizePt',{}).get('value') for r in observed]
                if not sizes or None in sizes or not actual_sizes or any(v is None for v in actual_sizes):
                    page_issues.append({'code':'FONT_SIZE_EVIDENCE_UNRESOLVED','objectId':sid})
                elif any(v not in sizes for v in actual_sizes):
                    page_issues.append({'code':'SOURCE_FONT_SIZE_CHANGED','objectId':sid})
            if c.has_table:
                if s is None or not s.has_table:
                    page_issues.append({'code':'TABLE_STYLE_SOURCE_UNRESOLVED','objectId':sid});continue
                for ri,row in enumerate(c.table.rows):
                    for ci,cell in enumerate(row.cells):
                        if not cell.text.strip(): continue
                        old=s.table.cell(ri,ci)
                        expected={r.font.size for p in old.text_frame.paragraphs for r in p.runs}
                        observed=[r.font.size for p in cell.text_frame.paragraphs for r in p.runs]
                        if not expected or None in expected or not observed or any(v is None or v not in expected for v in observed):
                            page_issues.append({'code':'TABLE_FONT_SIZE_EVIDENCE_UNRESOLVED_OR_CHANGED','objectId':sid,'cell':[ri,ci]})
        if not isinstance(page.get('notes'),str) or not page['notes'].strip():
            page_issues.append({'code':'OUTPUT_PAGE_NOTES_REQUIRED'})
        pages.append({'pageId':page['id'],'sourceSlideId':page['sourceSlideId'],'decorationsChecked':decorations,
                      'nativeTextObjectsChecked':text_checks,'status':'passed' if not page_issues else 'blocked','issues':page_issues})
        issues.extend({'pageId':page['id'],**i} for i in page_issues)
    return {'status':'passed' if not issues else 'blocked','pages':pages,'issues':issues,
            'webVisualFloor':'not_implemented','artisticHierarchyAssessment':'external_review_required',
            'scope':['immutable_source_style_parts','source_decoration_and_picture_retention','source_group_geometry_retention',
                     'source_native_text_font_size_retention','table_explicit_font_size_retention','output_page_notes_presence'],
            'limitations':['Conservative exact source retention; intentional decoration changes need an unsupported exception workflow',
                           'Not the complete Web visual floor; no artistic hierarchy or chart-label font assessment',
                           'No verification of notes truth, citations, topic coverage, business facts or reviewer identity']}
