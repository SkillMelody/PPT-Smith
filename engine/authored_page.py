"""Verification for model-authored high-fidelity pages.

The high-fidelity route gives a capable model full freedom over geometry,
shape composition, charts and visual language. In exchange, every piece of
*business text* it renders must declare which Presentation IR content it
comes from, using the shape name as a binding channel:

    bind:slide:<slide_id>:message
    bind:slide:<slide_id>:title
    bind:block:<slide_id>:<block_id>:<field>     field: text|label|value|item
    bind:chart:<slide_id>                        native chart data workbook
    bind:chart:<slide_id>:<chart_id>             one chart in a multi-chart page
    bind:source:<slide_id>:<block_id>            source/citation line
    decoration:<name>                            purely visual, no business text

Unbound business text, bindings to non-existent IR content, and text that
does not match the bound IR value are rejected. Geometry is never inspected
here: layout freedom is the point of this route.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

_BINDABLE_FIELDS = {
    "text", "label", "value", "display_value", "unit", "baseline",
    "delta", "caption", "item", "detail",
}


def _normalized(text: str) -> str:
    return " ".join((text or "").replace("\u00a0", " ").split())


def _slide_map(ir: dict) -> dict[str, dict]:
    return {slide.get("id"): slide for slide in ir.get("slides", []) if slide.get("id")}


def _block_map(slide: dict) -> dict[str, dict]:
    return {block.get("id"): block for block in slide.get("blocks", []) if block.get("id")}


def _chart_map(slide: dict) -> dict[str, dict]:
    charts = slide.get("charts", [])
    if not isinstance(charts, list):
        return {}
    return {
        chart.get("id"): chart
        for chart in charts
        if isinstance(chart, dict) and chart.get("id")
    }


def _iter_shapes(shapes) -> list:
    """Flatten group shapes so grouping cannot bypass content binding."""
    flattened = []
    for shape in shapes:
        flattened.append(shape)
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            flattened.extend(_iter_shapes(shape.shapes))
    return flattened


def _expected_values(slide: dict, block: dict, field: str) -> list[str]:
    if field == "text":
        return [_normalized(block.get("text", ""))]
    if field == "label":
        return [_normalized(block.get("label", ""))]
    if field == "value":
        return [_normalized(block.get("value", ""))]
    if field == "display_value":
        value = _normalized(block.get("value", ""))
        unit = _normalized(block.get("unit", ""))
        if not value:
            return []
        if not unit:
            return [value]
        return [value + unit, f"{value} {unit}"]
    if field in {"unit", "baseline", "delta", "caption"}:
        return [_normalized(block.get(field, ""))]
    if field == "detail":
        return [_normalized(item.get("detail", "")) for item in block.get("items", []) or []]
    if field == "item":
        return [_normalized(item.get("text", "")) for item in block.get("items", []) or []]
    return []


def _source_values(block: dict) -> list[str]:
    refs = []
    if block.get("source_ref"):
        refs.append(block["source_ref"])
    refs.extend(block.get("source_refs", []) or [])
    values = []
    for ref in refs:
        source_id = _normalized(ref.get("source_id", ""))
        loc = _normalized(ref.get("loc", ""))
        if source_id and loc:
            values.append(f"Source: {source_id} {loc}")
    return values


def _required_content_ids(ir: dict, slide_ids: list[str]) -> set[str]:
    required: set[str] = set()
    for slide_id in slide_ids:
        slide = _slide_map(ir).get(slide_id)
        if slide is None:
            continue
        if _normalized(slide.get("title", "")):
            required.add(f"slide:{slide_id}:title")
        if _normalized(slide.get("message", "")):
            required.add(f"slide:{slide_id}:message")
        if isinstance(slide.get("chart"), dict):
            required.add(f"chart:{slide_id}")
        for chart_id in _chart_map(slide):
            required.add(f"chart:{slide_id}:{chart_id}")
        for block in slide.get("blocks", []):
            if block.get("id"):
                required.add(f"{slide_id}:{block['id']}")
    return required


def _chart_data(shape) -> dict:
    """Read the editable data workbook embedded in a native chart shape."""
    chart = shape.chart
    categories = []
    if chart.plots:
        categories = [str(category.label) for category in chart.plots[0].categories]
    return {
        "categories": categories,
        "series": [
            {"name": str(series.name), "values": [float(value) for value in series.values]}
            for series in chart.series
        ],
    }


def _expected_chart_data(chart: dict) -> dict:
    return {
        "categories": [str(value) for value in chart["data"]["categories"]],
        "series": [
            {"name": str(series["name"]), "values": [float(value) for value in series["values"]]}
            for series in chart["data"]["series"]
        ],
    }


def _expected_table_rows(block: dict, source_docs: dict | None) -> list[list[str]] | None:
    ref = block.get("source_ref", {})
    doc = (source_docs or {}).get(ref.get("source_id"))
    element = doc.get(ref.get("loc", "")) if doc is not None else None
    if element is None or element.etype != "table":
        return None
    rows = [list(row) for row in element.rows]
    columns = block.get("column_filter")
    if columns:
        rows = [[row[index] for index in columns if index < len(row)] for row in rows]
    header, data_rows = rows[:1], rows[1:]
    offset = max(int(block.get("row_offset", 0)), 0)
    limit = min(int(block.get("row_limit", 12)), 30)
    return header + data_rows[offset:offset + limit]


def _table_rows(shape) -> list[list[str]]:
    return [[_normalized(cell.text) for cell in row.cells] for row in shape.table.rows]


def verify_authored_page(pptx_path: str | Path, ir: dict, *, slide_ids: list[str],
                         require_full_coverage: bool = False, source_docs: dict | None = None,
                         physical_slide_indices: set[int] | None = None,
                         declared_binding_names: set[str] | None = None) -> dict:
    """Verify declared IR bindings on a model-authored deck.

    Returns a machine-readable report. ``pass`` means every business text is
    bound to real IR content and matches it; it does not mean the page is
    visually approved.
    """
    slides = _slide_map(ir)
    prs = Presentation(str(pptx_path))
    unbound: list[str] = []
    invalid: list[dict] = []
    mismatched: list[dict] = []
    unbound_charts: list[str] = []
    unbound_tables: list[str] = []
    covered: set[str] = set()
    valid_bindings: set[str] = set()

    for physical_slide_index, slide in enumerate(prs.slides, 1):
        if physical_slide_indices is not None and physical_slide_index not in physical_slide_indices:
            continue
        for shape in _iter_shapes(slide.shapes):
            name = str(getattr(shape, "name", "") or "")
            # Strict template candidates begin with fully authored template
            # pages.  Only explicitly declared native operations belong to
            # this delivery's source-binding contract; inherited objects are
            # preserved rather than silently treated as model-authored text.
            if declared_binding_names is not None and name not in declared_binding_names:
                continue
            if getattr(shape, "shape_type", None) in {
                MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE,
            }:
                parts = name.split(":")
                if len(parts) != 4 or parts[:2] != ["bind", "image"]:
                    invalid.append({"binding": name, "reason": "unknown image binding"})
                    continue
                slide_id, block_id = parts[2], parts[3]
                target = slides.get(slide_id)
                block = _block_map(target).get(block_id) if target else None
                if block is None or block.get("role") != "image":
                    invalid.append({"binding": name, "reason": "unknown image block"})
                    continue
                crop = block.get("crop_audit")
                contaminated = (
                    block.get("asset_kind") == "page_screenshot"
                    or not isinstance(crop, dict)
                    or any(crop.get(field) is True for field in (
                        "includes_page_header", "includes_page_footer",
                        "includes_navigation", "includes_body_prose",
                    ))
                )
                if contaminated:
                    invalid.append({
                        "binding": name,
                        "reason": "image is not an audited standalone source asset",
                    })
                    continue
                covered.add(f"{slide_id}:{block_id}")
                valid_bindings.add(f"image:{slide_id}:{block_id}")
                continue
            if getattr(shape, "has_table", False):
                parts = name.split(":")
                if len(parts) != 4 or parts[:2] != ["bind", "table"]:
                    unbound_tables.append(name or "<unnamed native table>")
                    continue
                slide_id, block_id = parts[2], parts[3]
                target = slides.get(slide_id)
                block = _block_map(target).get(block_id) if target else None
                if block is None or block.get("role") != "table":
                    invalid.append({"binding": name, "reason": "unknown table binding"})
                    continue
                expected = _expected_table_rows(block, source_docs)
                if expected is None:
                    invalid.append({"binding": name, "reason": "table source data unavailable"})
                elif _table_rows(shape) != expected:
                    mismatched.append({"binding": name, "expected": expected, "actual": _table_rows(shape)})
                else:
                    covered.add(f"{slide_id}:{block_id}")
                    valid_bindings.add(f"table:{slide_id}:{block_id}")
                continue
            if getattr(shape, "has_chart", False):
                if name.startswith("decoration:"):
                    invalid.append({"binding": name, "reason": "decoration shapes may not contain charts"})
                    continue
                parts = name.split(":")
                if len(parts) not in {3, 4} or parts[:2] != ["bind", "chart"] or not all(parts[2:]):
                    unbound_charts.append(name or "<unnamed native chart>")
                    continue
                slide_id = parts[2]
                target = slides.get(slide_id)
                chart_id = parts[3] if len(parts) == 4 else None
                chart = (
                    _chart_map(target).get(chart_id)
                    if target is not None and chart_id is not None
                    else target.get("chart") if target is not None else None
                )
                if not isinstance(chart, dict):
                    invalid.append({"binding": name, "reason": "unknown chart binding"})
                    continue
                expected = _expected_chart_data(chart)
                actual = _chart_data(shape)
                if actual != expected:
                    mismatched.append({"binding": name, "expected": expected, "actual": actual})
                else:
                    content_id = f"chart:{slide_id}" + (f":{chart_id}" if chart_id is not None else "")
                    covered.add(content_id)
                    valid_bindings.add(content_id)
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            text = _normalized(shape.text or "")
            if not text:
                continue
            if name.startswith("decoration:"):
                invalid.append({
                    "binding": name,
                    "reason": "decoration shapes may not contain text",
                })
                continue
            if not name.startswith("bind:"):
                unbound.append(text)
                continue

            parts = name.split(":")
            kind = parts[1] if len(parts) > 1 else ""
            if kind == "slide" and len(parts) == 4:
                slide_id, field = parts[2], parts[3]
                target = slides.get(slide_id)
                if target is None or field not in {"message", "title"}:
                    invalid.append({"binding": name, "reason": "unknown slide binding"})
                    continue
                expected = _normalized(target.get(field, ""))
                if not expected:
                    invalid.append({"binding": name, "reason": f"slide has no {field}"})
                elif text != expected:
                    mismatched.append({"binding": name, "expected": expected, "actual": text})
                else:
                    covered.add(f"slide:{slide_id}:{field}")
                    valid_bindings.add(f"slide:{slide_id}:{field}")
            elif kind == "block" and len(parts) in {5, 6}:
                slide_id, block_id, field = parts[2], parts[3], parts[4]
                target = slides.get(slide_id)
                block = _block_map(target).get(block_id) if target else None
                if target is None or block is None:
                    invalid.append({"binding": name, "reason": "unknown block binding"})
                    continue
                if field not in _BINDABLE_FIELDS:
                    invalid.append({"binding": name, "reason": f"unsupported field '{field}'"})
                    continue
                if len(parts) == 6:
                    if field not in {"item", "detail"}:
                        invalid.append({"binding": name, "reason": "only item/detail bindings accept an index"})
                        continue
                    try:
                        item_index = int(parts[5])
                    except ValueError:
                        invalid.append({"binding": name, "reason": "item index must be an integer"})
                        continue
                    items = block.get("items", []) or []
                    if item_index < 0 or item_index >= len(items):
                        invalid.append({"binding": name, "reason": "item index is out of range"})
                        continue
                    item_field = "text" if field == "item" else "detail"
                    expected_values = [_normalized(items[item_index].get(item_field, ""))]
                else:
                    expected_values = [value for value in _expected_values(target, block, field) if value]
                expected_values = [value for value in expected_values if value]
                if not expected_values:
                    invalid.append({"binding": name, "reason": f"block has no {field}"})
                elif text not in expected_values:
                    mismatched.append({"binding": name, "expected": expected_values, "actual": text})
                else:
                    covered.add(f"{slide_id}:{block_id}")
                    suffix = f":{parts[5]}" if len(parts) == 6 else ""
                    valid_bindings.add(f"block:{slide_id}:{block_id}:{field}{suffix}")
            elif kind == "source" and len(parts) == 4:
                slide_id, block_id = parts[2], parts[3]
                target = slides.get(slide_id)
                block = _block_map(target).get(block_id) if target else None
                if target is None or block is None:
                    invalid.append({"binding": name, "reason": "unknown source binding"})
                elif not (block.get("source_ref") or block.get("source_refs")):
                    invalid.append({"binding": name, "reason": "bound block has no source reference"})
                else:
                    expected_values = _source_values(block)
                    if text not in expected_values:
                        mismatched.append({
                            "binding": name,
                            "expected": expected_values,
                            "actual": text,
                        })
                    else:
                        valid_bindings.add(f"source:{slide_id}:{block_id}")
            else:
                invalid.append({"binding": name, "reason": "malformed binding"})

    uncovered = sorted(_required_content_ids(ir, slide_ids) - covered) if require_full_coverage else []
    failed = bool(unbound or unbound_charts or unbound_tables or invalid or mismatched or uncovered)
    return {
        "status": "fail" if failed else "pass",
        "slide_ids": list(slide_ids),
        "unbound_text": unbound,
        "unbound_charts": unbound_charts,
        "unbound_tables": unbound_tables,
        "invalid_bindings": invalid,
        "mismatched_text": mismatched,
        "uncovered_content": uncovered,
        "bound_content_count": len(covered),
        "valid_bindings": sorted(valid_bindings),
    }
