"""Resolve reviewed component declarations against a native PPTX template.

The deterministic extractor deliberately does not guess that a cluster of
shapes is a funnel, timeline, or pyramid. A human or visual reviewer declares
the semantic family; this module proves that every declared native object
exists and records the geometry needed by later parameterized authoring.
"""
from __future__ import annotations

from copy import deepcopy
import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .topology import TOPOLOGIES


_GRANULARITIES = {"atomic", "micro", "composite", "section", "page_recipe"}
_NATIVE_FIDELITY = {"exact", "transformed", "style_authored"}
_COMPOSITION_ROLES = {"primary", "supporting", "context", "navigation", "decoration"}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return f"sha256:{digest.hexdigest()}"


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _kind(shape) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        return "chart"
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        return "table"
    if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
        return "image"
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return "connector"
    if getattr(shape, "has_text_frame", False):
        return "text"
    return "shape"


def _frame(shape, slide_width: int, slide_height: int) -> dict:
    return {
        "x": round(shape.left / slide_width, 6),
        "y": round(shape.top / slide_height, 6),
        "w": round(shape.width / slide_width, 6),
        "h": round(shape.height / slide_height, 6),
    }


def _normalized_placement(value: object, *, context: str) -> dict:
    if not isinstance(value, dict):
        raise ValueError(f"{context} requires a placement object")
    result: dict[str, float] = {}
    for key in ("x", "y", "w", "h"):
        item = value.get(key)
        if isinstance(item, bool) or not isinstance(item, (int, float)):
            raise ValueError(f"{context} placement {key} must be a number")
        result[key] = float(item)
    if result["x"] < 0 or result["y"] < 0 or result["w"] <= 0 or result["h"] <= 0:
        raise ValueError(f"{context} placement must have non-negative origin and positive size")
    if result["x"] + result["w"] > 1 or result["y"] + result["h"] > 1:
        raise ValueError(f"{context} placement must remain inside its parent")
    return result


def _resolve_group_declarations(
    declared_groups: object,
    *,
    shape_index: dict[str, object],
    slide_width: int,
    slide_height: int,
    context: str,
) -> list[dict]:
    if not isinstance(declared_groups, list):
        raise ValueError(f"{context} groups must be a list")
    resolved_groups: list[dict] = []
    for group in declared_groups:
        if not isinstance(group, dict):
            raise ValueError(f"{context} has an invalid group")
        role = group.get("role")
        names = group.get("shape_names")
        if not isinstance(role, str) or not role or not isinstance(names, list) or not names:
            raise ValueError(f"{context} has an invalid group")
        scope = group.get("scope")
        if scope is not None and scope not in {"item", "shared"}:
            raise ValueError(f"{context} group {role!r} has invalid scope")
        item_indices = group.get("item_indices")
        if item_indices is not None:
            if scope == "shared":
                raise ValueError(f"{context} shared group {role!r} cannot use item_indices")
            if (
                not isinstance(item_indices, list)
                or len(item_indices) != len(names)
                or any(
                    isinstance(item_index, bool)
                    or not isinstance(item_index, int)
                    or item_index < 0
                    for item_index in item_indices
                )
            ):
                raise ValueError(
                    f"{context} group {role!r} item_indices must align with shape_names"
                )
        bind_field = group.get("bind_field")
        if bind_field is not None:
            if role != "label" or not isinstance(bind_field, str) or not bind_field:
                raise ValueError(f"{context} group {role!r} has invalid bind_field")
        members: list[dict] = []
        for name in names:
            shape = shape_index.get(name)
            if shape is None:
                raise ValueError(f"{context} references missing template shape {name!r}")
            members.append({
                "shape_name": name,
                "shape_id": int(shape.shape_id),
                "kind": _kind(shape),
                "frame": _frame(shape, slide_width, slide_height),
            })
        resolved_group = {"role": role, "members": members}
        if scope is not None:
            resolved_group["scope"] = scope
        if item_indices is not None:
            resolved_group["item_indices"] = deepcopy(item_indices)
        if bind_field is not None:
            resolved_group["bind_field"] = bind_field
        resolved_groups.append(resolved_group)
    return resolved_groups


def _group_contract(groups: list[dict]) -> tuple:
    return tuple(
        (
            group.get("role"),
            group.get("scope"),
            group.get("bind_field"),
            tuple(group.get("item_indices", [])),
            tuple(member.get("kind") for member in group.get("members", [])),
        )
        for group in groups
    )


def build_component_atlas(template_pptx: str | Path, review: dict) -> dict:
    path = Path(template_pptx)
    if not path.is_file():
        raise FileNotFoundError(path)
    if not isinstance(review, dict) or review.get("schema_version") != "1.0.0":
        raise ValueError("component review schema_version must be 1.0.0")
    review_state = review.get("review")
    if not isinstance(review_state, dict) or review_state.get("state") != "reviewed":
        raise ValueError("component atlas requires a reviewed declaration")
    declarations = review.get("components")
    if not isinstance(declarations, list) or not declarations:
        raise ValueError("component review requires a non-empty components list")

    presentation = Presentation(str(path))
    resolved_components: list[dict] = []
    component_ids: set[str] = set()
    for declaration in declarations:
        component_id = declaration.get("component_id")
        if not isinstance(component_id, str) or not component_id:
            raise ValueError("component_id must be a non-empty string")
        if component_id in component_ids:
            raise ValueError(f"duplicate component_id {component_id!r}")
        component_ids.add(component_id)
        slide_index = declaration.get("slide_index")
        if not isinstance(slide_index, int) or not 1 <= slide_index <= len(presentation.slides):
            raise ValueError(f"component {component_id!r} has invalid slide_index")
        slide = presentation.slides[slide_index - 1]
        shape_index = {shape.name: shape for shape in _iter_shapes(slide.shapes)}
        granularity = declaration.get("granularity", "micro")
        if granularity not in _GRANULARITIES:
            raise ValueError(f"component {component_id!r} has invalid granularity")
        semantic_contract = declaration.get("semantic_contract", {})
        if not isinstance(semantic_contract, dict):
            raise ValueError(f"component {component_id!r} semantic_contract must be an object")
        renderer = declaration.get("renderer")
        if renderer is not None and (not isinstance(renderer, str) or not renderer):
            raise ValueError(f"component {component_id!r} renderer must be a non-empty string")
        children = declaration.get("children", [])
        if not isinstance(children, list):
            raise ValueError(f"component {component_id!r} children must be a list")
        resolved_children: list[dict] = []
        child_slots: set[str] = set()
        for child in children:
            if not isinstance(child, dict):
                raise ValueError(f"component {component_id!r} child declarations must be objects")
            slot_id = child.get("slot_id")
            child_component_id = child.get("component_id")
            if not isinstance(slot_id, str) or not slot_id or slot_id in child_slots:
                raise ValueError(f"component {component_id!r} has an invalid or duplicate child slot")
            if not isinstance(child_component_id, str) or not child_component_id:
                raise ValueError(f"component {component_id!r} child {slot_id!r} needs component_id")
            required = child.get("required", True)
            if not isinstance(required, bool):
                raise ValueError(f"component {component_id!r} child {slot_id!r} required must be boolean")
            child_slots.add(slot_id)
            resolved_children.append({
                "slot_id": slot_id,
                "component_id": child_component_id,
                "required": required,
                "placement": _normalized_placement(
                    child.get("placement"), context=f"component {component_id!r} child {slot_id!r}",
                ),
            })

        resolved_groups = _resolve_group_declarations(
            declaration.get("groups", []),
            shape_index=shape_index,
            slide_width=presentation.slide_width,
            slide_height=presentation.slide_height,
            context=f"component {component_id!r}",
        )
        if not resolved_groups and not resolved_children:
            raise ValueError(f"component {component_id!r} requires groups or child components")
        source_instances = declaration.get("source_instances", [])
        if not isinstance(source_instances, list):
            raise ValueError(f"component {component_id!r} source_instances must be a list")
        if source_instances and not resolved_groups:
            raise ValueError(f"component {component_id!r} without native groups cannot have source_instances")
        resolved_source_instances: list[dict] = []
        instance_ids: set[str] = set()
        occupied_members = {
            (slide_index, member["shape_id"])
            for group in resolved_groups
            for member in group["members"]
        }
        primary_contract = _group_contract(resolved_groups)
        for source_instance in source_instances:
            if not isinstance(source_instance, dict):
                raise ValueError(f"component {component_id!r} source instances must be objects")
            instance_id = source_instance.get("instance_id")
            if not isinstance(instance_id, str) or not instance_id or instance_id in instance_ids:
                raise ValueError(f"component {component_id!r} has an invalid or duplicate source instance")
            instance_ids.add(instance_id)
            instance_slide_index = source_instance.get("slide_index", slide_index)
            if (
                isinstance(instance_slide_index, bool)
                or not isinstance(instance_slide_index, int)
                or not 1 <= instance_slide_index <= len(presentation.slides)
            ):
                raise ValueError(
                    f"component {component_id!r} source instance {instance_id!r} has invalid slide_index"
                )
            instance_slide = presentation.slides[instance_slide_index - 1]
            instance_shape_index = {
                shape.name: shape for shape in _iter_shapes(instance_slide.shapes)
            }
            instance_groups = _resolve_group_declarations(
                source_instance.get("groups", []),
                shape_index=instance_shape_index,
                slide_width=presentation.slide_width,
                slide_height=presentation.slide_height,
                context=f"component {component_id!r} source instance {instance_id!r}",
            )
            if _group_contract(instance_groups) != primary_contract:
                raise ValueError(
                    f"component {component_id!r} source instance {instance_id!r} group contract "
                    "does not match the primary component"
                )
            instance_members = {
                (instance_slide_index, member["shape_id"])
                for group in instance_groups
                for member in group["members"]
            }
            if occupied_members & instance_members:
                raise ValueError(
                    f"component {component_id!r} source instance {instance_id!r} reuses an already "
                    "reviewed native shape"
                )
            occupied_members.update(instance_members)
            resolved_source_instances.append({
                "instance_id": instance_id,
                "slide_index": instance_slide_index,
                "groups": instance_groups,
            })
        declared_topologies = declaration.get("topologies", [])
        if not isinstance(declared_topologies, list) or any(
            not isinstance(topology, str) or topology not in TOPOLOGIES
            for topology in declared_topologies
        ):
            raise ValueError(
                f"component {component_id!r} has invalid reviewed topologies"
            )
        archetypes = declaration.get("archetypes", [])
        if not isinstance(archetypes, list) or any(
            not isinstance(archetype, str) or not archetype for archetype in archetypes
        ):
            raise ValueError(f"component {component_id!r} has invalid archetypes")
        composition_roles = declaration.get("composition_roles", [])
        if not isinstance(composition_roles, list) or any(
            role not in _COMPOSITION_ROLES for role in composition_roles
        ):
            raise ValueError(f"component {component_id!r} has invalid composition roles")
        text_capacity = declaration.get("text_capacity", {})
        if not isinstance(text_capacity, dict) or any(
            not isinstance(field, str)
            or isinstance(limit, bool)
            or not isinstance(limit, int)
            or limit < 1
            for field, limit in text_capacity.items()
        ):
            raise ValueError(f"component {component_id!r} has invalid text capacity")
        data_contract = declaration.get("data_contract", {})
        if not isinstance(data_contract, dict):
            raise ValueError(f"component {component_id!r} data contract must be an object")
        native_fidelity = declaration.get("native_fidelity")
        if native_fidelity is not None and native_fidelity not in _NATIVE_FIDELITY:
            raise ValueError(f"component {component_id!r} has invalid native fidelity")
        resolved_component = {
            "component_id": component_id,
            "family": declaration.get("family"),
            "granularity": granularity,
            "slide_index": slide_index,
            "semantic_uses": deepcopy(declaration.get("semantic_uses", [])),
            "topologies": deepcopy(declared_topologies),
            "groups": resolved_groups,
            "parameters": deepcopy(declaration.get("parameters", {})),
        }
        if archetypes:
            resolved_component["archetypes"] = deepcopy(archetypes)
        if composition_roles:
            resolved_component["composition_roles"] = deepcopy(composition_roles)
        if text_capacity:
            resolved_component["text_capacity"] = deepcopy(text_capacity)
        if data_contract:
            resolved_component["data_contract"] = deepcopy(data_contract)
        if native_fidelity is not None:
            resolved_component["native_fidelity"] = native_fidelity
        if semantic_contract:
            resolved_component["semantic_contract"] = deepcopy(semantic_contract)
        if renderer is not None:
            resolved_component["renderer"] = renderer
        if resolved_children:
            resolved_component["children"] = resolved_children
        if resolved_source_instances:
            resolved_component["source_instances"] = resolved_source_instances
        resolved_components.append(resolved_component)

    component_index = {component["component_id"]: component for component in resolved_components}
    for component in resolved_components:
        for child in component.get("children", []):
            if child["component_id"] not in component_index:
                raise ValueError(
                    f"component {component['component_id']!r} references missing child "
                    f"{child['component_id']!r}"
                )

    visiting: set[str] = set()
    visited: set[str] = set()

    def visit(component_id: str) -> None:
        if component_id in visiting:
            raise ValueError(f"component hierarchy contains a cycle at {component_id!r}")
        if component_id in visited:
            return
        visiting.add(component_id)
        for child in component_index[component_id].get("children", []):
            visit(child["component_id"])
        visiting.remove(component_id)
        visited.add(component_id)

    for component_id in component_index:
        visit(component_id)

    return {
        "schema_version": "1.0.0",
        "status": "reviewed",
        "source": {
            "filename": path.name,
            "sha256": _sha256(path),
            "slide_count": len(presentation.slides),
        },
        "review": deepcopy(review_state),
        "components": resolved_components,
    }


def validate_model_authoring_atlas(atlas: dict) -> dict:
    """Require semantic/capacity metadata before final model orchestration."""
    issues: list[dict] = []
    components = atlas.get("components") if isinstance(atlas, dict) else None
    if not isinstance(components, list) or not components:
        return {"status": "fail", "issues": [{"code": "COMPONENT_ATLAS_EMPTY"}]}
    for component in components:
        if not isinstance(component, dict):
            issues.append({"code": "COMPONENT_METADATA_INVALID"})
            continue
        component_id = component.get("component_id")
        required = {
            "topologies": component.get("topologies"),
            "semantic_contract": component.get("semantic_contract"),
            "archetypes": component.get("archetypes"),
            "composition_roles": component.get("composition_roles"),
            "text_capacity": component.get("text_capacity"),
            "native_fidelity": component.get("native_fidelity"),
        }
        for field, value in required.items():
            if value in (None, [], {}):
                issues.append({
                    "code": "COMPONENT_METADATA_INCOMPLETE",
                    "component_id": component_id,
                    "field": field,
                })
        has_chart = any(
            group.get("role") == "chart" for group in component.get("groups", [])
            if isinstance(group, dict)
        )
        if has_chart and not component.get("data_contract"):
            issues.append({
                "code": "CHART_COMPONENT_DATA_CONTRACT_MISSING",
                "component_id": component_id,
            })
    return {"status": "pass" if not issues else "fail", "issues": issues}


def select_component(atlas: dict, requirement: dict) -> dict:
    if not isinstance(atlas, dict) or atlas.get("status") != "reviewed":
        raise ValueError("component selection requires a reviewed atlas")
    if not isinstance(requirement, dict):
        raise ValueError("component requirement must be an object")
    semantic_use = requirement.get("semantic_use")
    element_count = requirement.get("element_count")
    family = requirement.get("family")
    topology = requirement.get("topology")
    required_slots = requirement.get("required_slots", [])
    archetype = requirement.get("archetype")
    composition_role = requirement.get("composition_role")
    text_requirements = requirement.get("text_requirements", {})
    data_shape = requirement.get("data_shape", {"kind": "none"})
    requested_component_id = requirement.get("component_id")
    if not isinstance(semantic_use, str) or not semantic_use:
        raise ValueError("component requirement needs semantic_use")
    if not isinstance(element_count, int) or isinstance(element_count, bool) or element_count < 1:
        raise ValueError("component requirement needs a positive element_count")
    if family is not None and (not isinstance(family, str) or not family):
        raise ValueError("component requirement family must be a non-empty string")
    if topology is not None and (not isinstance(topology, str) or topology not in TOPOLOGIES):
        raise ValueError("component requirement topology must be a supported topology")
    if (
        not isinstance(required_slots, list)
        or any(not isinstance(slot, str) or not slot for slot in required_slots)
    ):
        raise ValueError("component requirement required_slots must be a list of non-empty strings")
    if archetype is not None and (not isinstance(archetype, str) or not archetype):
        raise ValueError("component requirement archetype must be a string")
    if composition_role is not None and composition_role not in _COMPOSITION_ROLES:
        raise ValueError("component requirement composition role is invalid")
    if not isinstance(text_requirements, dict) or any(
        not isinstance(field, str)
        or isinstance(length, bool)
        or not isinstance(length, int)
        or length < 1
        for field, length in text_requirements.items()
    ):
        raise ValueError("component requirement text requirements are invalid")
    if not isinstance(data_shape, dict) or not isinstance(data_shape.get("kind"), str):
        raise ValueError("component requirement data shape is invalid")
    if requested_component_id is not None and (
        not isinstance(requested_component_id, str) or not requested_component_id
    ):
        raise ValueError("component requirement component_id must be a non-empty string")

    matches: list[tuple[int, str, dict, int, int]] = []
    for component in atlas.get("components", []):
        if requested_component_id is not None and component.get("component_id") != requested_component_id:
            continue
        uses = component.get("semantic_uses", [])
        if semantic_use not in uses:
            continue
        if family is not None and component.get("family") != family:
            continue
        if topology is not None and topology not in component.get("topologies", []):
            continue
        if archetype is not None and archetype not in component.get("archetypes", []):
            continue
        if (
            composition_role is not None
            and composition_role not in component.get("composition_roles", [])
        ):
            continue
        available_slots = set(component.get("semantic_contract", {}).get("required_fields", []))
        if not set(required_slots) <= available_slots:
            continue
        count = component.get("parameters", {}).get("element_count", {})
        minimum, maximum = count.get("minimum"), count.get("maximum")
        if not isinstance(minimum, int) or not isinstance(maximum, int):
            continue
        if not minimum <= element_count <= maximum:
            continue
        text_capacity = component.get("text_capacity", {})
        if any(text_capacity.get(field, 0) < length for field, length in text_requirements.items()):
            continue
        component_data = component.get("data_contract", {})
        data_kind = data_shape.get("kind", "none")
        if data_kind not in component_data.get("kinds", ["none"]):
            continue
        if data_kind == "chart":
            chart_type = data_shape.get("chart_type")
            if chart_type and chart_type not in component_data.get("chart_types", []):
                continue
            if data_shape.get("series_count", 0) > component_data.get("max_series", 0):
                continue
            if data_shape.get("category_count", 0) > component_data.get("max_categories", 0):
                continue
        if data_kind == "table":
            if data_shape.get("row_count", 0) > component_data.get("max_rows", 0):
                continue
            if data_shape.get("column_count", 0) > component_data.get("max_columns", 0):
                continue
        matches.append((maximum - element_count, component["component_id"], component, minimum, maximum))

    if not matches:
        return {
            "status": "no_match",
            "reason": (
                "no reviewed component satisfies topology, semantic use and element capacity"
                if topology is not None
                else "no reviewed component satisfies required semantic slots, semantic use and element capacity"
                if required_slots
                else "no reviewed component satisfies semantic use and element capacity"
            ),
        }
    _, _, selected, minimum, maximum = min(matches)
    return {
        "status": "selected",
        "component_id": selected["component_id"],
        "family": selected["family"],
        "reason": (
            f"semantic_use={semantic_use}; element_count={element_count} "
            f"within {minimum}..{maximum}"
        ),
    }


def find_feasible_components(atlas: dict, requirement: dict) -> list[dict]:
    """Return every reviewed component that satisfies the model requirement.

    Unlike ``select_component`` this is an audit surface: the model must see
    the complete feasible set before choosing reuse, composition, or a new
    component.
    """
    if not isinstance(atlas, dict) or atlas.get("status") != "reviewed":
        raise ValueError("component feasibility requires a reviewed atlas")
    if not isinstance(requirement, dict):
        raise ValueError("component requirement must be an object")
    semantic_use = requirement.get("semantic_use")
    element_count = requirement.get("element_count")
    family = requirement.get("family")
    topology = requirement.get("topology")
    required_slots = requirement.get("required_slots", [])
    archetype = requirement.get("archetype")
    composition_role = requirement.get("composition_role")
    text_requirements = requirement.get("text_requirements", {})
    data_shape = requirement.get("data_shape", {"kind": "none"})
    if not isinstance(semantic_use, str) or not semantic_use:
        raise ValueError("component requirement needs semantic_use")
    if (
        isinstance(element_count, bool)
        or not isinstance(element_count, int)
        or element_count < 1
    ):
        raise ValueError("component requirement needs a positive element_count")
    if family is not None and (not isinstance(family, str) or not family):
        raise ValueError("component requirement family must be a non-empty string")
    if topology is not None and topology not in TOPOLOGIES:
        raise ValueError("component requirement topology must be a supported topology")
    if (
        not isinstance(required_slots, list)
        or any(not isinstance(slot, str) or not slot for slot in required_slots)
    ):
        raise ValueError("component requirement required_slots must be strings")
    if archetype is not None and (not isinstance(archetype, str) or not archetype):
        raise ValueError("component requirement archetype must be a string")
    if composition_role is not None and composition_role not in _COMPOSITION_ROLES:
        raise ValueError("component requirement composition role is invalid")
    if not isinstance(text_requirements, dict) or any(
        not isinstance(field, str)
        or isinstance(length, bool)
        or not isinstance(length, int)
        or length < 1
        for field, length in text_requirements.items()
    ):
        raise ValueError("component requirement text requirements are invalid")
    if not isinstance(data_shape, dict) or not isinstance(data_shape.get("kind"), str):
        raise ValueError("component requirement data shape is invalid")

    feasible: list[dict] = []
    for component in atlas.get("components", []):
        if semantic_use not in component.get("semantic_uses", []):
            continue
        if family is not None and component.get("family") != family:
            continue
        if topology is not None and topology not in component.get("topologies", []):
            continue
        if archetype is not None and archetype not in component.get("archetypes", []):
            continue
        if (
            composition_role is not None
            and composition_role not in component.get("composition_roles", [])
        ):
            continue
        available_slots = set(
            component.get("semantic_contract", {}).get("required_fields", [])
        )
        if not set(required_slots) <= available_slots:
            continue
        count = component.get("parameters", {}).get("element_count", {})
        minimum, maximum = count.get("minimum"), count.get("maximum")
        if not isinstance(minimum, int) or not isinstance(maximum, int):
            continue
        if not minimum <= element_count <= maximum:
            continue
        text_capacity = component.get("text_capacity", {})
        if any(text_capacity.get(field, 0) < length for field, length in text_requirements.items()):
            continue
        component_data = component.get("data_contract", {})
        data_kind = data_shape.get("kind", "none")
        supported_kinds = component_data.get("kinds", ["none"])
        if data_kind not in supported_kinds:
            continue
        if data_kind == "chart":
            chart_type = data_shape.get("chart_type")
            if chart_type and chart_type not in component_data.get("chart_types", []):
                continue
            if data_shape.get("series_count", 0) > component_data.get("max_series", 0):
                continue
            if data_shape.get("category_count", 0) > component_data.get("max_categories", 0):
                continue
        if data_kind == "table":
            if data_shape.get("row_count", 0) > component_data.get("max_rows", 0):
                continue
            if data_shape.get("column_count", 0) > component_data.get("max_columns", 0):
                continue
        feasible.append({
            "component_id": component["component_id"],
            "family": component.get("family"),
            "capacity": {"minimum": minimum, "maximum": maximum},
            "capacity_slack": maximum - element_count,
            "granularity": component.get("granularity", "micro"),
            "native_fidelity": component.get("native_fidelity", "reviewed"),
        })
    return sorted(
        feasible,
        key=lambda item: (item["capacity_slack"], item["component_id"]),
    )


def resolve_chart_component_binding(atlas: dict, requirement: dict) -> dict:
    """Resolve one independently reusable native-chart semantic component."""
    selection = select_component(atlas, requirement)
    if selection["status"] != "selected":
        raise ValueError(selection["reason"])
    component = next(
        item for item in atlas.get("components", [])
        if item.get("component_id") == selection["component_id"]
    )
    chart_groups = [group for group in component.get("groups", []) if group.get("role") == "chart"]
    if len(chart_groups) != 1:
        raise ValueError(
            f"component {selection['component_id']!r} requires exactly one chart group"
        )
    chart_members = chart_groups[0].get("members", [])
    if len(chart_members) != 1 or chart_members[0].get("kind") != "chart":
        raise ValueError(
            f"component {selection['component_id']!r} chart group must contain one native chart"
        )
    label_fields: dict[str, list[str]] = {}
    decoration_names: list[str] = []
    for group in component.get("groups", []):
        if group is chart_groups[0]:
            continue
        names = [member.get("shape_name") for member in group.get("members", [])]
        if any(not isinstance(name, str) or not name for name in names):
            raise ValueError(f"component {selection['component_id']!r} has invalid native shape metadata")
        if group.get("role") == "label":
            field = group.get("bind_field", "text")
            if any(member.get("kind") != "text" for member in group.get("members", [])):
                raise ValueError(
                    f"component {selection['component_id']!r} label group must contain only text frames"
                )
            label_fields.setdefault(field, []).extend(names)
        else:
            decoration_names.extend(names)
    if not label_fields:
        raise ValueError(f"component {selection['component_id']!r} requires semantic label fields")
    return {
        "component_id": selection["component_id"],
        "component_type": selection["family"],
        "source_slide_index": component.get("slide_index"),
        "chart_names": [chart_members[0]["shape_name"]],
        "label_fields": label_fields,
        "decoration_names": decoration_names,
    }


def resolve_native_group_binding(atlas: dict, requirement: dict) -> dict:
    """Resolve one reviewed fixed native shape group and its text fields."""
    selection = select_component(atlas, requirement)
    if selection["status"] != "selected":
        raise ValueError(selection["reason"])
    component = next(
        item for item in atlas.get("components", [])
        if item.get("component_id") == selection["component_id"]
    )
    if component.get("renderer") != "native_group":
        raise ValueError("selected component is not a native_group renderer")
    label_fields: dict[str, list[str]] = {}
    decoration_names: list[str] = []
    for group in component.get("groups", []):
        members = group.get("members", [])
        names = [member.get("shape_name") for member in members]
        if any(not isinstance(name, str) or not name for name in names):
            raise ValueError(
                f"component {selection['component_id']!r} has invalid native shape metadata"
            )
        if group.get("role") == "label":
            if any(member.get("kind") != "text" for member in members):
                raise ValueError(
                    f"component {selection['component_id']!r} label groups require text frames"
                )
            field = group.get("bind_field", "text")
            label_fields.setdefault(field, []).extend(names)
        else:
            decoration_names.extend(names)
    if not label_fields:
        raise ValueError(f"component {selection['component_id']!r} requires semantic label fields")
    return {
        "component_id": selection["component_id"],
        "component_type": selection["family"],
        "source_slide_index": component.get("slide_index"),
        "label_fields": label_fields,
        "decoration_names": decoration_names,
    }


def resolve_component_binding(atlas: dict, requirement: dict) -> dict:
    """Resolve a semantic requirement into the native shape contract.

    Strict plans should describe why a component is needed and how many data
    elements it must hold.  The reviewed atlas owns the concrete template
    shape names.  Any same-length visual groups beside ``segment`` (for
    example a funnel's native rim/effect pieces) are paired with their segment
    so parameterized reflow keeps the whole native component intact.
    """
    selection = select_component(atlas, requirement)
    if selection["status"] != "selected":
        raise ValueError(selection["reason"])

    component = next(
        (
            item for item in atlas.get("components", [])
            if item.get("component_id") == selection["component_id"]
        ),
        None,
    )
    if component is None:
        raise ValueError(f"selected component {selection['component_id']!r} is missing from atlas")

    component_groups = component.get("groups", [])
    extended = any(
        any(key in group for key in ("scope", "item_indices", "bind_field"))
        for group in component_groups
    )
    if extended:
        segment_groups = [
            group for group in component_groups
            if group.get("role") == "segment" and group.get("scope", "item") == "item"
        ]
        if len(segment_groups) != 1:
            raise ValueError(
                f"component {selection['component_id']!r} requires exactly one item segment group"
            )

        def names_and_indices(group: dict) -> tuple[list[str], list[int]]:
            members = group.get("members")
            if not isinstance(members, list) or not members:
                raise ValueError(
                    f"component {selection['component_id']!r} has an empty {group.get('role')!r} group"
                )
            names = [member.get("shape_name") for member in members]
            if any(not isinstance(name, str) or not name for name in names):
                raise ValueError(
                    f"component {selection['component_id']!r} has invalid native shape metadata"
                )
            indices = group.get("item_indices")
            if indices is None:
                indices = list(range(len(names)))
            if (
                not isinstance(indices, list)
                or len(indices) != len(names)
                or any(
                    isinstance(item_index, bool)
                    or not isinstance(item_index, int)
                    or item_index < 0
                    for item_index in indices
                )
            ):
                raise ValueError(
                    f"component {selection['component_id']!r} has invalid item_indices"
                )
            return names, indices

        primary_names, primary_indices = names_and_indices(segment_groups[0])
        prototype_count = max(primary_indices) + 1
        if sorted(primary_indices) != list(range(prototype_count)):
            raise ValueError(
                f"component {selection['component_id']!r} segment item_indices must be contiguous and unique"
            )
        item_groups = [
            {"segment_names": [primary_names[index]], "label_fields": {}}
            for index in range(prototype_count)
        ]
        shared_names: list[str] = []

        for group in component_groups:
            if group is segment_groups[0]:
                continue
            role = group.get("role")
            scope = group.get("scope", "item")
            names, indices = names_and_indices(group)
            if scope == "shared":
                if group.get("item_indices") is not None:
                    raise ValueError(
                        f"component {selection['component_id']!r} shared group {role!r} cannot use item_indices"
                    )
                shared_names.extend(names)
                continue
            if any(item_index >= prototype_count for item_index in indices):
                raise ValueError(
                    f"component {selection['component_id']!r} group {role!r} references an unknown item"
                )
            if role == "label":
                field = group.get("bind_field", "text")
                for name, item_index in zip(names, indices):
                    item_groups[item_index]["label_fields"].setdefault(field, []).append(name)
            else:
                for name, item_index in zip(names, indices):
                    item_groups[item_index]["segment_names"].append(name)

        if any(not item["label_fields"] for item in item_groups):
            raise ValueError(
                f"component {selection['component_id']!r} requires labels for every item"
            )
        binding = {
            "component_id": selection["component_id"],
            "component_type": selection["family"],
            "source_slide_index": component.get("slide_index"),
            "binding_schema_version": "2.0.0",
            "binding_mode": "extended",
            "item_groups": item_groups,
            "shared_names": shared_names,
        }
        if component.get("renderer") is not None:
            binding["layout_mode"] = component["renderer"]
        return binding

    groups: dict[str, list[str]] = {}
    for group in component_groups:
        role = group.get("role")
        if not isinstance(role, str) or not role or role in groups:
            raise ValueError(
                f"component {selection['component_id']!r} has invalid or duplicate group roles"
            )
        members = group.get("members")
        if not isinstance(members, list) or not members:
            raise ValueError(f"component {selection['component_id']!r} has an empty {role!r} group")
        names = [member.get("shape_name") for member in members]
        if any(not isinstance(name, str) or not name for name in names):
            raise ValueError(
                f"component {selection['component_id']!r} has invalid native shape metadata"
            )
        groups[role] = names

    segment_names = groups.get("segment")
    label_names = groups.get("label")
    if not segment_names or not label_names:
        raise ValueError(
            f"component {selection['component_id']!r} requires segment and label groups"
        )
    if len(segment_names) != len(label_names):
        raise ValueError(
            f"component {selection['component_id']!r} segment and label groups must align"
        )

    companion_roles = [
        role for role, names in groups.items()
        if role not in {"segment", "label"} and len(names) == len(segment_names)
    ]
    segment_groups = [
        [segment_name, *(groups[role][index] for role in companion_roles)]
        for index, segment_name in enumerate(segment_names)
    ]
    binding = {
        "component_id": selection["component_id"],
        "component_type": selection["family"],
        "source_slide_index": component.get("slide_index"),
        "segment_groups": segment_groups,
        "label_names": label_names,
    }
    if component.get("renderer") is not None:
        binding["layout_mode"] = component["renderer"]
    return binding


def resolve_chart_dashboard_binding(atlas: dict, requirement: dict) -> dict:
    """Resolve a reviewed dashboard into ordered chart and text contracts.

    Unlike geometric components, a chart dashboard does not reflow arbitrary
    segments.  Its value is the reviewed coordination of several editable
    native charts and the explicit inventory of sample text that must be
    rebound or cleared before the page can pass strict delivery.
    """
    selection = select_component(atlas, requirement)
    if selection["status"] != "selected":
        raise ValueError(selection["reason"])
    component = next(
        item for item in atlas.get("components", [])
        if item.get("component_id") == selection["component_id"]
    )
    if component.get("family") != "chart_dashboard":
        raise ValueError("selected component is not a chart_dashboard")
    chart_groups = [group for group in component.get("groups", []) if group.get("role") == "chart"]
    if len(chart_groups) != 1:
        raise ValueError(
            f"component {selection['component_id']!r} requires exactly one chart group"
        )
    chart_group = chart_groups[0]
    members = chart_group.get("members", [])
    chart_names = [member.get("shape_name") for member in members]
    if not chart_names or any(member.get("kind") != "chart" for member in members):
        raise ValueError(
            f"component {selection['component_id']!r} chart group must contain only native charts"
        )
    indices = chart_group.get("item_indices", list(range(len(chart_names))))
    if sorted(indices) != list(range(len(chart_names))):
        raise ValueError(
            f"component {selection['component_id']!r} chart item_indices must be contiguous and unique"
        )
    ordered_charts = [name for _, name in sorted(zip(indices, chart_names))]
    text_names: list[str] = []
    decoration_names: list[str] = []
    for group in component.get("groups", []):
        if group is chart_group:
            continue
        names = [member.get("shape_name") for member in group.get("members", [])]
        if group.get("role") == "text":
            if any(member.get("kind") != "text" for member in group.get("members", [])):
                raise ValueError(
                    f"component {selection['component_id']!r} text group must contain only text frames"
                )
            text_names.extend(names)
        else:
            decoration_names.extend(names)
    return {
        "component_id": selection["component_id"],
        "component_type": "chart_dashboard",
        "source_slide_index": component.get("slide_index"),
        "chart_names": ordered_charts,
        "text_names": text_names,
        "decoration_names": decoration_names,
    }
