from __future__ import annotations

from pathlib import Path
from typing import Sequence

from pptx import Presentation


def write_slide_subset(
    source_pptx: Path,
    output_pptx: Path,
    *,
    slide_numbers: Sequence[int],
) -> dict[str, object]:
    """Write a standalone PPTX containing selected 1-based slides.

    Slides are removed from a copy of the source package so native charts,
    embedded workbooks, relationships, layouts, and editable shapes survive.
    """
    selected = list(slide_numbers)
    if not selected or selected != sorted(set(selected)):
        raise ValueError("slide_numbers must be a non-empty strictly increasing sequence")

    presentation = Presentation(source_pptx)
    source_slide_count = len(presentation.slides)
    if selected[0] < 1 or selected[-1] > source_slide_count:
        raise ValueError(
            f"slide_numbers must stay within 1..{source_slide_count}: {selected}"
        )

    keep_indexes = {number - 1 for number in selected}
    for slide_index in range(source_slide_count - 1, -1, -1):
        if slide_index in keep_indexes:
            continue
        slide_id = presentation.slides._sldIdLst[slide_index]
        presentation.part.drop_rel(slide_id.rId)
        del presentation.slides._sldIdLst[slide_index]

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_pptx)
    return {
        "source_slide_count": source_slide_count,
        "output_slide_count": len(presentation.slides),
        "slide_numbers": selected,
    }
