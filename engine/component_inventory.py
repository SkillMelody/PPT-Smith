"""Inventory complete native-template coverage by reviewed component contracts."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .component_atlas import _kind, _sha256


def _all_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _all_shapes(shape.shapes)


def _leaf_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _leaf_shapes(shape.shapes)
        else:
            yield shape


def _leaf_ids(shape) -> set[int]:
    if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
        return {int(shape.shape_id)}
    return {int(item.shape_id) for item in _leaf_shapes(shape.shapes)}


def _text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return " ".join(shape.text.split())


def build_template_component_inventory(template_pptx: str | Path, atlas: dict) -> dict:
    """Return a slide-by-slide matrix of reviewed and still-unreviewed content."""
    path = Path(template_pptx)
    if not path.is_file():
        raise FileNotFoundError(path)
    if not isinstance(atlas, dict) or atlas.get("status") != "reviewed":
        raise ValueError("component inventory requires a reviewed atlas")
    source_sha = _sha256(path)
    if atlas.get("source", {}).get("sha256") != source_sha:
        raise ValueError("component atlas source hash does not match template")

    presentation = Presentation(str(path))
    slides: list[dict] = []
    totals = {
        "native": 0, "covered": 0, "page_recipe_only": 0,
        "full": 0, "partial": 0, "uncovered": 0,
    }
    for slide_index, slide in enumerate(presentation.slides, 1):
        all_shapes = list(_all_shapes(slide.shapes))
        shape_index = {int(shape.shape_id): shape for shape in all_shapes}
        leaves = list(_leaf_shapes(slide.shapes))
        covered_by: dict[int, set[str]] = {int(shape.shape_id): set() for shape in leaves}
        page_recipe_covered_by: dict[int, set[str]] = {
            int(shape.shape_id): set() for shape in leaves
        }
        slide_components: list[str] = []
        for component in atlas.get("components", []):
            component_id = component.get("component_id")
            coverage_target = (
                page_recipe_covered_by
                if component.get("granularity") == "page_recipe"
                else covered_by
            )
            sources = [{
                "slide_index": component.get("slide_index"),
                "groups": component.get("groups", []),
            }, *component.get("source_instances", [])]
            matching_sources = [
                source for source in sources if source.get("slide_index") == slide_index
            ]
            if not matching_sources:
                continue
            slide_components.append(component_id)
            for source in matching_sources:
                for group in source.get("groups", []):
                    for member in group.get("members", []):
                        shape = shape_index.get(member.get("shape_id"))
                        if shape is None:
                            continue
                        for leaf_id in _leaf_ids(shape):
                            if leaf_id in coverage_target:
                                coverage_target[leaf_id].add(component_id)

        covered = []
        uncovered = []
        for shape in leaves:
            item = {
                "shape_id": int(shape.shape_id),
                "shape_name": shape.name,
                "kind": _kind(shape),
                "text": _text(shape),
            }
            component_ids = sorted(covered_by[int(shape.shape_id)])
            page_recipe_ids = sorted(page_recipe_covered_by[int(shape.shape_id)])
            if page_recipe_ids:
                item["page_recipe_covered_by"] = page_recipe_ids
            if component_ids:
                item["covered_by"] = component_ids
                covered.append(item)
            else:
                uncovered.append(item)
        page_recipe_only_count = sum(
            1 for shape in leaves
            if not covered_by[int(shape.shape_id)]
            and page_recipe_covered_by[int(shape.shape_id)]
        )
        if leaves and not uncovered:
            status = "full"
            totals["full"] += 1
        elif covered:
            status = "partial"
            totals["partial"] += 1
        else:
            status = "uncovered"
            totals["uncovered"] += 1
        totals["native"] += len(leaves)
        totals["covered"] += len(covered)
        totals["page_recipe_only"] += page_recipe_only_count
        slides.append({
            "slide_index": slide_index,
            "status": status,
            "component_ids": sorted(slide_components),
            "native_content_count": len(leaves),
            "covered_content_count": len(covered),
            "uncovered_content_count": len(uncovered),
            "page_recipe_only_content_count": page_recipe_only_count,
            "coverage_ratio": round(len(covered) / len(leaves), 6) if leaves else 1.0,
            "covered": covered,
            "uncovered": uncovered,
        })

    return {
        "schema_version": "1.0.0",
        "source": {
            "filename": path.name,
            "sha256": source_sha,
        },
        "summary": {
            "slide_count": len(slides),
            "native_content_count": totals["native"],
            "covered_content_count": totals["covered"],
            "uncovered_content_count": totals["native"] - totals["covered"],
            "page_recipe_only_content_count": totals["page_recipe_only"],
            "coverage_ratio": round(totals["covered"] / totals["native"], 6) if totals["native"] else 1.0,
            "fully_covered_slides": totals["full"],
            "partially_covered_slides": totals["partial"],
            "uncovered_slides": totals["uncovered"],
        },
        "slides": slides,
    }
