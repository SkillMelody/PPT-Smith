"""Execute a model-authored native component inside a strict Template deck.

The model owns the component's visual encoding and geometry. The runtime only
enforces the contract that makes this freedom deliverable: native objects in
the declared slot, IR-bound business content, no raster/page-screenshot
shortcuts, and the visual tokens observed in the user's template.
"""

from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path
from types import ModuleType
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .template_evidence import analyze_template


_EMU_PER_INCH = 914400


def _load_author_module(script_path: Path) -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        f"pptsmith_template_component_{abs(hash(script_path.resolve()))}", script_path,
    )
    if spec is None or spec.loader is None:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_SCRIPT_LOAD_FAILED")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    if not callable(getattr(module, "build", None)):
        raise ValueError("MODEL_TEMPLATE_COMPONENT_BUILD_REQUIRED")
    return module


def _iter_shapes(shapes: Any) -> list[Any]:
    flattened: list[Any] = []
    for shape in shapes:
        flattened.append(shape)
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            flattened.extend(_iter_shapes(shape.shapes))
    return flattened


def _normalized_placement(value: object) -> dict[str, float]:
    if not isinstance(value, dict):
        raise ValueError("MODEL_TEMPLATE_COMPONENT_PLACEMENT_REQUIRED")
    result: dict[str, float] = {}
    for key in ("x", "y", "w", "h"):
        item = value.get(key)
        if isinstance(item, bool) or not isinstance(item, (int, float)):
            raise ValueError("MODEL_TEMPLATE_COMPONENT_PLACEMENT_INVALID")
        result[key] = float(item)
    if result["x"] < 0 or result["y"] < 0 or result["w"] <= 0 or result["h"] <= 0:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_PLACEMENT_INVALID")
    if result["x"] + result["w"] > 1 or result["y"] + result["h"] > 1:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_PLACEMENT_INVALID")
    return result


def _rgb(color: Any) -> str | None:
    try:
        if color is not None and color.type is not None and color.rgb is not None:
            return f"#{color.rgb}"
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def _shape_colors(shape: Any) -> set[str]:
    colors: set[str] = set()
    for owner in (getattr(shape, "fill", None), getattr(shape, "line", None)):
        try:
            color = _rgb(owner.fore_color) if owner is not None else None
        except (AttributeError, TypeError, ValueError):
            color = None
        if color:
            colors.add(color)
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                color = _rgb(run.font.color)
                if color:
                    colors.add(color)
    return colors


def _shape_fonts(shape: Any) -> set[str]:
    if not getattr(shape, "has_text_frame", False):
        return set()
    return {
        run.font.name
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if isinstance(run.font.name, str) and run.font.name
    }


def _binding_slide_id(name: str) -> str | None:
    parts = name.split(":")
    if len(parts) < 3 or parts[0] != "bind":
        return None
    return parts[2]


def _sha256_bytes(payload: bytes) -> str:
    return "sha256:" + hashlib.sha256(payload).hexdigest()


def _approved_asset_bindings(
    value: object,
    *,
    slide_contract: dict,
    slide_id: str,
) -> dict[str, dict]:
    if value is None:
        value = []
    if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
        raise ValueError("MODEL_TEMPLATE_COMPONENT_ASSETS_INVALID")
    blocks = {
        block.get("id"): block
        for block in slide_contract.get("blocks", [])
        if isinstance(block, dict) and block.get("role") == "image"
    }
    approved: dict[str, dict] = {}
    for item in value:
        name = item.get("binding_name")
        asset_ref = item.get("asset_ref")
        asset_path = item.get("asset_path")
        expected_sha = item.get("sha256")
        if not all(isinstance(field, str) and field for field in (
            name, asset_ref, asset_path, expected_sha,
        )):
            raise ValueError("MODEL_TEMPLATE_COMPONENT_ASSETS_INVALID")
        parts = name.split(":")
        if len(parts) != 4 or parts[:2] != ["bind", "image"] or parts[2] != slide_id:
            raise ValueError("MODEL_TEMPLATE_COMPONENT_IMAGE_BINDING_INVALID")
        block = blocks.get(parts[3])
        if block is None or block.get("asset_ref") != asset_ref:
            raise ValueError("MODEL_TEMPLATE_COMPONENT_IMAGE_IR_MISMATCH")
        crop = block.get("crop_audit")
        if (
            block.get("asset_kind") == "page_screenshot"
            or not isinstance(crop, dict)
            or any(crop.get(field) is True for field in (
                "includes_page_header", "includes_page_footer",
                "includes_navigation", "includes_body_prose",
            ))
        ):
            raise ValueError("MODEL_TEMPLATE_COMPONENT_IMAGE_NOT_STANDALONE")
        path = Path(asset_path)
        if not path.is_file() or _sha256_bytes(path.read_bytes()) != expected_sha:
            raise ValueError("MODEL_TEMPLATE_COMPONENT_IMAGE_HASH_MISMATCH")
        if name in approved:
            raise ValueError("MODEL_TEMPLATE_COMPONENT_IMAGE_BINDING_DUPLICATED")
        approved[name] = {**item, "block_id": parts[3]}
    return approved


def author_model_template_component(
    pptx_path: str | Path,
    *,
    template_pptx: str | Path,
    script_path: str | Path,
    slide_index: int,
    slide_id: str,
    component_id: str,
    placement: dict,
    ir: dict,
    required_binding_names: list[str] | None = None,
    asset_bindings: list[dict] | None = None,
    author_context: dict | None = None,
) -> dict:
    """Run ``build(slide, context)`` and enforce native/style/binding gates."""
    target_path = Path(pptx_path)
    source_template = Path(template_pptx)
    author_script = Path(script_path)
    if not target_path.is_file() or not source_template.is_file():
        raise ValueError("MODEL_TEMPLATE_COMPONENT_PPTX_NOT_FOUND")
    if not author_script.is_file():
        raise ValueError("MODEL_TEMPLATE_COMPONENT_SCRIPT_NOT_FOUND")
    if not isinstance(slide_index, int) or slide_index < 1:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_SLIDE_INVALID")
    if not isinstance(slide_id, str) or not slide_id:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_SLIDE_ID_REQUIRED")
    if not isinstance(component_id, str) or not component_id:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_ID_REQUIRED")
    required = required_binding_names or []
    if not isinstance(required, list) or any(not isinstance(item, str) for item in required):
        raise ValueError("MODEL_TEMPLATE_COMPONENT_BINDINGS_INVALID")
    if author_context is None:
        author_context = {}
    if not isinstance(author_context, dict):
        raise ValueError("MODEL_TEMPLATE_COMPONENT_CONTEXT_INVALID")

    normalized = _normalized_placement(placement)
    presentation = Presentation(str(target_path))
    if slide_index > len(presentation.slides):
        raise ValueError("MODEL_TEMPLATE_COMPONENT_SLIDE_INVALID")
    slide = presentation.slides[slide_index - 1]
    before_ids = {int(shape.shape_id) for shape in _iter_shapes(slide.shapes)}
    evidence = analyze_template(source_template)
    tokens = evidence.get("tokens", {})
    slide_contract = next(
        (
            item for item in ir.get("slides", [])
            if isinstance(item, dict) and item.get("id") == slide_id
        ),
        None,
    )
    if slide_contract is None:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_IR_SLIDE_NOT_FOUND")
    approved_assets = _approved_asset_bindings(
        asset_bindings,
        slide_contract=slide_contract,
        slide_id=slide_id,
    )

    slot = {
        **normalized,
        "x_in": normalized["x"] * presentation.slide_width / _EMU_PER_INCH,
        "y_in": normalized["y"] * presentation.slide_height / _EMU_PER_INCH,
        "w_in": normalized["w"] * presentation.slide_width / _EMU_PER_INCH,
        "h_in": normalized["h"] * presentation.slide_height / _EMU_PER_INCH,
    }
    module = _load_author_module(author_script)
    context = {
        "ir": ir,
        "slide": slide_contract,
        "slide_id": slide_id,
        "component_id": component_id,
        "placement": slot,
        "template_tokens": tokens,
        "asset_bindings": list(approved_assets.values()),
        "author_context": author_context,
        "binding_prefix": f"bind:<kind>:{slide_id}",
        "policy": {
            "native_objects_only": True,
            "pictures": "approved_standalone_source_assets_only",
            "business_content_requires_binding": True,
            "decorations_must_not_contain_text": True,
        },
    }
    try:
        module.build(slide, context)
    except Exception as exc:
        raise ValueError(
            f"MODEL_TEMPLATE_COMPONENT_SCRIPT_FAILED: {type(exc).__name__}: {exc}"
        ) from exc

    all_shapes = _iter_shapes(slide.shapes)
    new_shapes = [shape for shape in all_shapes if int(shape.shape_id) not in before_ids]
    if not new_shapes:
        raise ValueError("MODEL_TEMPLATE_COMPONENT_EMPTY")

    issues: list[dict] = []
    binding_names: list[str] = []
    template_fonts = {
        item["value"] for item in tokens.get("fonts", [])
        if isinstance(item, dict) and isinstance(item.get("value"), str)
    }
    template_colors = {
        str(item["value"]).upper() for item in tokens.get("colors", [])
        if isinstance(item, dict) and isinstance(item.get("value"), str)
    }
    x0, y0 = normalized["x"], normalized["y"]
    x1, y1 = x0 + normalized["w"], y0 + normalized["h"]
    for shape in new_shapes:
        name = str(getattr(shape, "name", "") or "")
        shape_type = getattr(shape, "shape_type", None)
        if shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
            approved = approved_assets.get(name)
            if approved is None:
                issues.append({
                    "code": "MODEL_TEMPLATE_COMPONENT_IMAGE_NOT_APPROVED",
                    "shape": name,
                })
            else:
                actual_sha = _sha256_bytes(shape.image.blob)
                if actual_sha != approved["sha256"]:
                    issues.append({
                        "code": "MODEL_TEMPLATE_COMPONENT_IMAGE_HASH_MISMATCH",
                        "shape": name,
                    })
                else:
                    binding_names.append(name)
        text = " ".join(str(getattr(shape, "text", "") or "").split())
        has_business_content = bool(
            text or getattr(shape, "has_chart", False) or getattr(shape, "has_table", False)
        )
        if name.startswith("decoration:") and text:
            issues.append({"code": "MODEL_TEMPLATE_DECORATION_HAS_TEXT", "shape": name})
        elif has_business_content:
            bound_slide_id = _binding_slide_id(name)
            if bound_slide_id != slide_id:
                issues.append({
                    "code": "MODEL_TEMPLATE_COMPONENT_CONTENT_UNBOUND",
                    "shape": name,
                    "expected_slide_id": slide_id,
                })
            else:
                binding_names.append(name)
        elif not (name.startswith("decoration:") or name.startswith("bind:")):
            issues.append({"code": "MODEL_TEMPLATE_COMPONENT_SHAPE_UNDECLARED", "shape": name})
        if "..." in text or "…" in text:
            issues.append({"code": "MODEL_TEMPLATE_COMPONENT_ELLIPSIS", "shape": name})

        left = float(shape.left) / presentation.slide_width
        top = float(shape.top) / presentation.slide_height
        right = float(shape.left + shape.width) / presentation.slide_width
        bottom = float(shape.top + shape.height) / presentation.slide_height
        if left < x0 - 1e-5 or top < y0 - 1e-5 or right > x1 + 1e-5 or bottom > y1 + 1e-5:
            issues.append({
                "code": "MODEL_TEMPLATE_COMPONENT_OUTSIDE_SLOT",
                "shape": name,
                "frame": {"x": left, "y": top, "w": right - left, "h": bottom - top},
            })
        if template_fonts:
            unexpected_fonts = sorted(_shape_fonts(shape) - template_fonts)
            if unexpected_fonts:
                issues.append({
                    "code": "MODEL_TEMPLATE_COMPONENT_FONT_OUTSIDE_TEMPLATE",
                    "shape": name,
                    "fonts": unexpected_fonts,
                })
        if template_colors:
            unexpected_colors = sorted(
                color for color in _shape_colors(shape) if color.upper() not in template_colors
            )
            if unexpected_colors:
                issues.append({
                    "code": "MODEL_TEMPLATE_COMPONENT_COLOR_OUTSIDE_TEMPLATE",
                    "shape": name,
                    "colors": unexpected_colors,
                })

    missing = sorted(set(required) - set(binding_names))
    if missing:
        issues.append({"code": "MODEL_TEMPLATE_COMPONENT_BINDINGS_MISSING", "bindings": missing})
    if len(binding_names) != len(set(binding_names)):
        issues.append({"code": "MODEL_TEMPLATE_COMPONENT_BINDINGS_DUPLICATED"})
    if issues:
        codes = ",".join(sorted({issue["code"] for issue in issues}))
        raise ValueError(f"MODEL_TEMPLATE_COMPONENT_REJECTED: {codes}")

    presentation.save(str(target_path))
    return {
        "status": "pass",
        "component_id": component_id,
        "slide_id": slide_id,
        "slide_index": slide_index,
        "script": str(author_script),
        "new_shape_count": len(new_shapes),
        "binding_names": sorted(binding_names),
        "template_token_audit": {
            "status": "pass",
            "template_fonts": sorted(template_fonts),
            "template_colors": sorted(template_colors),
        },
        "approved_source_images": sorted(approved_assets),
    }
