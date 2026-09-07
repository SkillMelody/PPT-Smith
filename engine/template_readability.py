"""Portable readability normalization for strict Template delivery."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def enforce_bound_text_floor(pptx_path: str | Path, *, minimum_body_pt: float = 12.0) -> dict:
    prs = Presentation(str(pptx_path))
    adjusted = []
    for slide_index, slide in enumerate(prs.slides, 1):
        for shape in _iter_shapes(slide.shapes):
            name = str(getattr(shape, "name", "") or "")
            if not name.startswith("bind:") or not getattr(shape, "has_text_frame", False):
                continue
            exempt = name.startswith("bind:source:") or name.endswith(":caption") or ":page_no:" in name
            target = 9.0 if exempt else minimum_body_pt
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None and run.font.size.pt + 1e-6 < target:
                        adjusted.append({"slide_index": slide_index, "shape_name": name, "from_pt": round(run.font.size.pt, 2), "to_pt": target})
                        run.font.size = Pt(target)
    prs.save(str(pptx_path))
    return {"status": "adjusted" if adjusted else "unchanged", "adjusted_run_count": len(adjusted), "adjustments": adjusted}


def inspect_template_native_visual_floor(pptx_path: str | Path) -> dict:
    """Apply the shared native-object floor without coupling Template orchestration."""
    from .bespoke_quality import inspect_bespoke_visual_quality

    return inspect_bespoke_visual_quality(pptx_path)
