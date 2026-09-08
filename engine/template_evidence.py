"""Read-only PPTX template evidence extraction.

This module intentionally does not classify pages, create routing rules, or
connect to the deck compiler. Its output is descriptive evidence that remains
``unreviewed`` until a human or visual reviewer interprets it.
"""
from __future__ import annotations

import hashlib
import json
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
import zipfile

import jsonschema
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = ROOT / "schemas" / "v4" / "template-evidence.schema.json"
EMU_PER_IN = 914400
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return f"sha256:{digest.hexdigest()}"


def _rgb(color) -> str | None:
    try:
        if color is not None and color.type is not None and color.rgb is not None:
            return f"#{color.rgb}"
    except Exception:
        pass
    return None


def _fill(shape) -> str | None:
    try:
        if shape.fill.type is not None:
            return _rgb(shape.fill.fore_color)
    except Exception:
        pass
    return None


def _runs(shape) -> list:
    if not getattr(shape, "has_text_frame", False):
        return []
    return [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]


def _text(shape) -> str:
    return str(getattr(shape, "text", "") or "").strip()


def _font_size(shape) -> float:
    return max((run.font.size.pt for run in _runs(shape) if run.font.size), default=0.0)


def _kind(shape) -> str:
    shape_type = shape.shape_type
    if shape_type == MSO_SHAPE_TYPE.CHART:
        return "chart"
    if shape_type == MSO_SHAPE_TYPE.TABLE:
        return "table"
    if shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
        return "image"
    if shape_type == MSO_SHAPE_TYPE.LINE:
        return "connector"
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if _text(shape):
        return "text"
    return "shape"


def _frame(shape, width: int, height: int) -> dict:
    return {
        "x": round(shape.left / width, 5),
        "y": round(shape.top / height, 5),
        "w": round(shape.width / width, 5),
        "h": round(shape.height / height, 5),
    }


def _density(object_count: int) -> str:
    if object_count <= 8:
        return "low"
    if object_count <= 24:
        return "medium"
    return "high"


def _frequencies(counter: Counter, limit: int = 16) -> list[dict]:
    return [{"value": value, "count": count}
            for value, count in counter.most_common(limit) if value not in {None, ""}]


def _number_frequencies(counter: Counter, limit: int = 16) -> list[dict]:
    return [{"value": float(value), "count": count}
            for value, count in counter.most_common(limit)]


def _theme_colors(path: Path) -> list[dict]:
    """Read semantic colour roles from the first Office theme in the PPTX.

    Shape-level RGB extraction misses scheme colours, which caused authored
    Template pages to use a rare local maroon instead of the template's actual
    ``accent1`` red.  Keep the roles explicit so a model-authored component can
    consume the same visual grammar as cloned native components.
    """
    with zipfile.ZipFile(path) as archive:
        theme_names = sorted(
            name for name in archive.namelist()
            if name.startswith("ppt/theme/theme") and name.endswith(".xml")
        )
        if not theme_names:
            return []
        root = ET.fromstring(archive.read(theme_names[0]))
    scheme = root.find(f".//{{{DRAWING_NS}}}clrScheme")
    if scheme is None:
        return []
    result: list[dict] = []
    for role in list(scheme):
        color = next(iter(role), None)
        if color is None:
            continue
        color_kind = color.tag.rsplit("}", 1)[-1]
        value = color.get("lastClr") if color_kind == "sysClr" else color.get("val")
        if value and len(value) == 6 and all(character in "0123456789abcdefABCDEF" for character in value):
            result.append({"role": role.tag.rsplit("}", 1)[-1], "value": f"#{value.upper()}"})
    return result


def validate_template_evidence(report: dict) -> list[str]:
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    errors = sorted(
        jsonschema.Draft202012Validator(schema).iter_errors(report),
        key=lambda error: list(error.path),
    )
    return [f"/{'/'.join(map(str, error.path))}: {error.message}" for error in errors]


def analyze_template(pptx_path: str | Path) -> dict:
    """Return deterministic, private-safe evidence for a reference PPTX."""
    path = Path(pptx_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    presentation = Presentation(str(path))
    width, height = presentation.slide_width, presentation.slide_height
    color_counts: Counter[str] = Counter()
    font_counts: Counter[str] = Counter()
    size_counts: Counter[float] = Counter()
    slides: list[dict] = []

    for slide_index, slide in enumerate(presentation.slides, 1):
        elements: list[dict] = []
        geometry: Counter[str] = Counter()
        title_candidates: list[dict] = []
        occupied = 0.0
        for element_index, shape in enumerate(slide.shapes, 1):
            kind = _kind(shape)
            geometry[kind] += 1
            frame = _frame(shape, width, height)
            occupied += frame["w"] * frame["h"]
            text = _text(shape)
            max_font = round(_font_size(shape), 2)
            fill = _fill(shape)
            if fill:
                color_counts[fill] += 1
            for run in _runs(shape):
                if run.font.name:
                    font_counts[run.font.name] += 1
                if run.font.size:
                    size_counts[round(run.font.size.pt, 2)] += 1
                color = _rgb(run.font.color)
                if color:
                    color_counts[color] += 1
            elements.append({
                "element_id": f"s{slide_index:03d}-e{element_index:03d}",
                "kind": kind,
                "frame": frame,
                "text_length": len(text),
                "max_font_pt": max_font,
                "fill": fill,
            })
            if text and (frame["y"] < 0.25 or max_font >= 20):
                title_candidates.append({"text": text, "max_font_pt": max_font, "frame": frame})

        title_candidates.sort(key=lambda item: (item["frame"]["y"], -item["max_font_pt"], item["text"]))
        slides.append({
            "slide_index": slide_index,
            "object_count": len(elements),
            "density": _density(len(elements)),
            "occupied_area_sum": round(occupied, 5),
            "geometry": dict(sorted(geometry.items())),
            "title_candidates": title_candidates[:8],
            "elements": elements,
        })

    report = {
        "schema_version": "1.0.0",
        "status": "unreviewed",
        "source": {
            "filename": path.name,
            "sha256": _sha256(path),
            "slide_count": len(slides),
            "slide_size_in": {
                "width": round(width / EMU_PER_IN, 3),
                "height": round(height / EMU_PER_IN, 3),
            },
        },
        "tokens": {
            "colors": _frequencies(color_counts),
            "theme_colors": _theme_colors(path),
            "fonts": _frequencies(font_counts),
            "font_sizes_pt": _number_frequencies(size_counts),
        },
        "slides": slides,
        "rhythm": {
            "density_sequence": [slide["density"] for slide in slides],
            "object_count_sequence": [slide["object_count"] for slide in slides],
        },
        "review": {"required": True, "state": "unreviewed"},
        "limitations": [
            "Geometry evidence does not determine page intent or visual quality.",
            "This report is not connected to presentation generation or delivery.",
            "Human or visual-model review is required before deriving reusable design rules.",
        ],
    }
    errors = validate_template_evidence(report)
    if errors:
        raise ValueError("template evidence validation failed: " + "; ".join(errors[:8]))
    return report
