"""Actual-PPTX delivery checks that apply only to the strict Template route."""

from __future__ import annotations

from pathlib import Path
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


_FORBIDDEN_VISIBLE_RE = re.compile(
    r"(?i)(?:\b(?:undefined|null|nan|tbd|todo)\b|\[object\s+object\]|"
    r"lorem\s+ipsum|待补充|待完善|待定|待替换|占位(?:符|文本)?)"
)


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(
        paragraph.text.strip()
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text.strip()
    ).strip()


def _font_floor_exempt(name: str) -> bool:
    return (
        name.startswith("bind:source:")
        or ":page_no:" in name
        or name.endswith(":caption")
    )


def _footer_page_numbers(
    slide,
    slide_index: int,
    slide_width: int,
    slide_height: int,
) -> list[dict]:
    accepted = {str(slide_index), f"{slide_index:02d}"}
    values: list[dict] = []
    for shape in slide.shapes:
        text = _text(shape)
        if text not in accepted:
            continue
        if shape.left < slide_width * 0.75 or shape.top < slide_height * 0.88:
            continue
        values.append({
            "shape_name": str(getattr(shape, "name", "") or ""),
            "text": text,
            "x": round(shape.left / slide_width, 4),
            "y": round(shape.top / slide_height, 4),
            "w": round(shape.width / slide_width, 4),
            "h": round(shape.height / slide_height, 4),
        })
    return values


def inspect_template_delivery_quality(
    pptx_path: str | Path,
    *,
    ir: dict,
    minimum_body_pt: float = 12.0,
) -> dict:
    """Reject unresolved output, duplicate folios, and tiny body text."""
    presentation = Presentation(str(pptx_path))
    contracts = [slide for slide in ir.get("slides", []) if isinstance(slide, dict)]
    issues: list[dict] = []
    slide_reports: list[dict] = []

    for slide_index, slide in enumerate(presentation.slides, 1):
        contract = contracts[slide_index - 1] if slide_index <= len(contracts) else {}
        slide_id = str(contract.get("id") or f"S{slide_index:02d}")
        slide_role = str(contract.get("slide_role") or "content")
        placeholders: list[dict] = []
        small_fonts: list[dict] = []

        for shape in _iter_shapes(slide.shapes):
            text = _text(shape)
            name = str(getattr(shape, "name", "") or "")
            if text and _FORBIDDEN_VISIBLE_RE.search(text):
                placeholders.append({
                    "shape_name": name,
                    "text": text[:160],
                })
            if (
                slide_role in {"cover", "section", "closing"}
                or not text
                or _font_floor_exempt(name)
                or not getattr(shape, "has_text_frame", False)
            ):
                continue
            explicit_sizes = [
                run.font.size.pt
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.font.size is not None and run.text.strip()
            ]
            if explicit_sizes and min(explicit_sizes) + 1e-6 < minimum_body_pt:
                small_fonts.append({
                    "shape_name": name,
                    "minimum_font_pt": round(min(explicit_sizes), 2),
                    "required_font_pt": minimum_body_pt,
                    "text": text[:160],
                })

        page_numbers = _footer_page_numbers(
            slide, slide_index, presentation.slide_width, presentation.slide_height,
        )
        if placeholders:
            issues.append({
                "code": "TEMPLATE_RENDERED_PLACEHOLDER_TEXT",
                "slide_id": slide_id,
                "findings": placeholders,
            })
        if small_fonts:
            issues.append({
                "code": "TEMPLATE_BODY_FONT_BELOW_FLOOR",
                "slide_id": slide_id,
                "findings": small_fonts,
            })
        if len(page_numbers) > 1:
            issues.append({
                "code": "TEMPLATE_DUPLICATE_PAGE_NUMBER",
                "slide_id": slide_id,
                "slide_index": slide_index,
                "page_numbers": page_numbers,
            })
        slide_reports.append({
            "slide_id": slide_id,
            "slide_index": slide_index,
            "placeholder_count": len(placeholders),
            "small_body_text_count": len(small_fonts),
            "footer_page_number_count": len(page_numbers),
        })

    return {
        "status": "pass" if not issues else "fail",
        "minimum_body_pt": minimum_body_pt,
        "slides": slide_reports,
        "issues": issues,
    }
