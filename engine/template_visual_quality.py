"""Archetype-aware visual density checks for Template-route candidates."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


GENERIC_FAMILIES = frozenset({"card_grid", "icon_card_grid"})


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _is_semantic_content_name(name: str) -> bool:
    if name.startswith(("bind:chart:", "bind:table:", "bind:image:")):
        return True
    if not name.startswith("bind:block:"):
        return False
    return ":page_no:" not in name and not name.endswith(":page_no:value")


def _minimum_visible_text(page: dict) -> tuple[int, str]:
    archetype = page.get("archetype", "body")
    if archetype in {"cover", "section", "closing"}:
        return 30, archetype
    if page.get("chart_count", 0) >= 2:
        return 30, "data"
    families = {
        family for family in page.get("families", [])
        if isinstance(family, str)
    }
    if (
        page.get("semantic_element_count", 0) >= 2
        and families
        and not families <= GENERIC_FAMILIES
    ):
        return 50, "relationship"
    return 80, "body"


def evaluate_template_visual_quality(pages: list[dict]) -> dict:
    """Reject visually empty pages without treating charts as prose."""
    issues: list[dict] = []
    page_reports: list[dict] = []
    for index, page in enumerate(pages, 1):
        slide_id = str(page.get("slide_id") or f"S{index:02d}")
        text_chars = page.get("visible_text_chars")
        blank_score = page.get("blank_score")
        if isinstance(text_chars, bool) or not isinstance(text_chars, int) or text_chars < 0:
            issues.append({"code": "VISUAL_PAGE_METRICS_INVALID", "slide_id": slide_id})
            continue
        minimum, density_role = _minimum_visible_text(page)
        if text_chars < minimum:
            issues.append({
                "code": "VISUAL_CONTENT_DENSITY_BELOW_FLOOR",
                "slide_id": slide_id,
                "density_role": density_role,
                "visible_text_chars": text_chars,
                "minimum_visible_text_chars": minimum,
            })
        if isinstance(blank_score, (int, float)) and blank_score > 0.90:
            issues.append({
                "code": "VISUAL_PAGE_TOO_BLANK",
                "slide_id": slide_id,
                "blank_score": round(float(blank_score), 4),
                "maximum_blank_score": 0.90,
            })
        page_reports.append({
            "slide_id": slide_id,
            "density_role": density_role,
            "visible_text_chars": text_chars,
            "minimum_visible_text_chars": minimum,
            "chart_count": int(page.get("chart_count", 0)),
            "semantic_element_count": int(page.get("semantic_element_count", 0)),
            "blank_score": blank_score,
        })
    return {
        "status": "pass" if not issues else "fail",
        "pages": page_reports,
        "issues": issues,
    }


def inspect_template_visual_quality(
    pptx_path: str | Path,
    *,
    ir: dict,
    render_report: dict,
    page_composition: dict | None = None,
) -> dict:
    """Build density evidence from the actual PPTX and real-render report."""
    presentation = Presentation(str(pptx_path))
    ir_slides = [slide for slide in ir.get("slides", []) if isinstance(slide, dict)]
    rendered_by_index = {
        item.get("slide_index"): item
        for item in render_report.get("slides", [])
        if isinstance(item, dict)
    }
    composition_by_id = {
        item.get("slide_id"): item
        for item in (page_composition or {}).get("pages", [])
        if isinstance(item, dict) and isinstance(item.get("slide_id"), str)
    }
    pages: list[dict] = []
    for index, slide in enumerate(presentation.slides, 1):
        contract = ir_slides[index - 1] if index <= len(ir_slides) else {}
        shapes = list(_iter_shapes(slide.shapes))
        text = " ".join(
            str(getattr(shape, "text", "") or "").strip()
            for shape in shapes
            if getattr(shape, "has_text_frame", False)
        )
        intent = contract.get("component_intent", {})
        selected_ids = intent.get("selected_component_ids", [])
        families = (
            selected_ids
            if selected_ids
            else [f"model_authored:{intent.get('new_component_id')}"]
            if intent.get("mode") == "model_authored"
            else []
        )
        slide_role = contract.get("slide_role", "content")
        slide_id = contract.get("id", f"S{index:02d}")
        composition = composition_by_id.get(slide_id, {})
        pages.append({
            "slide_id": slide_id,
            "archetype": "body" if slide_role == "content" else slide_role,
            "visible_text_chars": len("".join(text.split())),
            "chart_count": sum(bool(getattr(shape, "has_chart", False)) for shape in shapes),
            "semantic_element_count": int(intent.get("element_count", 1)),
            "bound_semantic_object_count": sum(
                _is_semantic_content_name(str(getattr(shape, "name", "") or ""))
                for shape in shapes
            ),
            "declared_substantive_module_count": composition.get(
                "substantive_module_count"
            ),
            "declared_information_unit_count": composition.get(
                "information_unit_count"
            ),
            "families": families,
            "blank_score": rendered_by_index.get(index, {}).get("blank_score"),
        })
    report = evaluate_template_visual_quality(pages)
    for page in report["pages"]:
        source = next(item for item in pages if item["slide_id"] == page["slide_id"])
        page.update({
            "bound_semantic_object_count": source["bound_semantic_object_count"],
            "declared_substantive_module_count": source["declared_substantive_module_count"],
            "declared_information_unit_count": source["declared_information_unit_count"],
        })
        if (
            page["density_role"] not in {"cover", "section", "closing"}
            and source["bound_semantic_object_count"] < 2
        ):
            report["issues"].append({
                "code": "VISUAL_PAGE_SINGLE_SEMANTIC_OBJECT",
                "slide_id": page["slide_id"],
                "bound_semantic_object_count": source["bound_semantic_object_count"],
            })
    report["status"] = "pass" if not report["issues"] else "fail"
    return report
