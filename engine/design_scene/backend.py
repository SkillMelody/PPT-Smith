"""Fixed python-pptx backend; no task-supplied code, templates, or file paths."""
from __future__ import annotations

import io
import posixpath
from collections import Counter
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from .contract import validate, walk
from .store import DesignError, decode_image, digest

SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
          "ellipse": MSO_SHAPE.OVAL, "arrow": MSO_SHAPE.RIGHT_ARROW,
          "diamond": MSO_SHAPE.DIAMOND, "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE}
CHARTS = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
          "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}


def _xml(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(k, str(v))
    return e


def _font(font, spec):
    font.name = spec["family"]
    font.size = Pt(spec["size"])
    font.bold = spec.get("bold", False)
    font.italic = spec.get("italic", False)
    font.color.rgb = RGBColor.from_string(spec["color"])
    # Explicit East Asian and complex-script fonts avoid theme substitutions.
    for tag in ("a:ea", "a:cs"):
        q = font._rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
        if q is None:
            q = _xml(tag)
            font._rPr.append(q)
        q.set("typeface", spec["family"])


def _paint(shape, spec):
    style = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        shape._element.remove(style)
    if spec.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(spec["fill"])
    else:
        shape.fill.background()
    if spec.get("stroke"):
        shape.line.color.rgb = RGBColor.from_string(spec["stroke"])
        shape.line.width = Pt(spec.get("stroke_width", 1))
    else:
        shape.line.fill.background()
    sppr = shape._element.spPr
    for effect in list(sppr):
        if etree.QName(effect).localname in {"effectLst", "effectDag"}:
            sppr.remove(effect)
    sppr.append(_xml("a:effectLst"))


def _text(frame, text, node):
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = node.get("wrap", True)
    frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                             "bottom": MSO_ANCHOR.BOTTOM}[node.get("valign", "top")]
    margin = Pt(node.get("margin", 0))
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = margin
    offset = 0
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[node.get("align", "left")]
        p.line_spacing = node.get("line_spacing", 1.15)
        p.space_before = p.space_after = Pt(0)
        _font(p.font, node["font"])
        boundaries = {offset, offset + len(line)}
        for span in node.get("spans", []):
            boundaries.update(v for v in (span["start"], span["end"]) if offset < v < offset + len(line))
        marks = sorted(boundaries)
        for start, end in zip(marks, marks[1:]):
            r = p.add_run()
            r.text = text[start:end]
            selected = next((s["font"] for s in node.get("spans", []) if s["start"] <= start < s["end"]), node["font"])
            _font(r.font, selected)
            r.font._rPr.set("spc", str(round(node.get("letter_spacing", 0) * 100)))
        offset += len(line) + 1


def _path(shape, node):
    pr = shape._element.spPr
    for child in list(pr):
        if etree.QName(child).localname in {"prstGeom", "custGeom"}:
            pr.remove(child)
    geom = _xml("a:custGeom")
    for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
        geom.append(_xml(tag))
    geom.append(_xml("a:rect", l="0", t="0", r="r", b="b"))
    paths = _xml("a:pathLst")
    path = _xml("a:path", w=round(node["box"]["width"] * 1000), h=round(node["box"]["height"] * 1000))
    for command in node["commands"]:
        e = _xml({"M": "a:moveTo", "L": "a:lnTo", "Q": "a:quadBezTo", "C": "a:cubicBezTo", "Z": "a:close"}[command["op"]])
        for x, y in command["points"]:
            e.append(_xml("a:pt", x=round(x * 1000), y=round(y * 1000)))
        path.append(e)
    paths.append(path)
    geom.append(paths)
    pr.insert(1, geom)


def render(task, asset_bytes):
    """Return PPTX bytes. All input bytes have already passed asset ingestion."""
    validate(task, complete=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(task["canvas"]["width"]), Pt(task["canvas"]["height"])
    prs.core_properties.title = task["task_id"]
    prs.core_properties.author = "PPT Smith declarative renderer"
    content = {c["id"]: c for c in task["content"]["items"]}

    def add(shapes, nodes, ox=0, oy=0):
        for n in sorted(nodes, key=lambda n: n["z"]):
            b = n["box"]
            x, y, w, h = [Pt(v) for v in (ox + b["x"], oy + b["y"], b["width"], b["height"])]
            kind = n["type"]
            if kind == "group":
                shape = shapes.add_group_shape()
                add(shape.shapes, n["children"], ox + b["x"], oy + b["y"])
                xf = shape._element.grpSpPr.xfrm
                xf.off.x, xf.off.y = x, y
                xf.ext.cx, xf.ext.cy = w, h
                xf.chOff.x, xf.chOff.y = x, y
                xf.chExt.cx, xf.chExt.cy = w, h
            elif kind == "text":
                shape = shapes.add_textbox(x, y, w, h)
                _paint(shape, {})
                _text(shape.text_frame, content[n["content_id"]]["text"], n)
            elif kind in {"shape", "path"}:
                if kind == "shape" and n["shape"] == "line":
                    shape = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y + h)
                else:
                    shape = shapes.add_shape(SHAPES[n["shape"]] if kind == "shape" else MSO_SHAPE.RECTANGLE, x, y, w, h)
                if kind == "path":
                    _path(shape, n)
                _paint(shape, n["style"])
            elif kind == "image":
                data = asset_bytes[n["asset_id"]]
                im = decode_image(data)
                ratio = im.width / im.height
                if n["fit"] == "contain":
                    nw, nh = (w, int(w / ratio)) if ratio > w / h else (int(h * ratio), h)
                    shape = shapes.add_picture(io.BytesIO(data), x + (w - nw) // 2, y + (h - nh) // 2, nw, nh)
                else:
                    # Encode the actual crop once; never embed the whole reference repeatedly.
                    scale = max(w / im.width, h / im.height)
                    cw, ch = max(1, round(w / scale)), max(1, round(h / scale))
                    left, top = (im.width - cw) // 2, (im.height - ch) // 2
                    cropped = im.crop((left, top, left + cw, top + ch))
                    out = io.BytesIO()
                    cropped.save(out, format="PNG")
                    shape = shapes.add_picture(io.BytesIO(out.getvalue()), x, y, w, h)
            elif kind == "chart":
                c = content[n["content_id"]]
                cd = CategoryChartData()
                cd.categories = c["categories"]
                for series in c["series"]:
                    cd.add_series(series["name"], series["values"])
                shape = shapes.add_chart(CHARTS[n["chart_type"]], x, y, w, h, cd)
                chart = shape.chart
                chart.has_title = False
                chart.has_legend = n["legend"]
                _font(chart.font, n["font"])
                if n["legend"]:
                    chart.legend.position = (XL_LEGEND_POSITION.RIGHT if n.get("legend_position") == "right"
                                             else XL_LEGEND_POSITION.BOTTOM)
                    chart.legend.include_in_layout = False
                    _font(chart.legend.font, n["font"])
                for i, series in enumerate(chart.series):
                    if n["chart_type"] in {"pie", "doughnut"}:
                        targets = list(series.points)
                    else:
                        targets = [series]
                    for j, target in enumerate(targets):
                        color = RGBColor.from_string(n["colors"][(j if len(targets) > 1 else i) % len(n["colors"])])
                        target.format.fill.solid()
                        target.format.fill.fore_color.rgb = color
                        target.format.line.color.rgb = color
                plot = chart.plots[0]
                if n["chart_type"] == "doughnut" and "hole_size" in n:
                    holes = plot._element.xpath("c:holeSize")
                    hole = holes[0] if holes else _xml("c:holeSize")
                    hole.set("val", str(round(n["hole_size"])))
                    if not holes:
                        plot._element.insert_element_before(hole, "c:extLst")
                plot.has_data_labels = n["labels"]
                if n["labels"]:
                    _font(plot.data_labels.font, n["font"])
                    plot.data_labels.show_value = True
                    if "label_number_format" in n:
                        plot.data_labels.number_format = n["label_number_format"]
                        plot.data_labels.number_format_is_linked = False
                    if n["chart_type"] in {"pie", "doughnut"}:
                        plot.data_labels.position = XL_DATA_LABEL_POSITION.BEST_FIT
                        if "label_colors" in n:
                            for j, point in enumerate(chart.series[0].points):
                                _font(point.data_label.font, {**n["font"],
                                      "color": n["label_colors"][j % len(n["label_colors"])]})
                                # Per-point labels override plot defaults in LibreOffice.
                                if "label_number_format" in n:
                                    label = point.data_label._get_or_add_dLbl()
                                    fmt = _xml("c:numFmt", formatCode=n["label_number_format"], sourceLinked="0")
                                    label.insert_element_before(fmt, "c:spPr", "c:txPr", "c:dLblPos", "c:showLegendKey", "c:showVal")
            elif kind == "table":
                cells = content[n["content_id"]]["cells"]
                shape = shapes.add_table(len(cells), len(cells[0]), x, y, w, h)
                table = shape.table
                table.first_row = False
                table.horz_banding = table.vert_banding = False
                for i, width in enumerate(n.get("column_widths", [])):
                    table.columns[i].width = Pt(width)
                for r, row in enumerate(cells):
                    for col, value in enumerate(row):
                        cell = table.cell(r, col)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(n["header_fill"] if r == 0 else n["fill"])
                        font = {**n["font"], **({"color": n["header_color"], "bold": True} if r == 0 else {})}
                        _text(cell.text_frame, value, {"font": font, "margin": 4, "valign": "middle"})
            else:
                raise DesignError("UNSUPPORTED_NODE")
            shape.name = n["id"]

    for page in task["scene"]["pages"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(page["background"])
        add(slide.shapes, page["nodes"])
        notes = [n["text"] for n in task["content"]["notes"] if n["page_id"] == page["id"]]
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)
    out = io.BytesIO()
    prs.save(out)
    result = out.getvalue()
    inspect_package(result)
    return result


def inspect_package(data: bytes, *, workbook=False):
    """Reject active content, external relationships, zip bombs, and unapproved embeds."""
    if len(data) > 128 * 1024 * 1024:
        raise DesignError("PACKAGE_TOO_LARGE")
    try:
        with ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            names = {i.filename for i in infos}
            if len(infos) > 10000 or len(names) != len(infos) or sum(i.file_size for i in infos) > 256 * 1024 * 1024:
                raise DesignError("PACKAGE_RESOURCE_LIMIT")
            for info in infos:
                name = info.filename
                low = name.lower()
                if name.startswith("/") or ".." in name.split("/") or "\\" in name or info.flag_bits & 1:
                    raise DesignError("UNSAFE_PACKAGE_MEMBER")
                if any(x in low for x in ("vbaproject", "activex", "externaldata", "externallinks", "oleobject")):
                    raise DesignError("ACTIVE_PACKAGE_CONTENT")
                if info.file_size > 32 * 1024 * 1024:
                    raise DesignError("PACKAGE_MEMBER_TOO_LARGE")
                payload = z.read(info)
                if "/embeddings/" in low:
                    if workbook or not low.endswith(".xlsx"):
                        raise DesignError("UNAPPROVED_EMBEDDED_OBJECT")
                    inspect_package(payload, workbook=True)
                if low.endswith((".xml", ".rels")):
                    if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
                        raise DesignError("XML_ENTITY_FORBIDDEN")
                    root = etree.fromstring(payload, etree.XMLParser(resolve_entities=False, no_network=True))
                    if low.endswith(".rels"):
                        base = posixpath.dirname(posixpath.dirname(name))
                        for rel in root:
                            target = rel.get("Target", "")
                            if rel.get("TargetMode") == "External" or ":" in target or target.startswith(("/", "\\")):
                                raise DesignError("EXTERNAL_RELATIONSHIP_FORBIDDEN")
                            resolved = posixpath.normpath(posixpath.join(base, target))
                            if resolved.startswith("../") or resolved not in names:
                                raise DesignError("INVALID_PACKAGE_RELATIONSHIP")
                    for element in root.iter():
                        if etree.QName(element).localname in {"oleObj", "control", "videoFile", "audioFile"}:
                            raise DesignError("ACTIVE_PACKAGE_OBJECT")
                    if "macroEnabled" in payload.decode("utf-8"):
                        raise DesignError("MACRO_PACKAGE_FORBIDDEN")
            if workbook and "xl/workbook.xml" not in names:
                raise DesignError("INVALID_CHART_WORKBOOK")
    except DesignError:
        raise
    except Exception as exc:
        raise DesignError(f"INVALID_PACKAGE: {exc}") from exc
    return {"status": "passed", "external_relationships": 0, "macros": 0}


def inspect_objects(data: bytes, task):
    inspect_package(data)
    prs = Presentation(io.BytesIO(data))
    if len(prs.slides) != len(task["scene"]["pages"]):
        raise DesignError("PPTX_PAGE_COUNT_MISMATCH")
    if abs(prs.slide_width.pt - task["canvas"]["width"]) > .01 or abs(prs.slide_height.pt - task["canvas"]["height"]) > .01:
        raise DesignError("PPTX_CANVAS_MISMATCH")
    contents = {c["id"]: c for c in task["content"]["items"]}
    result = []

    def flatten(shapes):
        for shape in shapes:
            yield shape
            if hasattr(shape, "shapes"):
                yield from flatten(shape.shapes)

    for slide, page in zip(prs.slides, task["scene"]["pages"]):
        actual = {s.name: s for s in flatten(slide.shapes)}
        expected = [n for n, *_ in walk(page["nodes"])]
        if set(actual) != {n["id"] for n in expected}:
            raise DesignError("PPTX_OBJECT_MAPPING_MISMATCH")
        counts = Counter()
        for n in expected:
            shape = actual[n["id"]]
            kind = n["type"]
            c = contents.get(n.get("content_id"))
            if kind == "text" and (not shape.has_text_frame or shape.text != c["text"]):
                raise DesignError(f"PPTX_TEXT_MISMATCH: {n['id']}")
            if kind == "table":
                if not shape.has_table or [[cell.text for cell in row.cells] for row in shape.table.rows] != c["cells"]:
                    raise DesignError("PPTX_TABLE_MISMATCH")
            if kind == "chart":
                if not shape.has_chart:
                    raise DesignError("PPTX_CHART_NOT_NATIVE")
                chart = shape.chart
                if [x.label for x in chart.plots[0].categories] != c["categories"]:
                    raise DesignError("PPTX_CHART_CATEGORY_MISMATCH")
                series = [{"name": s.name, "values": list(s.values)} for s in chart.series]
                if series != c["series"]:
                    raise DesignError("PPTX_CHART_DATA_MISMATCH")
            counts[n["role"]] += 1
        result.append({"page_id": page["id"], "objects": dict(counts),
                       "bindings": [{"node_id": n["id"], "content_id": n.get("content_id"), "role": n["role"]} for n in expected]})
    return {"status": "passed", "pages": result, "pptx_sha256": digest(data)}
