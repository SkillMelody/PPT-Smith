"""Execution and delivery gates for the independent Bespoke route.

The author supplies a short Python module with ``build(prs, context)``.  It
starts from an empty ``python-pptx`` presentation; no Standard-route deck or
layout decision is injected.  The runtime locks the module to an approved
authoring manifest and source IR, then applies source-binding, structural,
native-editability and real-render gates.  Passing those gates deliberately
does *not* imply visual approval: that remains a separate human/model review.
"""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path
from types import ModuleType
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ppt_qa.render_report import write_render_report
from ppt_qa.verifier import run_render_report, run_structural_inspection

from .authored_page import verify_authored_page
from .bespoke import _canonical_sha256
from .bespoke_quality import inspect_bespoke_visual_quality


class TemplateReferenceError(ValueError):
    """A required reference template cannot safely reach the author."""


def _write_report(output_dir: Path, report: dict[str, Any]) -> str:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "authoring-report.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return str(path)


def _reject(output_dir: Path, code: str, **extra: Any) -> dict[str, Any]:
    report: dict[str, Any] = {"ok": False, "status": "rejected", "code": code, **extra}
    report["report"] = _write_report(output_dir, report)
    return report


def _load_author_module(script_path: Path) -> ModuleType:
    spec = importlib.util.spec_from_file_location("pptsmith_bespoke_author", script_path)
    if spec is None or spec.loader is None:
        raise ValueError("AUTHOR_SCRIPT_LOAD_FAILED")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    if not callable(getattr(module, "build", None)):
        raise ValueError("AUTHOR_SCRIPT_BUILD_REQUIRED")
    return module


def _iter_shapes(shapes: Any) -> list[Any]:
    all_shapes: list[Any] = []
    for shape in shapes:
        all_shapes.append(shape)
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            all_shapes.extend(_iter_shapes(shape.shapes))
    return all_shapes


def _native_editability(pptx_path: Path) -> dict[str, Any]:
    prs = Presentation(str(pptx_path))
    bindings: list[str] = []
    chart_bindings: list[str] = []
    table_bindings: list[str] = []
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            if not str(getattr(shape, "name", "") or "").startswith("bind:"):
                continue
            binding = str(shape.name)
            if getattr(shape, "has_chart", False):
                chart_bindings.append(binding)
            elif getattr(shape, "has_table", False):
                table_bindings.append(binding)
            elif getattr(shape, "has_text_frame", False):
                bindings.append(binding)
    return {
        "status": "pass" if bindings else "fail",
        "bound_native_text_shapes": len(bindings),
        "bindings": bindings,
        "bound_native_charts": len(chart_bindings),
        "chart_bindings": chart_bindings,
        "bound_native_tables": len(table_bindings),
        "table_bindings": table_bindings,
    }


def _load_source_docs(source_specs: list[str] | None) -> dict:
    """Parse author-provided source specs using the same anchor grammar as IR."""
    from .structural_parser import parse_source

    docs = {}
    for spec in source_specs or []:
        parts = spec.split(":", 2)
        if len(parts) != 3:
            raise ValueError(f"--source expects source_id:type:path, got '{spec}'")
        source_id, source_type, path = parts
        docs[source_id] = parse_source(Path(path).read_text(encoding="utf-8"), source_type, source_id)
    return docs


def _prepare_template_reference(
    manifest: dict,
    template_pptx: str | Path | None,
    output_dir: Path,
    render_engine: str,
) -> tuple[dict | None, dict | None]:
    """Verify and render a manifest-bound template for the author context."""
    template = manifest.get("template")
    if template is None:
        return None, None
    if template.get("use_mode") == "strict":
        # Starting from an empty Presentation cannot prove preservation of a
        # corporate master, theme or protected brand assets.  Refuse to make a
        # strict-reuse claim until a native master-cloning gate exists.
        raise TemplateReferenceError("TEMPLATE_STRICT_REUSE_NOT_IMPLEMENTED")
    if template_pptx is None:
        raise TemplateReferenceError("TEMPLATE_REFERENCE_REQUIRED")
    path = Path(template_pptx)
    if not path.is_file():
        raise TemplateReferenceError("TEMPLATE_REFERENCE_NOT_FOUND")

    from .template_evidence import analyze_template

    evidence = analyze_template(path)
    expected_sha = template.get("source", {}).get("sha256")
    if evidence["source"]["sha256"] != expected_sha:
        raise TemplateReferenceError("TEMPLATE_REFERENCE_HASH_MISMATCH")

    render = run_render_report(
        path, output_dir / "template-reference-render", engine=render_engine,
        expected_slides=evidence["source"]["slide_count"],
    )
    write_render_report(render, output_dir / "template-reference-render-report.json")
    # Reference decks legitimately contain sparse covers, section dividers or
    # near-empty master demonstrations. They remain usable visual evidence as
    # long as a real renderer produced the expected page images; blank-slide
    # heuristics are delivery gates for the new deck, not for a reference.
    if render.get("status") not in {"passed", "partial"}:
        raise TemplateReferenceError("TEMPLATE_REFERENCE_RENDER_FAILED")

    preferred = template.get("preferred_slide_indices", [])
    selected = preferred or list(range(1, evidence["source"]["slide_count"] + 1))
    rendered_by_index = {item.get("slide_index"): item for item in render.get("slides", [])}
    pages = []
    for slide_index in selected:
        rendered = rendered_by_index.get(slide_index)
        if not rendered or not rendered.get("image") or not Path(rendered["image"]).is_file():
            raise TemplateReferenceError("TEMPLATE_REFERENCE_IMAGE_MISSING")
        pages.append({"slide_index": slide_index, "image": rendered["image"]})

    reference = {
        "status": "rendered",
        "use_mode": template["use_mode"],
        "source_sha256": expected_sha,
        "reference_pptx": str(path),
        "selected_slide_indices": selected,
        "rendered_pages": pages,
        "render_report": str(output_dir / "template-reference-render-report.json"),
    }
    context = {
        **reference,
        "generation_policy": template["generation_policy"],
        "evidence": template["evidence"],
        "review": template["review"],
    }
    return context, reference


def run_bespoke_author(
    script_path: str | Path,
    manifest: dict,
    ir: dict,
    output_dir: str | Path,
    *,
    template_pptx: str | Path | None = None,
    source_specs: list[str] | None = None,
    render_engine: str = "auto",
) -> dict:
    """Run a Bespoke author script from blank and apply non-visual gates."""
    output = Path(output_dir)
    script = Path(script_path)
    if manifest.get("status") != "ready_for_authoring":
        return _reject(output, "AUTHORING_MANIFEST_NOT_READY")
    if manifest.get("route") != "bespoke":
        return _reject(output, "BESPOKE_ROUTE_REQUIRED")
    if manifest.get("source_ir_sha256") != _canonical_sha256(ir):
        return _reject(output, "AUTHORING_MANIFEST_IR_HASH_MISMATCH")
    locked_narrative = manifest.get("narrative_contract")
    if manifest.get("delivery_scope") == "complete_deck" and locked_narrative is None:
        return _reject(output, "AUTHORING_NARRATIVE_CONTRACT_REQUIRED")
    if locked_narrative is not None:
        if not isinstance(locked_narrative, dict):
            return _reject(output, "AUTHORING_NARRATIVE_CONTRACT_HASH_MISMATCH")
        narrative_payload = {
            key: value for key, value in locked_narrative.items() if key != "sha256"
        }
        if locked_narrative.get("sha256") != _canonical_sha256(narrative_payload):
            return _reject(output, "AUTHORING_NARRATIVE_CONTRACT_HASH_MISMATCH")
    if not script.is_file():
        return _reject(output, "AUTHOR_SCRIPT_NOT_FOUND", script=str(script))

    pages = manifest.get("pages", [])
    planned_pages = manifest.get("page_budget", {}).get("planned_pages")
    slide_ids = [page.get("slide_id") for page in pages if page.get("slide_id")]
    if not isinstance(planned_pages, int) or planned_pages != len(slide_ids):
        return _reject(output, "AUTHORING_MANIFEST_PAGE_PLAN_INVALID")
    try:
        source_docs = _load_source_docs(source_specs)
    except (OSError, ValueError) as exc:
        return _reject(output, "AUTHORING_SOURCE_LOAD_FAILED", error=str(exc))
    has_tables = any(block.get("role") == "table" for slide in ir.get("slides", [])
                     for block in slide.get("blocks", []))
    if has_tables and not source_docs:
        return _reject(output, "TABLE_SOURCE_DOCUMENTS_REQUIRED")

    try:
        template_context, template_reference = _prepare_template_reference(
            manifest, template_pptx, output, render_engine,
        )
    except TemplateReferenceError as exc:
        return _reject(output, str(exc))

    try:
        module = _load_author_module(script)
        presentation = Presentation()
        context = {
            "ir": ir,
            "manifest": manifest,
            "slide_ids": slide_ids,
            "output_dir": str(output),
            "started_from_blank": True,
            "template": template_context,
            "source_documents": {source_id: doc.to_dict() for source_id, doc in source_docs.items()},
        }
        module.build(presentation, context)
    except Exception as exc:  # preserve a concise, machine-readable delivery failure
        return _reject(output, "AUTHOR_SCRIPT_FAILED", error=f"{type(exc).__name__}: {exc}")

    output.mkdir(parents=True, exist_ok=True)
    deck = output / "deck.pptx"
    presentation.save(str(deck))
    actual_pages = len(Presentation(str(deck)).slides)
    binding = verify_authored_page(
        deck, ir, slide_ids=slide_ids, require_full_coverage=True, source_docs=source_docs,
    )
    required_bindings = {
        binding
        for page in pages
        for binding in page.get("required_bindings", [])
        if isinstance(binding, str)
    }
    missing_required_bindings = sorted(required_bindings - set(binding.get("valid_bindings", [])))
    inspection = run_structural_inspection(deck, ppt_ir=ir, route="bespoke")
    structural_blockers = [
        issue for issue in inspection.get("issues", [])
        if issue.get("severity") in {"fatal", "error"}
    ]
    native_editability = _native_editability(deck)
    visual_quality = inspect_bespoke_visual_quality(deck)
    page_count_ok = actual_pages == planned_pages

    if (
        binding["status"] == "pass"
        and not missing_required_bindings
        and not structural_blockers
        and page_count_ok
    ):
        render = run_render_report(
            deck, output / "render", engine=render_engine, expected_slides=actual_pages,
        )
        write_render_report(render, output / "render-report.json")
    else:
        render = {"status": "not_run", "reason": "pre_render_gate_failed"}

    ok = (
        binding["status"] == "pass"
        and not missing_required_bindings
        and not structural_blockers
        and page_count_ok
        and native_editability["status"] == "pass"
        and visual_quality["status"] == "pass"
        and render.get("status") == "passed"
    )
    result: dict[str, Any] = {
        "ok": ok,
        "status": "visual_unreviewed" if ok else "rejected",
        "started_from_blank": True,
        "deck": str(deck),
        "planned_pages": planned_pages,
        "actual_pages": actual_pages,
        "page_count_ok": page_count_ok,
        "binding": binding,
        "missing_required_bindings": missing_required_bindings,
        "structural_inspection": inspection,
        "structural_blockers": structural_blockers,
        "native_editability": native_editability,
        "visual_quality": visual_quality,
        "render": render,
        "visual_review": "required_not_run",
    }
    if template_reference is not None:
        result["template_reference"] = template_reference
    if missing_required_bindings:
        result["code"] = "AUTHORING_BINDINGS_INCOMPLETE"
    elif visual_quality["status"] != "pass":
        result["code"] = "BESPOKE_VISUAL_FLOOR_FAILED"
    result["report"] = _write_report(output, result)
    return result
